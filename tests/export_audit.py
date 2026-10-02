"""Checks what tests/export_audit.js collected: every PDF drawn to a picture
(first page) and searched for text running off the page or out of its
cell; every spreadsheet checked for print set-up, headings, widths and
numbers kept as numbers.
   python3 tests/export_audit.py /tmp/claude-0/exp
"""
import glob, os, sys, subprocess, json, re
from openpyxl import load_workbook
D = sys.argv[1] if len(sys.argv) > 1 else "/tmp/claude-0/exp"
bad = []
try:
    import pdfplumber
except ImportError:
    pdfplumber = None
for f in sorted(glob.glob(D + "/*.pdf")):
    n = os.path.basename(f)
    png = f[:-4]
    subprocess.run(["pdftoppm", "-r", "50", "-png", "-f", "1", "-l", "1", f, png], capture_output=True)
    if pdfplumber:
        with pdfplumber.open(f) as pdf:
            pages = len(pdf.pages)
            probs = []
            for pg in pdf.pages[:3]:
                W, H = pg.width, pg.height
                for w in pg.extract_words():
                    if w["x1"] > W - 8 or w["x0"] < 8 or w["bottom"] > H - 4:
                        probs.append(f"'{w['text'][:20]}' at edge")
                # characters overlapping other characters (text drawn over text)
                ch = sorted(pg.chars, key=lambda c: (round(c["top"]), c["x0"]))
                over = 0
                for a, b in zip(ch, ch[1:]):
                    if abs(a["top"] - b["top"]) < 1 and b["x0"] < a["x1"] - 1.2 and a["text"].strip() and b["text"].strip():
                        over += 1
                if over > 3:
                    probs.append(f"{over} overlapping characters")
            print(f"{'PDF ' if not probs else 'PDF!'} {n[:95]}  {pages}p {int(W)}x{int(H)} {'; '.join(sorted(set(probs))[:4])}")
            if probs: bad.append(n)
for f in sorted(glob.glob(D + "/*.xlsx")):
    n = os.path.basename(f)
    try:
        wb = load_workbook(f)
    except Exception as e:
        print("XLS! cannot open", n, e); bad.append(n); continue
    probs = []
    for ws in wb.worksheets:
        ps = ws.page_setup
        fit = ws.sheet_properties.pageSetUpPr.fitToPage if ws.sheet_properties.pageSetUpPr else False
        if not fit or (ps.fitToWidth not in (1, None)):
            probs.append(f"{ws.title}: not fitted to page width")
        if not ws.print_title_rows:
            probs.append(f"{ws.title}: heading row not repeated")
        # numbers kept as text
        txtnum = 0
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 300)):
            for c in row:
                if isinstance(c.value, str) and re.fullmatch(r"-?[\d,]+\.\d{2}", c.value.strip()):
                    txtnum += 1
        if txtnum > 2:
            probs.append(f"{ws.title}: {txtnum} amounts stored as text")
        widths = {k: v.width for k, v in ws.column_dimensions.items() if v.width}
        if not widths:
            probs.append(f"{ws.title}: no column widths set")
    print(f"{'XLS ' if not probs else 'XLS!'} {n[:95]}  {'; '.join(probs[:4])}")
    if probs: bad.append(n)
print(f"\n{len(bad)} file(s) with something to look at")
