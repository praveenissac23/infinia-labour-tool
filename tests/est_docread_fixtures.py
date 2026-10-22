"""One document of each kind an Estimation form can read, laid out the way offices send them.

    python3 tests/est_docread_fixtures.py /tmp/docs   -> prints {kind: {file, expect}} as JSON
"""
import json, os, sys
from docx import Document
import openpyxl
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, Spacer
from reportlab.lib.styles import getSampleStyleSheet

out = sys.argv[1] if len(sys.argv) > 1 else "/tmp/docs"
os.makedirs(out, exist_ok=True)
st = getSampleStyleSheet()
P = lambda t, s="Normal": Paragraph(t, st[s])
res = {}


def pdf(name, el):
    SimpleDocTemplate(os.path.join(out, name), pagesize=A4).build(el)
    return os.path.join(out, name)


def docx(name, paras, table=None):
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=0, cols=len(table[0]))
        for r in table:
            c = t.add_row().cells
            for i, v in enumerate(r):
                c[i].text = str(v)
    d.save(os.path.join(out, name))
    return os.path.join(out, name)


# letter of award -> project
res["project"] = {"file": pdf("LOA Al Barari villa.pdf", [
    P("LETTER OF AWARD", "Title"), P("Ref: RAM/LOA/2026/031"), P("Date: 15 September 2026"),
    P("To: M/s Infinia Contracting LLC, P.O. Box 85567, Dubai"),
    Table([["Project:", "Construction of G+1 Villa at Plot 6457, Al Barari"], ["Employer:", "Mr. Rashid Al Mansoori"],
           ["Consultant:", "Dimensions Engineering Consultants"], ["Location:", "Plot 6457, Al Barari, Dubai"],
           ["Contract Sum:", "AED 4,850,000.00 (excluding VAT)"], ["Commencement Date:", "01 October 2026"], ["Time for Completion:", "14 months"]]),
    Spacer(1, 10), P("We are pleased to award you the above works. An advance payment of 10% shall be paid against an advance payment bank guarantee. "
                     "Retention of 10% shall be deducted from each interim payment."), P("Yours faithfully, For the Employer")]),
    "expect": {"name": "Construction of G+1 Villa at Plot 6457, Al Barari", "client": "Mr. Rashid Al Mansoori", "consultant": "Dimensions Engineering Consultants",
               "contract_value": 4850000, "adv_pct": 10, "ret_pct": 10, "start_date": "2026-10-01", "end_date": "2027-12-01"}}

# tender invitation -> tender
res["tender"] = {"file": docx("ITT Jumeirah Golf villa.docx", [
    "INVITATION TO TENDER", "Date: 20 October 2026", "Project Name: Design & Build of B+G+1 Villa, Jumeirah Golf Estates",
    "Client: Emaar Development PJSC", "Consultant: Arcadia Engineering Consultants", "Location: Plot JGE-114, Jumeirah Golf Estates, Dubai",
    "You are invited to submit your best lump sum price. Tender submission date: 5th November 2026, 12:00 noon.", "Regards, Tendering Department"]),
    "expect": {"name": "Design & Build of B+G+1 Villa, Jumeirah Golf Estates", "client": "Emaar Development PJSC", "consultant": "Arcadia Engineering Consultants",
               "received": "2026-10-20", "due": "2026-11-05"}}

# a subcontractor's invoice -> invoice received (922: Al Noor Technical Manpower)
res["invoice"] = {"file": pdf("Al Noor invoice ANT-12.pdf", [
    P("AL NOOR TECHNICAL MANPOWER SUPPLY", "Title"), P("TAX INVOICE"),
    Table([["Invoice No:", "ANT/12", "Date:", "18/10/2026"], ["Bill to:", "Infinia Contracting LLC", "Project:", "922"]]), Spacer(1, 8),
    Table([["S.No", "Description", "Qty", "Unit", "Rate", "Amount"], ["1", "Steel fixers - September 2026", "640", "hrs", "25.00", "16,000.00"],
           ["2", "Shuttering carpenters - September 2026", "640", "hrs", "25.00", "16,000.00"],
           ["", "Total (excluding VAT)", "", "", "", "32,000.00"]]), Spacer(1, 6),
    Table([["VAT 5%", "1,600.00"], ["Total amount including VAT", "AED 33,600.00"]])]),
    "expect": {"sub": "Al Noor Technical Manpower", "inv_ref": "ANT/12", "inv_date": "2026-10-18", "inv_amount": 33600, "lines": 2}}

# a subcontractor's variation quotation -> variation to a subcontract
res["subvo"] = {"file": docx("Woodcraft VO quotation.docx", [
    "WOODCRAFT INTERIORS", "Date: 9 October 2026", "Our Ref: WD/VO/26-27/0042",
    "Subject: Additional glass handrail at staircase", "Further to the site instruction we quote as below.",
    "Total amount: AED 12,600.00 (excluding VAT)"],
    [["No", "Description", "Qty", "Unit", "Rate", "Amount"], ["1", "Toughened glass handrail 12mm with SS cap", "14", "m", "900.00", "12,600.00"]]),
    "expect": {"title": "Additional glass handrail at staircase", "amount": 12600, "on_date": "2026-10-09"}}

# a cost proposal -> client variation
res["variation"] = {"file": pdf("CP07 roof pergola.pdf", [
    P("Ref: IC/909/CP/07"), P("Date: 2 October 2026"), P("To: The Engineer, Dimensions Engineering Consultants"),
    P("Subject: Cost proposal - Additional pergola at roof terrace"), Spacer(1, 6),
    Table([["Item", "Description", "Qty", "Unit", "Rate", "Amount"], ["1", "Steel pergola frame, powder coated", "1", "LS", "36,500.00", "36,500.00"],
           ["2", "WPC louvres and fixing", "1", "LS", "12,250.00", "12,250.00"], ["", "Total (excluding VAT)", "", "", "", "48,750.00"]]),
    P("We request your approval to proceed.")]),
    "expect": {"title": "Additional pergola at roof terrace", "submitted_amount": 48750, "submitted": "2026-10-02"}}

# an interim payment certificate -> client payment certificate
wb = openpyxl.Workbook(); ws = wb.active
for r in [["INTERIM PAYMENT CERTIFICATE No. 05"], ["Project", "Villa 909"], ["Period ending", "25/09/2026"], [],
          ["Description", "Amount (AED)"], ["Total value of work done to date", 2345600], ["Material on site", 45000],
          ["Certified to date", 2210000], ["Less retention 10%", -221000], ["Less previous payments", -1650000]]:
    ws.append(r)
wb.save(os.path.join(out, "IPC 05.xlsx"))
res["clientpc"] = {"file": os.path.join(out, "IPC 05.xlsx"),
                   "expect": {"claimed_cum": 2345600, "certified_cum": 2210000, "mos_cum": 45000, "period_to": "2026-09-25"}}

# an RFI -> RFI log
res["log"] = {"file": pdf("RFI-014.pdf", [
    P("REQUEST FOR INFORMATION", "Title"), Table([["RFI No:", "RFI-909-014"], ["Date:", "12 October 2026"], ["Reply required by:", "19 October 2026"]]),
    P("Subject: Clarification on staircase handrail height"), P("Question: Please confirm the handrail height at the double-height void - drawing shows 900 mm, specification says 1100 mm.")]),
    "expect": {"title": "Clarification on staircase handrail height", "ref": "RFI-909-014", "sent": "2026-10-12", "due": "2026-10-19"}}

# a NOC -> approvals
res["approval"] = {"file": pdf("DEWA NOC.pdf", [
    P("Dubai Electricity & Water Authority (DEWA)", "Title"), P("NO OBJECTION CERTIFICATE"),
    Table([["NOC No:", "DEWA/2026/88412"], ["Date:", "7 October 2026"]]), P("Subject: NOC for power connection - Villa 909"),
    P("DEWA has no objection to the above subject to the conditions overleaf.")]),
    "expect": {"title": "NOC for power connection - Villa 909", "ref": "DEWA/2026/88412", "party": "DEWA", "replied": "2026-10-07"}}
print(json.dumps(res))
