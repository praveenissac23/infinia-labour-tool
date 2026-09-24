"""The monthly report: fixed columns, a note against each worker, one PDF.

Run: cd app && DATABASE_URL=sqlite:////tmp/mr.db python3 ../tests/monthly_report_test.py

Management sees the same report every month end. The figures are rebuilt
from whichever cycle is chosen, so an old month comes back on request;
the notes typed against each worker are stored against that cycle, so
it comes back with the words that went out with it. The PDF carries the
title and a Notes column after Adjusted Final Salary, with columns sized
so text wraps where it should and numbers do not.
"""
import sys, io
sys.path.insert(0, '.')
from fastapi.testclient import TestClient
import database, models, auth
models.Base.metadata.create_all(database.engine)
import main

db = database.SessionLocal()
db.add(models.User(username='admin', hashed_password=auth.hash_password('p'),
                   full_name='Praveen', role='admin'))
db.add(models.User(username='site1', hashed_password=auth.hash_password('p'),
                   full_name='Site Engineer', role='site', permissions='attendance'))
db.add(models.Employee(emp_no='D-03', name='AHMAD ALI', trade='DRIVER',
                       total_salary=2500, basic_salary=1500, active=True))
db.add(models.Employee(emp_no='F-722', name='MACHHAFA MONDAL', trade='MASON',
                       total_salary=1350, basic_salary=1000, active=True))
db.add(models.Site(code='914', active=True))
db.add(models.Engineer(name='AKHIL', active=True))
db.commit(); db.close()

c = TestClient(main.app)
def tok(u):
    return {'Authorization': 'Bearer ' + c.post('/auth/login',
            data={'username': u, 'password': 'p'}).json()['access_token']}
H, SITE = tok('admin'), tok('site1')
FAIL = []
def ck(l, ok, x=''):
    print(('PASS ' if ok else 'FAIL ') + l + ('' if ok else f'  [{x}]'))
    if not ok: FAIL.append(l)

AUG, SEP = 'August 2026', 'September 2026'
for cyc, mm, days in ((AUG, '08', range(1, 11)), (SEP, '09', range(1, 6))):
    c.post('/attendance/save', json={'month_year': cyc, 'rows': [
        {'emp_no': e, 'full_date': f'2026-{mm}-{d:02d}', 'am': 'Present', 'pm': 'Present',
         'site': '914', 'engineer': 'AKHIL', 'ot': 0, 'bh': 0, 'comments': ''}
        for e in ('D-03', 'F-722') for d in days]}, headers=H)
sums = {s['emp_no']: s['id'] for s in c.get(f'/summaries/{SEP}', headers=H).json()}
c.post(f"/summaries/{sums['D-03']}/adjustments", json={'description': 'Parking fine for D38023', 'amount': 170, 'is_deduction': True}, headers=H)
c.post(f"/summaries/{sums['D-03']}/adjustments", json={'description': 'Extra Allowance', 'amount': 50, 'is_deduction': False}, headers=H)

# ---- A note against each worker, kept against the cycle ---------------
ck('a cycle with no notes is empty', c.get(f'/reports/monthly-notes/{SEP}', headers=H).json()['notes'] == {})
r = c.post(f'/reports/monthly-notes/{SEP}', json={'emp_no': 'D-03', 'note': 'Fined for parking at 914; recovered this month.'}, headers=H)
ck('a note saves against the worker', r.status_code == 200 and r.json()['notes'].get('D-03', '').startswith('Fined'), r.text[:200])
c.post(f'/reports/monthly-notes/{SEP}', json={'emp_no': 'F-722', 'note': 'ILOE insurance renewed.'}, headers=H)
c.post(f'/reports/monthly-notes/{AUG}', json={'emp_no': 'D-03', 'note': 'Joined on the 1st.'}, headers=H)

sep = c.get(f'/reports/monthly-notes/{SEP}', headers=H).json()['notes']
aug = c.get(f'/reports/monthly-notes/{AUG}', headers=H).json()['notes']
ck('September holds both September notes', set(sep) == {'D-03', 'F-722'}, sep)
ck('August holds only its own', aug == {'D-03': 'Joined on the 1st.'}, aug)
ck('saving one worker leaves the other alone',
   c.post(f'/reports/monthly-notes/{SEP}', json={'emp_no': 'D-03', 'note': 'Rewritten.'}, headers=H).json()['notes'].get('F-722') == 'ILOE insurance renewed.')
ck('an empty note clears the line',
   'D-03' not in c.post(f'/reports/monthly-notes/{SEP}', json={'emp_no': 'D-03', 'note': '  '}, headers=H).json()['notes'])
c.post(f'/reports/monthly-notes/{SEP}', json={'emp_no': 'D-03', 'note': 'Fined for parking at 914; recovered this month.'}, headers=H)
whole = c.post(f'/reports/monthly-notes/{SEP}', json={'notes': {'D-03': 'Whole-page save.', 'F-722': ''}}, headers=H).json()['notes']
ck('the Save button sends the whole page and an emptied box clears', whole == {'D-03': 'Whole-page save.'}, whole)
c.post(f'/reports/monthly-notes/{SEP}', json={'emp_no': 'D-03', 'note': 'Fined for parking at 914; recovered this month.'}, headers=H)
c.post(f'/reports/monthly-notes/{SEP}', json={'emp_no': 'F-722', 'note': 'ILOE insurance renewed.'}, headers=H)
ck('a note with no worker is refused',
   c.post(f'/reports/monthly-notes/{SEP}', json={'note': 'x'}, headers=H).status_code == 400)
ck('a site engineer cannot read the notes',
   c.get(f'/reports/monthly-notes/{SEP}', headers=SITE).status_code in (401, 403))
ck('a site engineer cannot write them',
   c.post(f'/reports/monthly-notes/{SEP}', json={'emp_no': 'D-03', 'note': 'x'}, headers=SITE).status_code in (401, 403))

# ---- The fixed columns rebuild for any cycle ------------------------
MEAS = 'total_salary,ot_hours,bh_hours,adjustments,adjusted_final_salary'
def report(cyc):
    return c.get(f'/reports/custom?month_year={cyc}&data_source=summary&dimensions=emp_no,name&measures={MEAS}', headers=H).json()
sep_r, aug_r = report(SEP), report(AUG)
ck('report headers are the monthly set',
   [x['label'] for x in sep_r['columns']] == ['Employee No', 'Employee Name', 'Total Salary (AED)', 'OT Hours',
                                               'BH Hours', 'Adjustments', 'Adjusted Final Salary (AED)'],
   [x['label'] for x in sep_r['columns']])
d03 = next(r for r in sep_r['rows'] if r['dim_0'] == 'D-03')
ck('September row carries its adjustments', d03['adjustments'] == '-170 Parking fine for D38023, +50 Extra Allowance', d03)
ck('the two cycles give different figures',
   d03['adjusted_final_salary'] != next(r for r in aug_r['rows'] if r['dim_0'] == 'D-03')['adjusted_final_salary'])

# ---- The exports carry the Notes column against each worker ----------
t = c.post('/auth/download-token', headers=H).json()['token']
base = f'/export/{SEP}/custom-report?token={t}&data_source=summary&dimensions=emp_no,name&measures={MEAS}'
pdf = c.get(base + '&format=pdf&monthly=1')
ck('monthly PDF builds', pdf.status_code == 200 and pdf.content[:4] == b'%PDF', pdf.status_code)
ck('named as the monthly report', 'Monthly_Report' in pdf.headers.get('content-disposition', ''))
import warnings; warnings.simplefilter('ignore')
from pypdf import PdfReader
pages = PdfReader(io.BytesIO(pdf.content)).pages
# Wrapped lines come back with newlines inside a sentence, so read it flat.
text = ' '.join(' '.join((p.extract_text() or '').split()) for p in pages)
ck('PDF is titled Monthly Payroll Report', 'Monthly Payroll Report' in text)
ck('PDF has the combined column', 'Adjustments & Notes' in text)
ck('PDF has no separate Adjustments column', ' Adjustments ' not in text.replace('Adjustments & Notes', ''))
ck('no header is broken mid-word', 'Employ ee' not in text and 'H ours' not in text, text[:300])
ck("PDF carries D-03's note", 'recovered this month' in text)
ck("PDF carries F-722's note", 'ILOE insurance renewed' in text)
ck('PDF carries the adjustments', 'Parking fine' in text and 'Extra Allowance' in text)
ck('PDF carries the total', 'TOTAL' in text)
ck('two workers fit on one page', len(pages) == 1, len(pages))

import openpyxl
xl = c.get(base + '&format=excel&monthly=1')
ws = openpyxl.load_workbook(io.BytesIO(xl.content)).active
head = next(([v for v in row if v] for row in ws.iter_rows(values_only=True) if row and row[0] == 'Employee No'), [])
ck('Excel ends with the combined column', head and head[-1] == 'Adjustments & Notes', head)
cells = [str(v) for row in ws.iter_rows(values_only=True) for v in row if v]
ck('Excel carries the notes', any('ILOE insurance renewed' in v for v in cells))

plain = c.get(base + '&format=pdf')
ptext = ' '.join((p.extract_text() or '') for p in PdfReader(io.BytesIO(plain.content)).pages)
ck('the ordinary export has no Notes column', plain.status_code == 200 and 'ILOE insurance' not in ptext)

# ---- Preview: the same sheet, before it is a file ---------------------
# A report that can only be checked by downloading it, opening it and
# then deleting it again does not get checked.
import sys as _sys
_sys.path.insert(0, '.')
import export_web
pv = c.get(base.replace('/custom-report?', '/custom-report/view?') + '&monthly=1')
ck('the monthly report previews on screen', pv.status_code == 200, pv.status_code)
body = pv.text
ck('the preview carries the Infinia logo', 'data:image/png;base64,' in body)
ck('and the letterhead', 'INFINIA CONTRACTING LLC' in body)
ck('and the report title the file uses', 'Monthly Payroll Report' in body)
ck('with PDF, Excel and Print on it',
   'format=pdf' in body and 'format=excel' in body and 'window.print()' in body)
ck('the figures are the ones in the file',
   'ILOE insurance renewed' in body and 'Parking fine' in body)
ck('and the total is on it', 'TOTAL' in body)

# The preview and the file must agree which way up the page goes.
from pypdf import PdfReader as _PR
box = _PR(io.BytesIO(c.get(base + '&format=pdf&monthly=1').content)).pages[0].mediabox
pdf_turn = 'landscape' if box.width > box.height else 'portrait'
pv_turn = 'portrait' if 'width:210mm' in body else 'landscape'
ck('the preview is the same way up as the printed copy', pdf_turn == pv_turn,
   f'{pdf_turn} vs {pv_turn}')

# And the spreadsheet comes off a printer as a report, not a mess.
ws2 = openpyxl.load_workbook(io.BytesIO(c.get(base + '&format=excel&monthly=1').content)).active
ck('the spreadsheet is set to A4', int(ws2.page_setup.paperSize or 0) == 9,
   ws2.page_setup.paperSize)
ck('fitted to one page wide',
   bool(ws2.page_setup.fitToWidth) and bool(ws2.sheet_properties.pageSetUpPr.fitToPage))
ck('with the heading row repeated on later pages', bool(ws2.print_title_rows),
   ws2.print_title_rows)

print('\n' + ('ALL PASS' if not FAIL else f'{len(FAIL)} FAILED: ' + '; '.join(FAIL)))
sys.exit(1 if FAIL else 0)
