"""Total Salary is the figure in Master Data. Payable Salary is what
this cycle came to. They are never called by the same name.

Run: cd app && DATABASE_URL=sqlite:////tmp/st.db python3 ../tests/salary_terms_test.py

Master Data holds a worker's Total Salary - basic plus fixed allowance,
the contractual monthly figure. The live card, the printed card and the
reports then showed a second "Total Salary" underneath Basic Pay that
was really the share earned for the days paid this cycle, and the two
figures never matched. That second one is Payable Salary everywhere now,
and the reports offer the Master Data figure under its own name.
"""
import sys, io
sys.path.insert(0, '.')
from fastapi.testclient import TestClient
import database, models, auth
models.Base.metadata.create_all(database.engine)
import main

CYCLE = 'September 2026'
db = database.SessionLocal()
db.add(models.User(username='admin', hashed_password=auth.hash_password('p'),
                   full_name='Administrator', role='admin'))
db.add(models.Employee(emp_no='F-764', name='MOHAMMED ZUBAIR', trade='MASON',
                       total_salary=3000, basic_salary=2000, active=True))
db.add(models.Site(code='914', active=True))
db.add(models.Engineer(name='AKHIL', active=True))
db.commit(); db.close()

c = TestClient(main.app)
H = {'Authorization': 'Bearer ' + c.post('/auth/login',
     data={'username': 'admin', 'password': 'p'}).json()['access_token']}
FAIL = []
def ck(l, ok, x=''):
    print(('PASS ' if ok else 'FAIL ') + l + ('' if ok else f'  [{x}]'))
    if not ok: FAIL.append(l)

# Ten paid days of a 3,000 salary: payable = 10 x 100 = 1,000.
c.post('/attendance/save', json={'month_year': CYCLE, 'rows': [
    {'emp_no': 'F-764', 'full_date': f'2026-09-{d:02d}', 'am': 'Present', 'pm': 'Present',
     'site': '914', 'engineer': 'AKHIL', 'ot': 0, 'bh': 0, 'comments': ''}
    for d in range(1, 11)]}, headers=H)

card = c.get(f'/live-card/F-764/{CYCLE}', headers=H).json()['summary']
ck('Total Salary on the card is the Master Data figure', card['total_salary'] == 3000, card['total_salary'])
ck('Payable Salary is what ten days came to', card['total_salary_component'] == 1000,
   card['total_salary_component'])
ck('and they are different numbers', card['total_salary'] != card['total_salary_component'])

# The report offers both, under names that cannot be confused.
cat = c.get('/reports/builder-catalog', headers=H).json()['summary']['measures']
labels = {m['key']: m['label'] for m in cat}
ck('report offers Total Salary as the Master Data figure', labels.get('total_salary') == 'Total Salary (AED)', labels.get('total_salary'))
ck('report calls the cycle figure Payable Salary', labels.get('total_salary_component') == 'Payable Salary (AED)',
   labels.get('total_salary_component'))
ck('nothing else in the report is called Total Salary',
   [k for k, v in labels.items() if 'Total Salary' in v] == ['total_salary'],
   [k for k, v in labels.items() if 'Total Salary' in v])

r = c.get(f'/reports/custom?month_year={CYCLE}&data_source=summary&dimensions=emp_no'
          f'&measures=total_salary,total_salary_component,final_salary', headers=H).json()
row = r['rows'][0]
ck('report row: Total Salary 3,000', row['total_salary'] == 3000, row)
ck('report row: Payable Salary 1,000', row['total_salary_component'] == 1000, row)
ck('report headers read Total Salary / Payable Salary / Final Salary',
   [x['label'] for x in r['columns']][1:] == ['Total Salary (AED)', 'Payable Salary (AED)', 'Final Salary (AED)'],
   [x['label'] for x in r['columns']])

# The printed salary card and the Excel say the same.
tok = c.post('/auth/download-token', headers=H).json()['token']
import openpyxl
xl = c.get(f'/export/{CYCLE}/excel?token={tok}')
ck('salary cards export to Excel', xl.status_code == 200, xl.status_code)
wb = openpyxl.load_workbook(io.BytesIO(xl.content))
cells = [str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v]
ck('printed card says Payable Salary', any('Payable Salary' in v for v in cells))
ck('printed card no longer labels the cycle figure Total Salary',
   not any(v.strip() == 'Total Salary' for v in cells), [v for v in cells if 'Total Salary' in v][:3])

pdf = c.get(f'/export/{CYCLE}/pdf?token={tok}')
ck('salary cards export to PDF', pdf.status_code == 200 and pdf.content[:4] == b'%PDF', pdf.status_code)

rep = c.get(f'/export/{CYCLE}/custom-report?token={tok}&data_source=summary&dimensions=emp_no'
            f'&measures=total_salary,total_salary_component&format=excel')
ws = openpyxl.load_workbook(io.BytesIO(rep.content)).active
# The header row sits under a title block, so find it rather than assume row 1.
head = next(([v for v in row if v] for row in ws.iter_rows(values_only=True)
             if row and row[0] == 'Employee No'), [])
ck('report Excel headers carry both names', head[1:] == ['Total Salary (AED)', 'Payable Salary (AED)'], head)

print('\n' + ('ALL PASS' if not FAIL else f'{len(FAIL)} FAILED: ' + '; '.join(FAIL)))
sys.exit(1 if FAIL else 0)
