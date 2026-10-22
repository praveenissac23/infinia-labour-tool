"""Three agreements laid out nothing like each other (and nothing like Infinia's own) for the
"Read from agreement" test: a Word subcontract with a BOQ table, an Excel LPO and a PDF purchase order.

    python3 tests/est_agreement_fixtures.py /tmp/agr      -> prints the expected figures as JSON
"""
import json, os, sys
out = sys.argv[1] if len(sys.argv) > 1 else "/tmp/agr"
os.makedirs(out, exist_ok=True)
exp = {}

# ---- 1. Word: MEP subcontract, re-measurable, bills 1 and 2 with sub-totals ----
from docx import Document
d = Document()
d.add_heading("SUB-CONTRACT AGREEMENT FOR MEP WORKS", 0)
d.add_paragraph("Agreement No.: IC/912/SCA/04/26")
d.add_paragraph("Date: 05/09/2026")
d.add_paragraph("This agreement is made between M/s. Infinia Contracting LLC, P.O. Box 85567, Dubai (the Main Contractor) and "
                "M/s. Al Raida Electromechanical LLC, P.O. Box 1234, Sharjah (hereinafter called the Sub-Contractor).")
d.add_paragraph("Scope of Works: Supply and installation of electrical and plumbing works for the villa as per drawings.")
d.add_paragraph("Contract Sum: AED 1,250,000.00 (exclusive of VAT)")
d.add_paragraph("This is a re-measurable contract based on actual quantities executed at the agreed unit rates.")
d.add_paragraph("Payment terms: An advance payment of 10% shall be paid against a security cheque. Interim payments monthly against certified progress.")
d.add_paragraph("Retention money of 5% will be held from each interim payment until handover.")
rows = [("BILL 1 - ELECTRICAL WORKS", None, None, None, None, None),
        ("1.1", "Main distribution board MDB with meter panel", "No", "1", "85,000.00", "85,000.00"),
        ("1.2", "Power cabling 4C x 70 sq.mm XLPE", "m", "250", "420.00", "105,000.00"),
        ("1.3", "Lighting points complete with wiring", "Pt", "600", "350.00", "210,000.00"),
        ("1.4", "Power sockets 13A twin", "No", "400", "275.00", "110,000.00"),
        ("", "Sub-total Bill 1", "", "", "", "510,000.00"),
        ("BILL 2 - PLUMBING WORKS", None, None, None, None, None),
        ("2.1", "PPR water supply piping 25 mm", "m", "1,200", "95.00", "114,000.00"),
        ("2.2", "uPVC drainage piping 110 mm", "m", "800", "180.00", "144,000.00"),
        ("2.3", "Sanitary fixtures installation", "Set", "40", "6,000.00", "240,000.00"),
        ("2.4", "Water pumps and tanks complete", "LS", "1", "242,000.00", "242,000.00"),
        ("", "Sub-total Bill 2", "", "", "", "740,000.00"),
        ("", "GRAND TOTAL (excluding VAT)", "", "", "", "1,250,000.00")]
t = d.add_table(rows=1, cols=6)
for i, h in enumerate(["Item", "Description", "Unit", "Qty", "Rate (AED)", "Amount (AED)"]):
    t.rows[0].cells[i].text = h
for r in rows:
    c = t.add_row().cells
    if r[1] is None:
        c[0].merge(c[5]); c[0].text = r[0]
    else:
        for i, v in enumerate(r):
            c[i].text = v
d.add_paragraph("Signed for and on behalf of the Sub-Contractor: ________________")
d.save(os.path.join(out, "912 MEP subcontract.docx"))
exp["docx"] = {"name": "Al Raida Electromechanical LLC", "agreement": "IC/912/SCA/04/26", "date": "2026-09-05", "value": 1250000.0,
               "type": "Unit rate re-measurement contract", "adv": 10, "ret": 5, "items": 8, "bills": {"1": 510000.0, "2": 740000.0}}

# ---- 2. Excel: an LPO with its lines ----
import openpyxl
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "LPO"
for r in [["LOCAL PURCHASE ORDER"], [], ["LPO No.", "LPO/913/0045"], ["Date", "12-Oct-2026"], ["Supplier", "M/S Gulf Shoring Equipment Rental LLC"],
          ["Project", "913 - Villa"], ["Subject", "Shoring works"], [], ["S.No", "Description", "Qty", "Unit", "Rate", "Amount"],
          [1, "Sheet piles supply and driving", 120, "m", 650, 78000], [2, "Walers and struts", 1, "LS", 24500, 24500],
          [3, "Dewatering for 30 days", 30, "Day", 850, 25500], [None, "Total", None, None, None, 128000],
          [], ["Payment: 100% within 30 days after completion. Retention 10% for 6 months."]]:
    ws.append(r)
wb.save(os.path.join(out, "913 shoring LPO.xlsx"))
exp["xlsx"] = {"name": "Gulf Shoring Equipment Rental LLC", "lpo": "LPO/913/0045", "date": "2026-10-12", "value": 128000.0, "ret": 10, "items": 3}

# ---- 3. PDF: a two-column purchase order with no "Contract value" label ----
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, Spacer
from reportlab.lib.styles import getSampleStyleSheet
st = getSampleStyleSheet()
doc = SimpleDocTemplate(os.path.join(out, "915 painting work order.pdf"), pagesize=A4)
lines = [["1", "Internal walls - emulsion, 2 coats on primer", "1,850", "m2", "18.00", "33,300.00"],
         ["2", "Ceilings - emulsion, 2 coats", "920", "m2", "16.50", "15,180.00"],
         ["3", "External walls - textured paint system", "1,100", "m2", "32.00", "35,200.00"],
         ["4", "Wood doors - lacquer finish", "24", "Nos", "450.00", "10,800.00"]]
el = [Paragraph("WORK ORDER", st["Title"]),
      Table([["Work Order Ref:", "IC/915/SCA/03/26", "Issued on:", "October 3, 2026"],
             ["Subcontractor:", "M/S Royal Everest Painting Contracting", "Project:", "915 - Sobha Hartland villa"]]),
      Spacer(1, 12), Paragraph("Painting works for the villa as per the finishes schedule. Advance payment: nil. Retention of 5% from each payment.", st["Normal"]),
      Spacer(1, 12),
      Table([["Item", "Description", "Qty", "Unit", "Rate", "Amount"]] + lines + [["", "Total amount (excl. VAT)", "", "", "", "94,480.00"]]),
      Spacer(1, 12), Paragraph("Terms: progressive payment monthly against measured work.", st["Normal"])]
doc.build(el)
exp["pdf2"] = {"name": "Royal Everest Painting Contracting", "agreement": "IC/915/SCA/03/26", "date": "2026-10-03", "value": 94480.0, "ret": 5, "items": 4}
print(json.dumps(exp))
