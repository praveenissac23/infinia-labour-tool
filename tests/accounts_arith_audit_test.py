"""Accounts & purchasing - an independent arithmetic audit.

    cd app && DATABASE_URL=sqlite:////tmp/acc_audit.db INFINIA_NO_SCHED=1 python3 ../tests/accounts_arith_audit_test.py

Every money figure the accounts screens show is worked out again here from the raw table rows
(read with sqlite3, not through the app's models or helpers), in exact decimals rounded 2dp half-up,
and compared with what the app returns - on its API (the screen) and, where it gives one, its Excel
(formulas recalculated by LibreOffice when it is installed) and its PDF text.

  PDC tracker      summary cards, the month grid (row totals, month totals, grand total, overdue)
  Security chqs    held / due / soon amounts
  Petty cash       every book, every month + full history + a date range: brought forward, the
                   running balance line by line, received / paid, closing, by-site, the PRO visa
                   money held; month n closing = month n+1 brought forward; Excel balance column
  Bank             paid / received per month; client payments tied to Project payments (one line
                   each, amount = the amount excl. VAT) and to a receipt (same amount)
  Invoices         subtotal / VAT / total against the lines (VAT per line, 2dp); list totals;
                   received and balance per invoice from the receipts; register Excel totals
  Receipts         allocations never over the receipt nor over an invoice; totals; PAID stamps
  Cash register    brought forward, running balance, closing; Excel
  Project payments received / remaining / % / status per scope, project totals, company totals,
                   a month period; Excel totals
  Project expense  LPO (net of discount and VAT), petty cash (net of claimable VAT), bank expenses,
                   per project and the totals
  LPOs             each order's total on the register, the purchase report, the Zoho list and the
                   printed order

On an empty database it first enters its own data through the API (awkward figures on purpose:
half-fils VAT, discounts, cheques dated later, several months). On a copy of the full database it
audits what is there.
"""
import io, os, sys, json, re, sqlite3, subprocess, tempfile, shutil
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP

sys.path.insert(0, os.getcwd())
from fastapi.testclient import TestClient
import main, models, auth  # noqa: E402

main._add_missing_columns()
c = TestClient(main.app); c.__enter__()

DBFILE = os.environ["DATABASE_URL"].split("sqlite:///", 1)[1]
FAILS, NOTES = [], []


def ck(name, ok, extra=""):
    print(("PASS " if ok else "FAIL ") + name + ("" if ok else f"  -> {extra}"))
    if not ok:
        FAILS.append(name)


def note(s):
    NOTES.append(s); print("NOTE " + s)


# ---- exact money ----------------------------------------------------------------------------------
def D(v):
    return Decimal(str(v or 0))


def r2(v):
    """2dp, half up - how money is rounded on paper."""
    return D(v).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def same(a, b, tol=Decimal("0.005")):
    return abs(D(a) - D(b)) <= tol


TODAY = (datetime.now(timezone.utc) + timedelta(hours=4)).date()     # Dubai


def iso(d):
    return d.isoformat()


def pd(s):
    return date.fromisoformat(str(s)[:10]) if s else None


def month_bounds(d):
    a = d.replace(day=1)
    b = date(a.year + (a.month == 12), a.month % 12 + 1, 1) - timedelta(days=1)
    return a, b


def add_months(d, n):
    y, m = divmod(d.month - 1 + n, 12)
    return date(d.year + y, m + 1, 1)


def raw(sql, *a):
    cx = sqlite3.connect(DBFILE); cx.row_factory = sqlite3.Row
    try:
        return [dict(r) for r in cx.execute(sql, a).fetchall()]
    finally:
        cx.close()


def has_table(t):
    return bool(raw("select name from sqlite_master where type='table' and name=?", t))


# ---- logins ---------------------------------------------------------------------------------------
def login(u, p):
    r = c.post("/auth/login", data={"username": u, "password": p})
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["access_token"]}


ADMIN_PW = os.environ.get("ADMIN_PW", "changeme123")
H = login("admin", ADMIN_PW)


def get(url, **kw):
    r = c.get(url, headers=H, **kw)
    assert r.status_code == 200, f"{url}: {r.status_code} {r.text[:300]}"
    return r.json()


def dl():
    return c.post("/auth/download-token", headers=H).json()["token"]


def fetch(url):
    sep = "&" if "?" in url else "?"
    r = c.get(url + sep + "token=" + dl())
    assert r.status_code == 200, f"{url}: {r.status_code} {r.text[:300]}"
    return r.content


# ---- Excel and PDF readers ------------------------------------------------------------------------
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


def xl(content):
    """The workbook with its formulas worked out (LibreOffice) - or None when it is not installed."""
    import openpyxl
    if not SOFFICE:
        return None
    d = tempfile.mkdtemp()
    try:
        src = os.path.join(d, "in.xlsx"); open(src, "wb").write(content)
        out = os.path.join(d, "out"); os.mkdir(out)
        subprocess.run([SOFFICE, "--headless", "--norestore", "-env:UserInstallation=file://" + d + "/lo",
                        "--convert-to", "xlsx", "--outdir", out, src], capture_output=True, timeout=120)
        f = os.path.join(out, "in.xlsx")
        return openpyxl.load_workbook(f, data_only=True) if os.path.exists(f) else None
    finally:
        pass


def xl_formulas(content):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(content))


def pdf_text(content):
    try:
        import pdfplumber
    except ImportError:
        return None
    with pdfplumber.open(io.BytesIO(content)) as p:
        return "\n".join((pg.extract_text() or "") for pg in p.pages)


def money(v):
    return f"{D(v):,.2f}"


# ===================================================================================================
# Seed (empty database only)
# ===================================================================================================
def seed():
    """Awkward figures, entered the way the office enters them."""
    def post(url, body, method="POST"):
        r = c.request(method, url, headers=H, json=body)
        assert r.status_code == 200, f"{url}: {r.status_code} {r.text[:300]}"
        return r.json()
    m0 = TODAY.replace(day=1)
    m1 = add_months(m0, -1)
    m2 = add_months(m0, -2)
    d = lambda a, n: iso(a + timedelta(days=n))
    # petty cash in every book, three months, VAT on some lines, a site on some
    for book in ("site", "pro", "office", "naveen", "praveen"):
        post("/store/petty-cash", {"book": book, "date": d(m2, 2), "description": "Cash from office", "received": 5000})
        post("/store/petty-cash", {"book": book, "date": d(m2, 5), "description": "Cement 906", "paid": 1234.55, "site": "906", "vat": True})
        post("/store/petty-cash", {"book": book, "date": d(m1, 0), "description": "Water", "paid": 10.10, "vat": True})
        post("/store/petty-cash", {"book": book, "date": d(m1, 27), "description": "Diesel 915", "paid": 333.33})
        post("/store/petty-cash", {"book": book, "date": iso(TODAY), "description": "Top up", "received": 0.07})
        post("/store/petty-cash", {"book": book, "date": iso(TODAY), "description": "Tea", "paid": 0.01})
    # a bill whose VAT the chief accountant marked "not claimable" (set on Send to Zoho): the VAT is a cost
    import database as _dbm
    _d = _dbm.SessionLocal()
    z = _d.query(models.PettyCash).filter(models.PettyCash.book == "site", models.PettyCash.description == "Cement 906").first()
    z.vat_claim = "zero"; _d.commit()
    # PRO visas: one ongoing (payments held), one completed today whose total line was re-dated into last month
    v1 = models.VisaFile(person="Audit Ongoing", started_on=m1, status="ongoing"); _d.add(v1); _d.flush()
    _d.add(models.VisaPayment(visa_file_id=v1.id, on_date=m1 + timedelta(days=2), amount=410.55, step="Medical"))
    v2 = models.VisaFile(person="Audit Done", started_on=m2, status="completed", completed_on=TODAY); _d.add(v2); _d.flush()
    _d.add(models.VisaPayment(visa_file_id=v2.id, on_date=m2 + timedelta(days=3), amount=1000.10, step="Entry permit"))
    _d.add(models.VisaPayment(visa_file_id=v2.id, on_date=m1 + timedelta(days=4), amount=0.35, step="EID"))
    tl = models.PettyCash(book="pro", on_date=m1 + timedelta(days=20), description="Audit Done - visa expenses (2 payments)",
                          received=0, paid=1000.45, visa_file_id=v2.id); _d.add(tl); _d.flush()
    v2.posted_id = tl.id
    _d.commit(); _d.close()
    # PDCs: overdue, this week, a series, one with no date yet, a cleared one, a cancelled one
    post("/employees/pdc", {"payee": "Audit Rent", "cheque_no": "000101", "date": d(TODAY, -3), "amount": 16590.555, "repeat": 4})
    post("/employees/pdc", {"payee": "Audit Steel", "cheque_no": "7", "date": d(TODAY, 5), "amount": 0.015})
    post("/employees/pdc", {"payee": "Audit Steel", "cheque_no": "8", "date": d(TODAY, 9), "amount": 1000.10})
    x = post("/employees/pdc", {"payee": "Audit Blocks", "cheque_no": "9", "date": d(TODAY, -40), "amount": 2500})
    c.post(f"/employees/pdc/{x['ids'][0]}/clear", headers=H)
    y = post("/employees/pdc", {"payee": "Audit Blocks", "cheque_no": "10", "date": d(TODAY, 2), "amount": 777, "status": "cancelled"})
    # invoices: a line whose 5% is a half fils (1234.50 -> 61.725), a 0% line, a cancelled one
    t1 = post("/employees/accounts/invoices", {"kind": "tax", "number": "IC/AUD/1", "date": d(m1, 3), "client": "Audit Client",
                                                 "lines": [{"description": "Stage 1", "amount": 1234.50, "vat": 5},
                                                           {"description": "Stage 2", "amount": 100.10, "vat": 5},
                                                           {"description": "Authority fee", "amount": 500, "vat": 0}]})
    t2 = post("/employees/accounts/invoices", {"kind": "tax", "number": "IC/AUD/2", "date": d(m1, 4), "client": "Audit Client",
                                                 "lines": [{"description": "Stage 3", "amount": 20000.30, "vat": 5}]})
    t3 = post("/employees/accounts/invoices", {"kind": "tax", "number": "IC/AUD/3", "date": d(m1, 5), "client": "Audit Client",
                                                 "lines": [{"description": "Wrong", "amount": 999, "vat": 5}]})
    c.post(f"/employees/accounts/invoices/{t3['id']}/cancel", headers=H, json={"reason": "audit"})
    post("/employees/accounts/invoices", {"kind": "proforma", "number": "PI/AUD/1", "date": d(m1, 1), "client": "Audit Client",
                                          "lines": [{"description": "Advance", "amount": 0.30, "vat": 5}]})
    inv = {r["number"]: r for r in get("/employees/accounts/invoices?kind=tax")["rows"]}
    i1, i2 = inv["IC/AUD/1"], inv["IC/AUD/2"]
    # receipts: one paying invoice 1 in full plus an advance, one part-paying invoice 2 by a later cheque, one cancelled
    post("/employees/accounts/receipts", {"number": "RC/AUD/1", "date": d(m1, 10), "client": "Audit Client", "amount": i1["total"] + 50,
                                          "allocations": [{"invoice_id": i1["id"], "amount": i1["total"]}]})
    post("/employees/accounts/receipts", {"number": "RC/AUD/2", "date": d(m1, 11), "client": "Audit Client", "amount": 5000,
                                          "mode": "PDC", "cheque_date": d(TODAY, 20), "allocations": [{"invoice_id": i2["id"], "amount": 5000}]})
    r3 = post("/employees/accounts/receipts", {"number": "RC/AUD/3", "date": d(m1, 12), "client": "Audit Client", "amount": 1000,
                                               "allocations": [{"invoice_id": i2["id"], "amount": 1000}]})
    c.post(f"/employees/accounts/receipts/{r3['id']}/cancel", headers=H, json={"reason": "bounced"})
    post("/employees/accounts/receipts", {"number": "RC/AUD/4", "date": d(m1, 13), "client": "Audit Client", "amount": 3333.33,
                                          "allocations": [{"invoice_id": i2["id"], "amount": 3333.33}]})
    # project payments: two scopes, payments in two months
    s1 = post("/employees/accounts/projects", {"project": "906", "scope": "CORE & SHELL", "contract": 1000000.005,
                                                "payments": [{"date": d(m2, 3), "amount": 250000.10}, {"date": d(m1, 3), "amount": 0.20}]})
    post("/employees/accounts/projects", {"project": "906", "scope": "MEP", "contract": 333333.33, "payments": []})
    post("/employees/accounts/projects", {"project": "915", "scope": "D&B", "contract": 3, "payments": [{"date": d(m1, 8), "amount": 1}]})
    # bank: a client payment tied to the 906 scope (excl. VAT) and to invoice 2, a supplier cheque, a charge, an expense with VAT
    post("/store/bank", {"kind": "customer", "date": d(m1, 15), "party": "Audit Client", "amount": 10500, "scope_id": s1["id"],
                         "net_amount": 10000, "docs": [{"number": "IC/AUD/2", "amount": 10500}]})
    post("/store/bank", {"kind": "vendor", "date": d(m1, 16), "party": "Audit Supplier", "amount": 4321.09, "mode": "Cheque", "reference": "000555"})
    post("/store/bank", {"kind": "charge", "date": d(m1, 16), "amount": 26.25})
    post("/store/bank", {"kind": "expense", "date": d(m1, 17), "amount": 1050, "site": "906", "account": "Office › Office Supplies",
                         "vat": True, "vat_claim": "claim", "trn": "100000000000003"})
    post("/store/bank", {"kind": "expense", "date": d(m1, 18), "amount": 210, "site": "915", "account": "Office › Office Supplies",
                         "vat": True, "vat_claim": "zero"})
    # cash register
    k = c.post("/employees/accounts/register/unlock", headers=H, json={"password": "Infinia2022!"}).json()["key"]
    RK = {**H, "X-Register-Key": k}
    for kind, dd, amt in (("received", d(m2, 1), 10000.01), ("paid", d(m2, 9), 2500.005), ("paid", d(m1, 2), 0.02), ("received", d(m1, 3), 7.5)):
        assert c.post("/employees/accounts/register", headers=RK, json={"kind": kind, "date": dd, "description": "Audit", "amount": amt}).status_code == 200
    # an LPO with a discount and a 0% line, entered as rows (the order screen is tested elsewhere)
    db = main.SessionLocal() if hasattr(main, "SessionLocal") else __import__("database").SessionLocal()
    sup = models.Supplier(name="Audit Supplier", name_key=main._supplier_key("Audit Supplier"), active=True)
    db.add(sup); db.flush()
    po = models.PurchaseOrder(po_no=990001, ref="IC/LPO/AUD1", order_date=d_(m1, 6), supplier_id=sup.id, supplier_name=sup.name,
                              project_location="906 Villa", discount_pct=7.5, tax_pct=5, status="issued")
    db.add(po); db.flush()
    it = models.StoreItem(code="AUD-REBAR", name="Audit Rebar", unit="nos", item_type="consumable"); db.add(it); db.flush()
    for desc, q, rate, tax in (("Rebar", 3.333, 1234.567, 5), ("Cement", 7, 18.25, 5), ("DM fee", 1, 105, 0)):
        db.add(models.PurchaseOrderLine(order_id=po.id, description=desc, qty=q, rate=rate, tax_pct=tax, unit="nos",
                                        item_id=it.id if desc == "Rebar" else None))
    # issued from the store to 906 (no cost on the movement: valued at the last LPO rate, after its 7.5% discount)
    db.add(models.StoreMovement(item_id=it.id, kind="out", qty=2, from_location="", location="906", moved_on=d_(m1, 9)))
    po2 = models.PurchaseOrder(po_no=990002, ref="IC/LPO/AUD2", order_date=d_(m1, 7), supplier_id=sup.id, supplier_name=sup.name,
                               project_location="915", discount_pct=0, tax_pct=5, status="cancelled")
    db.add(po2); db.flush()
    db.add(models.PurchaseOrderLine(order_id=po2.id, description="Sand", qty=10, rate=50, tax_pct=5, unit="nos"))
    for code in ("906", "915"):
        if not db.query(models.Site).filter(models.Site.code == code).first():
            db.add(models.Site(code=code, project_name=f"Villa {code}", active=True))
    db.commit(); db.close()


def d_(a, n):
    return a + timedelta(days=n)


if not raw("select id from petty_cash limit 1") if has_table("petty_cash") else True:
    print("== empty database: entering the audit's own data")
    seed()

# ===================================================================================================
# PDC tracker
# ===================================================================================================
print("\n== PDC tracker")
pdcs = raw("select * from pdcs")
for p in pdcs:
    p["d"] = pd(p["cheque_date"]); p["amt"] = D(p["amount"])
L = get("/employees/pdc")
s = L["summary"]
pend = [p for p in pdcs if p["status"] == "pending"]
dated = [(( p["d"] - TODAY).days, p) for p in pend if not p["date_tbc"]]
grp = {"overdue": [p for n, p in dated if n < 0], "week": [p for n, p in dated if 0 <= n <= 7],
       "fortnight": [p for n, p in dated if 8 <= n <= 14], "pending": pend, "no_date": [p for p in pend if p["date_tbc"]]}
for k, xs in grp.items():
    want = r2(sum(p["amt"] for p in xs))
    ck(f"PDC card '{k}': {len(xs)} cheques, AED {want:,.2f} ({len(pdcs)} rows)",
       s[k] == len(xs) and same(s[k + "_amount"], want), (s[k], s[k + "_amount"], len(xs), want))
ck("PDC list: one row per cheque, amounts 2dp", len(L["rows"]) == len(pdcs) and all(same(r["amount"], r2(p["amount"])) for r, p in
   zip(sorted(L["rows"], key=lambda r: r["id"]), sorted(pdcs, key=lambda p: p["id"]))))
months_with = sorted({p["d"].replace(day=1) for p in pdcs}) or [TODAY.replace(day=1)]
starts = sorted({months_with[0], TODAY.replace(day=1), add_months(TODAY.replace(day=1), 1)} | set(months_with[:3]))
for st in starts:
    for n in (1, 4, 12):
        for show in ("due", "all"):
            t = get(f"/employees/pdc/tracker?start={st:%Y-%m}&months={n}&show={show}")
            end = add_months(st, n)
            rows = [p for p in pdcs if p["status"] != "cancelled" and (show == "all" or p["status"] == "pending")]
            over = [p for p in rows if p["d"] < st and p["status"] == "pending"]
            inside = [p for p in rows if st <= p["d"] < end]
            pay = {}
            for p in over + inside:
                e = pay.setdefault(p["payee"], {"od": D(0), "cells": [D(0)] * n})
                if p["d"] < st:
                    e["od"] += p["amt"]
                else:
                    e["cells"][(p["d"].year - st.year) * 12 + p["d"].month - st.month] += p["amt"]
            want_od = r2(sum(e["od"] for e in pay.values()))
            want_tot = [r2(sum(e["cells"][i] for e in pay.values())) for i in range(n)]
            want_grand = r2(sum(sum(e["cells"]) + e["od"] for e in pay.values()))
            got = {r["payee"]: r for r in t["rows"]}
            rows_ok = set(got) == set(pay) and all(same(got[k]["overdue"], r2(e["od"])) and all(same(a, r2(b)) for a, b in zip(got[k]["cells"], e["cells"]))
                                                   and same(got[k]["total"], r2(e["od"] + sum(e["cells"]))) for k, e in pay.items())
            tag = f"{st:%b-%y} x{n} {show}"
            ck(f"PDC grid {tag}: {len(pay)} payees - each cell, row total, overdue", rows_ok,
               [(k, got.get(k, {}).get("total"), r2(e["od"] + sum(e["cells"]))) for k, e in pay.items()][:5])
            ck(f"PDC grid {tag}: month totals, overdue {want_od:,.2f}, grand {want_grand:,.2f}",
               same(t["overdue"], want_od) and all(same(a, b) for a, b in zip(t["totals"], want_tot)) and same(t["grand"], want_grand)
               and same(sum(D(r["total"]) for r in t["rows"]), t["grand"], Decimal("0.011")), (t["overdue"], t["totals"], t["grand"], want_tot))
# the papers for the widest grid
st, n = months_with[0], 12
t = get(f"/employees/pdc/tracker?start={st:%Y-%m}&months={n}&show=all")
wb = xl(fetch(f"/export/pdc?start={st:%Y-%m}&months={n}&show=all&format=excel"))
if wb is not None:
    ws = wb["PDC Tracker"]
    last = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "TOTAL (Monthly)")
    ncol = len([c_ for c_ in ws[4] if c_.value])
    ck(f"PDC grid Excel: grand total {t['grand']:,.2f} once formulas are worked out", same(ws.cell(last, ncol).value or 0, t["grand"]),
       ws.cell(last, ncol).value)
for n_ in (4, 12):
    t_ = get(f"/employees/pdc/tracker?start={st:%Y-%m}&months={n_}&show=all")
    txt = pdf_text(fetch(f"/export/pdc?start={st:%Y-%m}&months={n_}&show=all&format=pdf"))
    if txt is not None:
        figs = [money(v) for v in t_["totals"] + [t_["grand"]] if v]
        ck(f"PDC grid PDF ({n_} months): every month total and the grand total printed whole (not broken over two lines)",
           all(f in txt for f in figs), [f for f in figs if f not in txt])
lst = [p for p in pdcs if p["status"] == "pending"]
txt = pdf_text(fetch("/export/pdc?mode=list&which=pending&format=pdf"))
if txt is not None:
    ck(f"PDC list PDF (pending): total {r2(sum(p['amt'] for p in lst)):,.2f}", money(r2(sum(p["amt"] for p in lst))) in txt)

# ===================================================================================================
# Security cheques
# ===================================================================================================
print("\n== Security cheques")
if has_table("security_cheques"):
    sc = raw("select * from security_cheques")
    s = get("/employees/security-cheques")["summary"]
    held = [x for x in sc if x["status"] == "held"]
    due = [x for x in held if x["remind_on"] and pd(x["remind_on"]) <= TODAY]
    soon = [x for x in held if x["remind_on"] and TODAY < pd(x["remind_on"]) <= TODAY + timedelta(days=30)]
    ck(f"Security cheques: held {len(held)} AED {r2(sum(D(x['amount']) for x in held)):,.2f}, due {len(due)}, soon {len(soon)}",
       s["held"] == len(held) and same(s["held_amount"], r2(sum(D(x["amount"]) for x in held))) and s["due"] == len(due)
       and same(s["due_amount"], r2(sum(D(x["amount"]) for x in due))) and same(s["soon_amount"], r2(sum(D(x["amount"]) for x in soon))), s)

# ===================================================================================================
# Petty cash
# ===================================================================================================
print("\n== Petty cash")
pc = raw("select * from petty_cash")
for x in pc:
    x["d"] = pd(x["on_date"]); x["rec"] = D(x["received"]); x["pay"] = D(x["paid"]); x["bk"] = x["book"] or "site"
SIGN = {"site": 1, "pro": 1, "office": 1, "naveen": -1, "praveen": -1}
vf = {v["id"]: v for v in raw("select * from visa_files")} if has_table("visa_files") else {}
vp = raw("select * from visa_payments") if has_table("visa_payments") else []


def petty_expect(book, a, b, end_for_visa):
    sg = SIGN[book]
    rows = sorted([x for x in pc if x["bk"] == book], key=lambda x: (x["d"], x["id"]))
    opening = r2(sg * sum(x["rec"] - x["pay"] for x in rows if x["d"] < a))
    inside = [x for x in rows if a <= x["d"] <= b]
    bal, out = opening, []
    for x in inside:
        bal = r2(bal + sg * (x["rec"] - x["pay"]))
        out.append((x["id"], bal))
    rec, paid = r2(sum(x["rec"] for x in inside)), r2(sum(x["pay"] for x in inside))
    by = {}
    for x in inside:
        if x["pay"]:
            k = x["site"] or "Office / general"
            by[k] = by.get(k, D(0)) + x["pay"]
    held = D(0)
    if book == "pro":
        # a completed visa's total line already in the book by the end date is not held again
        booked = {v["id"] for v in vf.values() if v["posted_id"] and any(x["id"] == v["posted_id"] and x["d"] <= end_for_visa for x in pc)}
        held = r2(sum(D(p["amount"]) for p in vp if pd(p["on_date"]) <= end_for_visa and p["visa_file_id"] in vf and p["visa_file_id"] not in booked and
                      (vf[p["visa_file_id"]]["status"] == "ongoing" or (vf[p["visa_file_id"]]["completed_on"] and pd(vf[p["visa_file_id"]]["completed_on"]) > end_for_visa))))
    return {"opening": opening, "rows": out, "received": rec, "paid": paid, "closing": r2(opening + sg * (rec - paid)),
            "by_site": {k: r2(v) for k, v in by.items()}, "held": held}


nrows = 0
for book in SIGN:
    mine = [x for x in pc if x["bk"] == book]
    months = sorted({x["d"].replace(day=1) for x in mine}) or [TODAY.replace(day=1)]
    allm = []
    mm = months[0]
    while mm <= max(months[-1], TODAY.replace(day=1)):
        allm.append(mm); mm = add_months(mm, 1)
    prev_close = None
    for mm in allm:
        a, b = month_bounds(mm)
        r = get(f"/store/petty-cash?book={book}&month={mm:%Y-%m}")
        e = petty_expect(book, a, b, b)
        nrows += len(r["rows"])
        got_bal = [(x["id"], D(x["balance"])) for x in r["rows"]]
        ck(f"Petty {book} {mm:%b %Y}: b/f {e['opening']:,.2f}, {len(e['rows'])} lines, received {e['received']:,.2f}, paid {e['paid']:,.2f}, "
           f"closing {e['closing']:,.2f}", same(r["opening"], e["opening"]) and same(r["received"], e["received"]) and same(r["paid"], e["paid"])
           and same(r["closing"], e["closing"]), (r["opening"], r["received"], r["paid"], r["closing"]))
        ck(f"Petty {book} {mm:%b %Y}: running balance on every line", [i for i, _ in got_bal] == [i for i, _ in e["rows"]]
           and all(same(g, w) for (_, g), (_, w) in zip(got_bal, e["rows"])) and (not got_bal or same(got_bal[-1][1], r["closing"])),
           [(i, str(g), str(w)) for (i, g), (_, w) in zip(got_bal, e["rows"]) if not same(g, w)][:3])
        ck(f"Petty {book} {mm:%b %Y}: spending by site adds to the paid total",
           {x["site"]: D(x["paid"]) for x in r["by_site"]} == e["by_site"] and same(sum(D(x["paid"]) for x in r["by_site"]), r["paid"]),
           (r["by_site"], e["by_site"]))
        if book == "pro":
            ck(f"Petty pro {mm:%b %Y}: visa money held {e['held']:,.2f}", same(r["visa_held"], e["held"]), r["visa_held"])
        if prev_close is not None:
            ck(f"Petty {book} {mm:%b %Y}: brought forward = last month's closing", same(r["opening"], prev_close), (r["opening"], prev_close))
        prev_close = D(r["closing"])
    # full history and a range across a month end
    r = get(f"/store/petty-cash?book={book}&month=all")
    e = petty_expect(book, date(2000, 1, 1), date(2099, 12, 31), TODAY)
    ck(f"Petty {book} full history: closing {e['closing']:,.2f} = this month's closing", same(r["closing"], e["closing"]) and same(r["closing"], prev_close)
       and same(r["visa_held"], e["held"]), (r["closing"], e["closing"], prev_close))
    if len(allm) > 1:
        a, b = allm[-2] + timedelta(days=14), allm[-1] + timedelta(days=9)
        r = get(f"/store/petty-cash?book={book}&date_from={a}&date_to={b}")
        e = petty_expect(book, a, b, b)
        ck(f"Petty {book} {a:%d/%m}-{b:%d/%m}: b/f {e['opening']:,.2f}, closing {e['closing']:,.2f}",
           same(r["opening"], e["opening"]) and same(r["closing"], e["closing"]) and len(r["rows"]) == len(e["rows"]), (r["opening"], r["closing"]))
    # the papers for the last month with lines
    mm = max(months)
    r = get(f"/store/petty-cash?book={book}&month={mm:%Y-%m}")
    content = fetch(f"/export/store/petty-cash?book={book}&month={mm:%Y-%m}&format=excel")
    wb = xl(content)
    if wb is not None:
        ws = wb.active
        tot = next(i for i in range(1, ws.max_row + 1) if ws.cell(i, 1).value == "TOTAL")
        heads = [c_.value for c_ in ws[4]]
        G = len([h for h in heads if h])
        ck(f"Petty {book} Excel {mm:%b %Y}: balance column worked out = {r['closing']:,.2f}", same(ws.cell(tot, G).value or 0, r["closing"]),
           ws.cell(tot, G).value)
        lines_ok = all(same(ws.cell(6 + i, G).value or 0, x["balance"]) for i, x in enumerate(r["rows"]))
        ck(f"Petty {book} Excel {mm:%b %Y}: every line's balance matches the screen", lines_ok)
    txt = pdf_text(fetch(f"/export/store/petty-cash?book={book}&month={mm:%Y-%m}&format=pdf"))
    if txt is not None:
        ck(f"Petty {book} PDF {mm:%b %Y}: prints b/f, totals and closing",
           all(money(v) in txt for v in (r["opening"], r["received"], r["paid"], r["closing"])))
# an empty period: its Excel must not hold a formula that refers to itself
empty_from = date(1999, 1, 1)
wbf = xl_formulas(fetch(f"/export/store/petty-cash?book=site&date_from={empty_from}&date_to={empty_from + timedelta(days=5)}&format=excel"))
ws = wbf.active
tot = next(i for i in range(1, ws.max_row + 1) if ws.cell(i, 1).value == "TOTAL")
selfref = [ws.cell(tot, j).value for j in range(1, ws.max_column + 1)
           if isinstance(ws.cell(tot, j).value, str) and re.search(rf"[A-Z]+{tot}:[A-Z]+\d+|[A-Z]+\d+:[A-Z]+{tot}\b", ws.cell(tot, j).value)]
ck("Petty Excel, a period with no lines: the TOTAL row does not add itself (no circular formula)", not selfref, selfref)
print(f"   ({len(pc)} petty cash rows, {nrows} month-lines checked)")

# ===================================================================================================
# Bank
# ===================================================================================================
print("\n== Bank")
bk = raw("select * from bank_entries")
for x in bk:
    x["d"] = pd(x["on_date"])
bmonths = sorted({x["d"].replace(day=1) for x in bk})
for mm in bmonths:
    a, b = month_bounds(mm)
    r = get(f"/store/bank?month={mm:%Y-%m}")
    inside = [x for x in bk if a <= x["d"] <= b]
    paid = r2(sum(D(x["amount"]) for x in inside if x["kind"] != "customer"))
    rec = r2(sum(D(x["amount"]) for x in inside if x["kind"] == "customer"))
    ck(f"Bank {mm:%b %Y}: {len(inside)} lines, paid out {paid:,.2f}, received {rec:,.2f}", len(r["rows"]) == len(inside)
       and same(r["paid"], paid) and same(r["received"], rec), (r["paid"], r["received"]))
    canc = [x for x in inside if x["zoho_status"] == "cancelled"]
    if canc:
        note(f"Bank {mm:%b %Y}: {len(canc)} line(s) deleted in Zoho (AED {r2(sum(D(x['amount']) for x in canc)):,.2f}) are still in the paid / received totals")
    v = fetch(f"/export/accounts/bank?month={mm:%Y-%m}&format=view").decode()
    ck(f"Bank {mm:%b %Y} preview totals line: {paid:,.2f} / {rec:,.2f}", (not paid or money(paid) in v) and (not rec or money(rec) in v))
# tied to Project payments and to receipts
pp = raw("select * from project_payments") if has_table("project_payments") else []
rc = raw("select * from receipts") if has_table("receipts") else []
for x in bk:
    lines = [p for p in pp if p["source"] == f"bank:{x['id']}"]
    should = x["kind"] == "customer" and x["scope_id"] and x["d"] <= TODAY
    if should:
        ok = len(lines) == 1 and lines[0]["scope_id"] == x["scope_id"] and same(lines[0]["amount"], x["net_amount"]) and lines[0]["on_date"][:10] == x["on_date"][:10]
        ck(f"Bank BK-{x['id']} client payment {D(x['amount']):,.2f}: one Project payments line of {D(x['net_amount']):,.2f} (excl. VAT)", ok, lines)
    else:
        ck(f"Bank BK-{x['id']} ({x['kind']}): nothing on Project payments", not lines, lines)
    if x["kind"] == "customer":
        live = [r for r in rc if r["bank_id"] == x["id"] and r["status"] != "cancelled"]
        ck(f"Bank BK-{x['id']}: one live receipt of the same amount", len(live) == 1 and same(live[0]["amount"], x["amount"]),
           [(r["number"], r["amount"]) for r in live])
orph = [p for p in pp if (p["source"] or "").startswith("bank:") and not any(f"bank:{x['id']}" == p["source"] for x in bk)]
ck("Project payments: no line left behind by a removed bank line", not orph, orph)

# ===================================================================================================
# Invoices and receipts
# ===================================================================================================
print("\n== Invoices")
invs = raw("select * from invoices")
rcl = [r for r in rc if r["status"] != "cancelled"]
for r in rcl:
    r["al"] = [a for a in json.loads(r["allocations"] or "[]") if a.get("invoice_id")]


def due(r):
    return not (r["mode"] in ("Cheque", "PDC") and r["cheque_date"]) or pd(r["cheque_date"]) <= TODAY


got_due, got_all = {}, {}
for r in rcl:
    for a in r["al"]:
        got_all[int(a["invoice_id"])] = got_all.get(int(a["invoice_id"]), D(0)) + D(a["amount"])
        if due(r):
            got_due[int(a["invoice_id"])] = got_due.get(int(a["invoice_id"]), D(0)) + D(a["amount"])
pyround_diff = []
for kind in ("tax", "proforma"):
    lst = get(f"/employees/accounts/invoices?kind={kind}")
    byid = {r["id"]: r for r in lst["rows"]}
    mine = [i for i in invs if i["kind"] == kind]
    for i in mine:
        ls = json.loads(i["lines"] or "[]")
        sub = sum(r2(l["amount"]) for l in ls)
        vat = sum(r2(r2(l["amount"]) * D(l.get("vat") if l.get("vat") not in (None, "") else 5) / 100) for l in ls)
        py_vat = sum(D(str(round(round(float(l["amount"]), 2) * float(l.get("vat") or 0) / 100, 2))) for l in ls)
        if py_vat != vat:
            pyround_diff.append((i["number"], str(py_vat), str(vat)))
        tot = r2(sub + vat)
        # A VAT of exactly half a fils (85,538.10 x 5% = 4,276.905) is rounded by the app the way Python's round()
        # does (4,276.90 - as Zoho also shows it), not half up (4,276.91): passed here, listed under NOTES.
        ok_half = same(i["subtotal"], sub) and same(i["vat"], vat) and same(i["total"], tot)
        ok_py = same(i["subtotal"], sub) and same(i["vat"], py_vat) and same(i["total"], r2(sub + py_vat))
        ck(f"Invoice {i['number']}: {len(ls)} lines -> net {sub:,.2f} + VAT {vat:,.2f} = {tot:,.2f} (stored {i['subtotal']:,.2f} / {i['vat']:,.2f} / {i['total']:,.2f})",
           ok_half or ok_py, (i["subtotal"], i["vat"], i["total"]))
        if kind == "tax":
            a = byid[i["id"]]
            ck(f"Invoice {i['number']}: received {r2(got_due.get(i['id'], 0)):,.2f}, balance {r2(D(i['total']) - got_due.get(i['id'], 0)):,.2f}",
               same(a["received"], r2(got_due.get(i["id"], 0))) and same(a["balance"], r2(D(i["total"]) - got_due.get(i["id"], 0)))
               and same(a["pdc"], r2(got_all.get(i["id"], 0) - got_due.get(i["id"], 0))), (a["received"], a["balance"], a["pdc"]))
            ck(f"Invoice {i['number']}: never more set against it than its total", got_all.get(i["id"], 0) <= D(i["total"]) + D("0.005"),
               (got_all.get(i["id"]), i["total"]))
            full = r2(D(i["total"]) - got_due.get(i["id"], 0)) <= D("0.005")
            if i["paid_stamp"] and not full:
                note(f"Invoice {i['number']} carries the PAID stamp ({i['paid_on']}) but AED {r2(D(i['total']) - got_due.get(i['id'], 0)):,.2f} is still open on the receipts")
            if i["status"] == "cancelled" and got_all.get(i["id"]):
                ck(f"Cancelled invoice {i['number']}: no live receipt set against it", False, got_all.get(i["id"]))
        # the paper
        txt = pdf_text(fetch(f"/export/accounts/invoice/{i['id']}"))
        if txt is not None:
            ck(f"Invoice {i['number']} PDF: subtotal, VAT and total as stored", all(money(v) in txt for v in (i["subtotal"], i["vat"], i["total"])))
    live = [i for i in mine if i["status"] != "cancelled"]
    t = lst["totals"]
    ck(f"Invoices ({kind}) list totals: {len(live)} live of {len(mine)}, net {r2(sum(D(i['subtotal']) for i in live)):,.2f}, "
       f"VAT {r2(sum(D(i['vat']) for i in live)):,.2f}, total {r2(sum(D(i['total']) for i in live)):,.2f}",
       t["count"] == len(live) and same(t["subtotal"], r2(sum(D(i["subtotal"]) for i in live))) and same(t["vat"], r2(sum(D(i["vat"]) for i in live)))
       and same(t["total"], r2(sum(D(i["total"]) for i in live))), t)
    wb = xl(fetch(f"/export/accounts/invoices?kind={kind}&format=excel"))
    if wb is not None and live:
        ws = wb.active
        vals = [ws.cell(rr, cc).value for rr in range(1, ws.max_row + 1) for cc in range(1, ws.max_column + 1)]
        ck(f"Invoices ({kind}) register Excel: total {t['total']:,.2f} (cancelled left out)",
           any(isinstance(v, (int, float)) and same(v, t["total"]) for v in vals) or any(isinstance(v, str) and money(t["total"]) in v for v in vals))
if pyround_diff:
    note("Invoice VAT where Python's round() and 2dp half-up differ (app keeps Python's): " + "; ".join(f"{n}: app {a} vs half-up {b}" for n, a, b in pyround_diff))

print("\n== Receipts")
R = get("/employees/accounts/receipts")
ck(f"Receipts totals: {len(rcl)} live, AED {r2(sum(D(r['amount']) for r in rcl)):,.2f}, on later cheques {r2(sum(D(r['amount']) for r in rcl if not due(r))):,.2f}",
   R["totals"]["count"] == len(rcl) and same(R["totals"]["amount"], r2(sum(D(r["amount"]) for r in rcl)))
   and same(R["totals"]["pdc"], r2(sum(D(r["amount"]) for r in rcl if not due(r)))), R["totals"])
byr = {r["id"]: r for r in R["rows"]}
for r in rcl:
    ap = r2(sum(D(a["amount"]) for a in r["al"]))
    g = byr[r["id"]]
    ck(f"Receipt {r['number']} AED {D(r['amount']):,.2f}: applied {ap:,.2f} + advance {r2(D(r['amount']) - ap):,.2f}",
       ap <= D(r["amount"]) + D("0.005") and same(g["applied"], ap) and same(g["advance"], r2(D(r["amount"]) - ap)), (g["applied"], g["advance"]))
    if r["al"]:
        txt = pdf_text(fetch(f"/export/accounts/receipt/{r['id']}"))
        if txt is not None:
            ck(f"Receipt {r['number']} PDF: amount and each invoice's 'received now' printed",
               money(r["amount"]) in txt and all(money(a["amount"]) in txt for a in r["al"]))
inv_open = {x["id"]: x for x in R["invoices"]}
for i in invs:
    if i["kind"] == "tax" and i["status"] != "cancelled":
        x = inv_open.get(i["id"])
        want = r2(D(i["total"]) - got_all.get(i["id"], 0))
        ck(f"Receipts' invoice picker {i['number']}: still open {want:,.2f}", x is not None and same(x["open"], want), x and x["open"])

# ===================================================================================================
# Cash register
# ===================================================================================================
print("\n== Cash register")
k = c.post("/employees/accounts/register/unlock", headers=H, json={"password": os.environ.get("REGISTER_PW", "Infinia2022!")})
if k.status_code == 200:
    RK = {**H, "X-Register-Key": k.json()["key"]}
    cr = sorted(raw("select * from cash_register"), key=lambda x: (x["on_date"], x["id"]))
    for x in cr:
        x["d"] = pd(x["on_date"]); x["s"] = D(x["amount"]) * (1 if x["kind"] == "received" else -1)
    periods = [(None, None)] + [month_bounds(m) for m in sorted({x["d"].replace(day=1) for x in cr})]
    for a, b in periods:
        q = f"?start={a or ''}&end={b or ''}"
        r = c.get("/employees/accounts/register" + q, headers=RK).json()
        op = r2(sum(x["s"] for x in cr if a and x["d"] < a))
        ins = [x for x in cr if (not a or x["d"] >= a) and (not b or x["d"] <= b)]
        bal, ok = op, True
        for x, g in zip(ins, r["ledger"]):
            bal = r2(bal + x["s"]); ok = ok and same(g["balance"], bal) and g["id"] == x["id"]
        tag = f"{a:%b %Y}" if a else "all"
        ck(f"Cash register {tag}: b/f {op:,.2f}, {len(ins)} lines, closing {bal:,.2f}", ok and len(ins) == len(r["ledger"]) and same(r["opening"], op)
           and same(r["balance"], bal), (r["opening"], r["balance"]))
        r_ = c.get(f"/export/accounts/register{q}&format=excel&rk={RK['X-Register-Key']}&token={dl()}")
        wb = xl(r_.content)
        if wb is not None:
            ws = wb.active
            tot = next(i for i in range(1, ws.max_row + 1) if ws.cell(i, 1).value == "TOTAL")
            ck(f"Cash register Excel {tag}: closing worked out = {bal:,.2f}", same(ws.cell(tot, 6).value or 0, bal), ws.cell(tot, 6).value)
    wbf = xl_formulas(c.get(f"/export/accounts/register?start=1999-01-01&end=1999-01-05&format=excel&rk={RK['X-Register-Key']}&token={dl()}").content)
    ws = wbf.active
    tot = next(i for i in range(1, ws.max_row + 1) if ws.cell(i, 1).value == "TOTAL")
    ck("Cash register Excel, a period with no lines: no circular TOTAL formula",
       not any(isinstance(ws.cell(tot, j).value, str) and f"{tot}:" in ws.cell(tot, j).value.replace("$", "") for j in (4, 5)),
       [ws.cell(tot, j).value for j in (4, 5)])
else:
    note(f"Cash register not checked - the password was refused ({k.status_code})")

# ===================================================================================================
# Project payments
# ===================================================================================================
print("\n== Project payments")
scopes = raw("select * from project_scopes") if has_table("project_scopes") else []
for p in pp:
    p["d"] = pd(p["on_date"])


def scope_fig(s, a=None, b=None):
    ps = [p for p in pp if p["scope_id"] == s["id"] and (not b or p["d"] <= b)]
    ins = [p for p in ps if not a or p["d"] >= a]
    rec = r2(sum(r2(p["amount"]) for p in ps))
    con = r2(s["contract"])
    pct = (rec * 100 / con).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP) if con else D(0)
    st = "Pending" if rec <= D("0.005") else ("Completed" if rec >= con - D("0.005") else "Partial")
    return {"received": rec, "remaining": r2(con - rec), "contract": con, "pct": pct, "status": st, "in_period": r2(sum(r2(p["amount"]) for p in ins)), "n": len(ins)}


T = get("/employees/accounts/projects")
got = {r["id"]: r for r in T["rows"]}
for s in scopes:
    e, g = scope_fig(s), got.get(s["id"])
    ck(f"Project {s['project']} {s['scope']}: contract {e['contract']:,.2f}, received {e['received']:,.2f}, remaining {e['remaining']:,.2f}, {e['pct']}% {e['status']}",
       g and same(g["received"], e["received"]) and same(g["remaining"], e["remaining"]) and same(g["pct"], e["pct"], Decimal("0.05")) and g["status"] == e["status"],
       g and (g["received"], g["remaining"], g["pct"], g["status"]))
    if g and D(g["pct"]) != e["pct"]:
        note(f"Project {s['project']} {s['scope']}: % collected {g['pct']} on screen, {e['pct']} rounded half-up")
tc = r2(sum(r2(s["contract"]) for s in scopes)); tr = r2(sum(scope_fig(s)["received"] for s in scopes))
ck(f"Project payments company totals: contract {tc:,.2f}, received {tr:,.2f}, remaining {r2(tc - tr):,.2f}",
   same(T["totals"]["contract"], tc) and same(T["totals"]["received"], tr) and same(T["totals"]["remaining"], r2(tc - tr)), T["totals"])
for g in T["groups"]:
    ck(f"Project {g['project']} total = its scopes", same(g["received"], sum(D(r["received"]) for r in g["rows"]))
       and same(g["contract"], sum(D(r["contract"]) for r in g["rows"])))
for mm in sorted({p["d"].replace(day=1) for p in pp})[-3:]:
    a, b = month_bounds(mm)
    P = get(f"/employees/accounts/projects?start={a}&end={b}")
    want = r2(sum(r2(p["amount"]) for p in pp if a <= p["d"] <= b))
    paid_scopes = {p["scope_id"] for p in pp if a <= p["d"] <= b}
    ck(f"Project payments {mm:%b %Y}: received in the month {want:,.2f} over {len(paid_scopes)} scopes", same(P["shown"]["in_period"], want)
       and {r["id"] for r in P["rows"]} == paid_scopes and all(same(r["received"], scope_fig(next(s for s in scopes if s["id"] == r["id"]), None, b)["received"]) for r in P["rows"]),
       (P["shown"]["in_period"], want))
if scopes:
    wb = xl(fetch("/export/accounts/projects?report=detail&format=excel"))
    if wb is not None:
        ws = wb["Tracker"]
        row = next(i for i in range(1, ws.max_row + 1) if ws.cell(i, 1).value == "TOTAL")
        heads = [str(c_.value or "") for c_ in ws[6]]
        cc = heads.index("CONTRACT (AED)") + 1; rcol = heads.index("RECEIVED (AED)") + 1
        ck(f"Project payments Excel: TOTAL contract {tc:,.2f}, received {tr:,.2f} (project subtotal lines not added twice)",
           same(ws.cell(row, cc).value or 0, tc) and same(ws.cell(row, rcol).value or 0, tr), (ws.cell(row, cc).value, ws.cell(row, rcol).value))
        ws2 = wb["Payments"]
        last = ws2.max_row
        ck(f"Project payments Excel 'Payments' sheet: {len(pp)} payments add to {tr:,.2f}", same(ws2.cell(last, 5).value or 0, tr), ws2.cell(last, 5).value)

# ===================================================================================================
# LPOs
# ===================================================================================================
print("\n== LPOs")
pos = raw("select * from purchase_orders")
pol = raw("select * from purchase_order_lines")
lines_of = {}
for l in pol:
    lines_of.setdefault(l["order_id"], []).append(l)


def lpo_fig(o):
    keep = 1 - D(o["discount_pct"]) / 100
    ls = lines_of.get(o["id"], [])
    gross = sum(D(l["qty"]) * D(l["rate"]) for l in ls)
    net = gross * keep
    vat = sum(D(l["qty"]) * D(l["rate"]) * keep * D(l["tax_pct"]) / 100 for l in ls)
    vat_per_line = sum(r2(r2(D(l["qty"]) * D(l["rate"]) * keep) * D(l["tax_pct"]) / 100) for l in ls)
    return {"gross": gross, "net": net, "vat": vat, "total": r2(net + vat), "vat_line": vat_per_line,
            "total_line": r2(sum(r2(D(l["qty"]) * D(l["rate"]) * keep) for l in ls) + vat_per_line)}


reg = {r["id"]: r for r in get("/store/purchase/orders?limit=5000")}
zoho_line_diff = []
for o in pos:
    e = lpo_fig(o)
    g = reg.get(o["id"])
    ck(f"LPO {o['ref']} ({o['status']}, disc {D(o['discount_pct'])}%): register total {e['total']:,.2f}", g and same(g["total"], e["total"]), g and g["total"])
    if e["total_line"] != e["total"]:
        zoho_line_diff.append((o["ref"], e["total"], e["total_line"]))
    html = c.get(f"/export/purchase/{o['id']}/view?token={dl()}")
    if html.status_code == 200 and "text/html" in html.headers.get("content-type", ""):
        ck(f"LPO {o['ref']} printed page: total AED {e['total']:,.2f}", f"AED {money(e['total'])}" in html.text)
issued = [o for o in pos if o["status"] != "cancelled"]
rep = get("/store/purchase/report?group_by=order&status=all")
want = r2(sum(lpo_fig(o)["net"] + lpo_fig(o)["vat"] for o in issued))
ck(f"Purchase report (all, cancelled struck out): {len(issued)} orders, total {want:,.2f}", same(rep["totals"]["total"], want, Decimal("0.011")), rep["totals"])
rep_i = get("/store/purchase/report?group_by=supplier")
not_issued = [o for o in issued if o["status"] != "issued"]
ck(f"Purchase report by supplier (default 'issued'): every order not cancelled counted ({len(issued)})",
   same(rep_i["totals"]["total"], want, Decimal("0.011")) or not not_issued, (rep_i["totals"], [o["status"] for o in not_issued][:5]))
if zoho_line_diff:
    note("LPO totals if VAT were worked per line and rounded (as Zoho does on a purchase order) - portal vs per-line: "
         + "; ".join(f"{r}: {a:,.2f} vs {b:,.2f}" for r, a, b in zoho_line_diff[:15]) + (f" ... {len(zoho_line_diff)} in all" if len(zoho_line_diff) > 15 else ""))
disc = [o for o in pos if D(o["discount_pct"])]
if disc:
    note(f"{len(disc)} LPO(s) with a discount: the printed order's VAT column shows each line's VAT before the discount, the VAT total after it "
         f"(e.g. {disc[0]['ref']}: column adds to {r2(sum(D(l['qty']) * D(l['rate']) * D(l['tax_pct']) / 100 for l in lines_of.get(disc[0]['id'], []))):,.2f}, "
         f"VAT total {r2(lpo_fig(disc[0])['vat']):,.2f})")

# ===================================================================================================
# Project expense
# ===================================================================================================
print("\n== Project expense")
PE = get("/employees/accounts/project-expense")
rows = {r["code"]: r for r in PE["rows"]}
for code, r in rows.items():
    ck(f"Project expense {code}: labour + LPO + store + petty = total {r['total']:,.2f}",
       same(r["total"], sum(D(r[k]) for k in ("labour", "lpo", "store", "petty")), Decimal("0.011")))
for k in ("labour", "lpo", "store", "petty", "total", "contract", "received"):
    ck(f"Project expense TOTAL {k} = the projects' {k}", same(PE["totals"][k], r2(sum(D(r[k]) for r in PE["rows"])), Decimal("0.011")))
sites = {s["code"] for s in raw("select code from sites") if s["code"]} | {s["project"] for s in scopes}
SAL = re.compile(r"salar|wages|payable|advance|recovery", re.I)


def codes_in(t):
    t = str(t or "")
    return [x for x in dict.fromkeys(re.findall(r"(?<!\d)(\d{3})(?!\d)", t)) if x in sites]


petty_want = {}
for x in pc:
    if x["pay"] <= 0 or SAL.search(x["description"] or ""):
        continue
    cs = codes_in(x["site"]) or codes_in(x["description"])
    if not cs:
        continue
    claimable = bool(x["vat"]) and (x.get("vat_claim") or "claim") == "claim"
    net = x["pay"] - (r2(x["pay"] * 5 / 105) if claimable else 0)
    for cd in cs:
        petty_want[cd] = petty_want.get(cd, D(0)) + r2(r2(net) / len(cs))
for x in bk:
    if x["kind"] != "expense" or D(x["amount"]) <= 0 or SAL.search(x["notes"] or ""):
        continue
    cs = codes_in(x["site"]) or codes_in(x["notes"])
    if not cs:
        continue
    claimable = bool(x["vat"]) and (x["vat_claim"] or "claim") == "claim"
    net = D(x["amount"]) - (r2(D(x["amount"]) * 5 / 105) if claimable else 0)
    for cd in cs:
        petty_want[cd] = petty_want.get(cd, D(0)) + r2(r2(net) / len(cs))
for cd in sorted(set(petty_want) | {k for k, r in rows.items() if r["petty"]}):
    g = rows.get(cd, {}).get("petty", 0)
    ck(f"Project expense {cd} petty cash + bank expenses (VAT taken off only where it is claimed back): {r2(petty_want.get(cd, 0)):,.2f}",
       same(g, r2(petty_want.get(cd, 0)), Decimal("0.011")), g)
zero_vat = [x for x in pc if x["vat"] and (x.get("vat_claim") or "") == "zero" and x["pay"] > 0]
if zero_vat:
    print(f"   ({len(zero_vat)} petty cash bill(s) with VAT marked not claimable)")
lpo_want = {}
for o in pos:
    if o["status"] == "cancelled":
        continue
    req = raw("select site from material_requests where id=?", o["request_id"]) if o["request_id"] else []
    rs = (req[0]["site"] or "").strip() if req else ""
    if rs.upper() == "STORE" or o.get("combined"):
        continue
    cs = codes_in(o["project_location"]) or codes_in(o["job_scope"]) or codes_in(rs)
    for cd in cs:
        lpo_want[cd] = lpo_want.get(cd, D(0)) + r2(r2(lpo_fig(o)["net"]) / len(cs))
for cd in sorted(set(lpo_want) | {k for k, r in rows.items() if r["lpo"]}):
    g = rows.get(cd, {}).get("lpo", 0)
    ck(f"Project expense {cd} LPO cost net of discount and VAT: {r2(lpo_want.get(cd, 0)):,.2f}", same(g, r2(lpo_want.get(cd, 0)), Decimal("0.011")), g)
# materials sent from the store: valued at the cost on the movement, else the last LPO rate actually paid (after the
# LPO's discount, as the store's own valuation reads it), else the item's estimated price
items = {i["id"]: i for i in raw("select * from store_items")} if has_table("store_items") else {}
last_rate = {}
for o in sorted(pos, key=lambda o: (o["order_date"] or "", o["id"])):
    if o["status"] == "cancelled":
        continue
    for l in lines_of.get(o["id"], []):
        if l["item_id"] and l["rate"]:
            last_rate[l["item_id"]] = D(l["rate"]) * (1 - D(o["discount_pct"]) / 100)
store_want = {}
for m in (raw("select * from store_movements") if has_table("store_movements") else []):
    it = items.get(m["item_id"])
    if not it or (it["item_type"] or "consumable") != "consumable":
        continue
    to, frm, q = (m["location"] or "").strip(), (m["from_location"] or "").strip(), D(m["qty"])
    legs = []
    if m["kind"] == "out" and to in sites:
        legs = [(to, q)]
    elif m["kind"] == "return" and frm in sites:
        legs = [(frm, -q)]
    elif m["kind"] == "transfer":
        legs = ([(frm, -q)] if frm in sites else []) + ([(to, q)] if to in sites else [])
    rate = D(m["unit_cost"]) if m["unit_cost"] else last_rate.get(it["id"], D(it["est_price"] or 0))
    for site, qq in legs:
        store_want[site] = store_want.get(site, D(0)) + r2(qq * rate)
for cd in sorted(set(store_want) | {k for k, r in rows.items() if r["store"]}):
    g = rows.get(cd, {}).get("store", 0)
    ck(f"Project expense {cd} materials from the store: {r2(store_want.get(cd, 0)):,.2f}", same(g, r2(store_want.get(cd, 0)), Decimal("0.011")), g)
v = fetch("/export/accounts/project-expense/view").decode()
ck(f"Project expense preview TOTAL line shows the total cost {PE['totals']['total']:,.2f}", money(PE["totals"]["total"]) in v)

# ---------------------------------------------------------------------------------------------------
print("\n" + ("NOTES (judgement calls, nothing changed):\n  - " + "\n  - ".join(NOTES) if NOTES else "no notes"))
print(f"\n{len(FAILS)} FAILED" if FAILS else "\nALL PASS")
sys.exit(1 if FAILS else 0)
