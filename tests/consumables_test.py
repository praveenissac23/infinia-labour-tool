"""A consumable is stock in the store and used up at the site.

Run: cd app && DATABASE_URL=sqlite:////tmp/cn.db python3 ../tests/consumables_test.py

Fifty bags of cement sent to 901 are in the wall within a day or two.
Carrying them as stock at 901 forever makes every site look like a
warehouse and buries the things that really are there - the drill that
must come back, the scaffolding somebody is paying for by the week. So
a site holds tools, owned equipment and hired equipment, and never a
consumable; what was sent shows as what was last sent, which is the
question asked before sending more. The central store still holds
everything, because there it is genuinely on a shelf.
"""
import sys, io
sys.path.insert(0, '.')
from fastapi.testclient import TestClient
import database, models, auth
models.Base.metadata.create_all(database.engine)
import main

db = database.SessionLocal()
db.add(models.User(username='admin', hashed_password=auth.hash_password('p'),
                   full_name='Administrator', role='admin'))
for code in ('901', '902'):
    db.add(models.Site(code=code, active=True))
db.commit(); db.close()

c = TestClient(main.app)
H = {'Authorization': 'Bearer ' + c.post('/auth/login',
     data={'username': 'admin', 'password': 'p'}).json()['access_token']}

FAIL = []
def ck(l, ok, x=''):
    print(('PASS ' if ok else 'FAIL ') + l + ('' if ok else f'  [{x}]'))
    if not ok: FAIL.append(l)

def item(name, kind, unit):
    body = {'name': name, 'unit': unit, 'item_type': kind}
    if kind == 'rental':
        body['rental_supplier'] = 'Al Raha Scaffolding'
    return c.post('/store/items', json=body, headers=H).json()

cement = item('Cement OPC 50kg', 'consumable', 'bag')
drill = item('Hilti Drill', 'returnable', 'pcs')
mixer = item('Concrete Mixer', 'asset', 'pcs')
scaff = item('Scaffold Ledger', 'rental', 'pcs')

def move(it, kind, qty, frm='', to='', on='2026-09-10'):
    return c.post('/store/movements', json={'item_id': it['id'], 'kind': kind, 'qty': qty,
                  'from_location': frm, 'location': to, 'incharge': 'Amal',
                  'notes': '', 'moved_on': on}, headers=H)

for it, q in ((cement, 200), (drill, 4), (mixer, 2), (scaff, 500)):
    ck(f'{it["name"]} received into the store', move(it, 'in', q, to='').status_code == 200)

# Everything goes out to 901.
ck('cement sent to 901', move(cement, 'out', 50, frm='', to='901', on='2026-09-12').status_code == 200)
ck('drill sent to 901', move(drill, 'out', 2, frm='', to='901').status_code == 200)
ck('mixer sent to 901', move(mixer, 'out', 1, frm='', to='901').status_code == 200)
ck('scaffolding sent to 901', move(scaff, 'out', 300, frm='', to='901').status_code == 200)

at = c.get('/store/at-site?site=901', headers=H).json()
held = {r['name'] for r in at['rows']}
ck('the drill is held at the site', 'Hilti Drill' in held, held)
ck('the mixer is held at the site', 'Concrete Mixer' in held, held)
ck('the hired scaffolding is held at the site', 'Scaffold Ledger' in held, held)
ck('the cement is NOT held at the site', 'Cement OPC 50kg' not in held, held)

recent = {r['name']: r for r in at['recent']}
ck('the cement shows as last sent instead', 'Cement OPC 50kg' in recent, list(recent))
ck('with the quantity that went', recent.get('Cement OPC 50kg', {}).get('qty') == 50,
   recent.get('Cement OPC 50kg'))
ck('and the day it went', recent.get('Cement OPC 50kg', {}).get('on') == '2026-09-12',
   recent.get('Cement OPC 50kg'))
ck('nothing that is held is repeated in last-sent',
   not (held & set(recent)), held & set(recent))

# A second delivery moves the date on, and only the latest is shown.
move(cement, 'out', 30, frm='', to='901', on='2026-09-16')
at = c.get('/store/at-site?site=901', headers=H).json()
cem_recent = [r for r in at['recent'] if r['name'] == 'Cement OPC 50kg']
ck('only the most recent sending is shown', len(cem_recent) == 1, cem_recent)
ck('and it is the newest one', cem_recent[0]['on'] == '2026-09-16' and cem_recent[0]['qty'] == 30,
   cem_recent[0])

# The store still holds everything, and the figures still add up there.
stock = {s['name']: s for s in c.get('/store/stock', headers=H).json()}
cem = stock['Cement OPC 50kg']
ck('the store still shows the cement it holds', cem['central'] == 120, cem)
ck('the cement is not counted as out at sites', cem['out_at_sites'] == 0, cem)
ck('and no site is listed against it', cem['by_site'] == {}, cem)
ck("the cement's total is what is in the store", cem['total'] == 120, cem)

dr = stock['Hilti Drill']
ck('the drill is still counted as out at a site', dr['out_at_sites'] == 2, dr)
ck('against the site holding it', dr['by_site'] == {'901': 2}, dr)
ck('and its total covers both', dr['total'] == 4, dr)

# The reports say the same thing.
by_site = c.get('/store/report?kind=by_site', headers=H).json()['rows']
names = {r['name'] for r in by_site}
ck('the sites report lists the drill', 'Hilti Drill' in names, names)
ck('the sites report leaves the cement out', 'Cement OPC 50kg' not in names, names)

stock_rep = {r['name']: r for r in c.get('/store/report?kind=stock', headers=H).json()['rows']}
ck('the stock report shows no cement at sites',
   stock_rep['Cement OPC 50kg']['at_sites'] == 0 and stock_rep['Cement OPC 50kg']['by_site'] == {},
   stock_rep['Cement OPC 50kg'])
ck('the stock report still shows the drill at its site',
   stock_rep['Hilti Drill']['by_site'] == {'901': 2}, stock_rep['Hilti Drill'])

# A site that has never had anything says so cleanly.
empty = c.get('/store/at-site?site=902', headers=H).json()
ck('a fresh site holds nothing', empty['rows'] == [], empty['rows'])
ck('and has nothing sent to it yet', empty['recent'] == [], empty['recent'])

# The things that do come back still can.
ck('the drill can be returned', move(drill, 'return', 2, frm='901', to='').status_code == 200)
at = c.get('/store/at-site?site=901', headers=H).json()
ck('and then it is no longer at the site',
   'Hilti Drill' not in {r['name'] for r in at['rows']}, at['rows'])

# The ledger is untouched by any of this - it still knows where the
# cement went, which is what the usage report is built on.
usage = c.get('/store/report?kind=usage', headers=H).json()['rows']
ck('the usage report still accounts for the cement',
   any('Cement' in str(r.get('name', '')) for r in usage), usage[:2])

# ---- The asset register: no empty lines, no "where" column ------------
item('Jack Hammer', 'asset', 'pcs')            # owned on paper, none held
assets = c.get('/store/report?kind=assets', headers=H).json()['rows']
names = {r['name'] for r in assets}
ck('the asset register lists equipment that is held', 'Concrete Mixer' in names, names)
ck('a machine with none anywhere is not a line', 'Jack Hammer' not in names, names)
ck('no "where" column', all('where' not in r for r in assets), assets[:1])

# ---- The download is the screen: the search box narrows it too -------
t = c.post('/auth/download-token', headers=H).json()['token']
import openpyxl
xl = c.get(f'/export/store/report?kind=stock&format=excel&q=mixer&token={t}')
ck('store export with a search builds', xl.status_code == 200, xl.status_code)
cells = [str(v) for row in openpyxl.load_workbook(io.BytesIO(xl.content)).active.iter_rows(values_only=True) for v in row if v]
ck('the export holds only what matched', any('Concrete Mixer' in v for v in cells)
   and not any('Cement OPC' in v for v in cells), [v for v in cells if 'OPC' in v or 'Mixer' in v])
full = c.get(f'/export/store/report?kind=stock&format=excel&token={t}')
cells = [str(v) for row in openpyxl.load_workbook(io.BytesIO(full.content)).active.iter_rows(values_only=True) for v in row if v]
ck('with no search the export is the whole list', any('Cement OPC' in v for v in cells))

print('\n' + ('ALL PASS' if not FAIL else f'{len(FAIL)} FAILED: ' + '; '.join(FAIL)))
sys.exit(1 if FAIL else 0)
