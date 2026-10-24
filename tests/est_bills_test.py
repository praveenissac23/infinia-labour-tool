"""Estimation > Bills: the consultant's own Excel bill as the format; each month's bill rolled over, typed, worked out
exactly as Excel does, and out as the same workbook, a preview and a PDF.
    cp /tmp/full_base.db /tmp/eb.db && cd app && DATABASE_URL=sqlite:////tmp/eb.db INFINIA_NO_SCHED=1 python3 ../tests/est_bills_test.py
Needs the sample bills in BILLS_DIR (default: the folder they were attached in).
"""
import io, os, sys, json, glob, zipfile, warnings
warnings.filterwarnings("ignore")
sys.path.insert(0, os.getcwd())
from fastapi.testclient import TestClient
import openpyxl
import main as M
c = TestClient(M.app); c.__enter__()
D = os.environ.get("BILLS_DIR", "/mnt/user-data/attachments/80e871d8-70e2-5c6a-a360-434939dea359/")
fails = []
def ck(n, ok, x=""):
    print(("PASS " if ok else "FAIL ") + n + ("" if ok else f"  -> {str(x)[:300]}"))
    if not ok: fails.append(n)
H = {"Authorization": "Bearer " + c.post("/auth/login", data={"username": "admin", "password": "changeme123"}).json()["access_token"]}
projects = {p["code"]: p["id"] for p in c.get("/store/est/meta", headers=H).json()["projects"]}
f913 = glob.glob(D + "*913-IPA_09.xlsx")[0]
pid = projects["913"]

# ---- the format
r = c.post("/store/est/ipa/format", headers=H, data={"project_id": str(pid), "name": "913 - client format"},
           files={"file": ("913-IPA_09.xlsx", open(f913, "rb").read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
ck("format uploaded", r.status_code == 200, r.text)
fm = r.json()
ck("every formula gives Excel's own figure", fm["formulas"] > 500 and fm["same"] == fm["formulas"], (fm["same"], fm["formulas"], fm.get("different")))
ck("the % lines, bill no., issue date and total due are found", fm["lines"] > 50 and fm["bill_no"] and fm["issue_date"] and fm["total_due"], fm)
xls = glob.glob(D + "*914_MEP*.xls")
if xls:
    r = c.post("/store/est/ipa/format", headers=H, data={"project_id": str(pid)}, files={"file": ("914_MEP.xls", open(xls[0], "rb").read(), "application/vnd.ms-excel")})
    ck("an old .xls is refused with how to save it as .xlsx", r.status_code == 400 and "Save As" in r.text, r.text)
r = c.post("/store/est/ipa/format", headers=H, data={"project_id": str(pid)}, files={"file": ("x.xlsx", b"not a workbook", "application/octet-stream")})
ck("a file that is not a workbook is refused", r.status_code == 400, r.text)

# ---- the first bill (the month after the uploaded IPA 9)
b = c.post("/store/est/ipa/new", headers=H, json={"project_id": pid, "format_id": fm["id"]}).json()
bill = c.get(f"/store/est/ipa/bill/{b['id']}", headers=H).json()
ck("new bill is IPA 10, dated today", bill["label"] == "10" and bill["issue_date"] == M._dubai_today().isoformat(), (bill["label"], bill["issue_date"]))
lines = [ln for t in bill["progress"] for ln in t["lines"]]
ck("Previous % = last month's Total % on every line", all(ln["prev"] == ln["total"] for ln in lines if ln["tot_typed"] and ln["prev"] is not None), [(l["desc"][:20], l["prev"], l["total"]) for l in lines[:5]])
pay = bill["payments"][0]
ck("IPA 9 put on the previous payments, amount to fill", any((row["desc"]["raw"] or "") == "Payment Application No 9" for row in pay["rows"]), [row["desc"]["raw"] for row in pay["rows"]][:20])
# this month: MC Scope - RCC superstructure from 30% to 45%
t = next(t for t in bill["progress"] if t["sheet"] == "MC Scope")
ln = next(l for l in t["lines"] if l["tot_typed"] and abs((l["total"] or 0) - 0.3) < 1e-9)
before = bill["total_due"]
r = c.put(f"/store/est/ipa/bill/{b['id']}", headers=H, json={"set": {"MC Scope": {ln["tot_key"]: 0.45}}})
ck("Total % saved", r.status_code == 200, r.text)
after = r.json()["total_due"]
# 15% of 1,315,621 more work, less advance recovery at the job's %, plus 5% VAT (the sheet's own formulas)
ck("total due worked out again", after > before, (before, after))
ln2 = next(l for t2 in r.json()["progress"] if t2["sheet"] == "MC Scope" for l in t2["lines"] if l["r"] == ln["r"])
ck("the line shows 45% and is marked changed", abs(ln2["total"] - 0.45) < 1e-9 and ln2["changed"], ln2)
# payment for IPA 9 typed as the office does: (received - authority fees) / 1.05
row = next(row for row in pay["rows"] if (row["desc"]["raw"] or "") == "Payment Application No 9")
r = c.put(f"/store/est/ipa/bill/{b['id']}", headers=H, json={"set": {pay["sheet"]: {row["amount"]["key"]: "=(989148-280)/1.05", row["received"]["key"]: "20.10.2026"}}})
ck("payment typed as a formula is worked out", r.status_code == 200 and abs(next(x for x in r.json()["payments"][0]["rows"] if x["r"] == row["r"])["amount"]["value"] - (989148 - 280) / 1.05) < 0.01, r.text[:300])
due2 = r.json()["total_due"]
ck("...and comes off the total due", due2 < after, (after, due2))
r = c.put(f"/store/est/ipa/bill/{b['id']}", headers=H, json={"set": {"Summary": {"18,7": 5}}})
ck("a figure the sheet works out cannot be typed over", r.status_code == 400 and "worked out" in r.text, r.text)
r = c.put(f"/store/est/ipa/bill/{b['id']}", headers=H, json={"set": {"MC Scope": {ln["tot_key"]: "=SUM("}}})
ck("a formula that cannot be read is refused", r.status_code == 400, r.text)

# ---- out: Excel, preview, PDF
tok = c.post("/auth/download-token", headers=H).json()["token"]
x = c.get(f"/export/est/ipa/{b['id']}?token={tok}&format=excel")
ck("Excel comes out", x.status_code == 200 and x.content[:2] == b"PK", x.status_code)
wb = openpyxl.load_workbook(io.BytesIO(x.content), data_only=True)
ws = wb["IPA"]
ck("the workbook carries this month's figures (bill no., total due)", ws["C8"].value == 10 and abs(ws["E22"].value - due2) < 0.5, (ws["C8"].value, ws["E22"].value, due2))
ck("...and the amount in words", "DIRHAMS" in str(ws["B23"].value) and "ONLY" in str(ws["B23"].value), ws["B23"].value)
ck("...with the logo, stamps and signatures kept", sum(1 for n in zipfile.ZipFile(io.BytesIO(x.content)).namelist() if "/media/" in n) >= 4)
wbf = openpyxl.load_workbook(io.BytesIO(x.content))
was_hidden = openpyxl.load_workbook(f913)["APP A"].row_dimensions[row["r"]].hidden
ck("the payments row hidden in the format is shown in the workbook once filled", wbf["APP A"].row_dimensions[row["r"]].hidden is False or not wbf["APP A"].row_dimensions[row["r"]].hidden, (was_hidden, wbf["APP A"].row_dimensions[row["r"]].hidden))
ck("...and its formulas kept", str(wbf["IPA"]["E22"].value).startswith("="), wbf["IPA"]["E22"].value)
p = c.get(f"/export/est/ipa/{b['id']}?token={tok}&format=preview")
ck("preview shows the bill", p.status_code == 200 and "PAYMENT APPLICATION" in p.text and "Payment Application No 9" in p.text, p.status_code)
pdf = c.get(f"/export/est/ipa/{b['id']}?token={tok}&format=pdf")
ck("PDF comes out", pdf.status_code == 200 and pdf.content[:4] == b"%PDF", pdf.status_code)
sh = c.get(f"/store/est/ipa/bill/{b['id']}/sheet/0", headers=H).json()
ck("a sheet for typing in comes up with its typed cells marked", 'class="ed' in sh["html"] and sh["width"] > 300, sh.get("name"))

# ---- the next month starts from this one
c.post(f"/store/est/ipa/bill/{b['id']}/status", headers=H, json={"status": "Submitted"})
b2 = c.post("/store/est/ipa/new", headers=H, json={"project_id": pid, "format_id": fm["id"]}).json()
bill2 = c.get(f"/store/est/ipa/bill/{b2['id']}", headers=H).json()
l3 = next(l for t2 in bill2["progress"] if t2["sheet"] == "MC Scope" for l in t2["lines"] if l["r"] == ln["r"])
ck("next bill is IPA 11 and its Previous % is this month's 45%", bill2["label"] == "11" and abs(l3["prev"] - 0.45) < 1e-9, (bill2["label"], l3["prev"]))
ck("...IPA 10 put on its previous payments", any((row["desc"]["raw"] or "") == "Payment Application No 10" for row in bill2["payments"][0]["rows"]))
ck("...and IPA 9's payment carried on", any((row["desc"]["raw"] or "") == "Payment Application No 9" and abs((row["amount"]["value"] or 0) - (989148 - 280) / 1.05) < 0.01 for row in bill2["payments"][0]["rows"]))
lst = c.get(f"/store/est/ipa/list?project_id={pid}", headers=H).json()
ck("the list shows both bills, newest first", [x["label"] for x in lst["bills"]][:2] == ["11", "10"], [x["label"] for x in lst["bills"]])
r = c.delete(f"/store/est/ipa/bill/{b['id']}", headers=H)
ck("admin can remove a submitted bill", r.status_code == 200, r.text)

# ---- other consultants' formats
for pat, code in (("*914_-_IPA_-_07_-_R2.xlsx", "914"), ("*909-IPA09.xlsx", "909")):
    g = glob.glob(D + pat)
    if not g:
        continue
    r = c.post("/store/est/ipa/format", headers=H, data={"project_id": str(projects[code])}, files={"file": (os.path.basename(g[0]), open(g[0], "rb").read(), "application/octet-stream")})
    fx = r.json()
    ck(f"{code}: format read, every formula as Excel", r.status_code == 200 and fx["same"] == fx["formulas"], (fx.get("same"), fx.get("formulas")))
    bx = c.post("/store/est/ipa/new", headers=H, json={"project_id": projects[code], "format_id": fx["id"]}).json()
    bb = c.get(f"/store/est/ipa/bill/{bx['id']}", headers=H).json()
    ck(f"{code}: next bill numbered on ({bb['label']})", bb["label"] in ("8", "10"), bb["label"])
# ---- Total = Previous + Current (Current typed): a new month starts Current at 0, Previous = last Total
g = glob.glob(D + "*911-IPA_Pre_Final_Bill.xlsx")
if g:
    pid911 = projects.get("911") or pid
    r = c.post("/store/est/ipa/format", headers=H, data={"project_id": str(pid911)}, files={"file": ("911.xlsx", open(g[0], "rb").read(), "application/octet-stream")})
    fx = r.json()
    import est_bills_engine as BE2
    m = BE2.read_format(open(g[0], "rb").read(), "911.xlsx"); v0 = BE2.calculate(m, {})
    bx = c.post("/store/est/ipa/new", headers=H, json={"project_id": pid911, "format_id": fx["id"]}).json()
    bb = c.get(f"/store/est/ipa/bill/{bx['id']}", headers=H).json()
    cl = [l for t in bb["progress"] for l in t["lines"] if l["how"] == "cur"]
    ck("911: lines where Current % is typed are found", len(cl) > 0, len(cl))
    t0 = next(t for t in m["map"]["progress"] if any(l.get("how") == "cur" for l in t["lines"]))
    l0 = next(l for l in t0["lines"] if l.get("how") == "cur")
    was = v0[t0["sheet"]].get(f"{l0['r']},{t0['tot_c']}")
    got = next(l for t in bb["progress"] if t["sheet"] == t0["sheet"] for l in t["lines"] if l["r"] == l0["r"])
    ck("...new month: Current 0, Previous = last Total", abs((got["cur"] or 0)) < 1e-9 and abs(got["prev"] - was) < 1e-9, (got["cur"], got["prev"], was))
    r = c.put(f"/store/est/ipa/bill/{bx['id']}", headers=H, json={"set": {t0["sheet"]: {got["cur_key"]: 0.1}}})
    got2 = next(l for t in r.json()["progress"] if t["sheet"] == t0["sheet"] for l in t["lines"] if l["r"] == l0["r"])
    ck("...typing Current 10% makes Total = Previous + 10%", abs(got2["total"] - (was + 0.1)) < 1e-9, got2)
    ck("911: cover names the client", "KARIM" in (r.json()["cover"]["facts"].get("client") or ""), r.json()["cover"]["facts"])
r = c.put(f"/store/est/ipa/bill/{b2['id']}", headers=H, json={"issue_date": "2026-11-05"})
ck("issue date set from a date, written as the format writes it", r.status_code == 200 and r.json()["issue_date"] == "2026-11-05", r.text[:200])
# ---- a bill that comes out below zero (Total % lowered) saves, with "MINUS" in the words - not an error
ln0 = next(l for t in bill2["progress"] for l in t["lines"] if l["tot_typed"] and (l["total"] or 0) > 0.5)
sh0 = next(t["sheet"] for t in bill2["progress"] if any(l is ln0 for l in t["lines"]))
r = c.put(f"/store/est/ipa/bill/{b2['id']}", headers=H, json={"set": {sh0: {ln0["tot_key"]: 0.0}}})
ck("Total % lowered so the bill is below zero: saved, no error", r.status_code == 200 and r.json()["total_due"] < 0, (r.status_code, r.text[:200]))
ck("...the words say MINUS", "MINUS" in (r.json()["cover"]["words"] or "").upper(), r.json()["cover"]["words"])
# ---- the amount in words is written exactly as each consultant writes it
import est_bills_engine as BE3, est_bills_render as BR3
for fpath in sorted(glob.glob(D + "*.xlsx")):
    mm = BE3.read_format(open(fpath, "rb").read(), fpath)
    vv = BE3.calculate(mm, {})
    for w in mm["map"]["words"]:
        like = BE3._text(next(x for x in mm["sheets"] if x["name"] == w["sheet"]), w["r"], w["c"])
        got = BE3.amount_words(vv[w["of"]["sheet"]][f'{w["of"]["r"]},{w["of"]["c"]}'], w.get("upper"), like)
        ck(f"words as the consultant writes them: {os.path.basename(fpath)[:24]}", got == like, (like, got))
# ---- pages break where the consultant's do (914: 36 printed pages, Division 01 broken after rows 36, 65, 104, 148)
g = glob.glob(D + "*914_-_IPA_-_07_-_R2.xlsx")
if g:
    mm = BE3.read_format(open(g[0], "rb").read(), g[0])
    plans = {s_["name"]: BR3.page_plan(mm, s_, {}) for s_ in BR3.printed_sheets(mm)}
    ck("914 prints on 36 pages", sum(len(p_["pages"]) for p_ in plans.values()) == 36, sum(len(p_["pages"]) for p_ in plans.values()))
    ck("...Division 01 breaks where the consultant broke it", [p_[1][-1] for p_ in plans["Division 01-General Requirement"]["pages"]] == [36, 65, 104, 148, 197])
    ck("...heading rows repeated on every page", all(p_[0][:3] == [1, 2, 3] for p_ in plans["Division 01-General Requirement"]["pages"]))
    ck("...APP C keeps its hidden rows hidden", all(r_ not in mm["sheets"][[x["name"] for x in mm["sheets"]].index("APP C")]["hidden_rows"] for p_ in plans["APP C"]["pages"] for r_ in p_[0]))
    ck("...footer Page x of y, and no file-name header", plans["IPA"]["hf"].get("fc") == "Page &P of &N" and not any("&F" in v for v in plans["IPA"]["hf"].values()), plans["IPA"]["hf"])
ck("the Excel file carries no file-name header either", all(b"&amp;F" not in zipfile.ZipFile(io.BytesIO(x.content)).read(n) for n in zipfile.ZipFile(io.BytesIO(x.content)).namelist() if n.startswith("xl/worksheets/sheet")))
# ---- a big workbook in pieces (the web server takes ~1 MB a request)
g = glob.glob(D + "*909-IPA09.xlsx")
if g:
    data = open(g[0], "rb").read(); uid = "ab12cd34ef56ab12cd34"; n = (len(data) + 699_999) // 700_000
    for i in range(n):
        c.post("/store/est/ipa/format/part", headers=H, data={"upload_id": uid, "index": str(i)}, files={"file": ("part", data[i * 700_000:(i + 1) * 700_000], "application/octet-stream")})
    r = c.post("/store/est/ipa/format", headers=H, data={"project_id": str(projects["909"]), "upload_id": uid, "parts": str(n), "filename": "909-IPA09.xlsx"})
    ck("2.6 MB workbook sent in pieces is read whole", r.status_code == 200 and r.json()["same"] == r.json()["formulas"], r.text[:200])
# ---- who signs: full name, designation and an uploaded signature (white paper taken out)
from PIL import Image as _I, ImageDraw as _D
im = _I.new("RGB", (600, 220), (250, 250, 248)); dr = _D.Draw(im)
for i in range(0, 520, 6):
    dr.line((40 + i, 120 + (i % 60) - 30, 46 + i, 120 - (i % 50) + 25), fill=(20, 40, 160), width=5)
buf = io.BytesIO(); im.save(buf, "JPEG")
r = c.post("/store/est/ipa/signature", headers=H, data={"name": "Rahul Menon", "title": "Quantity Surveyor"}, files={"file": ("sig.jpg", buf.getvalue(), "image/jpeg")})
ck("a signature is uploaded and cleaned", r.status_code == 200 and r.json()["sig"].startswith("data:image/png"), r.text[:200])
sg = _I.open(io.BytesIO(__import__("base64").b64decode(r.json()["sig"].split(",", 1)[1])))
ck("...the paper is see-through, cropped to the ink", sg.mode == "RGBA" and sg.getpixel((1, 1))[3] == 0 and sg.width < 600, (sg.mode, sg.size))
r = c.post("/store/est/ipa/signature", headers=H, data={"name": "x"}, files={"file": ("x.txt", b"hello", "text/plain")})
ck("...a file that is not a picture is refused", r.status_code == 400, r.text)
g = glob.glob(D + "*909-IPA09.xlsx")[0]
r = c.post("/store/est/ipa/format", headers=H, data={"project_id": str(projects["909"])}, files={"file": ("909.xlsx", open(g, "rb").read(), "x")})
bx = c.post("/store/est/ipa/new", headers=H, json={"project_id": projects["909"], "format_id": r.json()["id"]}).json()
bb = c.get(f"/store/est/ipa/bill/{bx['id']}", headers=H).json()
neethu = next((x for x in bb["signers"] if "neethu" in x["key"]), None)
ck("909: the preparer is found (Neethu Treesa Jose, Quantity Surveyor, on IPA and APP B) with her signature", neethu and neethu["title"] == "Quantity Surveyor" and set(neethu["sheets"]) >= {"IPA", "APP B"} and neethu["sig"].startswith("data:image"), neethu)
r = c.put(f"/store/est/ipa/bill/{bx['id']}", headers=H, json={"signs": {neethu["key"]: {"name": "Rahul Menon", "title": "Quantity Surveyor"}}})
nb = next(x for x in r.json()["signers"] if x["key"] == neethu["key"])
ck("changed to Rahul Menon: his own signature is on file, not yet on the bill", r.status_code == 200 and nb["name"] == "Rahul Menon" and nb["own"] and not nb["signed_at"], nb)
tok2 = c.post("/auth/download-token", headers=H).json()["token"]
pv = c.get(f"/export/est/ipa/{bx['id']}?token={tok2}&format=preview").text
ck("...the bill prints 'Name: Rahul Menon' on the cover and the appendix", pv.count("Rahul Menon") >= 2 and "Neethu" not in pv, pv.count("Rahul Menon"))
x = c.get(f"/export/est/ipa/{bx['id']}?token={tok2}&format=excel")
zz = zipfile.ZipFile(io.BytesIO(x.content))
ck("...no signature picture in the Excel before anyone signs (the format's own taken off too)", not any("infsig" in n for n in zz.namelist()), [n for n in zz.namelist() if "media" in n][:12])
pdf0 = c.get(f"/export/est/ipa/{bx['id']}?token={tok2}&format=pdf").content
ck("...PDF comes out", pdf0[:4] == b"%PDF")
# ---- the chain: Prepared by (Rahul) -> Checked by (Shafeeq) -> Approved by (Anoop); each in turn, named logins only
sg = r.json()["signers"]
ck("the chain is in order: prepared, checked, approved", [x["step"] for x in sg] == sorted(x["step"] for x in sg) and sg[0]["turn"] and not sg[1]["turn"], [(x["labels"][0], x["name"], x["turn"]) for x in sg])
r = c.post(f"/store/est/ipa/bill/{bx['id']}/sign", headers=H, json={"key": sg[1]["key"]})
ck("the checker cannot sign before the preparer", r.status_code == 400 and "signs first" in r.text, r.text)
# a login for Rahul
# the staff register gives the names: a QS in the register, a login with only his first name
import models as _Mo
from database import SessionLocal as _SL
_db = _SL()
if not _db.query(_Mo.Employee).filter(_Mo.Employee.name == "RAHUL MENON").first():
    _db.add(_Mo.Employee(emp_no="T-RAHUL", name="RAHUL MENON", staff=True, designation="QS", active=True)); _db.commit()
pp = c.get("/store/est/ipa/people", headers=H).json()["people"]
ck("people to sign come from the staff register, designation written out", any(x["name"] == "Rahul Menon" and x["designation"] == "Quantity Surveyor" for x in pp), [x["name"] for x in pp][:8])
ur = c.post("/users", headers=H, json={"username": "rahul", "password": "rahul12345", "full_name": "Rahul", "role": "office"})
HR = None
if ur.status_code in (200, 201):
    c.post(f"/users/{ur.json()['id']}/permissions", headers=H, json={"permissions": "estimation_projects"})
    HR = {"Authorization": "Bearer " + c.post("/auth/login", data={"username": "rahul", "password": "rahul12345"}).json().get("access_token", "")}
    bell = c.get("/notifications", headers=HR).json()["notifications"]
    ck("Rahul's bell shows nothing yet (preparer: his own bill)", not any(n["kind"] == "bill" for n in bell))
    r = c.post(f"/store/est/ipa/bill/{bx['id']}/sign", headers=HR, json={"key": sg[1]["key"]})
    ck("Rahul cannot sign for Shafeeq", r.status_code in (400, 403), r.text)
    pp = c.get("/store/est/ipa/people", headers=H).json()["people"]
    ck("...the register shows his login (first-name login 'Rahul' matched to RAHUL MENON)", any(x["name"] == "Rahul Menon" and x["login"] == "rahul" for x in pp))
r = c.post(f"/store/est/ipa/bill/{bx['id']}/sign", headers=HR or H, json={"key": sg[0]["key"]})
ck("Rahul signs as prepared; now awaiting approval, next is the checker", r.status_code == 200 and r.json()["status"] == "Awaiting approval" and r.json()["next"]["name"] == sg[1]["name"], r.text[:300])
pv = c.get(f"/export/est/ipa/{bx['id']}?token={tok2}&format=excel")
zz = zipfile.ZipFile(io.BytesIO(pv.content))
ck("...his signature is now on the Excel (cover and appendix), the others' still not", sum(1 for n in zz.namelist() if "infsig" in n) == len(sg[0]["sheets"]))
bell = c.get("/notifications", headers=H).json()["notifications"]
ck("admin's bell does not carry it (not their turn)", not any(n["kind"] == "bill" for n in bell))
us = c.post("/users", headers=H, json={"username": "shafiq2", "password": "shafeeq12345", "full_name": "Shafiq", "role": "office"})
if us.status_code in (200, 201):
    c.post(f"/users/{us.json()['id']}/permissions", headers=H, json={"permissions": "estimation_projects"})
    HS = {"Authorization": "Bearer " + c.post("/auth/login", data={"username": "shafiq2", "password": "shafeeq12345"}).json()["access_token"]}
    bell = c.get("/notifications", headers=HS).json()["notifications"]
    ck("login 'Shafiq' gets the bell for 'Shafeeq' on the bill (spelt differently, same person)", any(n["kind"] == "bill" and n["bill"] == bx["id"] for n in bell), [n["title"] for n in bell][:3])
r = c.post(f"/store/est/ipa/bill/{bx['id']}/sign", headers=H, json={"key": sg[1]["key"]})
ck("admin signs on the checker's behalf", r.status_code == 200 and r.json()["signers"][1]["signed_at"], r.text[:200])
for _ in range(6):
    cur = c.get(f"/store/est/ipa/bill/{bx['id']}", headers=H).json()
    nxt = next((x for x in cur["signers"] if x["turn"]), None)
    if not nxt:
        break
    r = c.post(f"/store/est/ipa/bill/{bx['id']}/sign", headers=H, json={"key": nxt["key"]})
ck("...and the rest in turn, the approver last: the bill is Approved", r.status_code == 200 and r.json()["status"] == "Approved" and all(x["signed_at"] for x in r.json()["signers"]), r.text[:200])
r = c.post(f"/store/est/ipa/bill/{bx['id']}/sign", headers=H, json={"key": sg[1]["key"], "undo": True})
ck("taking the checker's signing back also takes back the approver's; status back to awaiting", r.status_code == 200 and r.json()["status"] == "Awaiting approval" and not r.json()["signers"][-1]["signed_at"] and r.json()["signers"][0]["signed_at"], r.text[:200])
r = c.put(f"/store/est/ipa/bill/{bx['id']}", headers=H, json={"signs": {neethu["key"]: {"name": "Someone New", "title": "QS"}}})
nb = next(x for x in r.json()["signers"] if x["key"] == neethu["key"])
ck("a different preparer named: they sign afresh, no signature on the bill", nb["sig"] == "" and not nb["own"] and not nb["signed_at"], nb)
b3 = c.post("/store/est/ipa/new", headers=H, json={"project_id": projects["909"], "format_id": bx and r.json()["format_id"]}).json()
n3 = next(x for x in c.get(f"/store/est/ipa/bill/{b3['id']}", headers=H).json()["signers"] if x["key"] == neethu["key"])
ck("next month's bill keeps the same people, unsigned", n3["name"] == "Someone New" and not n3["signed_at"], n3)
# ---- payments to authorities (913 APP B): this bill's lines typed, lines added, last month's hidden, sub-total this bill only
g = glob.glob(D + "*913-IPA_09.xlsx")[0]
r = c.post("/store/est/ipa/format", headers=H, data={"project_id": str(projects["913"]), "name": "913 fees"}, files={"file": ("913.xlsx", open(g, "rb").read(), "x")})
bf = c.post("/store/est/ipa/new", headers=H, json={"project_id": projects["913"], "format_id": r.json()["id"]}).json()
bb = c.get(f"/store/est/ipa/bill/{bf['id']}", headers=H).json()
fe = bb["fees"][0]
ck("913: the authority fees table is found, with this bill's heading and an empty line, last month's lines hidden", fe["rows"][0]["head"] and fe["rows"][0]["desc"]["value"] == "IPA 10" and fe["hidden"] >= 30 and fe["total"] == 0, [(x["r"], x["desc"]["value"]) for x in fe["rows"]])
ln = fe["rows"][1]
r = c.put(f"/store/est/ipa/bill/{bf['id']}", headers=H, json={"set": {fe["sheet"]: {ln["desc"]["key"]: "DM: Pergola permit fee", ln["ref"]["key"]: "WIFBP2030001 dtd 20.10.2026", ln["amount"]["key"]: 150}}})
ck("a fee typed: sub-total 150, on the cover 150", r.json()["fees"][0]["total"] == 150 and any(x["amount"] == 150 for x in r.json()["cover"]["rows"] if "uthorit" in x["label"]), r.json()["fees"][0]["total"])
r = c.post(f"/store/est/ipa/bill/{bf['id']}/rows", headers=H, json={"sheet": fe["sheet"], "at": ln["r"] + 1, "n": 1, "like": ln["r"]})
ck("a line added (a row put in, as in Excel)", r.status_code == 200 and len(r.json()["fees"][0]["rows"]) == len(fe["rows"]) + 1, r.text[:200])
nr = next(x for x in r.json()["fees"][0]["rows"] if x["r"] == ln["r"] + 1)
r = c.put(f"/store/est/ipa/bill/{bf['id']}", headers=H, json={"set": {fe["sheet"]: {nr["desc"]["key"]: "DM: Second fee", nr["amount"]["key"]: 130}}})
ck("...typed: sub-total 280 and the cover follows", r.json()["fees"][0]["total"] == 280 and any(x["amount"] == 280 for x in r.json()["cover"]["rows"] if "uthorit" in x["label"]), r.json()["fees"][0]["total"])
x = c.get(f"/export/est/ipa/{bf['id']}?token={tok2}&format=excel")
wbf = openpyxl.load_workbook(io.BytesIO(x.content)); ws = wbf["APP B"]
tr = next(rr for rr in range(1, 60) if "sub - total" in str(ws.cell(rr, 2).value or "").lower())
ck("Excel: the sub-total sums this bill's rows only, last month's rows hidden", str(ws.cell(tr, 4).value).startswith("=SUM(") and ws.row_dimensions[tr - 5].hidden and not ws.row_dimensions[tr - 1].hidden, (ws.cell(tr, 4).value, tr))
ck("...the cover refers to the moved sub-total cell", f"D{tr}" in str(wbf["IPA"].cell(19, 5).value), wbf["IPA"].cell(19, 5).value)
pv = c.get(f"/export/est/ipa/{bf['id']}?token={tok2}&format=preview").text
ck("preview prints this bill's fees, not last month's", "Pergola permit" in pv and "Second fee" in pv and "Pergola Drawing Submission" not in pv)
c.post(f"/store/est/ipa/bill/{bf['id']}/status", headers=H, json={"status": "Submitted"})
bn = c.post("/store/est/ipa/new", headers=H, json={"project_id": projects["913"], "format_id": bf and bb["format_id"]}).json()
fn = c.get(f"/store/est/ipa/bill/{bn['id']}", headers=H).json()["fees"][0]
ck("next month: IPA 11 heading, this month's fees hidden, sub-total 0", fn["rows"][0]["desc"]["value"] == "IPA 11" and fn["total"] == 0 and fn["hidden"] >= 33, [(x["r"], x["desc"]["value"]) for x in fn["rows"]])
# ---- the company stamp: bottom right of every page (in the footer), never dropped on the tables; logos keep their shape
import re as _re, base64 as _b64
from PIL import Image as _I
pv = c.get(f"/export/est/ipa/{bn['id']}?token={c.post('/auth/download-token', headers=H).json()['token']}&format=preview").text
pages = pv.count('<section class="pg')
ck("stamp at the bottom right of every page", pv.count('class="stp"') == pages and pages >= 9, (pv.count('class="stp"'), pages))
imgs = _re.findall(r'<img class="xi" src="data:image/(\w+);base64,([^"]+)" style="left:[\d.]+px;top:[\d.]+px;width:(\d+)px;height:(\d+)px"', pv)
bad = []
for ext, data, w, h in imgs:
    iw, ih = _I.open(io.BytesIO(_b64.b64decode(data))).size
    if int(w) >= 12 and int(h) >= 12 and abs((int(w) / int(h)) / (iw / ih) - 1) > 0.06:
        bad.append((w, h, iw, ih))
ck("every logo / picture keeps its own shape (not stretched)", imgs and not bad, bad[:5])
ck("long text runs over without moving its cell (no positioned cells over the grid)", 'td.ov{overflow:visible}' in pv and "td.ov>span" in pv)
x = c.get(f"/export/est/ipa/{bn['id']}?token={c.post('/auth/download-token', headers=H).json()['token']}&format=excel")
z = zipfile.ZipFile(io.BytesIO(x.content))
sx = z.read("xl/worksheets/sheet3.xml").decode()
ck("Excel: the stamp is in the footer of each sheet (&G with its picture)", "&amp;R&amp;G" in sx and "legacyDrawingHF" in sx and "xl/media/infstamp.png" in z.namelist())
wbx = openpyxl.load_workbook(io.BytesIO(x.content))
ck("Excel: the workbook still opens", "APP B" in wbx.sheetnames)
print(f"\n{len(fails)} FAILED: {fails}" if fails else "\nALL PASS"); sys.exit(1 if fails else 0)
