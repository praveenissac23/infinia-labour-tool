"""Internal BOQ workbook -> Overview > BOQ & budget.

    cd app && DATABASE_URL=sqlite:////tmp/boqimp.db INFINIA_NO_SCHED=1 python3 ../tests/est_boqimport_test.py [internal_boq.xlsx ...]

A small workbook in the internal BOQ's layout is made here; any real internal BOQ files given on the
command line are read too and must add up to their own summary sheet to the fils.
"""
import io, os, sys
sys.path.insert(0, os.getcwd())
import openpyxl
from fastapi.testclient import TestClient
import main, estimation as E, est_boqimport as B
from database import SessionLocal, engine

main._add_missing_columns()
E.create_tables(engine)
c = TestClient(main.app); c.__enter__()
fails = []
def ck(name, ok, extra=""):
    print(("PASS " if ok else "FAIL ") + name + ("" if ok else f"  -> {extra}"))
    if not ok: fails.append(name)


def book(extra_line=False, elec_summary=None):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "COVER PAGE"
    p = wb.create_sheet("1.PRELIM")
    for r in [["1. PRELIMINARIES INCLUDING MOBILIZATION"], ["S/NO.", "ITEMS", "LUMPSUM AMOUNT"], [1, "Insurance", 6000],
              [2, "Municipality permit", "Authority payment"], [3, "Allowance for site administrator", None],
              [None, "1 Engineer (part time)", 145000], [None, "1 Foreman", 60000], [4, "Watchman", None], [None, "TOTAL", 211000]]:
        p.append(r)
    s = wb.create_sheet("2.SUBSTR")
    rows = [["2. SUBSTRUCTURE"], ["A. EARTHWORKS"], ["S/NO.", "ITEMS", "CONS. QTY", "IC QTY", "UNIT", "RATE", "LUMPSUM AMOUNT"],
            [1, "Excavation", None, 100, "m3", 30, 3000], [2, "Imported fill", None, 10.5, "m3", 25, 262.5],
            [3, "Others", None, None, "m2", None, 0], [None, "TOTAL", 3262.5], [],
            ["B. DEWATERING WORKS"], ["S/NO.", "ITEMS", "CONS. QTY", "IC QTY", "UNIT", "RATE", "LUMPSUM AMOUNT"],
            [4, "Dewatering", None, 30, "Days", 3200, 96000], [None, "TOTAL", "EXCLUDED"], [],
            ["C. CONCRETE"], ["S/NO.", "ITEMS", "CONS. QTY", "IC QTY", "UNIT", "RATE", "LUMPSUM AMOUNT"],
            [5, "Raft C40", None, 597, "m3", 1050, 626850], [6, "Curved column shutter", None, 5, "m3", "LS", 8000],
            [7, "Pool walls", None, 0, "m3", 1100, "Included in beams"]]
    if extra_line:
        rows.append([8, "Extra item added by the QS", None, 2, "No.", 500, 1000])
    rows += [[None, "TOTAL", 634850], [], ["SUMMARY"], ["A.", "EARTHWORKS", 3262.5], ["B.", "DEWATERING WORKS", "Excluded"],
             ["C.", "CONCRETE", 634850 + (1000 if extra_line else 0)], [None, "SUBSTRUCTURE TOTAL", 0]]
    for r in rows:
        s.append(r)
    j = wb.create_sheet("7. JOINERY")
    for r in [["7. CARPENTRY AND JOINERY"], ["A. DOORS"], ["S/NO.", "ITEM", "SIZE", None, "AREA", "TYPE", "CONS QTY", "IC QTY", "UNIT", "RATE", "AMOUNT"],
              [None, None, "L (m)", "H (m)"], [None, "Supply and fixing of all wooden doors"], [1, "D1", 1.2, 2.2, 2.64, "Single", None, 3, "No.", 2000, 6000]]:
        j.append(r)
    e = wb.create_sheet("10. ELEC.")
    for r in [["10. ELECTRICAL WORKS"], ["S/NO.", "ITEMS", "CONS. QTY", "IC QTY", "UNIT", "RATE", "AMOUNT"], [1, "Builders work", None, 1, "LS", 25000, 25000],
              [2, "OHP", None, None, 0.1, None, None], [None, "ELECTRICAL TOTAL", 25000], [], ["PRICE COMPARISON - ELECTRICAL"],
              ["SL.NO", "Item Description", "Subcon 1"], [None, "Electrical", 175000]]:
        e.append(r)
    m = wb.create_sheet("SUMMARY")
    sub = 3262.5 + 634850 + (1000 if extra_line else 0)
    elec = 25000 if elec_summary is None else elec_summary
    for r in [["TOTAL SUMMARY"], ["S/NO.", "DESCRIPTION", "AMOUNT", "DISCOUNT"], ["01.", "Preliminaries", 211000], ["02.", "Sub structure", sub],
              ["07.", "Carpentry & Joinery", "Phase 02"], ["10.", "Electrical", elec, "Yet to Confirm" if elec_summary else None],
              [None, "TOTAL", 211000 + sub + elec - 100]]:
        m.append(r)
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()


# ---- reading ----
r = B.read_internal_boq(book())
L = {x["description"]: x for x in r["lines"]}
ck("prelim lump sums come in as 1 LS x amount", L["Insurance"]["qty"] == 1 and L["Insurance"]["unit"] == "LS" and L["Insurance"]["amount"] == 6000)
ck("an amount in words is kept as the remark, no money", L["Municipality permit"]["remark"] == "Authority payment" and L["Municipality permit"]["amount"] == 0)
ck("sub-items carry their lead-in line", "Allowance for site administrator - 1 Engineer (part time)" in L and L["Allowance for site administrator - 1 Engineer (part time)"]["amount"] == 145000)
ck("unpriced template rows and OHP / Others left out", "Watchman" not in L and "Others" not in L and not any(d.startswith("OHP") for d in L))
ck("bill and section names", L["Excavation"]["bill"] == "02. Substructure" and L["Excavation"]["section"] == "A. EARTHWORKS")
ck("decimals kept (10.5 m3 x 25 = 262.50)", L["Imported fill"]["qty"] == 10.5 and L["Imported fill"]["amount"] == 262.5)
ck("a section the sheet's summary excludes is kept but left out", L["Dewatering"]["scope"] == "Excluded" and L["Dewatering"]["amount"] == 96000)
ck("a lump sum whose rate is 'LS' keeps its amount", L["Curved column shutter"]["amount"] == 8000 and L["Curved column shutter"]["rate"] == 0)
ck("a bill the summary marks Phase 02 is kept but left out", L["Supply and fixing of all wooden doors - D1 (1.2 x 2.2 m, Single)"]["scope"] == "Phase 02")
ck("price comparison tables are not BOQ lines", "Electrical" not in L)
ck("total = summary total, the discount added as a line", r["total"] == r["summary_total"] and any(x["description"].startswith("Discount") and x["amount"] == -100 for x in r["lines"]), (r["total"], r["summary_total"]))
r2 = B.read_internal_boq(book(extra_line=True))
ck("an extra line added in the workbook comes in", any(x["description"] == "Extra item added by the QS" and x["amount"] == 1000 for x in r2["lines"]))
r3 = B.read_internal_boq(book(elec_summary=650000))
adj = [x for x in r3["lines"] if x["description"].startswith("Difference")]
ck("a bill whose summary figure differs gets the difference as a line, with the note", len(adj) == 1 and adj[0]["amount"] == 625000 and "Yet to Confirm" in adj[0]["description"] and r3["total"] == r3["summary_total"])
# ---- the real workbooks, if given: every one adds up to its own summary ----
for f in sys.argv[1:]:
    rr = B.read_internal_boq(open(f, "rb").read())
    ck(f"{os.path.basename(f)}: {len(rr['lines'])} lines, {rr['total']:,.2f} = summary {rr['summary_total']:,.2f}", rr["summary_total"] is not None and abs(rr["total"] - rr["summary_total"]) < 0.005)

# ---- through the app ----
tok = c.post("/auth/login", data={"username": "admin", "password": "changeme123"}).json()["access_token"]
H = {"Authorization": "Bearer " + tok}
db = SessionLocal()
p = db.query(E.EstProject).filter_by(code="BOQT").first() or E.EstProject(code="BOQT", name="BOQ test")
db.add(p); db.commit(); pid = p.id
db.query(E.EstBoq).filter_by(project_id=pid).delete(); db.commit()
db.add(E.EstBoq(project_id=pid, section="OLD", code="1", description="Excavation", unit="m3", qty=90, rate=30, cost_rate=22, done_pct=40)); db.commit()
up = lambda data, **kw: c.post(f"/store/est/projects/{pid}/boq-import", files={"file": ("922 Internal BOQ.xlsx", data)}, data={k: str(v) for k, v in kw.items()}, headers=H)
x = up(book())
ck("preview reads, saves nothing", x.status_code == 200 and x.json()["existing"] == 1 and db.query(E.EstBoq).filter_by(project_id=pid).count() == 1, x.text[:200])
x = up(book(), save=1, mode="replace").json()
db.expire_all()
rows = db.query(E.EstBoq).filter_by(project_id=pid).all()
ck("replace: lines in, the old one out", len(rows) == x["saved"] and not any(b.section == "OLD" for b in rows))
ex = next(b for b in rows if b.description == "Excavation")
ck("replace: cost rate and done % typed before stay on the line that came back", ex.cost_rate == 22 and ex.done_pct == 40 and x["kept_budget"] == 1)
bq = c.get(f"/store/est/projects/{pid}", headers=H).json()
ck("BOQ value counts only the awarded scope, to the fils", bq["summary"]["boq_value"] == r["total"], (bq["summary"]["boq_value"], r["total"]))
ck("budget cost from the cost rate (100 x 22)", bq["summary"]["boq_cost"] == 2200)
line = next(b for b in bq["boq"] if b["description"] == "Curved column shutter")
ck("a given amount is shown as given", line["amount"] == 8000 and line["given"])
c.patch(f"/store/est/boq/{line['id']}", json={"qty": 6, "rate": 1500}, headers=H)
line = next(b for b in c.get(f"/store/est/projects/{pid}", headers=H).json()["boq"] if b["id"] == line["id"])
ck("typing qty / rate works the amount out again", line["amount"] == 9000 and not line["given"])
c.post(f"/store/est/projects/{pid}/boq-scope", json={"bill": "07. Carpentry and Joinery", "scope": ""}, headers=H)
v2 = c.get(f"/store/est/projects/{pid}", headers=H).json()["summary"]["boq_value"]
ck("bringing the Phase 02 bill into scope counts it", abs(v2 - (r["total"] + 1000 + 6000)) < 0.005, v2)
c.post(f"/store/est/projects/{pid}/boq-scope", json={"bill": "02. Substructure", "section": "B. DEWATERING WORKS", "scope": ""}, headers=H)
v3 = c.get(f"/store/est/projects/{pid}", headers=H).json()["summary"]["boq_value"]
ck("one section brought in on its own", abs(v3 - v2 - 96000) < 0.005, v3)
n0 = db.query(E.EstBoq).filter_by(project_id=pid).count()
up(book(extra_line=True), save=1, mode="add")
ck("add: lines go after the ones there", db.query(E.EstBoq).filter_by(project_id=pid).count() > n0)
ck("not a workbook is refused", up(b"hello").status_code == 400)
t = c.post("/auth/download-token", headers=H).json()["token"]
ck("BOQ exports with the bill and remarks", c.get(f"/export/est/boq?pid={pid}&token={t}&format=xlsx").status_code == 200)
ck("no file kept", not any(f.kind == "boq" for f in db.query(E.EstFile).filter_by(project_id=pid)) if hasattr(E.EstFile, "kind") else True)
db.query(E.EstBoq).filter_by(project_id=pid).delete(); db.delete(db.get(E.EstProject, pid)); db.commit()
print(f"\n{len(fails)} FAILED: {fails}" if fails else "\nALL PASS")
sys.exit(1 if fails else 0)
