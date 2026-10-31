"""Estimation arithmetic audit - every money figure worked out again here, independently of the app's code (Decimal,
half-up to the fils, written from the rules as the certificates print them), and compared with what the app gives:
the project page, the dashboard, the Accounts payment list, every Preview / Excel export, every printed subcontractor
certificate, and the bill workbooks (Excel's "precision as displayed", the bill saved = the bill exported = the next
bill's previous).

    cp /tmp/claude-0/full.db /tmp/arith.db
    cd app && DATABASE_URL=sqlite:////tmp/arith.db INFINIA_NO_SCHED=1 INFINIA_NO_PUSH=1 python3 ../tests/est_arith_audit_test.py

It adds one project of its own (code ARITH-T, two contracts, a lump-sum and a re-measured subcontract) so every rule
is exercised on any database, and then audits every project in the database. Nothing that was there is changed.

The rules, as worked out here:
  line (qty)    = qty x rate x stage% (100% when blank) - deduction                     to the fils
  line (pct)    = (line amount - deduction) x % done                                     to the fils
  work / variations to date = sum of the lines (schedule lines marked variation go to variations)
  net to date   = work + variations + material on site + advance paid to date - advance recovered - retention
                  + retention released - other deductions
  this PC       = net to date - net to date of the PC before (on the same subcontract / client contract)
  VAT           = this PC x VAT%;  payable = this PC + VAT
  retention (auto)  = (work + variations) x retention%
  recovery (auto)   = min(advance paid to date, work x recovery%)   (recovery% = the fixed % / the advance %)
  revised subcontract = value + variations - cancelled;  a re-measured one (no value) = the larger of that and the work measured
  to pay        = certified to date + VAT on all PCs - paid;   balance = revised - (work + variations)
  committed     = revised of every subcontract not cancelled;  margin = contract + approved VOs - committed
"""
import io, os, sys, json, re, zipfile, glob, warnings
from decimal import Decimal as D, ROUND_HALF_UP
warnings.filterwarnings("ignore")
sys.path.insert(0, os.getcwd())
from fastapi.testclient import TestClient
import openpyxl
import main, models, estimation as E, est_pcsheet as S, est_bills as EB, est_bills_engine as BE
from database import SessionLocal, engine

main._add_missing_columns()
E.create_tables(engine)
c = TestClient(main.app); c.__enter__()
TOK = c.post("/auth/login", data={"username": "admin", "password": "changeme123"}).json()["access_token"]
H = {"Authorization": "Bearer " + TOK}
db = SessionLocal()
fails, n = [], [0]


def m2(x):
    return D(str(x or 0)).quantize(D("0.01"), ROUND_HALF_UP)


def same(name, app, want):
    n[0] += 1
    if app is None or abs(D(str(app)) - D(str(want))) > D("0.005"):
        fails.append(f"{name}: app {app} / worked out {want}")


def ck(name, ok, extra=""):
    n[0] += 1
    if not ok:
        fails.append(f"{name} {extra}")


def post(path, body):
    r = c.post(path, json=body, headers=H)
    assert r.status_code == 200, (path, r.status_code, r.text)
    return r.json()


# ======================================================================== a project of its own
for old in db.query(E.EstProject).filter(E.EstProject.code == "ARITH-T").all():
    c.delete(f"/store/est/projects/{old.id}", headers=H)
db.expire_all()
P = E.EstProject(code="ARITH-T", name="Arithmetic audit villa", status="Running", adv_pct=0, ret_pct=10, vat_pct=5)
db.add(P); db.commit()
PID = P.id
# the second contract makes the project's own agreement its first one ("Shell & core")
k2 = post("/store/est/contracts", {"project_id": PID, "name": "MEP", "value": 400000, "adv_pct": 20, "ret_pct": 5, "vat_pct": 5, "first_name": "Shell & core"})["id"]
K1 = next(k.id for k in db.query(E.EstContract).filter_by(project_id=PID).all() if k.id != k2)
rp = c.patch(f"/store/est/contracts/{K1}", json={"value": 1000000, "adv_pct": 10, "ret_pct": 10, "vat_pct": 5}, headers=H)
assert rp.status_code == 200, rp.text
# client: an advance and three interim bills on contract 1, two on contract 2 - retention / recovery worked out by the app
post("/store/est/clientpcs", {"project_id": PID, "contract_id": K1, "pc_type": "Advance", "status": "Certified"})
for cum, st in ((123456.78, "Paid"), (333333.33, "Certified"), (501234.57, "Submitted")):
    post("/store/est/clientpcs", {"project_id": PID, "contract_id": K1, "claimed_cum": cum, "certified_cum": cum if st != "Submitted" else 0, "status": st,
                                  "adv_auto": True, "ret_auto": True})
for cum in (45678.91, 98765.43):
    post("/store/est/clientpcs", {"project_id": PID, "contract_id": k2, "claimed_cum": cum, "status": "Draft", "adv_auto": False, "adv_rec_cum": round(cum * 0.2, 2), "ret_auto": True})
post("/store/est/variations", {"project_id": PID, "contract_id": K1, "title": "Extra pergola", "status": "Approved", "submitted_amount": 25000, "approved_amount": 21345.67})
post("/store/est/variations", {"project_id": PID, "contract_id": k2, "title": "Omitted spot lights", "kind": "Omission", "status": "Approved", "approved_amount": -3210.5})
post("/store/est/variations", {"project_id": PID, "contract_id": k2, "title": "Pending", "status": "Submitted", "submitted_amount": 9999})
for q, r_, a in ((12.5, 333.33, None), (3, 6539.45, 19618.34), (7.25, 41.1, None)):
    post("/store/est/boq", {"project_id": PID, "description": "line", "qty": q, "rate": r_, **({"amount": a} if a else {})})
post("/store/est/boq", {"project_id": PID, "description": "by others", "qty": 1, "rate": 5000, "scope": "By client"})
# subcontract 1: lump sum by %, advance 20% recovered at a fixed 10%, retention 10%, a variation line and a deduction
s1 = post("/store/est/subs", {"project_id": PID, "name": "AUDIT MEP LLC", "value": 568800, "adv_pct": 20, "ret_pct": 10, "vat_pct": 5,
                              "adv_rec_mode": "pct", "adv_rec_pct": 10})["id"]
post("/store/est/subvos", {"sub_id": s1, "title": "Extra DB", "amount": 12345.6})
items = c.put(f"/store/est/subs/{s1}/items", headers=H, json={"items": [
    {"description": "Electrical", "qty": 1, "unit": "LS", "amount": 172800, "measure": "pct"},
    {"description": "Plumbing", "qty": 1, "unit": "LS", "amount": 153000, "measure": "pct"},
    {"description": "VRF", "qty": 1, "unit": "LS", "amount": 243000, "measure": "pct"},
    {"description": "Extra DB", "qty": 2, "unit": "No", "rate": 6172.8, "measure": "qty", "is_var": True}]}).json()["items"]
IT = {i["description"]: i["id"] for i in items}
post("/store/est/subpcs", {"sub_id": s1, "pc_type": "Advance"})
for lines, deds in ((((IT["Electrical"], 12.5, 0), (IT["Plumbing"], 7.33, 0)), []),
                    (((IT["Electrical"], 33.33, 0), (IT["Plumbing"], 21.5, 0), (IT["VRF"], 5, 0), (IT["Extra DB"], 50, 1)), [("Material", 2, 117.35)]),
                    (((IT["Electrical"], 61.7, 0), (IT["Plumbing"], 21.5, 0), (IT["VRF"], 18.25, 0), (IT["Extra DB"], 100, 2)), [])):
    pc = post("/store/est/subpcs", {"sub_id": s1, "inv_ref": "INV"})["id"]
    body = {"lines": [{"item_id": i, "pct": p, "qty": q} for i, p, q in lines]}
    if deds:
        body["deds"] = [{"cat": ct, "description": "supplied", "qty": q, "rate": rt} for ct, q, rt in deds]
    r = c.put(f"/store/est/subpcs/{pc}/sheet", headers=H, json=body)
    assert r.status_code == 200, r.text
# subcontract 2: re-measured quantities, a typed first certificate, retention released, an other deduction
s2 = post("/store/est/subs", {"project_id": PID, "name": "AUDIT SHORING", "value": 0, "adv_pct": 0, "ret_pct": 10, "vat_pct": 5})["id"]
items = c.put(f"/store/est/subs/{s2}/items", headers=H, json={"items": [
    {"description": "Footing", "unit": "m3", "rate": 235, "measure": "qty"}, {"description": "Slab", "unit": "m3", "rate": 120.5, "measure": "qty"}]}).json()["items"]
IT2 = {i["description"]: i["id"] for i in items}
for q1, q2 in ((88.32, 0), (141.5, 20.33), (141.5, 48.07)):
    pc = post("/store/est/subpcs", {"sub_id": s2, "inv_ref": "RE"})["id"]
    c.put(f"/store/est/subpcs/{pc}/sheet", headers=H, json={"lines": [{"item_id": IT2["Footing"], "qty": q1}, {"item_id": IT2["Slab"], "qty": q2}]})
last2 = pc
c.put(f"/store/est/subpcs/{last2}/sheet", headers=H, json={"pc": {"ret_rel_cum": 1000}, "deds": [{"cat": "Other", "description": "water", "amount": 550}]})
db.expire_all()

# ======================================================================== the rules, written out here
def line_val(it, ln):
    if (it.measure or "pct") == "qty":
        stage = D(str(ln.pct)) / 100 if ln.pct else D(1)
        return m2(D(str(ln.qty or 0)) * D(str(it.rate or 0)) * stage - D(str(ln.deduct or 0)))
    return m2((D(str(it.amount or 0)) - D(str(ln.deduct or 0))) * D(str(ln.pct or 0)) / 100)


def chain(rows, vat_pct, work_of):
    out, prev, adv = {}, D(0), D(0)
    for p in sorted(rows, key=lambda p: (p.seq or 0, p.id)):
        adv = m2(adv + D(str(p.adv_amount or 0)))
        work, var = m2(work_of(p)), m2(getattr(p, "var_cum", 0) or 0)
        net = m2(work + var + D(str(p.mos_cum or 0)) + adv - D(str(p.adv_rec_cum or 0)) - D(str(p.ret_cum or 0))
                 + D(str(p.ret_rel_cum or 0)) - D(str(p.other_cum or 0)))
        this = m2(net - prev)
        vat = m2(this * D(str(vat_pct or 0)) / 100)
        out[p.id] = dict(work=work, var=var, adv=adv, net=net, prev=prev, this=this, vat=vat, payable=m2(this + vat))
        prev = net
    return out


def client_work(x):
    return D(str((x.certified_cum if (x.certified_cum or 0) > 0 or x.status in ("Certified", "Invoiced", "Paid") else x.claimed_cum) or 0))


def dl():
    return c.post("/auth/download-token", headers=H).json()["token"]


def xlsx_rows(path_query):
    r = c.get(path_query + f"&token={dl()}&format=excel")
    assert r.status_code == 200, (path_query, r.status_code, r.text[:200])
    ws = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, row in enumerate(rows) if row and any(isinstance(v, str) and v.strip() in ("PC No.", "Subcontractor", "Bill", "Section", "Item", "Description", "Amount") for v in row))
    head = [str(v).strip() if v is not None else "" for v in rows[hi]]
    body = [dict(zip(head, row)) for row in rows[hi + 1:] if any(v not in (None, "") for v in row)
            and not any(isinstance(v, str) and v.strip().lower() in ("total", "totals") for v in row)]   # the totals line under the table
    return body


dash = c.get("/store/est/dash", headers=H).json()
drow = {r["id"]: r for r in dash["projects"]}
dtot = {k: D(0) for k in dash["totals"]}
pay_rows = c.get("/store/est/payqueue", headers=H).json().get("rows", [])
all_payable = {}
counts = dict(projects=0, subs=0, subpcs=0, lines=0, clientpcs=0, docs=0, exports=0)
for p in db.query(E.EstProject).order_by(E.EstProject.code).all():
    counts["projects"] += 1
    b = c.get(f"/store/est/projects/{p.id}", headers=H).json()
    sm, tag = b["summary"], p.code
    committed = cert = paid = to_pay = D(0)
    for sb in b["subs"]:
        counts["subs"] += 1
        s = db.get(E.EstSub, sb["id"])
        vos = db.query(E.EstSubVO).filter_by(sub_id=s.id).all()
        pcs = db.query(E.EstSubPC).filter_by(sub_id=s.id).all()
        items = {i.id: i for i in db.query(S.EstSubItem).filter_by(sub_id=s.id).all()}
        ch = chain(pcs, s.vat_pct, lambda x: D(str(x.work_cum or 0)))
        last = ch[max(ch, key=lambda k: (db.get(E.EstSubPC, k).seq or 0, k))] if ch else None
        revised = m2(D(str(s.value or 0)) + sum(D(str(v.amount or 0)) for v in vos) - D(str(s.cancelled or 0)))
        if not (s.value or 0) and last:
            revised = max(revised, m2(last["work"] + last["var"]))          # re-measured: what has been measured so far
        same(f"{tag} {s.name}: revised value", sb["revised"], revised)
        rule = E.sub_rec_rule(s)
        locked_upto = max([i for i, x in enumerate(sorted(pcs, key=lambda x: (x.seq or 0, x.id))) if x.status in E.LOCKED_STATUS] or [-1])
        for i, pc in enumerate(sorted(pcs, key=lambda x: (x.seq or 0, x.id))):
            counts["subpcs"] += 1
            w, a = ch[pc.id], next(x for x in sb["pcs"] if x["id"] == pc.id)
            nm = f"{tag} {s.name} PC{pc.seq:02d}"
            for k, kk in (("work", "work"), ("var", "var"), ("adv_paid_cum", "adv"), ("net_cum", "net"), ("prev_net", "prev"), ("this", "this"), ("vat", "vat"), ("payable", "payable")):
                same(f"{nm} {k}", a[k], w[kk])
            all_payable[pc.id] = w["payable"]
            lines = db.query(S.EstPCLine).filter_by(pc_id=pc.id).all()
            if lines and items:
                counts["lines"] += len(lines)
                same(f"{nm} work to date = its lines", pc.work_cum, sum((line_val(items[l.item_id], l) for l in lines if l.item_id in items and not items[l.item_id].is_var), D(0)))
                same(f"{nm} variations to date = its lines", pc.var_cum or 0, sum((line_val(items[l.item_id], l) for l in lines if l.item_id in items and items[l.item_id].is_var), D(0)))
            if i > locked_upto:
                if pc.ret_auto:
                    same(f"{nm} retention {s.ret_pct}% auto", pc.ret_cum, m2((D(str(pc.work_cum or 0)) + D(str(pc.var_cum or 0))) * D(str(s.ret_pct or 0)) / 100))
                if pc.adv_auto and rule["mode"] in ("", "pct") and (pc.seq or 1) >= rule["from"] and rule["pct"]:
                    same(f"{nm} advance recovered {rule['pct']}% auto", pc.adv_rec_cum, min(w["adv"], m2(D(str(pc.work_cum or 0)) * D(str(rule["pct"])) / 100)))
            if items:                                     # the printed certificate
                counts["docs"] += 1
                d = S.build_doc(db, pc)
                same(f"{nm} printed: A = work + variations", d["A_total"], w["work"] + w["var"])
                same(f"{nm} printed: E = previous certificates", d["E_total"], w["prev"])
                same(f"{nm} printed: E lines add up to E", sum((D(str(x[1])) for x in d["E"]), D(0)), w["prev"])
                same(f"{nm} printed: net", d["net"], w["this"])
                same(f"{nm} printed: VAT", d["vat"], w["vat"])
                same(f"{nm} printed: payable", d["payable"], w["payable"])
                same(f"{nm} printed: A + B - C + D + MOS - E = net",
                     m2(D(str(d["A_total"])) + sum(D(str(x[2])) for x in d["B"]) - D(str(d["C_total"])) + sum(D(str(x[2])) for x in d["D"]) + D(str(d["mos"])) - D(str(d["E_total"]))), w["this"])
                same(f"{nm} printed: C lines add up to C", sum((D(str(x[2])) for x in d["C"]), D(0)), d["C_total"])
                if lines:
                    same(f"{nm} printed: Appendix A = A", sum((D(str(bl["cum"])) for bl in d["bills"]), D(0)), d["A_total"])
                same(f"{nm} printed: history = earlier certificates", d["hist_total"]["amount"], w["prev"])
        s_cert = last["net"] if last else D(0)
        s_paid = m2(sum((D(str(x.paid_amount or 0)) for x in pcs), D(0)))
        s_vat = m2(sum((v["vat"] for v in ch.values()), D(0)))
        same(f"{tag} {s.name}: certified to date", sb["certified"], s_cert)
        same(f"{tag} {s.name}: paid", sb["paid"], s_paid)
        same(f"{tag} {s.name}: to pay", sb["to_pay"], m2(s_cert + s_vat - s_paid))
        same(f"{tag} {s.name}: balance", sb["balance"], m2(revised - ((last["work"] + last["var"]) if last else 0)))
        if s.status != "Cancelled":
            committed += revised
        cert += s_cert; paid += s_paid; to_pay += m2(s_cert + s_vat - s_paid)
        if pcs:                                           # the Excel export of the certificates: each line adds up across
            counts["exports"] += 1
            for row in xlsx_rows(f"/export/est/subpcs?pid={p.id}&sub_id={s.id}"):
                if not str(row.get("PC No.") or "").startswith("PC No."):
                    continue
                g = lambda k: D(str(row.get(k) or 0))
                same(f"{tag} {s.name} Excel {row['PC No.']}: the line adds up to Net to date",
                     g("Work done to date") + g("Variations to date") + g("Material on site") + g("Advance paid") + g("Adv. recovery") + g("Retention")
                     + g("Other ded.") + g("Ret. released"), g("Net to date"))
                same(f"{tag} {s.name} Excel {row['PC No.']}: payable = this + VAT", g("This certificate") + g("VAT"), g("Payable"))
    same(f"{tag}: committed", sm["committed"], m2(committed))
    same(f"{tag}: subcontractors certified", sm["sub_certified"], m2(cert))
    same(f"{tag}: subcontractors paid", sm["sub_paid"], m2(paid))
    same(f"{tag}: subcontractors to pay", sm["sub_to_pay"], m2(to_pay))
    # ---- client side: each contract's certificates on their own (or the project's when it has none)
    vars_ = db.query(E.EstVariation).filter_by(project_id=p.id).all()
    appr = m2(sum((D(str(v.approved_amount or 0)) for v in vars_ if v.status == "Approved"), D(0)))
    same(f"{tag}: approved variations", sm["approved_vo"], appr)
    same(f"{tag}: revised contract", sm["revised"], m2(D(str(sm["contract"])) + appr))
    same(f"{tag}: margin = contract + approved VOs - committed", sm["margin"], m2(D(str(sm["contract"])) + appr - committed))
    boq = db.query(E.EstBoq).filter_by(project_id=p.id).all()
    bv = m2(sum(((m2(x.amount) if x.amount is not None else m2(D(str(x.qty or 0)) * D(str(x.rate or 0)))) for x in boq if not (x.scope or "")), D(0)))
    same(f"{tag}: BOQ value (in scope)", sm["boq_value"], bv)
    ks = db.query(E.EstContract).filter_by(project_id=p.id).all()
    if ks:
        same(f"{tag}: project contract = sum of its contracts", p.contract_value, m2(sum((D(str(k.value or 0)) for k in ks), D(0))))
    cpcs = db.query(E.EstClientPC).filter_by(project_id=p.id).all()
    groups = [([x for x in cpcs if x.contract_id == k.id], k) for k in ks] if ks else [(cpcs, p)]
    cc = {}
    for rows, k in groups:
        cc.update(chain(rows, k.vat_pct, client_work))
        srt = sorted(rows, key=lambda x: (x.seq or 0, x.id))
        lu = max([i for i, x in enumerate(srt) if x.status in E.LOCKED_STATUS] or [-1])
        for i, x in enumerate(srt):
            if i > lu and x.ret_auto:
                same(f"{tag} client PC{x.seq:02d} retention {k.ret_pct}% auto", x.ret_cum, m2(client_work(x) * D(str(k.ret_pct or 0)) / 100))
    if any(x.id not in cc for x in cpcs):
        ck(f"{tag}: every client certificate belongs to a contract", False, [x.id for x in cpcs if x.id not in cc])
    for x in b["clientpcs"]:
        counts["clientpcs"] += 1
        if x["id"] not in cc:
            continue
        w = cc[x["id"]]
        for kk in ("net_cum", "this", "vat", "payable"):
            same(f"{tag} client PC{x['seq']:02d} (contract {x.get('contract_id')}) {kk}", x.get(kk), w[{"net_cum": "net"}.get(kk, kk)])
    if b["clientpcs"]:
        counts["exports"] += 1
        rows = [r for r in xlsx_rows(f"/export/est/clientpcs?pid={p.id}") if isinstance(r.get("PC No."), (int, float))]
        same(f"{tag} client certificates Excel: total of This certificate", sum((D(str(r.get("This certificate") or 0)) for r in rows), D(0)),
             sum((cc[x.id]["this"] for x in cpcs if x.id in cc), D(0)))
        same(f"{tag} client certificates Excel: total with VAT", sum((D(str(r.get("Total") or 0)) for r in rows), D(0)),
             sum((cc[x.id]["payable"] for x in cpcs if x.id in cc), D(0)))
    if boq:
        counts["exports"] += 1
        rows = xlsx_rows(f"/export/est/boq?pid={p.id}")
        same(f"{tag} BOQ Excel: Amount column = BOQ value", sum((D(str(r.get("Amount") or 0)) for r in rows if isinstance(r.get("Amount"), (int, float))), D(0)), bv)
    for k in ("contract", "revised", "client_certified", "committed", "sub_certified", "sub_paid", "sub_to_pay", "approved_vo", "margin"):
        same(f"{tag}: dashboard {k} = project page", drow[p.id][k], sm[k])
    for k in dtot:
        dtot[k] += D(str(drow[p.id][k] or 0))
for k, v in dtot.items():
    same(f"dashboard total {k} = sum of the projects", dash["totals"][k], m2(v))
for r in pay_rows:
    if r.get("pc_id") in all_payable:
        same(f"Accounts payment list {r.get('sub')} PC{r.get('seq')}: payable", r.get("payable"), all_payable[r["pc_id"]])

# ---- the audit project's own figures, worked out by hand (not from the rows) - the rules above really ran
ab = c.get(f"/store/est/projects/{PID}", headers=H).json()
sa = next(x for x in ab["subs"] if x["id"] == s1)
pc3 = sa["pcs"][-1]
# Electrical 61.7% of 172,800 = 106,617.60; Plumbing 21.5% of 153,000 = 32,895.00; VRF 18.25% of 243,000 = 44,347.50
same("ARITH-T sub 1 last PC: work to date by hand", pc3["work"], D("183860.10"))
same("ARITH-T sub 1 last PC: variations by hand (2 x 6,172.80)", pc3["var"], D("12345.60"))
same("ARITH-T sub 1 last PC: retention 10% of work + variations", pc3["ret_cum"], D("19620.57"))
same("ARITH-T sub 1 last PC: recovery 10% of work (capped at the advance 113,760)", pc3["adv_rec_cum"], D("18386.01"))
same("ARITH-T sub 1 last PC: material deduction 2 x 117.35 still counted", pc3["other_cum"], D("234.70"))
same("ARITH-T sub 1 last PC: net to date", pc3["net_cum"], D("183860.10") + D("12345.60") + D("113760") - D("18386.01") - D("19620.57") - D("234.70"))
s2b = next(x for x in ab["subs"] if x["id"] == s2)
same("ARITH-T re-measured sub: revised = measured (141.5 x 235 + 48.07 x 120.5)", s2b["revised"], m2(D("141.5") * 235 + D("48.07") * D("120.5")))
k1b = next(x for x in ab["contracts"] if x["id"] == K1)
ck("ARITH-T contract 1: advance certificate = 10% of 1,000,000", abs(next(x for x in ab["clientpcs"] if x["pc_type"] == "Advance")["adv_amount"] - 100000) < 0.005)
same("ARITH-T contract 1: certified = last certified PC's net (Submitted one not counted)", k1b["certified"],
     next(x for x in ab["clientpcs"] if x["contract_id"] == K1 and x["seq"] == 3)["net_cum"])

# ======================================================================== bill workbooks: Excel's "precision as displayed"
def precision_book(full_precision_off=True):
    """A small bill workbook as Excel saves it with "precision as displayed": B2 = A2/3 shows 0.00 and Excel kept 33.33;
    B3 = B2*3 is then 99.99 (not 100.00); B5 = VAT 5% of B4. Cached figures written as Excel writes them."""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "IPA"
    ws["A1"] = "Description"; ws["B1"] = "Amount"
    ws["A2"] = 100; ws["B2"] = "=A2/3"; ws["B3"] = "=B2*3"; ws["A4"] = 1234.567; ws["B4"] = "=A4*1"; ws["B5"] = "=B4*5%"; ws["B6"] = "=B3+B4+B5"
    for k in ("B2", "B3", "B4", "B5", "B6"):
        ws[k].number_format = "#,##0.00"
    buf = io.BytesIO(); wb.save(buf)
    shown = {"B2": "33.33", "B3": "99.99", "B4": "1234.57", "B5": "61.73", "B6": "1396.29"}
    full = {"B2": repr(100 / 3), "B3": "100", "B4": "1234.567", "B5": repr(1234.567 * 0.05), "B6": repr(100 + 1234.567 + 1234.567 * 0.05)}
    vals = shown if full_precision_off else full
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue())); out = io.BytesIO(); zo = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        raw = zin.read(it.filename).decode("utf-8") if it.filename.endswith(".xml") else zin.read(it.filename)
        if it.filename == "xl/worksheets/sheet1.xml":
            for ref, v in vals.items():
                raw = re.sub(r'(<c r="%s"[^>]*>\s*<f>[^<]*</f>)\s*(?:<v\s*/>|<v>[^<]*</v>)?' % ref, lambda m: m.group(1) + f"<v>{v}</v>", raw)
        if it.filename == "xl/workbook.xml" and full_precision_off:
            raw = re.sub(r"<calcPr([^>]*)/>", r'<calcPr\1 fullPrecision="0"/>', raw) if "<calcPr" in raw else raw.replace("</workbook>", '<calcPr fullPrecision="0"/></workbook>')
        zo.writestr(it, raw)
    zo.close()
    return out.getvalue()


pad = precision_book(True)
m = BE.read_format(pad, "pad.xlsx")
ck("precision as displayed: the workbook's setting is read", m.get("precision_shown") is True)
chk = BE.check_against_excel(m)
ck("precision as displayed: every formula gives Excel's saved figure", chk["mismatch"] == 0 and chk["same"] == chk["formulas"] == 5, chk)
same("precision as displayed: B3 = 33.33 x 3 = 99.99", BE.Book(m).value("IPA", 3, 2), D("99.99"))
# the same workbook kept without the setting (an .xls converted before the setting was carried over) - recovered
lost = json.loads(json.dumps(m)); lost.pop("precision_shown")
ck("kept without the setting: figures are not Excel's (the old way)", BE.check_against_excel(lost)["mismatch"] > 0)
again = BE.recover_precision_shown(lost) if hasattr(BE, "recover_precision_shown") else None
ck("kept without the setting: recovered - every formula is Excel's again", again and again["mismatch"] == 0 and lost.get("precision_shown") is True, again)
full = BE.read_format(precision_book(False), "full.xlsx")
ck("a full-precision workbook is never switched to precision as displayed", not full.get("precision_shown") and BE.check_against_excel(full)["mismatch"] == 0)
x = BE.write_xlsx(pad, lost, {}, BE.calculate(lost, {}))
ck("the bill's Excel keeps 'precision as displayed' so Excel works it out the same on opening",
   'fullPrecision="0"' in zipfile.ZipFile(io.BytesIO(x)).read("xl/workbook.xml").decode())
x2 = BE.write_xlsx(precision_book(False), full, {}, BE.calculate(full, {}))
ck("a full-precision bill's Excel is not set to precision as displayed", 'fullPrecision="0"' not in zipfile.ZipFile(io.BytesIO(x2)).read("xl/workbook.xml").decode())
# a format already kept that way: the bills made from it are worked out as Excel does, and the list says so
lost.pop("precision_shown", None)
fm = EB.EstBillFormat(project_id=PID, contract_id=K1, name="pad", filename="pad.xlsx", data=pad, model=json.dumps(lost),
                      checked=json.dumps({"formulas": 5, "same": 1, "different": []}), created_by=1)
db.add(fm); db.commit()
EB._MODEL_CACHE.pop(fm.id, None); EB._LIST_CACHE.pop(fm.id, None); getattr(EB, "_RECHECK", {}).pop(fm.id, None)
lst = c.get(f"/store/est/ipa/list?project_id={PID}", headers=H).json()
row = next(f for f in lst["formats"] if f["id"] == fm.id)
ck("kept format: the list shows every formula as Excel's", row["same"] == row["formulas"] == 5, row)
mm = EB._model(db.get(EB.EstBillFormat, fm.id))
ck("kept format: worked out with precision as displayed", mm.get("precision_shown") is True)
same("kept format: B6 as Excel (99.99 + 1,234.57 + 61.73)", BE.Book(mm).value("IPA", 6, 2), D("1396.29"))
db.expire_all()
ck("nothing stored was changed", json.loads(db.get(EB.EstBillFormat, fm.id).model).get("precision_shown") is None)
rd = c.delete(f"/store/est/ipa/format/{fm.id}", headers=H)
ck("format removed (and its worked-out copies forgotten)", rd.status_code == 200 and fm.id not in EB._MODEL_CACHE and fm.id not in getattr(EB, "_RECHECK", {}), rd.text[:200])
# the real 914 MEP IPA02 (if this computer has it): 593 / 593 formulas as Excel
for f in glob.glob("/mnt/user-data/attachments/*/*914_MEP_IPA02*.xls") + glob.glob("/tmp/claude-0/bills/914 MEP IPA02*.xls"):
    mr = BE.read_format(BE.xls_to_xlsx(open(f, "rb").read()), "914 MEP IPA02 - R2.xlsx")
    r = BE.check_against_excel(mr)
    ck(f"real {os.path.basename(f)}: every formula as Excel", r["mismatch"] == 0 and r["formulas"] == r["same"], (r["formulas"], r["same"]))
    break

# ======================================================================== bills: saved = exported = next bill's previous
mk = os.path.join(os.path.dirname(os.path.abspath(__file__)), "make_ipa_like_914.py")
import subprocess, tempfile
tmp = os.path.join(tempfile.mkdtemp(), "914-like.xlsx")
subprocess.run([sys.executable, mk, tmp], check=True, capture_output=True)
r = c.post("/store/est/ipa/format", headers=H, data={"project_id": str(PID), "contract_id": str(K1), "name": "audit format", "force": "1"},
           files={"file": ("914-like.xlsx", open(tmp, "rb").read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
ck("bill format uploaded", r.status_code == 200, r.text[:300])
if r.status_code == 200:
    fid = r.json()["id"]
    b1 = post("/store/est/ipa/new", {"project_id": PID, "format_id": fid})["id"]
    bill = c.get(f"/store/est/ipa/bill/{b1}", headers=H).json()
    sheets = bill.get("progress") or bill.get("tables") or []
    ln = [l for t in sheets for l in t["lines"] if l.get("amount")]
    ck("bill has % lines with amounts", len(ln) >= 2, len(ln))
    typed = {}
    for i, l in enumerate(ln):
        typed.setdefault(next(t["sheet"] for t in sheets if l in t["lines"]), {})[l["tot_key"]] = round(min(1.0, (l["prev"] or 0) + 0.1 + i * 0.0137), 4)
    rr = c.put(f"/store/est/ipa/bill/{b1}", headers=H, json={"set": typed})
    ck("bill saved", rr.status_code == 200, rr.text[:200])
    bill = c.get(f"/store/est/ipa/bill/{b1}", headers=H).json()
    eb = db.get(EB.EstBill, b1); db.refresh(eb)
    fmt = db.get(EB.EstBillFormat, fid)
    model = EB._bill_model(fmt, EB._inputs(eb))
    vals = BE.calculate(model, BE.with_words(model, EB._inputs(eb)))
    mp = model["map"]
    td = vals[mp["total_due"]["sheet"]].get(f"{mp['total_due']['r']},{mp['total_due']['c']}")
    same("bill: total due in the database = the workbook worked out", eb.total_due, m2(td))
    same("bill: total due on screen = database", bill["total_due"], m2(eb.total_due))
    wd = sum(((l["amount"] or 0) * (l["total"] or 0) for t in EB._progress(model, vals, {}) for l in t["lines"]), 0)
    same("bill: work done in the database = sum of amount x Total %", eb.work_done, m2(wd))
    xr = c.get(f"/export/est/ipa/{b1}?token={dl()}&format=excel")
    ck("bill Excel exported", xr.status_code == 200, xr.status_code)
    if xr.status_code == 200:
        wbx = openpyxl.load_workbook(io.BytesIO(xr.content), data_only=True)
        same("bill: total due in the Excel = database", wbx[mp["total_due"]["sheet"]].cell(mp["total_due"]["r"], mp["total_due"]["c"]).value, m2(eb.total_due))
        for t in mp["progress"]:
            for l in t["lines"][:50]:
                v = wbx[t["sheet"]].cell(l["r"], t["tot_c"]).value
                if isinstance(v, (int, float)):
                    same(f"bill Excel {t['sheet']} row {l['r']}: Total % = database", v, D(str(vals[t["sheet"]].get(f"{l['r']},{t['tot_c']}") or 0)).quantize(D("0.000001")))
    pv = c.get(f"/export/est/ipa/{b1}?token={dl()}&format=preview")
    import est_bills_render as BR
    tcell = next(x for x in model["sheets"] if x["name"] == mp["total_due"]["sheet"])["cells"].get(f"{mp['total_due']['r']},{mp['total_due']['c']}") or {}
    shown = BR.fmt(eb.total_due, model["styles"][tcell.get("s", 0)].get("nf") or "General")      # as the cell's own format shows it
    ck("bill preview shows the total due", pv.status_code == 200 and shown in pv.text, (shown, re.findall(r"[-(]?[\d,]+\.\d\d\)?", pv.text)[-12:]))
    b2 = post("/store/est/ipa/new", {"project_id": PID, "format_id": fid})["id"]
    e2 = db.get(EB.EstBill, b2)
    inp2 = EB._inputs(e2)
    for t in mp["progress"]:
        for l in t["lines"]:
            tot1 = vals[t["sheet"]].get(f"{l['r']},{t['tot_c']}")
            if isinstance(tot1, (int, float)):
                same(f"next bill {t['sheet']} row {l['r']}: Previous % = last bill's Total %", inp2.get(t["sheet"], {}).get(f"{l['r']},{t['prev_c']}"), D(str(tot1)))

# ======================================================================== the live figures (if this database has them)
# 31 Oct 2026, read from the live app: the payment recorded on each paid subcontractor certificate = its payable
for pc in db.query(E.EstSubPC).filter(E.EstSubPC.status == "Paid").all():
    s = db.get(E.EstSub, pc.sub_id)
    ch = chain(db.query(E.EstSubPC).filter_by(sub_id=s.id).all(), s.vat_pct, lambda x: D(str(x.work_cum or 0)))
    if pc.paid_amount and abs(D(str(pc.paid_amount)) - ch[pc.id]["payable"]) > D("0.005"):
        print(f"NOTE {s.name} PC{pc.seq:02d}: paid {pc.paid_amount:,.2f} - payable {ch[pc.id]['payable']:,.2f}")

c.delete(f"/store/est/projects/{PID}", headers=H)
print(f"{n[0]} figures checked - {counts}")
for f in fails[:80]:
    print("FAIL", f)
print(f"\n{len(fails)} FAILED" if fails else "\nALL PASS")
sys.exit(1 if fails else 0)
