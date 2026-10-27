"""Checks every file tests/audit_exports.js saved: PDFs (pages, A4, text on them), Excel (A4 paper, fitted to one
page wide, heading rows repeated on every printed page, heading frozen, columns not too narrow), previews (an A4
page set for printing). Prints one line per file and a summary of what is off.
    python3 tests/check_exports.py [/tmp/claude-0/exports]"""
import json, os, re, subprocess, sys, zipfile
from collections import defaultdict

D = sys.argv[1] if len(sys.argv) > 1 else "/tmp/claude-0/exports"
J = json.load(open(os.path.join(D, "exports.json")))
issues = defaultdict(list)


def pdf_check(f):
    out = subprocess.run(["pdfinfo", f], capture_output=True, text=True).stdout
    pages = int((re.search(r"Pages:\s+(\d+)", out) or [0, 0])[1])
    m = re.search(r"Page size:\s+([\d.]+) x ([\d.]+)", out)
    w, h = (float(m[1]), float(m[2])) if m else (0, 0)
    a4 = {round(w), round(h)} <= {595, 596, 841, 842} or (abs(w - 595) < 3 and abs(h - 842) < 3) or (abs(w - 842) < 3 and abs(h - 595) < 3)
    txt = subprocess.run(["pdftotext", "-l", "2", f, "-"], capture_output=True, text=True).stdout.strip()
    bad = []
    if not pages: bad.append("no pages")
    if not a4: bad.append(f"page {w:.0f}x{h:.0f}pt, not A4")
    if not txt: bad.append("no text on it")
    return f"PDF {pages}p {'A4' if a4 else f'{w:.0f}x{h:.0f}'}", bad


def xlsx_check(f):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(f)
    except Exception as e:
        return "XLSX unreadable", [f"cannot open: {e}"]
    bad, notes = [], []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        ps = ws.page_setup
        fit = ws.sheet_properties.pageSetUpPr is not None and ws.sheet_properties.pageSetUpPr.fitToPage
        rows, cols = ws.max_row, ws.max_column
        if rows <= 1 and cols <= 1:
            continue
        nm = ws.title[:20]
        if str(ps.paperSize) not in ("9", "None") and ps.paperSize not in (9, None):
            bad.append(f"[{nm}] paper size {ps.paperSize}, not A4")
        if ps.paperSize is None:
            bad.append(f"[{nm}] no paper size set (prints on the printer's default)")
        if not fit or (ps.fitToWidth not in (1, "1", None)):
            bad.append(f"[{nm}] not fitted to one page wide")
        form = ws.title in ("Final Settlement", "Combined Cards")      # a document laid out as a form, not a list
        if rows > 30 and not ws.print_title_rows and not form:
            bad.append(f"[{nm}] {rows} rows but no heading rows repeated on each printed page")
        if rows > 25 and not ws.freeze_panes and not form:
            bad.append(f"[{nm}] {rows} rows but the heading is not frozen")
        narrow = [c for c in range(1, min(cols, 30) + 1)
                  if (ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width or 8.43) < 4 and any(ws.cell(r, c).value not in (None, "") for r in range(1, min(rows, 50) + 1))]
        if narrow:
            bad.append(f"[{nm}] columns too narrow: {narrow[:5]}")
        notes.append(f"{nm}:{rows}x{cols}{' fit' if fit else ''}{' A4' if ps.paperSize in (9, '9') else ''}{' ' + ps.orientation[:4] if ps.orientation else ''}")
    return "XLSX " + " ".join(notes)[:90], bad


def html_check(f):
    s = open(f, encoding="utf-8", errors="ignore").read()
    bad = []
    if not re.search(r"@page[^{]*\{[^}]*size\s*:\s*A4", s, re.I) and "size: A4" not in s and "size:A4" not in s:
        bad.append("no A4 page set for printing")
    if not re.search(r"<table|<div", s, re.I):
        bad.append("empty page")
    return "HTML preview", bad


seen = 0
for r in J["rows"]:
    where = f"{r['screen']} / {r['button']}"
    files = [x["file"] for x in r.get("fetched") or []] + ([r["file"]] if r.get("file") else [])
    if not files:
        if r.get("dup"):
            continue
        if r.get("sharebox") or any(u.startswith("SHARE:") for u in r.get("urls") or []):
            print(f"{where:60s} share: {(r.get('sharebox') or [u for u in r['urls'] if u.startswith('SHARE:')][0])[:60]}")
            continue
        issues["nothing happened"].append(where)
        print(f"{where:60s} NOTHING opened or downloaded {r.get('urls')}")
        continue
    for fx in r.get("fetched") or []:
        if fx["status"] != 200:
            issues["server error"].append(f"{where}: {fx['status']} {fx['url']}")
    for f in files:
        if not os.path.exists(f) or os.path.getsize(f) == 0:
            issues["empty file"].append(where); continue
        seen += 1
        head = open(f, "rb").read(5)
        kind = "pdf" if head.startswith(b"%PDF") else "xlsx" if head.startswith(b"PK") else "html"
        desc, bad = {"pdf": pdf_check, "xlsx": xlsx_check, "html": html_check}[kind](f)
        for b in bad:
            issues[f"{kind}: " + re.sub(r"\[.*?\] |\d+ rows but |: \[.*", "", b)].append(f"{where}: {b}")
        print(f"{where:60s} {desc} {'  !! ' + '; '.join(bad) if bad else ''}")

print(f"\n{seen} files checked. Page errors: {len(J['errs'])}")
for k, v in sorted(issues.items(), key=lambda kv: -len(kv[1])):
    print(f"\n== {k} ({len(v)})")
    for x in v[:40]:
        print("   ", x)
