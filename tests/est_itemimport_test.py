"""Subcontract item schedule out of Excel (agreement BOQ or a payment certificate).

    cd app && DATABASE_URL=sqlite:////tmp/itimp.db INFINIA_NO_SCHED=1 python3 ../tests/est_itemimport_test.py [certificate.xlsx ...]

A workbook in the 908 MEP certificate layout is made here; real certificates given on the command line
must read with every sheet equal to its own total row and every line's done x amount = its certified figure.
"""
import io, os, sys
sys.path.insert(0, os.getcwd())
import openpyxl
from fastapi.testclient import TestClient
import main, estimation as E, est_pcsheet as S, est_itemimport as I
from database import SessionLocal, engine

main._add_missing_columns()
E.create_tables(engine)
c = TestClient(main.app); c.__enter__()
fails = []
def ck(name, ok, extra=""):
    print(("PASS " if ok else "FAIL ") + name + ("" if ok else f"  -> {extra}"))
    if not ok: fails.append(name)

HEAD = ["Bill No", "Description", "Approx. Qty", "Unit", "Rate", "Amount", "Previous % Work Done", "Current % Work Done", "Total % Work Done", "Certified Amount", "Remarks"]


def mep_book():
    wb = openpyxl.Workbook(); wb.active.title = "Certificate"
    hid = wb.create_sheet("Old"); hid.append(HEAD); hid.append([1, "Hidden line", 1, "Item", "LS", 999]); hid.sheet_state = "hidden"
    ws = wb.create_sheet("Bill 9 -Electrical Works")
    for r in [[], [None, "Appendix - B1", None, None, None, None, None, None, None, None, "Payment Certificate No.: 08/B1"],
              [None, "Value of Work done & Variations"], [None, "Electrical Works"], [], HEAD, [],
              [9, "ELECTRICAL WORKS"], [None, "Contractor to include for supply of all materials as per drawings and specifications."],
              [None, "1) The contractor to include cost of supply and installation of main distribution boards"],
              [None, "a. M.D.B (IP.55) with meter panel", "LS", "NO's", "LS", 10000, 0.1, 0.35, 0.45, 4500],
              [None, "b. SMDB-RF", None, "NO's"],
              [None, "c. DB Ground Floor", None, "NO's", None, 6500, 0.1, 0.35, 0.45, 2925],
              [None, "2) Providing and installing main power cables", 1, "Item", "LS", 15000, 0.07, 0.73, 0.8, 12000],
              [None, "10) TELEPHONE SYSTEM:"],
              [None, "Supply, install, test and commission telephone system", "LS", "Item", "LS", 5000, 0, 0.5, 0.5, 2500],
              [None, "11) Contractor to include for fixing of light fixtures", 1, "Item", "LS", 6500, 0, 0, 0, 0],
              [None, "14) Contractor to include for strip lighting", 1, "Item", "LS", "Included"],
              [None, "Total To-date carried over to Certificate", None, None, None, 43000, None, None, 0.5099, 21925],
              [], ["Subcontractor: M/S. AL RAIDA", None, None, None, None, "Contractor: M/s. INFINIA"], ["Date:"]]:
        ws.append(r)
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()


r = I.read_schedule(mep_book())
L = r["items"]
D = {x["description"]: x for x in L}
ck("hidden sheets are not read", "Hidden line" not in D)
ck("bill code and name from 'Appendix - B1' and the line under it", L[0]["bill"] == "B1" and L[0]["bill_title"] == "Electrical Works", (L[0]["bill"], L[0]["bill_title"]))
k = next(x for x in D if x.endswith("a. M.D.B (IP.55) with meter panel"))
ck("a lettered line carries the sentence that introduces it", k.startswith("1) The contractor to include cost") and D[k]["amount"] == 10000 and D[k]["measure"] == "pct" and D[k]["unit"] == "NO's")
ck("an unpriced lettered line is left out (counted in the sheet's 'unpriced')", not any("SMDB" in x for x in D) and r["sheets"][0]["unpriced"] == 1)
ck("% done read as a percentage (0.45 -> 45%)", D[k]["done_pct"] == 45)
ck("a numbered heading is a heading", D["Supply, install, test and commission telephone system"]["section"] == "10) TELEPHONE SYSTEM:")
ck("the next numbered item goes back to the bill's own heading", D["11) Contractor to include for fixing of light fixtures"]["section"] == "9 ELECTRICAL WORKS")
ck("an amount in words is kept, with no money", D["14) Contractor to include for strip lighting"]["words"] == "Included" and D["14) Contractor to include for strip lighting"]["amount"] == 0)
ck("sheet total = lines (contract 43,000.00, certified 21,925.00), every line checked", r["contract"] == 43000 and r["certified"] == 21925 and r["all_checked"], (r["contract"], r["certified"], r["notes"]))
ck("the signature block is not a line", not any("Subcontractor" in x or "INFINIA" in x for x in D))

# a wrong figure in the file is caught
wb = openpyxl.load_workbook(io.BytesIO(mep_book())); ws = wb["Bill 9 -Electrical Works"]
for row in ws.iter_rows():
    if row[1].value == "2) Providing and installing main power cables":
        row[9].value = 11000
bio = io.BytesIO(); wb.save(bio)
bad = I.read_schedule(bio.getvalue())
ck("a certified figure that doesn't match its % is flagged, not passed", not bad["all_checked"] and any("not the certified figure" in n for n in bad["notes"]), bad["notes"])

for f in sys.argv[1:]:
    rr = I.read_schedule(open(f, "rb").read())
    hard = [n for n in rr["notes"] if "no rate" not in n]
    ck(f"{os.path.basename(f)}: {len(rr['items'])} lines, contract {rr['contract']:,.2f}, certified {rr['certified']:,.2f} - every sheet = its total, every line = its certified figure",
       rr["all_checked"] and not hard, rr["notes"])

# ---- through the app ----
tok = c.post("/auth/login", data={"username": "admin", "password": "changeme123"}).json()["access_token"]
H = {"Authorization": "Bearer " + tok}
db = SessionLocal()
p = db.query(E.EstProject).first()
sid = c.post("/store/est/subs", json={"project_id": p.id, "name": "IMPORT TEST LLC", "scope": "MEP", "value": 43000}, headers=H).json()["id"]
up = lambda data, **kw: c.post(f"/store/est/subs/{sid}/items/import", files={"file": ("908 PC08.xlsx", data)}, data={k: str(v) for k, v in kw.items()}, headers=H)
x = up(mep_book()).json()
ck("preview: every line shown, nothing saved", x["lines"] == len(L) and len(x["items"]) == len(L) and db.query(S.EstSubItem).filter_by(sub_id=sid).count() == 0)
x = up(mep_book(), save=1).json()
items = c.get(f"/store/est/subs/{sid}/items", headers=H).json()
ck("saved: schedule total = 43,000.00 = subcontract value", x["saved"] == len(L) and abs(items["total"] - 43000) < .005, items.get("total"))
ck("words kept on the line (Type / mark)", any(i["ref"] == "Included" for i in items["items"]))
ck("the lines can be edited afterwards", c.put(f"/store/est/subs/{sid}/items", json={"items": [dict(i, rate=0) for i in items["items"]]}, headers=H).status_code == 200)
pc = c.post("/store/est/subpcs", json={"sub_id": sid, "inv_ref": "INV 1"}, headers=H).json()["id"]
it0 = items["items"][0]["id"]
c.put(f"/store/est/subpcs/{pc}/sheet", json={"lines": [{"item_id": it0, "pct": 45}]}, headers=H)
x = up(mep_book(), save=1, mode="replace")
ck("replace refused once lines are on a certificate", x.status_code == 400 and "certificates" in x.text)
ck("add still works", up(mep_book(), save=1, mode="add").status_code == 200)
ck("not a workbook is refused", up(b"nope").status_code == 400)
# 911: quantity x rate x stage %  (1,023.14 m2 x 16.00 x 70% = 11,459.17)
sid2 = c.post("/store/est/subs", json={"project_id": p.id, "name": "NAIF TEST", "scope": "Paint", "value": 0}, headers=H).json()["id"]
c.post(f"/store/est/subs/{sid2}/items/paste", json={"text": "1\tInternal wall\t1345\tM2\t16\t21520", "measure": "qty"}, headers=H)
it = c.get(f"/store/est/subs/{sid2}/items", headers=H).json()["items"][0]
pc2 = c.post("/store/est/subpcs", json={"sub_id": sid2, "inv_ref": "PC 2"}, headers=H).json()["id"]
c.put(f"/store/est/subpcs/{pc2}/sheet", json={"lines": [{"item_id": it["id"], "qty": 1023.14, "pct": 70}]}, headers=H)
w = db.get(E.EstSubPC, pc2); db.refresh(w)
ck("quantity x rate x stage %: 11,459.17", abs((w.work_cum or 0) - 11459.17) < .005, w.work_cum)
c.put(f"/store/est/subpcs/{pc2}/sheet", json={"lines": [{"item_id": it["id"], "qty": 849.86, "pct": 0}]}, headers=H)
db.refresh(w)
ck("no stage % = the whole quantity (849.86 x 16 = 13,597.76)", abs((w.work_cum or 0) - 13597.76) < .005, w.work_cum)
for s_ in (sid, sid2):
    c.delete(f"/store/est/subs/{s_}", headers=H)
print(f"\n{len(fails)} FAILED: {fails}" if fails else "\nALL PASS")
sys.exit(1 if fails else 0)
