"""Every list whose screen can be narrowed (a Show box, a picker, a search)
prints - Preview, PDF and Excel - only what the screen shows."""
import os, sys, io
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))
os.environ.setdefault("INFINIA_NO_PUSH", "1")
os.environ.setdefault("INFINIA_CRED_KEY_FILE", "/tmp/export_filters_test.key")
from datetime import date, timedelta
from fastapi.testclient import TestClient
import openpyxl
import main, models, auth, credentials as CR, seccheque as SC, accounts as AC
from database import SessionLocal
FAIL = []
def ck(l, ok, x=""):
    print(("PASS " if ok else "FAIL ") + l + ("" if ok else f"  [{str(x)[:300]}]"))
    if not ok: FAIL.append(l)
T = main._dubai_today()
db = SessionLocal()
db.add(models.User(username="a", hashed_password=auth.hash_password("p"), full_name="A", role="admin"))
c1 = models.Company(name="INFINIA CONTRACTING L.L.C.", short_name="Infinia"); c2 = models.Company(name="PRIME INFINIA", short_name="Prime")
db.add_all([c1, c2]); db.commit()
s1 = models.Employee(emp_no="IC501", name="ALPHA STAFF", staff=True, pay_group="staff", company_id=c1.id, basic_salary=2000, allowance=3000, total_salary=5000, active=True, joined_on=date(2024, 1, 1))
s2 = models.Employee(emp_no="IC502", name="BRAVO STAFF", staff=True, pay_group="staff", company_id=c2.id, basic_salary=2000, allowance=3000, total_salary=5000, active=True, joined_on=date(2024, 1, 1))
l1 = models.Employee(emp_no="F-901", name="CHARLIE LAB", staff=False, basic_salary=1000, total_salary=1500, active=True, joined_on=date(2024, 1, 1))
db.add_all([s1, s2, l1]); db.commit()
db.add_all([models.StaffLeave(employee_id=s1.id, on_date=T - timedelta(days=3), portion=1, kind="absent", paid=False),
            models.StaffLeave(employee_id=s2.id, on_date=T - timedelta(days=2), portion=1, kind="absent", paid=False),
            models.PayItem(employee_id=s1.id, month_year="October 2026", direction="add", category="taxi", amount=111),
            models.PayItem(employee_id=s2.id, month_year="October 2026", direction="add", category="taxi", amount=222),
            models.EmployeeDocument(employee_id=s1.id, kind="eid", expires_on=T + timedelta(days=5)),
            models.EmployeeDocument(employee_id=s2.id, kind="eid", expires_on=T + timedelta(days=20)),
            models.EmployeeDocument(employee_id=l1.id, kind="eid", expires_on=T + timedelta(days=6)),
            models.CompanyExpiry(category="Vehicle", item="Hilux ZZ1", kind="Registration", expires_on=T + timedelta(days=4)),
            models.CompanyExpiry(category="Licence", item="Trade licence", kind="Licence", expires_on=T + timedelta(days=4)),
            models.CompanyExpiry(category="Vehicle", item="Canter YY2", kind="Insurance", expires_on=T + timedelta(days=25)),
            SC.SecurityCheque(kind="cheque", party="EMAAR", amount=1000, status="held"),
            SC.SecurityCheque(kind="cheque", party="NAKHEEL", amount=2000, status="held"),
            AC.Invoice(kind="tax", number="INV-001", inv_date=T, client="SOBHA", subtotal=100, vat=5, total=105),
            AC.Invoice(kind="tax", number="INV-002", inv_date=T, client="DAMAC", subtotal=100, vat=5, total=105),
            CR.Credential(name="DEWA", category="Utilities", username="dewa_user", secret=CR.lock("x")),
            CR.Credential(name="RTA", category="Utilities", username="rta_user", secret=CR.lock("y")),
            models.Pdc(payee="ACME STEEL", cheque_date=T + timedelta(days=3), amount=500, status="pending"),
            models.Pdc(payee="BETA CEMENT", cheque_date=T + timedelta(days=60), amount=700, status="pending"),
            models.Pdc(payee="GAMMA OLD", cheque_date=T - timedelta(days=60), amount=900, status="cleared"),
            models.MaterialRequest(ref="MR-0001", site="912", requested_by="Amal", status="pending", requested_on=T, urgency="normal"),
            models.MaterialRequest(ref="MR-0002", site="905", requested_by="Amal", status="delivered", requested_on=T, urgency="normal"),
            models.MaterialRequest(ref="MR-0003", site="905", requested_by="Amal", status="approved", requested_on=T, urgency="normal")])
C1 = c1.id; db.commit(); db.close()
c = TestClient(main.app)
h = {"Authorization": "Bearer " + c.post("/auth/login", data={"username": "a", "password": "p"}).json()["access_token"]}
r = c.post("/auth/download-token", headers=h); tok = r.json().get("token") or r.json().get("download_token")
def view(path): r = c.get(path + f"&token={tok}"); return r.status_code, r.text
def xl(path):
    r = c.get(path.replace("/view?", "?") + f"&format=excel&token={tok}")
    if r.status_code != 200: return r.status_code, r.text[:200]
    try: wb = openpyxl.load_workbook(io.BytesIO(r.content))
    except Exception: return 0, r.text[:300]
    return 200, " ".join(str(x.value) for row in wb.active.iter_rows() for x in row if x.value)
def both(label, path, yes, no):
    for kind, (st, t) in (("preview", view(path)), ("excel", xl(path))):
        ck(f"{label} - {kind}", st == 200 and all(y in t for y in yes) and not any(n in t for n in no), (st, t[:400]))
    p = c.get(path.replace("/view?", "?") + f"&format=pdf&token={tok}"); ck(f"{label} - pdf", p.status_code == 200 and p.content[:4] == b"%PDF", p.status_code)

both("requests: still open", "/export/store/report/view?kind=mr_history&status=open", ["MR-0001", "MR-0003"], ["MR-0002"])
both("requests: search 912", "/export/store/report/view?kind=mr_history&status=all&q=912", ["MR-0001"], ["MR-0002", "MR-0003"])
both("staff: one company", f"/export/payroll/staff/view?company_id={C1}", ["ALPHA"], ["BRAVO"])
both("staff: search", "/export/payroll/staff/view?q=bravo", ["BRAVO"], ["ALPHA"])
both("leave: one person", "/export/payroll/leave/view?month_year=&emp_no=IC501", ["ALPHA"], ["BRAVO"])
both("additions: one person", "/export/payroll/items/view?month_year=&emp_no=IC502", ["BRAVO"], ["ALPHA"])
both("documents: due in 7 days", "/export/payroll/documents/view?within=w", ["ALPHA", "CHARLIE"], ["BRAVO"])
both("documents: 8-30 days", "/export/payroll/documents/view?within=m", ["BRAVO"], ["ALPHA", "CHARLIE"])
both("documents: office only + search", "/export/payroll/documents/view?within=-1&who=staff&q=alpha", ["ALPHA"], ["BRAVO", "CHARLIE"])
both("documents: missing", "/export/payroll/documents/view?within=missing&who=all", ["Missing"], [])
both("expiries: due in 7, vehicle", "/export/payroll/expiries/view?within=w&cat=Vehicle", ["Hilux"], ["Trade licence", "Canter"])
both("expiries: search", "/export/payroll/expiries/view?within=-1&q=canter", ["Canter"], ["Hilux"])
both("security cheques: search", "/export/security-cheques/view?show=held&q=emaar", ["EMAAR"], ["NAKHEEL"])
both("invoices: search", "/export/accounts/invoices?kind=tax&q=damac".replace("?", "?format=view&", 0), ["INV-002"], ["INV-001"]) if False else None
st, t = c.get(f"/export/accounts/invoices?kind=tax&q=damac&format=view&token={tok}").status_code, c.get(f"/export/accounts/invoices?kind=tax&q=damac&format=view&token={tok}").text
ck("invoices: search - preview", st == 200 and "INV-002" in t and "INV-001" not in t, t[:300])
st, t = xl("/export/accounts/invoices?kind=tax&q=damac")
ck("invoices: search - excel", st == 200 and "INV-002" in t and "INV-001" not in t, t[:300])
st, t = view("/export/accounts/credentials/view?cat=&q=rta")
ck("credentials: search - preview", st == 200 and "rta_user" in t and "dewa_user" not in t, t[:300])
st, t = xl("/export/accounts/credentials?cat=&q=rta")
ck("credentials: search - excel", st == 200 and "rta_user" in t and "dewa_user" not in t, t[:300])
both("pdc list: due in 7 days", "/export/pdc/view?mode=list&which=week", ["ACME"], ["BETA", "GAMMA"])
both("pdc list: search", "/export/pdc/view?mode=list&which=all&q=gamma", ["GAMMA"], ["ACME", "BETA"])
st, t = view("/export/pdc/view?months=4")
ck("pdc tracker still works", st == 200, t[:200])
print("\nALL PASS" if not FAIL else f"\n{len(FAIL)} FAILED: {FAIL}")
