"""PPE register: the Excel log loaded once, item limits, too soon only with
a reason, 'approved' lines not counted, quick repeats, stock, rights."""
import os, sys, tempfile, json, asyncio, io
from datetime import date, timedelta
os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mktemp(suffix='.db')}"
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "app"))
import main as M, models, ppe, auth
from database import SessionLocal, engine
from fastapi import HTTPException
models.Base.metadata.create_all(bind=engine); M._add_missing_columns()
s = SessionLocal(); fails = 0
def ok(c, m):
    global fails
    print(("PASS " if c else "FAIL ") + m); fails += 0 if c else 1
def refused(fn, *a):
    try:
        fn(*a); return ""
    except HTTPException as e:
        return e.detail or "refused"
seed = json.load(open(os.path.join(os.path.dirname(__file__), "..", "app", "ppe_seed.json")))
for i in sorted({r["id"] for r in seed["issues"]} - {"PI005"}):
    s.add(models.Employee(emp_no=f"{i[0]}-{i[1:]}" if i[0] == "F" else i, name=f"Worker {i}", trade="Helper", active=True, staff=False))
s.commit()
E = {e.emp_no: e for e in s.query(models.Employee).all()}
jomon = models.User(username="jomon", full_name="Jomon", hashed_password="x", role="office", permissions="reports,ppe")
viewer = models.User(username="v", full_name="V", hashed_password="x", role="office", permissions="reports")
s.add_all([jomon, viewer]); s.commit()
n = ppe.seed(s)
ok(n == 293, f"Excel register loaded: 293 issues ({n})")
ok(ppe.seed(s) == 0, "loaded once only")
matched = s.query(models.PpeIssue).filter(models.PpeIssue.employee_id.isnot(None)).count()
ok(matched >= 290, f"issues matched to the staff list by ID ({matched}/293)")
d = ppe.get_ppe(s, jomon)
ok(d["may_edit"] and not ppe.get_ppe(s, viewer)["may_edit"], "PPE right can record; Labour reports alone can only look")
ok(refused(ppe._editor, viewer), "viewer cannot record")
W = {w["no"]: w for w in d["workers"]}
# a worker's cells
w = W["F-782"]
ok(w["cells"]["Coverall"]["last"] and w["cells"]["Safety Shoes"]["last"], "F782's coverall and shoes shown with dates")
# rules on a fresh worker
s.add(models.Employee(emp_no="F-900", name="Test Man", active=True, staff=False)); s.commit()
t9 = s.query(models.Employee).filter_by(emp_no="F-900").first()
today = M._dubai_today()
ppe.add({"employee_id": t9.id, "item": "Safety Shoes", "size": "42", "qty": 1, "date": (today - timedelta(days=7)).isoformat()}, s, jomon)
c = ppe.check(t9.id, "", "Safety Shoes", 1, today.isoformat(), 0, s, jomon)
ok(c["too_soon"] and c["days_since"] == 7 and c["last_size"] == "42" and c["next_due"], f"shoes again after 7 days: too soon, last size 42, next due {c['next_due']}")
from datetime import date as _d
due = _d.fromisoformat(c["next_due"])
ok(not ppe.check(t9.id, "", "Safety Shoes", 1, due.isoformat(), 0, s, jomon)["too_soon"] and ppe.check(t9.id, "", "Safety Shoes", 1, (due - timedelta(days=1)).isoformat(), 0, s, jomon)["too_soon"], f"due date is exact: allowed on {due}, not the day before")
r = refused(ppe.add, {"employee_id": t9.id, "item": "Safety Shoes", "qty": 1}, s, jomon)
ok("Too soon" in r and "7 days ago" in r and "next due" in r, "saving without a reason is refused with the last issue shown")
ppe.add({"employee_id": t9.id, "item": "Safety Shoes", "qty": 1, "early_reason": "torn on site - seen by Shaiju"}, s, jomon)
ok(s.query(models.PpeIssue).filter_by(employee_id=t9.id, item="Safety Shoes").count() == 2, "with a reason it is saved")
e2 = s.query(models.PpeIssue).filter_by(employee_id=t9.id).order_by(models.PpeIssue.id.desc()).first()
ok(e2.early_reason.startswith("torn"), "the reason is kept on the entry")
# coverall: 2 in 6 months allowed
ppe.add({"employee_id": t9.id, "item": "Coverall", "qty": 2, "date": today.isoformat()}, s, jomon)
ok("Too soon" in refused(ppe.add, {"employee_id": t9.id, "item": "Coverall", "qty": 1}, s, jomon), "coverall: 2 given (new joiner) - a 3rd is too soon")
# approved, not given: does not count
s.add(models.Employee(emp_no="F-901", name="Other Man", active=True, staff=False)); s.commit()
t1 = s.query(models.Employee).filter_by(emp_no="F-901").first()
ppe.add({"employee_id": t1.id, "item": "Helmet", "qty": 1, "condition": "Eligible"}, s, jomon)
c = ppe.check(t1.id, "", "Helmet", 1, today.isoformat(), 0, s, jomon)
ok(not c["too_soon"], "'approved - give later' does not count until given")
d = ppe.get_ppe(s, jomon); W = {w["no"]: w for w in d["workers"]}
ok(W["F-900"]["cells"]["Safety Shoes"]["state"] == "full" and W["F-900"]["early"] == 1, "grid: F900 shoes red (at limit), 1 early flag")
ok(W["F-901"]["pending"] == 1 and d["tiles"]["pending"] >= 1, "approved-not-given counted on the tile")
ok(any(q["no"] == "F-900" and q["item"] == "Safety Shoes" and q["days"] == 7 for q in d["quick"]), "quick repeat listed: same item within 30 days")
# helmet limit 1 a year
s.add(models.Employee(emp_no="F-902", name="Helmet Man", active=True, staff=False)); s.commit()
t2 = s.query(models.Employee).filter_by(emp_no="F-902").first()
ppe.add({"employee_id": t2.id, "item": "Helmet", "qty": 1, "date": (today - timedelta(days=200)).isoformat()}, s, jomon)
ok(ppe.check(t2.id, "", "Helmet", 1, today.isoformat(), 0, s, jomon)["too_soon"], "helmet after 200 days: still too soon (1 a year)")
ok(not ppe.check(t2.id, "", "Safety Vest", 1, today.isoformat(), 0, s, jomon)["too_soon"], "other items not affected")
# stock
ppe.set_stock({"item": "Safety Shoes", "counted": 20, "reorder": 5}, s, jomon)
st = {x["item"]: x for x in ppe.get_ppe(s, jomon)["stock"]}
ok(st["Safety Shoes"]["counted"] == 20 and st["Safety Shoes"]["available"] == 20 - st["Safety Shoes"]["issued_since"], "stock: counted today, given since taken off")
ppe.add({"employee_id": t1.id, "item": "Safety Shoes", "qty": 1}, s, jomon)
st2 = {x["item"]: x for x in ppe.get_ppe(s, jomon)["stock"]}
ok(st2["Safety Shoes"]["available"] == st["Safety Shoes"]["available"] - 1, "giving a pair takes one off the store")
# delete a wrong entry
ppe.delete(e2.id, s, jomon)
ok(not ppe.check(t9.id, "", "Safety Shoes", 1, (today + timedelta(days=200)).isoformat(), 0, s, jomon)["too_soon"], "wrong entry removed; after 6 months he is due again")
# future date refused
ok("future" in refused(ppe.add, {"employee_id": t9.id, "item": "Jacket", "date": (today + timedelta(days=3)).isoformat()}, s, jomon), "future date refused")
# exports
async def body(r):
    b = b""
    async for c in r.body_iterator: b += c if isinstance(c, bytes) else c.encode()
    return b
tok = auth.create_download_token("v")
import openpyxl
x = asyncio.run(body(ppe.export(tok, "excel", "grid", s)))
rows = list(openpyxl.load_workbook(io.BytesIO(x)).active.iter_rows(values_only=True))
ok(any(r and "F-900" in [str(v) for v in r] for r in rows), "by-worker Excel lists the workers")
x = asyncio.run(body(ppe.export(tok, "excel", "log", s)))
rows = [r for r in openpyxl.load_workbook(io.BytesIO(x)).active.iter_rows(values_only=True) if r and r[1] and str(r[1]).startswith(("F", "D", "P"))]
ok(len(rows) == s.query(models.PpeIssue).count(), f"issue register Excel has every entry ({len(rows)})")
ok(asyncio.run(body(ppe.export(tok, "pdf", "grid", s)))[:4] == b"%PDF", "PDF made")
print("fails=", fails)
