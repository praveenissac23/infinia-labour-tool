"""A varied synthetic corpus for the document reader: bills / IPAs, subcontractor PCs, variation logs and contracts
laid out many different ways, rendered as Excel, typed PDF, scanned PDF (rasterised with skew, noise, blur, JPEG)
and Word - with the true figures kept in truth.json.

    python3 tests/docread/synth.py OUT_DIR [--n 24] [--seed 7] [--scans 0.5]
"""
import argparse
import io
import json
import os
import random
import subprocess
import sys
import tempfile
from datetime import date, timedelta

# ---- figures ---------------------------------------------------------------------------------------------

R2 = lambda x: round(x + 1e-9, 2)


def bill_figures(rng, sub=False):
    contract = rng.choice([rng.randrange(250, 990) * 1000, rng.randrange(1000, 9000) * 1000 + rng.choice([0, 500, 750])])
    done = rng.uniform(0.15, 0.95)
    this_share = rng.uniform(0.06, 0.35)
    work_c = R2(contract * done + rng.random())
    work_p = R2(work_c * (1 - this_share))
    has_var = rng.random() < 0.55
    var_c = R2(contract * rng.uniform(0.01, 0.08)) if has_var else 0.0
    var_p = R2(var_c * rng.uniform(0.3, 1.0)) if has_var else 0.0
    ret_pct = rng.choice([5, 10, 10])
    adv_pct = rng.choice([0, 10, 10, 20]) if not sub else rng.choice([0, 10, 10])
    gross_c, gross_p = R2(work_c + var_c), R2(work_p + var_p)
    ret_c, ret_p = R2(gross_c * ret_pct / 100), R2(gross_p * ret_pct / 100)
    adv_c, adv_p = R2(work_c * adv_pct / 100), R2(work_p * adv_pct / 100)
    contra = []
    if rng.random() < (0.6 if sub else 0.3):
        names = rng.sample(["Material supplied by main contractor", "Manpower supplied", "Back charges - cleaning", "Contra charges",
                            "Plant & equipment supplied", "Penalty for delay", "Other deductions"], rng.choice([1, 1, 2, 3]))
        for nm in names:
            c = R2(rng.uniform(500, 25000))
            p = R2(c * rng.choice([0, 0.5, 1.0]))
            contra.append({"label": nm, "c": c, "p": p})
    con_c, con_p = R2(sum(x["c"] for x in contra)), R2(sum(x["p"] for x in contra))
    advpay = R2(contract * adv_pct / 100) if sub and adv_pct else 0.0
    net_c = R2(gross_c + advpay - ret_c - adv_c - con_c)
    net_p = R2(gross_p + advpay - ret_p - adv_p - con_p)
    net_t = R2(net_c - net_p)
    vat = R2(net_t * 0.05)
    auth = R2(rng.choice([0, 0, 0, 280, 360, 1250])) if not sub else 0.0
    total = R2(net_t + vat + auth)
    rounded = rng.random() < 0.25
    if rounded:
        total = float(round(total))
    seq = rng.randrange(2, 24)
    d = date(2026, rng.randrange(1, 11), rng.randrange(1, 28))
    return dict(contract=contract, work_c=work_c, work_p=work_p, var_c=var_c, var_p=var_p, gross_c=gross_c, gross_p=gross_p,
                ret_pct=ret_pct, adv_pct=adv_pct, ret_c=ret_c, ret_p=ret_p, adv_c=adv_c, adv_p=adv_p, contra=contra, con_c=con_c, con_p=con_p,
                advpay=advpay, net_c=net_c, net_p=net_p, net_t=net_t, vat=vat, auth=auth, total=total, seq=seq, date=d, has_var=has_var,
                rounded=rounded)


def truth_of(F, kind, var_shown=True):
    if kind == "ipa":
        t = {"seq": F["seq"], "submitted": F["date"].isoformat(), "claimed_cum": F["gross_c"], "adv_rec_cum": F["adv_c"] if F["adv_pct"] else None,
             "ret_cum": F["ret_c"], "previous": F["net_p"], "vat": F["vat"], "total": F["total"]}
        if F["auth"]:
            t["authority"] = F["auth"]
        if F["has_var"] and var_shown:
            t["variation"] = F["var_c"]
    else:
        t = {"seq": F["seq"], "work_cum": F["work_c"], "var_cum": F["var_c"] if F["has_var"] else None, "adv_cum": F["advpay"] or None,
             "adv_rec_cum": F["adv_c"] or None, "ret_cum": F["ret_c"], "other_cum": F["con_c"] or None, "prev_cert": F["net_p"], "net": F["net_t"],
             "vat": F["vat"], "total": F["total"]}
    return {k: v for k, v in t.items() if v is not None}


# ---- number styles ----------------------------------------------------------------------------------------

def style(rng):
    return {"cur": rng.choice(["", "", "AED ", "Dhs. ", "AED"]), "neg": rng.choice(["-", "()", "()", "less"]), "dec": rng.random() < 0.9,
            "sep": rng.random() < 0.92}


def fmt(v, st, neg=False):
    s = f"{abs(v):,.2f}" if st["dec"] else f"{abs(v):,.0f}"
    if not st["sep"]:
        s = s.replace(",", "")
    if not st["dec"] and abs(v - round(v)) > 0.004:
        s = f"{abs(v):,.2f}" if st["sep"] else f"{abs(v):.2f}"
    s = (st["cur"] + s) if st["cur"] else s
    if neg and v != 0:
        if st["neg"] == "-":
            s = "-" + s
        elif st["neg"] == "()":
            s = "(" + s + ")"
    return s


def xfmt(st, neg_style):
    base = "#,##0.00" if st["dec"] else "#,##0"
    if not st["sep"]:
        base = base.replace("#,##", "")
    if st["cur"].strip():
        base = f'"{st["cur"].strip()} "' + base
    if neg_style == "()":
        return base + ";(" + base + ")"
    return base


# ---- the page model: rows of cells on a column grid (rendered to Excel, PDF and scans) --------------------

class Page:
    def __init__(self, widths, title=""):
        self.widths = widths          # mm
        self.rows = []
        self.title = title

    def row(self, *cells, h=6.0, bold=False, border=False, hidden=False, gap=0):
        """cells: (col, span, value, kind) - kind: text | num | negnum | pct | date"""
        self.rows.append({"cells": list(cells), "h": h, "bold": bold, "border": border, "hidden": hidden, "gap": gap})


def T(col, text, span=1, align="l"):
    return {"col": col, "span": span, "v": text, "k": "text", "align": align}


def N(col, v, span=1, neg=False):
    return {"col": col, "span": span, "v": v, "k": "num", "neg": neg, "align": "r"}


def P(col, v, span=1):
    return {"col": col, "span": span, "v": v, "k": "pct", "align": "r"}


def D(col, d, span=1, f="%d/%m/%Y"):
    return {"col": col, "span": span, "v": d, "k": "date", "f": f, "align": "c"}


def render_xlsx(pages, path, st, hidden_sheet=None):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin")
    for pi, pg in enumerate(pages):
        ws = wb.create_sheet(pg.title or f"Sheet{pi + 1}")
        for i, w in enumerate(pg.widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w / 2.0
        r = 1
        for row in pg.rows:
            r += row.get("gap", 0)
            for c in row["cells"]:
                cell = ws.cell(r, c["col"] + 1)
                if c["k"] == "text":
                    cell.value = c["v"]
                elif c["k"] == "num":
                    v = c["v"]
                    neg = c.get("neg")
                    if neg and st["neg"] in ("-", "()"):
                        v = -abs(v)
                    cell.value = v
                    cell.number_format = xfmt(st, st["neg"])
                elif c["k"] == "pct":
                    cell.value = c["v"] / 100.0
                    cell.number_format = "0.00%"
                elif c["k"] == "date":
                    cell.value = c["v"]
                    cell.number_format = "dd/mm/yyyy"
                if c["span"] > 1:
                    ws.merge_cells(start_row=r, start_column=c["col"] + 1, end_row=r, end_column=c["col"] + c["span"])
                cell.alignment = Alignment(horizontal={"l": "left", "r": "right", "c": "center"}[c["align"]], wrap_text=False)
                if row["bold"]:
                    cell.font = Font(bold=True)
                if row["border"]:
                    cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if row["hidden"]:
                ws.row_dimensions[r].hidden = True
            r += 1
    if hidden_sheet:
        ws = wb.create_sheet("old rev")
        for i, rr in enumerate(hidden_sheet):
            for j, v in enumerate(rr):
                ws.cell(i + 1, j + 1).value = v
        ws.sheet_state = "hidden"
    wb.save(path)


def _txt(c, st):
    if c["k"] == "text":
        return c["v"]
    if c["k"] == "num":
        if c["v"] < 0 and not c.get("neg"):
            return fmt(c["v"], dict(st, neg=st["neg"] if st["neg"] in ("-", "()") else "-"), True)   # a minus figure shows its sign
        return fmt(c["v"], st, c.get("neg") and st["neg"] in ("-", "()"))
    if c["k"] == "pct":
        return f"{c['v']:.2f}%" if c["v"] != int(c["v"]) else f"{int(c['v'])}%"
    if c["k"] == "date":
        return c["v"].strftime(c.get("f", "%d/%m/%Y"))
    return str(c["v"])


def render_pdf(pages, path, st, font_size=8.5, rules=True, arabic=False):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    fn, fb = "Helvetica", "Helvetica-Bold"
    try:
        pdfmetrics.registerFont(TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
        pdfmetrics.registerFont(TTFont("DejaVuB", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))
        if arabic:
            fn, fb = "DejaVu", "DejaVuB"
    except Exception:
        pass
    from reportlab.lib.pagesizes import landscape
    size = landscape(A4) if max(sum(pg.widths) for pg in pages) > 200 else A4
    c = canvas.Canvas(path, pagesize=size)
    W, H = size
    for pg in pages:
        left = 14 * mm
        xs = [left]
        tot = sum(pg.widths)
        k = (W - 28 * mm) / (tot * mm) if tot * mm > (W - 28 * mm) else 1.0
        for w in pg.widths:
            xs.append(xs[-1] + w * mm * k)
        y = H - 16 * mm
        for row in pg.rows:
            if row["hidden"]:
                continue
            y -= row.get("gap", 0) * row["h"] * mm
            h = row["h"] * mm
            # long words in a table cell wrap onto more lines (as Excel's "wrap text" prints them): a taller row
            plan = []
            nl = 1
            for cell in row["cells"]:
                x0, x1 = xs[cell["col"]], xs[min(len(xs) - 1, cell["col"] + cell["span"])]
                t = _txt(cell, st)
                f = fb if row["bold"] else fn
                size = font_size
                while c.stringWidth(t, f, size) > (x1 - x0 - 3) and size > 5.5 and cell["k"] != "text":
                    size -= 0.5
                lines = [t]
                if cell["k"] == "text" and row["border"] and c.stringWidth(t, f, size) > (x1 - x0 - 3):
                    if len(t) > 28 and (x1 - x0) > 60:
                        lines, cur = [], ""
                        for wd in t.split():
                            if cur and c.stringWidth(cur + " " + wd, f, size) > (x1 - x0 - 3):
                                lines.append(cur)
                                cur = wd
                            else:
                                cur = (cur + " " + wd).strip()
                        lines.append(cur)
                    else:
                        while c.stringWidth(t, f, size) > (x1 - x0 - 3) and size > 5:
                            size -= 0.5
                if cell["k"] == "text" and not row["border"] and c.stringWidth(t, f, size) > (x1 - x0 - 3):
                    while c.stringWidth(t, f, size) > (x1 - x0 - 3) * 1.6 and size > 6:
                        size -= 0.5
                nl = max(nl, len(lines))
                plan.append((cell, x0, x1, lines, f, size))
            lead = font_size + 2
            h = max(h, nl * lead + 4)
            if y - h < 14 * mm:
                c.showPage()
                y = H - 16 * mm
            for cell, x0, x1, lines, f, size in plan:
                c.setFont(f, size)
                top = y - (h - len(lines) * lead) / 2 - size
                for li, t in enumerate(lines):
                    ty = top - li * lead + 1.2 if len(lines) > 1 else y - h + (h - size) / 2 + 1.2
                    if cell["align"] == "r":
                        c.drawRightString(x1 - 1.5, ty, t)
                    elif cell["align"] == "c":
                        c.drawCentredString((x0 + x1) / 2, ty, t)
                    else:
                        c.drawString(x0 + 1.5, ty, t)
                if row["border"] and rules:
                    c.setLineWidth(0.5)
                    c.rect(x0, y - h, x1 - x0, h)
            y -= h
        c.showPage()
    c.save()


def scan_pdf(pdf_path, out_path, rng, dpi=200, jpeg=True):
    """A typed PDF made to look scanned: rasterised, turned a little, noise, blur, JPEG, back into a PDF of images."""
    from PIL import Image, ImageFilter
    import numpy as np
    with tempfile.TemporaryDirectory() as d:
        subprocess.run(["pdftoppm", "-r", str(dpi), "-gray", "-png", pdf_path, os.path.join(d, "p")], check=True, capture_output=True)
        imgs = []
        for fn in sorted(os.listdir(d)):
            im = Image.open(os.path.join(d, fn)).convert("L")
            ang = rng.uniform(-1.8, 1.8)
            im = im.rotate(ang, resample=Image.BICUBIC, fillcolor=255, expand=False)
            a = np.asarray(im, dtype=np.float32)
            a = a * rng.uniform(0.85, 1.0) + rng.uniform(0, 25)            # greyer paper, lighter ink
            a += np.random.default_rng(rng.randrange(1 << 30)).normal(0, rng.uniform(4, 14), a.shape)
            im = Image.fromarray(np.clip(a, 0, 255).astype(np.uint8))
            im = im.filter(ImageFilter.GaussianBlur(rng.uniform(0.3, 0.9)))
            if jpeg:
                buf = io.BytesIO()
                im.save(buf, format="JPEG", quality=rng.randrange(45, 80))
                im = Image.open(io.BytesIO(buf.getvalue())).convert("L")
            imgs.append(im)
        imgs[0].save(out_path, "PDF", resolution=float(dpi), save_all=True, append_images=imgs[1:])


# ---- bill layouts -----------------------------------------------------------------------------------------

SYN = {
    "work": ["Total Value of Work done", "Gross value of work done to date", "Value of Work Executed to Date", "Work done up to date",
             "Original contract works completed", "Measured works", "Value of works completed", "Gross Valuation of Work"],
    "var": ["Approved Variations", "Variations executed", "Value of Variation Orders", "Variation works (VOs)", "Change orders approved"],
    "gross": ["Total value of work executed (A+B)", "Gross total including variations", "Total gross value", "Total (A + B)"],
    "ret": ["Less: Retention @ {r}%", "Retention money ({r}%)", "Less Retention {r}%", "Retention held", "Retension {r}%"],
    "adv": ["Less: Advance Recovery", "Recovery of advance payment ({a}%)", "Advance payment recovery", "Less: Mobilisation advance recovered"],
    "prev": ["Less: Previously certified", "Less Previous Progressive Payments", "Amount paid up to last certificate", "Less: Total of previous payments",
             "Previous IPAs certified"],
    "net": ["Net amount due (excl. VAT)", "Net amount for this certificate", "Amount due this period before VAT", "Net payable excluding VAT"],
    "netc": ["Net certified to date", "Total certified to date (net)", "Net value to date"],
    "vat": ["Add: VAT 5%", "VAT @ 5%", "Value Added Tax (5%)", "Add. 5% Tax Payment"],
    "total": ["Total Amount due for Payment", "Total payable including VAT", "Amount certified for payment", "Net payable incl. VAT", "TOTAL DUE"],
    "auth": ["Add: Payment to Authorities", "Authority fees (DEWA / DM)"],
    "advpay": ["Advance Payment", "Mobilisation advance paid"],
    "rel": ["Retention release"],
}
COLH = {"p": ["Previous", "Upto Last", "Previous Certified", "Up to previous", "Prev. Cumulative"],
        "t": ["This Period", "This Bill", "Current", "Present", "This Certificate"],
        "c": ["Cumulative", "To Date", "Up to date", "Total to date", "Gross Certified"]}
AR = {"Previous": "السابق", "This Period": "الحالي", "Cumulative": "التراكمي", "Description": "البيان", "Amount": "المبلغ"}

PROJECTS = ["Construction of G+1 Villa at Plot 6457, Al Barari", "Proposed B+G+1+R Residential Villa, Sobha Hartland, Plot 347-3106",
            "Fit-out works for Villa WS-46, Sobha Hartland", "Construction of Private Villa at Keturah Reserve VL-88"]
CLIENTS = ["Mr. Rashid Al Mansoori", "Emkaan Design L.L.C", "Nenaya Holdings Ltd.", "Mr. Asif Akbar Ali"]
SUBS = ["Al Raida Electromechanical Services LLC", "Royal Everst Technical Services", "Woodcraft Interiors LLC", "OneGB Piling & Foundation Contracting LLC",
        "Naif Crescent Technical Services LLC", "Blue Line Aluminium & Glass LLC"]


def head_block(pg, rng, F, kind, ncols, ls=1):
    title = rng.choice(["INTERIM PAYMENT CERTIFICATE", "PAYMENT APPLICATION", "INTERIM PAYMENT APPLICATION", "PROGRESS PAYMENT CERTIFICATE"]) \
        if kind == "ipa" else rng.choice(["PAYMENT CERTIFICATE", "SUBCONTRACTOR PAYMENT CERTIFICATE", "INTERIM PAYMENT CERTIFICATE (SUBCONTRACT)"])
    pg.row(T(0, title, span=ncols, align="c"), bold=True, h=8)
    pg.row(T(0, "Project:", span=ls), T(ls, rng.choice(PROJECTS), span=ncols - ls))
    if kind == "ipa":
        pg.row(T(0, "Client:", span=ls), T(ls, rng.choice(CLIENTS), span=ncols - ls))
        lab = rng.choice(["Interim Bill No:", "IPA No.", "Payment Application No:", "Certificate No."])
        no = rng.choice([f"{F['seq']}", f"IPA-{F['seq']:02d}", f"914/IPA/{F['seq']:02d} R1", f"{F['seq']:02d}"])
    else:
        sub = rng.choice(SUBS)
        F["sub"] = sub
        pg.row(T(0, "Subcontractor:", span=ls), T(ls, "M/s. " + sub, span=ncols - ls))
        lab = rng.choice(["IPC No.", "PC No.", "Payment Certificate No."])
        no = rng.choice([f"{F['seq']}", f"PC-{F['seq']:02d}", f"{F['seq']:02d}"])
    if ncols - 2 > ls:
        pg.row(T(0, lab, span=ls), T(ls, no), T(ncols - 2, rng.choice(["Date:", "Issue Date:", "Date of issue:"]), align="r"), D(ncols - 1, F["date"], f=rng.choice(["%d/%m/%Y", "%d.%m.%Y", "%d-%b-%Y"])))
    else:
        pg.row(T(0, lab, span=ls), T(ls, no))
        pg.row(T(0, rng.choice(["Date:", "Issue Date:"]), span=ls), D(ls, F["date"]))
    pg.row(T(0, rng.choice(["Contract Value", "Original Contract Sum", "Contract Sum"]), span=ls), N(ls, F["contract"]), gap=1)


def lab(rng, k, F):
    return rng.choice(SYN[k]).replace("{r}", str(F["ret_pct"])).replace("{a}", str(F["adv_pct"]))


def layout_single(rng, F, kind, arabic=False):
    """The office's own single-column payment summary (label | amount)."""
    pg = Page([70, 20, 25, 35], "IPA" if kind == "ipa" else "PC")
    head_block(pg, rng, F, kind, 4)
    if rng.random() < 0.6:
        pg.row(T(0, "Description" + (" / البيان" if arabic else "")), T(3, "Amount (AED)" + (" / المبلغ" if arabic else ""), align="r"), bold=True, border=True, gap=1)
    var_sep = F["has_var"] and rng.random() < 0.7
    neg = lambda v: N(3, v, neg=True)
    if var_sep:
        pg.row(T(0, lab(rng, "work", F)), N(3, F["work_c"]), border=True)
        pg.row(T(0, lab(rng, "var", F)), N(3, F["var_c"]), border=True)
        if rng.random() < 0.7:
            pg.row(T(0, lab(rng, "gross", F)), N(3, F["gross_c"]), border=True, bold=True)
    else:
        pg.row(T(0, lab(rng, "work", F)), P(2, round(F["gross_c"] / F["contract"] * 100, 2)) if rng.random() < 0.4 else T(2, ""), N(3, F["gross_c"]), border=True)
    if kind == "subpc" and F["advpay"]:
        pg.row(T(0, lab(rng, "advpay", F)), P(2, F["adv_pct"]), N(3, F["advpay"]), border=True)
    items = []
    if F["adv_pct"]:
        items.append((lab(rng, "adv", F), F["adv_c"]))
    items.append((lab(rng, "ret", F), F["ret_c"]))
    for x in F["contra"]:
        items.append(("Less: " + x["label"] if rng.random() < 0.5 else x["label"], x["c"]))
    rng.shuffle(items)
    prev_first = rng.random() < 0.4
    if prev_first:
        items.insert(0, (lab(rng, "prev", F), F["net_p"]))
    else:
        items.append((lab(rng, "prev", F), F["net_p"]))
    for l, v in items:
        pg.row(T(0, l), neg(v), border=True)
    if rng.random() < 0.5:
        pg.row(T(0, lab(rng, "net", F)), N(3, F["net_t"]), border=True, bold=True)
    pg.row(T(0, lab(rng, "vat", F)), N(3, F["vat"]), border=True)
    if F["auth"]:
        pg.row(T(0, lab(rng, "auth", F)), N(3, F["auth"]), border=True)
    pg.row(T(0, lab(rng, "total", F)), N(3, F["total"]), border=True, bold=True)
    if rng.random() < 0.3:
        # a figure to the side of the total (a distractor): the previous revision's total
        pg.rows[-1]["cells"].append(N(2, R2(F["total"] * rng.uniform(0.7, 0.95))))
    if rng.random() < 0.5:
        pg.row(T(0, f"(Dirhams {words(F['total'])} Only)", span=4))
    return [pg], var_sep


def layout_three(rng, F, kind, arabic=False):
    """A consultant's certificate: Previous | This period | Cumulative columns (in any order), maybe a contract
    value column, maybe a two-row heading."""
    order = rng.sample(["p", "t", "c"], 3) if rng.random() < 0.35 else ["p", "t", "c"]
    with_contract = rng.random() < 0.35
    widths = [62] + ([26] if with_contract else []) + [27, 27, 27]
    pg = Page(widths, "Certificate")
    head_block(pg, rng, F, kind, len(widths))
    off = 2 if with_contract else 1
    hdr = [T(0, "Description" + (" / البيان" if arabic else ""))]
    if with_contract:
        hdr.append(T(1, "Contract Value", align="c"))
    names = {k: rng.choice(COLH[k]) for k in "ptc"}
    if arabic:
        names = {k: (names[k] + " / " + AR.get(["Previous", "This Period", "Cumulative"]["ptc".index(k)], "")) for k in "ptc"}
    two_row = rng.random() < 0.4
    if two_row:
        top = [T(0, "")] + ([T(1, "")] if with_contract else []) + [T(off, "Amount (AED)", span=3, align="c")]
        pg.row(*top, bold=True, border=True, gap=1)
    for j, k in enumerate(order):
        hdr.append(T(off + j, names[k], align="c"))
    pg.row(*hdr, bold=True, border=True, gap=0 if two_row else 1)
    V = {"p": {}, "t": {}, "c": {}}
    def put(key, c, p):
        V["c"][key], V["p"][key], V["t"][key] = c, p, R2(c - p)
    put("work", F["work_c"], F["work_p"])
    if F["has_var"]:
        put("var", F["var_c"], F["var_p"])
        put("gross", F["gross_c"], F["gross_p"])
    if kind == "subpc" and F["advpay"]:
        put("advpay", F["advpay"], F["advpay"])
    if F["adv_pct"]:
        put("adv", F["adv_c"], F["adv_p"])
    put("ret", F["ret_c"], F["ret_p"])
    for i, x in enumerate(F["contra"]):
        put(f"contra{i}", x["c"], x["p"])
    put("net", F["net_c"], F["net_p"])
    def line(label, key, neg=False, bold=False):
        cells = [T(0, label)]
        if with_contract:
            cells.append(N(1, F["contract"]) if key == "work" else T(1, ""))
        for j, k in enumerate(order):
            cells.append(N(off + j, V[k][key], neg=neg))
        pg.row(*cells, border=True, bold=bold)
    line(lab(rng, "work", F), "work")
    if F["has_var"]:
        line(lab(rng, "var", F), "var")
        line(lab(rng, "gross", F), "gross", bold=True)
    if "advpay" in V["c"]:
        line(lab(rng, "advpay", F), "advpay")
    if F["adv_pct"]:
        line(lab(rng, "adv", F), "adv", neg=True)
    line(lab(rng, "ret", F), "ret", neg=True)
    for i, x in enumerate(F["contra"]):
        line(x["label"], f"contra{i}", neg=True)
    line(lab(rng, "netc", F), "net", bold=True)
    last = len(widths) - 1
    pg.row(T(0, lab(rng, "net", F), span=last), N(last, F["net_t"]), border=True, gap=1)
    pg.row(T(0, lab(rng, "vat", F), span=last), N(last, F["vat"]), border=True)
    if F["auth"]:
        pg.row(T(0, lab(rng, "auth", F), span=last), N(last, F["auth"]), border=True)
    pg.row(T(0, lab(rng, "total", F), span=last), N(last, F["total"]), border=True, bold=True)
    return [pg], F["has_var"]


def layout_atoe(rng, F, kind, arabic=False):
    """Infinia's own subcontractor certificate: A work, B advance, C deductions, D additions, E previously certified."""
    pg = Page([10, 78, 22, 40], "PC")
    head_block(pg, rng, F, kind, 4, ls=2)
    pg.row(T(0, "A"), T(1, "ADD: VALUE OF WORK EXECUTED"), bold=True, gap=1)
    pg.row(T(1, "Value of Work Executed to Date"), D(2, F["date"]), N(3, F["work_c"]))
    pg.row(T(1, "Value of Variations Executed to Date"), N(3, F["var_c"]) if F["has_var"] else T(3, "-", align="r"))
    pg.row(T(1, "Total Value of Work Executed"), N(3, F["gross_c"]), bold=True)
    pg.row(T(0, "B"), T(1, "ADD: ADVANCE PAYMENT"), bold=True)
    pg.row(T(1, "Advance Payment"), P(2, F["adv_pct"]) if F["adv_pct"] else T(2, "%"), N(3, F["advpay"]) if F["advpay"] else T(3, "-", align="r"))
    pg.row(T(0, "C"), T(1, "LESS: DEDUCTIONS"), bold=True)
    pg.row(T(1, "Advance payment Recovery"), P(2, F["adv_pct"]) if F["adv_pct"] else T(2, "%"), N(3, F["adv_c"], neg=True) if F["adv_c"] else T(3, "-", align="r"))
    pg.row(T(1, "Retention"), P(2, F["ret_pct"]), N(3, F["ret_c"], neg=True))
    for x in F["contra"]:
        pg.row(T(1, x["label"]), N(3, x["c"], neg=True))
    pg.row(T(1, "Total Deductions"), N(3, R2(F["adv_c"] + F["ret_c"] + F["con_c"]), neg=True), bold=True)
    pg.row(T(0, "D"), T(1, "ADD: ADDITIONS"), bold=True)
    pg.row(T(1, "Retention Release"), T(3, "-", align="r"))
    pg.row(T(0, "E"), T(1, "LESS: PREVIOUSLY CERTIFIED AMOUNT"), bold=True)
    n = rng.randrange(1, 4)
    parts = [R2(F["net_p"] / n) for _ in range(n - 1)]
    parts.append(R2(F["net_p"] - sum(parts)))
    for i, x in enumerate(parts):
        pg.row(T(1, f"IPC - {i + 1:02d}"), N(3, x, neg=True))
    pg.row(T(1, "Total Previously Certified Amounts"), N(3, F["net_p"], neg=True), bold=True)
    pg.row(T(0, "Net Amount Due for Payment (A + B - C + D - E)", span=3), N(3, F["net_t"]), bold=True, gap=1)
    pg.row(T(0, "VAT (5%)", span=3), N(3, F["vat"]))
    pg.row(T(0, "Total Amount Due for Payment", span=3), N(3, F["total"]), bold=True)
    return [pg], True


def layout_kv(rng, F, kind, arabic=False):
    """Label: value written in one cell, two pairs to a row (a letter-style certificate)."""
    pg = Page([95, 95], "Cert")
    head_block(pg, rng, F, kind, 2)
    st = F["_st"]
    pairs = [(lab(rng, "work", F) if not F["has_var"] else lab(rng, "gross", F), F["gross_c"], False)]
    if F["has_var"]:
        pairs.insert(0, (lab(rng, "var", F), F["var_c"], False))
        pairs.insert(0, (lab(rng, "work", F), F["work_c"], False))
    if kind == "subpc" and F["advpay"]:
        pairs.append((lab(rng, "advpay", F), F["advpay"], False))
    if F["adv_pct"]:
        pairs.append((lab(rng, "adv", F), F["adv_c"], True))
    pairs.append((lab(rng, "ret", F), F["ret_c"], True))
    for x in F["contra"]:
        pairs.append((x["label"], x["c"], True))
    pairs.append((lab(rng, "prev", F), F["net_p"], True))
    pairs.append((lab(rng, "net", F), F["net_t"], False))
    pairs.append((lab(rng, "vat", F), F["vat"], False))
    if F["auth"]:
        pairs.append((lab(rng, "auth", F), F["auth"], False))
    pairs.append((lab(rng, "total", F), F["total"], False))
    sep = rng.choice([": ", " : ", " - "])
    txt = [f"{l.rstrip(':')}{sep}{fmt(v, st, n and st['neg'] in ('-', '()'))}" for l, v, n in pairs]
    if rng.random() < 0.5:
        for t in txt:
            pg.row(T(0, t, span=2))
    else:
        for i in range(0, len(txt), 2):
            pg.row(T(0, txt[i]), T(1, txt[i + 1]) if i + 1 < len(txt) else T(1, ""))
    return [pg], F["has_var"]


def layout_breakdown(rng, F, kind, arabic=False):
    """A long breakdown (BOQ items with previous / current / cumulative) before the summary - the summary at the end,
    the breakdown full of figures that look like bill lines."""
    pg = Page([12, 70, 26, 26, 26], "IPA")
    head_block(pg, rng, F, kind, 5, ls=2)
    pg.row(T(0, "No."), T(1, "Description"), T(2, rng.choice(COLH["p"]), align="c"), T(3, rng.choice(COLH["t"]), align="c"),
           T(4, rng.choice(COLH["c"]), align="c"), bold=True, border=True, gap=1)
    n = rng.randrange(8, 28)
    shares = [rng.random() for _ in range(n)]
    s = sum(shares)
    cum = [R2(F["work_c"] * x / s) for x in shares]
    cum[-1] = R2(F["work_c"] - sum(cum[:-1]))
    prv = [R2(c * (F["work_p"] / F["work_c"])) for c in cum]
    prv[-1] = R2(F["work_p"] - sum(prv[:-1]))
    items = ["Excavation and backfilling", "Blinding concrete", "Raft foundation", "Columns and shear walls", "Slabs and beams", "Block work",
             "Internal plaster", "External render", "Waterproofing", "Electrical first fix", "Plumbing first fix", "HVAC ducting", "Tiling",
             "Gypsum ceiling", "Painting", "Joinery", "Aluminium windows", "Landscaping", "Swimming pool", "Retention of earth - shoring",
             "Variation in retaining wall", "Advance works - mobilisation", "Preliminaries", "Site facilities", "Testing and commissioning",
             "Lift installation", "Staircase railings", "External works"]
    for i in range(n):
        pg.row(T(0, str(i + 1)), T(1, items[i % len(items)]), N(2, prv[i]), N(3, R2(cum[i] - prv[i])), N(4, cum[i]), border=True)
    pg.row(T(1, "TOTAL CARRIED TO SUMMARY"), N(2, F["work_p"]), N(3, R2(F["work_c"] - F["work_p"])), N(4, F["work_c"]), bold=True, border=True)
    pages = [pg]
    pg2 = Page([70, 20, 25, 35], "Summary")
    pg2.row(T(0, "SUMMARY OF PAYMENT", span=4, align="c"), bold=True)
    var_sep = F["has_var"]
    pg2.row(T(0, lab(rng, "work", F)), N(3, F["work_c"]), border=True)
    if var_sep:
        pg2.row(T(0, lab(rng, "var", F)), N(3, F["var_c"]), border=True)
        pg2.row(T(0, lab(rng, "gross", F)), N(3, F["gross_c"]), border=True)
    if kind == "subpc" and F["advpay"]:
        pg2.row(T(0, lab(rng, "advpay", F)), N(3, F["advpay"]), border=True)
    if F["adv_pct"]:
        pg2.row(T(0, lab(rng, "adv", F)), N(3, F["adv_c"], neg=True), border=True)
    pg2.row(T(0, lab(rng, "ret", F)), N(3, F["ret_c"], neg=True), border=True)
    for x in F["contra"]:
        pg2.row(T(0, x["label"]), N(3, x["c"], neg=True), border=True)
    pg2.row(T(0, lab(rng, "prev", F)), N(3, F["net_p"], neg=True), border=True)
    pg2.row(T(0, lab(rng, "net", F)), N(3, F["net_t"]), border=True)
    pg2.row(T(0, lab(rng, "vat", F)), N(3, F["vat"]), border=True)
    if F["auth"]:
        pg2.row(T(0, lab(rng, "auth", F)), N(3, F["auth"]), border=True)
    pg2.row(T(0, lab(rng, "total", F)), N(3, F["total"]), border=True, bold=True)
    pages.append(pg2)
    return pages, var_sep


BILL_LAYOUTS = {"single": layout_single, "three": layout_three, "atoe": layout_atoe, "kv": layout_kv, "breakdown": layout_breakdown}


def words(v):
    ones = "Zero One Two Three Four Five Six Seven Eight Nine Ten Eleven Twelve Thirteen Fourteen Fifteen Sixteen Seventeen Eighteen Nineteen".split()
    tens = "Twenty Thirty Forty Fifty Sixty Seventy Eighty Ninety".split()

    def w(n):
        if n < 20:
            return ones[n]
        if n < 100:
            return tens[n // 10 - 2] + ("" if n % 10 == 0 else " " + ones[n % 10])
        if n < 1000:
            return ones[n // 100] + " Hundred" + ("" if n % 100 == 0 else " and " + w(n % 100))
        if n < 1000000:
            return w(n // 1000) + " Thousand" + ("" if n % 1000 == 0 else " " + w(n % 1000))
        return w(n // 1000000) + " Million" + ("" if n % 1000000 == 0 else " " + w(n % 1000000))
    d = int(v)
    f = int(round((v - d) * 100))
    return w(d) + (f" and {w(f)} Fils" if f else "")


# ---- variation logs ---------------------------------------------------------------------------------------

VO_TITLES = ["Additional boundary wall height", "Change of floor tiles in majlis", "Extra power points in kitchen", "Omission of pergola",
             "Upgrade of sanitary fittings", "Additional waterproofing to roof", "Revised staircase design", "Swimming pool heater",
             "Change of AC system to VRF", "Additional lighting in garden", "Deletion of gypsum cornice", "Car porch shade structure",
             "Extra excavation due to rock", "Façade cladding change", "Smart home wiring", "Additional manhole", "Lift pit modification",
             "Change of kitchen worktop", "Extra plastering due to design change", "Additional steel for columns"]
VOH = {"ref": ["Ref.No", "VO No.", "Variation No", "CP No.", "Ref", "VO Ref"],
       "title": ["Description", "Description of Variation", "Item Description", "Particulars", "Details"],
       "submitted_amount": ["Submitted Amount", "Amount Submitted", "Claimed Amount", "Proposed Value", "Quoted Amount"],
       "revised_amount": ["Revised Amount", "Revised Value"],
       "approved_amount": ["Accepted Amount", "Approved Amount", "Certified Value", "Agreed Amount", "Approved Value"],
       "absorbed_amount": ["Absorbed Amount"],
       "submitted": ["Date of Submission", "Submitted On", "Submission Date"],
       "approved_on": ["Date of Acceptance", "Approval Date", "Approved On"],
       "remarks": ["Remarks", "Comments", "Notes"]}


def vo_rows(rng, n):
    rows = []
    d0 = date(2025, rng.randrange(1, 12), rng.randrange(1, 28))
    style_ref = rng.choice(["V{:02d}", "VO-{:03d}", "CP{:02d}", "SI-{}", "V.O. {}"])
    for i in range(n):
        omission = rng.random() < 0.12
        sub = R2(rng.uniform(1500, 180000)) * (-1 if omission else 1)
        st = rng.choices(["Approved", "Submitted", "Rejected", "Absorbed", "Identified"], [5, 3, 1, 1, 1])[0]
        rev = R2(sub * rng.uniform(0.85, 1.0)) if rng.random() < 0.25 and st in ("Approved", "Submitted") else 0.0
        appr = R2((rev or sub) * rng.uniform(0.8, 1.0)) if st == "Approved" else 0.0
        absd = R2(sub * rng.uniform(0.5, 1.0)) if st == "Absorbed" else 0.0
        ds = d0 + timedelta(days=rng.randrange(0, 300))
        da = ds + timedelta(days=rng.randrange(5, 60)) if st == "Approved" else None
        rem = {"Approved": rng.choice(["Accepted", "Approved by consultant", ""]), "Rejected": "Rejected by client", "Submitted": rng.choice(["Under review", ""]),
               "Absorbed": "Absorbed", "Identified": "To be submitted"}[st]
        if st == "Identified":
            sub = 0.0
        rows.append({"ref": style_ref.format(i + 1), "title": VO_TITLES[(i + rng.randrange(0, 20)) % len(VO_TITLES)], "submitted_amount": sub,
                     "revised_amount": rev, "approved_amount": appr, "absorbed_amount": absd, "submitted": ds if st != "Identified" else None,
                     "approved_on": da, "remarks": rem, "status": st})
    return rows


def vo_layout(rng, rows, arabic=False):
    cols = ["ref", "title", "submitted_amount"] + (["revised_amount"] if rng.random() < 0.5 else []) + ["approved_amount"] + \
           (["absorbed_amount"] if any(r["absorbed_amount"] for r in rows) or rng.random() < 0.3 else []) + ["submitted", "approved_on", "remarks"]
    if rng.random() < 0.4:                       # money columns after the dates
        mid = [c for c in cols if c.endswith("amount")]
        rest = [c for c in cols if not c.endswith("amount") and c not in ("ref", "title", "remarks")]
        cols = ["ref", "title"] + rest + mid + ["remarks"]
    widths = {"ref": 16, "title": 62, "remarks": 30}
    W = [widths.get(c, 24) for c in cols]
    pg = Page(W, "VO Log")
    pg.row(T(0, rng.choice(["VARIATION LOG", "VARIATION ORDER REGISTER", "SUMMARY OF VARIATIONS"]), span=len(cols), align="c"), bold=True, h=8)
    pg.row(T(0, "Project: " + rng.choice(PROJECTS), span=len(cols)))
    two_row = rng.random() < 0.4
    names = {c: rng.choice(VOH[c]) for c in cols}
    if arabic:
        names["title"] += " / الوصف"
    money = [i for i, c in enumerate(cols) if c.endswith("amount")]
    if two_row and money == list(range(money[0], money[0] + len(money))):
        top = []
        for i, c in enumerate(cols):
            if i == money[0]:
                top.append(T(i, "Amount (AED)", span=len(money), align="c"))
            elif i not in money:
                top.append(T(i, names[c] if c not in money else "", align="c"))
        pg.row(*top, bold=True, border=True, gap=1)
        sub = [T(i, names[c].replace(" Amount", "").replace("Amount ", "").replace(" Value", "") if i in money else "", align="c") for i, c in enumerate(cols)]
        pg.row(*sub, bold=True, border=True)
    else:
        pg.row(*[T(i, names[c], align="c") for i, c in enumerate(cols)], bold=True, border=True, gap=1)
    absorbed_section = rng.random() < 0.3 and any(r["status"] == "Absorbed" for r in rows)
    body = [r for r in rows if not (absorbed_section and r["status"] == "Absorbed")]
    tail = [r for r in rows if absorbed_section and r["status"] == "Absorbed"]
    dfmt = rng.choice(["%d/%m/%Y", "%d.%m.%Y", "%d-%b-%y", "%d %B %Y"])

    def line(r):
        cells = []
        for i, c in enumerate(cols):
            v = r.get(c)
            if c.endswith("amount"):
                cells.append(N(i, abs(v), neg=v < 0) if v else T(i, "-", align="r"))
            elif c in ("submitted", "approved_on"):
                cells.append(D(i, v, f=dfmt) if v else T(i, "-", align="c"))
            else:
                cells.append(T(i, v or ""))
        pg.row(*cells, border=True)
    for r in body:
        line(r)
    if tail:
        pg.row(T(1, "ABSORBED VARIATIONS"), bold=True)
        for r in tail:
            line(r)
    tot = [T(1, rng.choice(["TOTAL", "Grand Total", "Total"]))]
    for i, c in enumerate(cols):
        if c.endswith("amount"):
            s = R2(sum(r[c] for r in rows))
            tot.append(N(i, abs(s), neg=s < 0))
    pg.row(*tot, bold=True, border=True)
    return [pg], cols


# ---- contracts ----------------------------------------------------------------------------------------------

def contract_doc(rng, i):
    value = rng.randrange(400, 9000) * 1000 + rng.choice([0, 0, 500, 250])
    adv, ret = rng.choice([10, 15, 20, 0]), rng.choice([5, 10])
    start = date(2026, rng.randrange(1, 10), rng.randrange(1, 28))
    months = rng.choice([8, 10, 12, 14, 18])
    y, m = start.year + (start.month - 1 + months) // 12, (start.month - 1 + months) % 12 + 1
    end = date(y, m, min(start.day, 28))
    proj, client = rng.choice(PROJECTS), rng.choice(CLIENTS)
    cons = rng.choice(["Dimensions Engineering Consultants", "Renaissance Engineering Consultants", "Emkaan Architecture & Engineering", "Al Gurg Consultants"])
    loc = rng.choice(["Plot 6457, Al Barari, Dubai", "Plot 347-3106, Sobha Hartland, Dubai", "VL-88, Keturah Reserve, Meydan, Dubai"])
    truth = {"name": proj, "client": client.replace("Mr. ", "Mr. "), "consultant": cons, "contract_value": float(value),
             "ret_pct": float(ret), "start_date": start.isoformat(), "end_date": end.isoformat()}
    if adv:
        truth["adv_pct"] = float(adv)
    style_ = rng.choice(["table", "letter", "mixed"])
    paras, table = [], []
    paras.append(rng.choice(["LETTER OF AWARD", "CONTRACT AGREEMENT", "NOTIFICATION OF AWARD", "LETTER OF ACCEPTANCE"]))
    paras.append(f"Ref: {rng.choice(['RAM', 'DEC', 'EMK'])}/LOA/2026/{rng.randrange(10, 99):03d}")
    paras.append(f"Date: {(start - timedelta(days=rng.randrange(5, 30))).strftime(rng.choice(['%d %B %Y', '%d/%m/%Y']))}")
    paras.append("To: M/s Infinia Contracting LLC, P.O. Box 85567, Dubai")
    vat_in = R2(value * 1.05)
    bond = R2(value * 0.10)
    vtxt = f"AED {value:,.2f}" if rng.random() < 0.7 else f"AED {value:,.0f}/-"
    if style_ in ("table", "mixed"):
        table = [["Project:", proj], ["Employer:" if rng.random() < 0.5 else "Client:", client], ["Consultant:", cons], ["Location:", loc]]
        if style_ == "table":
            table.append([rng.choice(["Contract Sum:", "Contract Value:", "Lump Sum Price:"]), vtxt + " (exclusive of VAT)"])
            table.append(["Commencement Date:", start.strftime("%d/%m/%Y")])
            table.append(["Completion Date:", end.strftime("%d %B %Y")])
    if style_ != "table":
        paras.append(f"We are pleased to award you the works for the {proj} for the {rng.choice(['Employer', 'Client'])} {client}, "
                     f"under the supervision of {cons}.")
        paras.append(f"The Contract Sum is {vtxt} (Dirhams {words(value)} Only) exclusive of VAT. "
                     f"The total including 5% VAT amounts to AED {vat_in:,.2f}.")
        paras.append(f"The works shall commence on {start.strftime('%d %B %Y')} and the time for completion is {months} months, "
                     f"i.e. the works shall be completed by {end.strftime('%d %B %Y')}.")
    paras.append(f"A performance bond of 10% of the contract value (AED {bond:,.2f}) shall be submitted within 14 days.")
    if adv:
        paras.append(f"An advance payment of {adv}% shall be paid against an advance payment guarantee, recovered pro-rata from interim payments.")
    paras.append(f"Retention of {ret}% will be deducted from each interim payment, half released at handover.")
    paras.append("Insurance: Contractor's All Risk policy for AED 5,000,000.00 minimum.")
    return truth, paras, table


def render_contract(paras, table, path, fmt_):
    if fmt_ == "docx":
        from docx import Document
        d = Document()
        for p in paras[:4]:
            d.add_paragraph(p)
        if table:
            t = d.add_table(rows=0, cols=2)
            for r in table:
                cs = t.add_row().cells
                cs[0].text, cs[1].text = r
        for p in paras[4:]:
            d.add_paragraph(p)
        d.save(path)
        return
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    st = getSampleStyleSheet()
    el = [Paragraph(paras[0], st["Title"])] + [Paragraph(p, st["Normal"]) for p in paras[1:4]] + [Spacer(1, 8)]
    if table:
        tb = Table(table, colWidths=[120, 330])
        tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, (0, 0, 0)), ("FONTSIZE", (0, 0), (-1, -1), 9)]))
        el.append(tb)
        el.append(Spacer(1, 8))
    el += [Paragraph(p, st["Normal"]) for p in paras[4:]]
    SimpleDocTemplate(path, pagesize=A4).build(el)


def team_vo_logs(out, rng, scans=True):
    """The QS team's own variation-log layout, filled with the real rows of six projects (app/est_vo_seed.json,
    read from the team's files and checked against their totals) - Excel, typed PDF and a scan."""
    seed = json.load(open(os.path.join(os.path.dirname(__file__), "..", "..", "app", "est_vo_seed.json")))
    out_c = []
    for code, x in seed.items():
        rows = x["rows"]
        if len(rows) < 3:
            continue
        cols = ["ref", "title", "submitted_amount", "revised_amount", "approved_amount", "absorbed_amount", "submitted", "approved_on", "pc_ref", "remarks"]
        names = ["Ref.No", "Description", "Submitted Amount", "Revised Amount", "Accepted Amount", "Absorbed Amount", "Date of Submission",
                 "Date of Acceptance", "Claimed in PC", "Remarks"]
        pg = Page([14, 60, 22, 22, 22, 22, 22, 22, 16, 30], "Variation Log")
        pg.row(T(0, f"{code} - VARIATION LOG", span=10, align="c"), bold=True, h=8)
        pg.row(*[T(i, n, align="c") for i, n in enumerate(names)], bold=True, border=True, gap=1)
        truth = []
        seen = set()
        for r in rows:
            key = (r["ref"], r["title"])
            if key in seen or not r["ref"] or r["status"] not in ("Approved", "Submitted", "Rejected", "Absorbed", "Identified") or \
                    r["status"] == "Absorbed" and not r["absorbed_amount"]:
                continue
            seen.add(key)
            cells = []
            for i, c in enumerate(cols):
                v = r.get(c)
                if c.endswith("amount"):
                    cells.append(N(i, abs(v), neg=v < 0) if v else T(i, "-", align="r"))
                elif c in ("submitted", "approved_on"):
                    cells.append(D(i, date.fromisoformat(v)) if v else T(i, "-", align="c"))
                elif c == "remarks":
                    cells.append(T(i, (r.get("notes") or "").split("\n")[-1][:40]))
                else:
                    cells.append(T(i, v or ""))
            pg.row(*cells, border=True)
            truth.append({k: r.get(k) for k in ("ref", "title", "submitted_amount", "revised_amount", "approved_amount", "absorbed_amount",
                                                "submitted", "approved_on", "status")})
        tot = [T(1, "TOTAL")] + [N(i, abs(sum(t[c] for t in truth)), neg=sum(t[c] for t in truth) < 0) for i, c in enumerate(cols) if c.endswith("amount")]
        pg.row(*tot, bold=True, border=True)
        st = {"cur": "", "neg": "()", "dec": True, "sep": True}
        base = f"teamvo_{code}"
        render_xlsx([pg], os.path.join(out, base + ".xlsx"), st)
        out_c.append({"file": base + ".xlsx", "kind": "vo", "format": "xlsx", "layout": "team", "truth": truth, "real_data": True})
        render_pdf([pg], os.path.join(out, base + ".pdf"), st, font_size=7)
        out_c.append({"file": base + ".pdf", "kind": "vo", "format": "pdf", "layout": "team", "truth": truth, "real_data": True})
        if scans:
            scan_pdf(os.path.join(out, base + ".pdf"), os.path.join(out, base + "_scan.pdf"), rng, dpi=300)
            out_c.append({"file": base + "_scan.pdf", "kind": "vo", "format": "scan", "layout": "team", "dpi": 300, "truth": truth, "real_data": True})
    return out_c


# ---- the corpus -------------------------------------------------------------------------------------------

def dumps(o):
    return json.dumps(o, default=lambda x: x.isoformat() if hasattr(x, "isoformat") else str(x), indent=1)


def build(out, n=24, seed=7, scans=0.5, vo_n=16, contract_n=14):
    rng = random.Random(seed)
    os.makedirs(out, exist_ok=True)
    corpus = []
    for kind in ("ipa", "subpc"):
        for i in range(n):
            F = bill_figures(rng, sub=kind == "subpc")
            st = style(rng)
            F["_st"] = st
            lays = ["single", "three", "kv", "breakdown"] + (["atoe", "atoe"] if kind == "subpc" else [])
            lay = lays[i % len(lays)]
            arabic = rng.random() < 0.2
            pages, var_shown = BILL_LAYOUTS[lay](rng, F, kind, arabic)
            truth = truth_of(F, kind, var_shown)
            base = f"{kind}_{i:02d}_{lay}"
            # Excel: sometimes a hidden row / hidden sheet with an older revision's figures
            hidden_sheet = None
            if rng.random() < 0.3:
                pages[-1].row(T(0, lab(rng, "total", F)), N(len(pages[-1].widths) - 1, R2(F["total"] * 0.8)), hidden=True)
            if rng.random() < 0.3:
                hidden_sheet = [["Total Amount due for Payment", R2(F["total"] * 1.1)], ["Retention", R2(F["ret_c"] * 0.9)]]
            render_xlsx(pages, os.path.join(out, base + ".xlsx"), st, hidden_sheet)
            corpus.append({"file": base + ".xlsx", "kind": kind, "format": "xlsx", "layout": lay, "arabic": arabic, "truth": truth})
            pdfp = os.path.join(out, base + ".pdf")
            render_pdf(pages, pdfp, st, rules=rng.random() < 0.7, arabic=arabic)
            corpus.append({"file": base + ".pdf", "kind": kind, "format": "pdf", "layout": lay, "arabic": arabic, "truth": truth})
            if rng.random() < scans:
                dpi = rng.choice([150, 200, 200, 300])
                scan_pdf(pdfp, os.path.join(out, base + "_scan.pdf"), rng, dpi=dpi)
                corpus.append({"file": base + "_scan.pdf", "kind": kind, "format": "scan", "layout": lay, "dpi": dpi, "truth": truth})
            if kind == "subpc" and F.get("sub"):
                for c in corpus[-3:]:
                    if c["file"].startswith(base):
                        c["sub"] = F["sub"]
    for i in range(vo_n):
        rows = vo_rows(rng, rng.randrange(6, 30))
        st = style(rng)
        st["cur"] = ""
        st["neg"] = rng.choice(["-", "()"])          # an omission always shows its minus
        arabic = rng.random() < 0.2
        pages, cols = vo_layout(rng, rows, arabic)
        base = f"vo_{i:02d}"
        for r in rows:                      # what is not printed is not in the truth
            for k in ("revised_amount", "absorbed_amount"):
                if k not in cols:
                    r[k] = 0.0
        truth = [{k: (v.isoformat() if hasattr(v, "isoformat") else v) for k, v in r.items()} for r in rows]
        render_xlsx(pages, os.path.join(out, base + ".xlsx"), st)
        corpus.append({"file": base + ".xlsx", "kind": "vo", "format": "xlsx", "layout": "vo", "truth": truth})
        pdfp = os.path.join(out, base + ".pdf")
        render_pdf(pages, pdfp, st, font_size=7.5, arabic=arabic)
        corpus.append({"file": base + ".pdf", "kind": "vo", "format": "pdf", "layout": "vo", "truth": truth})
        if rng.random() < scans * 0.6:
            scan_pdf(pdfp, os.path.join(out, base + "_scan.pdf"), rng, dpi=300)
            corpus.append({"file": base + "_scan.pdf", "kind": "vo", "format": "scan", "layout": "vo", "dpi": 300, "truth": truth})
    corpus += team_vo_logs(out, rng, scans=scans > 0)
    for i in range(contract_n):
        truth, paras, table = contract_doc(rng, i)
        base = f"contract_{i:02d}"
        for f_ in ("pdf", "docx"):
            render_contract(paras, table, os.path.join(out, f"{base}.{f_}"), f_)
            corpus.append({"file": f"{base}.{f_}", "kind": "project", "format": f_, "layout": "contract", "truth": truth})
        if rng.random() < scans:
            scan_pdf(os.path.join(out, base + ".pdf"), os.path.join(out, base + "_scan.pdf"), rng, dpi=300)
            corpus.append({"file": base + "_scan.pdf", "kind": "project", "format": "scan", "layout": "contract", "dpi": 300, "truth": truth})
    with open(os.path.join(out, "truth.json"), "w") as fh:
        fh.write(dumps(corpus))
    return corpus


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("out")
    ap.add_argument("--n", type=int, default=24)
    ap.add_argument("--seed", type=int, default=7)
    ap.add_argument("--scans", type=float, default=0.5)
    ap.add_argument("--vo", type=int, default=16)
    ap.add_argument("--contracts", type=int, default=14)
    a = ap.parse_args()
    c = build(a.out, a.n, a.seed, a.scans, a.vo, a.contracts)
    print(len(c), "documents in", a.out)
