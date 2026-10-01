"""Accounts: invoices (tax / proforma) and the password-locked cash register.

    cd app && DATABASE_URL=sqlite:////tmp/acc.db python3 ../tests/accounts_test.py

Checks: who can reach what (admin, a login with the two rights, a login
with only invoices, a login with nothing); numbering; VAT arithmetic;
edit, cancel, proforma -> tax invoice; the PDF and the Excel; and the
register: nothing without a fresh key, wrong passwords lock, the key is
per user, exports need the key, the activity log carries no figures.
"""
import io, sys, os
sys.path.insert(0, os.getcwd())
from fastapi.testclient import TestClient
import main, models, auth, accounts
from database import SessionLocal

main._add_missing_columns()
c = TestClient(main.app)
fails = []
def ck(name, ok, extra=""):
    print(("PASS " if ok else "FAIL ") + name + ("" if ok else f"  -> {extra}"))
    if not ok: fails.append(name)

db = SessionLocal()
for un, pw, perms in [("acc_chief", "chief12345", "settings,accounts_invoices,accounts_register"),
                      ("acc_invonly", "inv12345", "settings,accounts_invoices"),
                      ("acc_site", "site12345", "attendance,settings")]:
    u = db.query(models.User).filter_by(username=un).first()
    if not u:
        u = models.User(username=un, full_name=un, role="office", hashed_password=auth.hash_password(pw), permissions=perms)
        db.add(u)
    else:
        u.permissions = perms; u.hashed_password = auth.hash_password(pw)
db.commit()
db.query(accounts.Invoice).delete(); db.query(accounts.CashEntry).delete(); db.commit()
for s in db.query(models.Setting).filter(models.Setting.key.like("register_%")).all():
    db.delete(s)
db.commit()

def tok(u, p):
    r = c.post("/auth/login", data={"username": u, "password": p}); assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["access_token"]}
A, CH, IO, SI = tok("admin", "changeme123"), tok("acc_chief", "chief12345"), tok("acc_invonly", "inv12345"), tok("acc_site", "site12345")
def dl(h):
    return c.post("/auth/download-token", headers=h).json()["token"]

# ---- rights -------------------------------------------------------------
ck("site login refused invoices", c.get("/employees/accounts/invoices", headers=SI).status_code == 403)
ck("site login refused register unlock", c.post("/employees/accounts/register/unlock", json={"password": "Infinia2022!"}, headers=SI).status_code == 403)
ck("invoices-only login sees invoices", c.get("/employees/accounts/invoices", headers=IO).status_code == 200)
ck("invoices-only login refused register", c.post("/employees/accounts/register/unlock", json={"password": "Infinia2022!"}, headers=IO).status_code == 403)
ck("chief sees invoices", c.get("/employees/accounts/invoices", headers=CH).status_code == 200)
me = c.get("/auth/me", headers=CH).json()
ck("accounts rights in ALL_SCREENS", {"accounts_invoices", "accounts_register"} <= set(main.ALL_SCREENS))
ck("no role has accounts by default", not any("accounts_invoices" in v for k, v in main.ROLE_DEFAULTS.items() if k != "admin"))

# ---- invoices ----------------------------------------------------------
n = c.get("/employees/accounts/invoices/next?kind=tax&project_no=906", headers=CH).json()["number"]
yy = f"{main._dubai_today():%y}"
ck("first number IC/yy/906/01", n == f"IC/{yy}/906/01", n)
lines = [{"description": "Phase 2 works - 30% progress", "amount": "56,339.00", "vat": 5},
         {"description": "Variation - extra boundary wall", "amount": 11138, "vat": 0},
         {"description": "", "amount": ""}]
body = {"kind": "tax", "number": n, "date": "2026-09-30", "client": "Mr. Mohammed Ahmed", "client_trn": "100000000000003",
        "client_address": "Dubai, UAE", "project": "(B+G+1+R) Villa", "project_no": "906", "plot": "6457380",
        "location": "Al Barsha South", "work": "Phase 2 Works", "lines": lines}
r = c.post("/employees/accounts/invoices", json=body, headers=CH); ck("tax invoice saved", r.status_code == 200, r.text)
iid = r.json()["id"]
lst = c.get("/employees/accounts/invoices?kind=tax", headers=CH).json()
x = lst["rows"][0]
ck("blank line dropped", len(x["lines"]) == 2, x["lines"])
ck("VAT 5% only on the 5% line", x["vat"] == round(56339 * .05, 2), x["vat"])
ck("totals", x["subtotal"] == 67477.0 and x["total"] == round(67477 + 2816.95, 2), (x["subtotal"], x["total"]))
ck("next number follows", c.get("/employees/accounts/invoices/next?kind=tax&project_no=906", headers=CH).json()["number"] == f"IC/{yy}/906/02")
r = c.post("/employees/accounts/invoices", json=body, headers=CH); ck("duplicate number refused", r.status_code == 400 and "already used" in r.text, r.text)
bad = dict(body, number="X1", lines=[{"description": "a", "amount": "abc"}]); r = c.post("/employees/accounts/invoices", json=bad, headers=CH)
ck("bad amount refused", r.status_code == 400, r.text)
bad = dict(body, number="X1", lines=[{"description": "a", "amount": 10, "vat": 7}]); ck("VAT 7% refused", c.post("/employees/accounts/invoices", json=bad, headers=CH).status_code == 400)
bad = dict(body, number="X1", client=""); ck("no client refused", c.post("/employees/accounts/invoices", json=bad, headers=CH).status_code == 400)
bad = dict(body, number="X1", lines=[]); ck("no lines refused", c.post("/employees/accounts/invoices", json=bad, headers=CH).status_code == 400)
# custom pattern continues
r = c.post("/employees/accounts/invoices", json=dict(body, number=f"IC/{yy}/906/PHASE2-10"), headers=CH)
ck("custom number saved", r.status_code == 200, r.text)
ck("custom number continues", c.get("/employees/accounts/invoices/next?kind=tax&project_no=906", headers=CH).json()["number"] == f"IC/{yy}/906/PHASE2-11")
ck("no project: own series", c.get("/employees/accounts/invoices/next?kind=tax", headers=CH).json()["number"] == f"IC/{yy}/01")
# edit
r = c.put(f"/employees/accounts/invoices/{iid}", json=dict(body, lines=[{"description": "Phase 2", "amount": 1000, "vat": 5}]), headers=CH)
ck("edit saved", r.status_code == 200, r.text)
x = [r for r in c.get("/employees/accounts/invoices?kind=tax", headers=CH).json()["rows"] if r["id"] == iid][0]
ck("edit recalculates", x["total"] == 1050.0, x["total"])
# PDF
t = dl(CH)
r = c.get(f"/export/accounts/invoice/{iid}?token={t}")
ck("PDF opens", r.status_code == 200 and r.content[:4] == b"%PDF", r.status_code)
ts = dl(SI)
ck("PDF refused to site login", c.get(f"/export/accounts/invoice/{iid}?token={ts}").status_code == 403)
r = c.get(f"/export/accounts/invoices?kind=tax&token={t}")
ck("Excel list", r.status_code == 200 and r.content[:2] == b"PK", r.status_code)
r = c.get(f"/export/accounts/invoices?kind=tax&format=excel&token={t}"); ck("Export Excel", r.content[:2] == b"PK")
r = c.get(f"/export/accounts/invoices?kind=tax&format=view&token={t}")
ck("Preview list = PDF inline", r.content[:4] == b"%PDF" and "inline" in r.headers.get("content-disposition", ""), r.headers.get("content-disposition"))
r = c.get(f"/export/accounts/invoices?kind=tax&format=pdf&token={t}")
ck("Export PDF list = download", r.content[:4] == b"%PDF" and "attachment" in r.headers.get("content-disposition", ""))
ck("list PDF refused to site login", c.get(f"/export/accounts/invoices?kind=tax&format=pdf&token={ts}").status_code == 403)
# cancel
r = c.post(f"/employees/accounts/invoices/{iid}/cancel", json={"reason": "wrong client"}, headers=CH); ck("cancel", r.status_code == 200)
ck("cancelled cannot be edited", c.put(f"/employees/accounts/invoices/{iid}", json=body, headers=CH).status_code == 400)
lst = c.get("/employees/accounts/invoices?kind=tax", headers=CH).json()
ck("totals leave cancelled out", lst["totals"]["count"] == 1, lst["totals"])
ck("cancelled PDF still opens", c.get(f"/export/accounts/invoice/{iid}?token={t}").status_code == 200)
# proforma -> tax
pn = c.get("/employees/accounts/invoices/next?kind=proforma&project_no=920", headers=CH).json()["number"]
ck("proforma number PI/", pn == f"PI/{yy}/920/01", pn)
r = c.post("/employees/accounts/invoices", json=dict(body, kind="proforma", number=pn, project_no="920"), headers=CH)
pid = r.json()["id"]
ck("proforma saved, separate list", len(c.get("/employees/accounts/invoices?kind=proforma", headers=CH).json()["rows"]) == 1)
ck("proforma PDF", c.get(f"/export/accounts/invoice/{pid}?token={t}").content[:4] == b"%PDF")
r = c.post(f"/employees/accounts/invoices/{pid}/convert", headers=CH); ck("convert", r.status_code == 200, r.text)
ck("converted number in tax series", r.json()["number"] == f"IC/{yy}/920/01", r.json())
ck("convert twice refused", c.post(f"/employees/accounts/invoices/{pid}/convert", headers=CH).status_code == 400)
pr = c.get("/employees/accounts/invoices?kind=proforma", headers=CH).json()["rows"][0]
ck("proforma linked", pr["converted_to_id"] == r.json()["id"])
# signature: the one kept in Settings (purchase orders) is used
import export_web, shutil
had = export_web.signature_file()
if not had and os.path.exists("/tmp/claude-0/inv-sig.png"):
    shutil.copy("/tmp/claude-0/inv-sig.png", export_web.SIG_PATHS[0])
ck("invoices use the Settings signature", c.get("/employees/accounts/invoices", headers=CH).json()["signature"] == bool(export_web.signature_file()))
ck("PDF with signature", c.get(f"/export/accounts/invoice/{iid}?token={t}").content[:4] == b"%PDF")
ck("no separate invoice signature upload", c.post("/employees/accounts/invoice-signature", files={"file": ("a.png", b"x", "image/png")}, headers=CH).status_code in (404, 405))

# ---- cash register -----------------------------------------------------
ck("register refused without key", c.get("/employees/accounts/register", headers=CH).status_code == 423)
ck("register refused with junk key", c.get("/employees/accounts/register", headers={**CH, "X-Register-Key": "junk"}).status_code == 423)
r = c.post("/employees/accounts/register/unlock", json={"password": "wrong"}, headers=CH)
ck("wrong password 401", r.status_code == 401 and "4 more" in r.text, r.text)
r = c.post("/employees/accounts/register/unlock", json={"password": "Infinia2022!"}, headers=CH)
ck("right password gives key", r.status_code == 200 and r.json().get("key"), r.text)
K = {**CH, "X-Register-Key": r.json()["key"]}
ck("register opens with key", c.get("/employees/accounts/register", headers=K).status_code == 200)
# key of chief does not work for admin
ck("key is per user", c.get("/employees/accounts/register", headers={**A, "X-Register-Key": r.json()["key"]}).status_code == 423)
# a normal login token is not a register key
ck("login token is not a key", c.get("/employees/accounts/register", headers={**CH, "X-Register-Key": CH["Authorization"][7:]}).status_code == 423)
for e in [("received", "2026-09-01", "Opening cash", 50000, ""), ("received", "2026-09-10", "Mr. Ahmed - 906", 20000, ""),
          ("paid", "2026-09-05", "Salaries advance", 12000.5, "Ramesh"), ("paid", "2026-09-20", "Cement", 3000, "")]:
    r = c.post("/employees/accounts/register", json=dict(zip(("kind", "date", "description", "amount", "remarks"), e)), headers=K)
    ck(f"entry {e[2]}", r.status_code == 200, r.text)
ck("entry without key refused", c.post("/employees/accounts/register", json={"kind": "paid", "date": "2026-09-01", "description": "x", "amount": 1}, headers=CH).status_code == 423)
ck("zero amount refused", c.post("/employees/accounts/register", json={"kind": "paid", "date": "2026-09-01", "description": "x", "amount": 0}, headers=K).status_code == 400)
g = c.get("/employees/accounts/register", headers=K).json()
ck("register totals", g["total_received"] == 70000 and g["total_paid"] == 15000.5 and g["balance"] == 54999.5, g)
ck("as on = last date", g["as_on"] == "2026-09-20", g["as_on"])
g2 = c.get("/employees/accounts/register?start=2026-09-06&end=2026-09-30", headers=K).json()
ck("period: brought forward", g2["opening"] == 50000 - 12000.5 and g2["balance"] == 54999.5, (g2["opening"], g2["balance"]))
ck("period: rows inside only", len(g2["received"]) == 1 and len(g2["paid"]) == 1)
eid = g["paid"][1]["id"]
ck("edit entry", c.put(f"/employees/accounts/register/{eid}", json={"kind": "paid", "date": "2026-09-20", "description": "Cement 50 bags", "amount": 3500}, headers=K).status_code == 200)
ck("delete entry", c.delete(f"/employees/accounts/register/{eid}", headers=K).status_code == 200)
ck("balance after delete", c.get("/employees/accounts/register", headers=K).json()["balance"] == 57999.5)
# exports
t = dl(CH)
r = c.get(f"/export/accounts/register?token={t}&rk={K['X-Register-Key']}&format=pdf"); ck("register PDF", r.content[:4] == b"%PDF", r.status_code)
r = c.get(f"/export/accounts/register?token={t}&rk={K['X-Register-Key']}&format=excel"); ck("register Excel", r.content[:2] == b"PK", r.status_code)
if r.content[:2] == b"PK":
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.content)).active
    vals = [c2.value for row in ws.iter_rows() for c2 in row if c2.value is not None]
    ck("Excel has balance formula", any(isinstance(v, str) and v.startswith("=C") for v in vals), vals[-5:])
ck("register export without key refused", c.get(f"/export/accounts/register?token={t}&rk=bad&format=pdf").status_code == 423)
ck("register export refused to invoices-only", c.get(f"/export/accounts/register?token={dl(IO)}&rk={K['X-Register-Key']}&format=pdf").status_code == 403)
# activity log: no figures
logs = [l.details or "" for l in db.query(models.AuditLog).filter(models.AuditLog.action.like("register%")).all()] if True else []
ck("activity log has no amounts", logs and not any(ch in " ".join(logs) for ch in ("50000", "20000", "Cement", "Ahmed")), logs[:5])
# password change: admin only
ka = c.post("/employees/accounts/register/unlock", json={"password": "Infinia2022!"}, headers=A).json()["key"]
KA = {**A, "X-Register-Key": ka}
ck("chief cannot change password", c.post("/employees/accounts/register/password", json={"current": "Infinia2022!", "new": "Another2026!"}, headers=K).status_code == 403)
ck("admin: wrong current refused", c.post("/employees/accounts/register/password", json={"current": "x", "new": "Another2026!"}, headers=KA).status_code == 400)
ck("admin: short refused", c.post("/employees/accounts/register/password", json={"current": "Infinia2022!", "new": "short"}, headers=KA).status_code == 400)
ck("admin changes password", c.post("/employees/accounts/register/password", json={"current": "Infinia2022!", "new": "Another2026!"}, headers=KA).status_code == 200)
ck("old password no longer opens", c.post("/employees/accounts/register/unlock", json={"password": "Infinia2022!"}, headers=CH).status_code == 401)
ck("new password opens", c.post("/employees/accounts/register/unlock", json={"password": "Another2026!"}, headers=CH).status_code == 200)
c.post("/employees/accounts/register/password", json={"current": "Another2026!", "new": "Infinia2022!"}, headers=KA)
# lockout
codes = [c.post("/employees/accounts/register/unlock", json={"password": "nope"}, headers=CH).status_code for _ in range(5)]
r = c.post("/employees/accounts/register/unlock", json={"password": "Infinia2022!"}, headers=CH)
ck("5 wrong -> locked even with right password", codes[-1] == 401 and r.status_code == 429 and "Try again" in r.text, (codes, r.status_code, r.text))
ck("lockout is per user (admin still opens)", c.post("/employees/accounts/register/unlock", json={"password": "Infinia2022!"}, headers=A).status_code == 200)
main.put_setting(db, f"register_tries_{db.query(models.User).filter_by(username='acc_chief').first().id}", "0|0")
ck("password not stored in plain text", "Infinia2022!" not in open(accounts.__file__).read())

print(f"\n{len(fails)} failure(s)" if fails else "\nALL PASS")
sys.exit(1 if fails else 0)
