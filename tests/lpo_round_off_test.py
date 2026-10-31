"""LPO round off (31 Oct 2026) - like Zoho's "Adjustment": a few fils on the total so the LPO equals the
supplier's invoice. Live cases: IC/LPO/20260130 13,660.50 vs invoice 13,661.00 (+0.50); IC/LPO/20260144
4,649.40 vs 4,649.00 (-0.40). Total = net after discount + VAT + round off; no VAT on the round off.

    cp /tmp/claude-0/full.db /tmp/ro.db; cd app && DATABASE_URL=sqlite:////tmp/ro.db INFINIA_NO_SCHED=1 INFINIA_NO_PUSH=1 \
        python3 ../tests/lpo_round_off_test.py

First, on the copy of the full database as it is: every LPO's total worked out from the raw rows (the old rule,
net + VAT) before the app starts; after the start-up migration adds purchase_orders.round_off, every one of them
must come out identical everywhere (register, report, Zoho list, printed page, PDF) and no row may be touched.
Zoho is never called: its calls are replaced here.
"""
import io, os, re, sqlite3, subprocess, sys, tempfile
from decimal import Decimal, ROUND_HALF_UP

sys.path.insert(0, '.')
DBFILE = os.environ["DATABASE_URL"].split("sqlite:///", 1)[1]
FAIL = []


def ck(l, ok, x=''):
    print(('PASS ' if ok else 'FAIL ') + l + ('' if ok else f'  [{str(x)[:500]}]'))
    if not ok:
        FAIL.append(l)


def r2(v):
    return Decimal(str(v or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


# ---- before: the totals on file, from the raw rows, and every row as stored ----------------------------------
con = sqlite3.connect(DBFILE); con.row_factory = sqlite3.Row
had_col = any(r[1] == "round_off" for r in con.execute("pragma table_info(purchase_orders)"))
before_rows = {r["id"]: dict(r) for r in con.execute("select * from purchase_orders")}
lines_raw = {}
for l in con.execute("select * from purchase_order_lines"):
    lines_raw.setdefault(l["order_id"], []).append(dict(l))
con.close()


def old_total(o):
    keep = 1 - (o["discount_pct"] or 0) / 100.0
    ls = lines_raw.get(o["id"], [])
    net = sum((l["qty"] or 0) * (l["rate"] or 0) for l in ls) * keep
    vat = sum((l["qty"] or 0) * (l["rate"] or 0) * keep * (l["tax_pct"] or 0) / 100.0 for l in ls)
    return round(net + vat, 2)


before_tot = {i: old_total(o) for i, o in before_rows.items()}

from fastapi.testclient import TestClient  # noqa: E402
import database, models, main, pozoho, projexpense, zoholog, export_web  # noqa: E402
import zohoapi as Z  # noqa: E402

c = TestClient(main.app, raise_server_exceptions=False); c.__enter__()     # start-up: the migration
H = {'Authorization': 'Bearer ' + c.post('/auth/login', data={'username': 'admin', 'password': 'changeme123'}).json()['access_token']}
DL = lambda: c.post('/auth/download-token', headers=H).json()['token']

# ---- the migration: additive, nothing rewritten -------------------------------------------------------------
con = sqlite3.connect(DBFILE); con.row_factory = sqlite3.Row
cols = [r[1] for r in con.execute("pragma table_info(purchase_orders)")]
ck("start-up adds purchase_orders.round_off" + (" (it was already there)" if had_col else ""), "round_off" in cols, cols)
after_rows = {r["id"]: dict(r) for r in con.execute("select * from purchase_orders")}
con.close()
# (columns other start-up steps add to an older copy are left out: only what was there before is compared)
changed = [(i, [k for k in o if after_rows[i].get(k) != o[k]]) for i, o in before_rows.items()
           if any(after_rows[i].get(k) != o[k] for k in o)]
ck(f"no existing LPO row rewritten by the migration ({len(before_rows)} LPOs; updated_at untouched too)", not changed, changed[:5])
ck("existing LPOs: round off 0 (or empty)", all((after_rows[i].get("round_off") or 0) == 0 for i in before_rows),
   [(i, after_rows[i].get("round_off")) for i in before_rows if after_rows[i].get("round_off")][:5])

# ---- every existing total identical everywhere --------------------------------------------------------------
reg = {r["id"]: r for r in c.get('/store/purchase/orders?limit=5000', headers=H).json()}
bad = [(before_rows[i]["ref"], t, reg.get(i, {}).get("total")) for i, t in before_tot.items() if reg.get(i, {}).get("total") != t]
ck(f"register: all {len(before_tot)} LPO totals identical before / after", not bad, bad[:5])
rep = c.get('/store/purchase/report?group_by=order&status=all', headers=H).json()
rt = {r["id"]: r["total"] for r in rep["rows"]}
bad = [(before_rows[i]["ref"], t, rt.get(i)) for i, t in before_tot.items() if abs((rt.get(i) or 0) - t) > 0.005]
ck("purchase report: every order's total identical", not bad, bad[:5])
ck("purchase report: no Round off column when no order has one", "round_off" not in [m["key"] for m in rep["measures"]], rep["measures"])
db = database.SessionLocal()
bad = []
for i, t in before_tot.items():
    o = db.get(models.PurchaseOrder, i)
    if main.lpo_total(o) != t or pozoho._total(o) != t or main._lpo_snap(o)["total"] != t:
        bad.append((o.ref, t, main.lpo_total(o), pozoho._total(o), main._lpo_snap(o)["total"]))
    if f"AED {t:,.2f}" not in main._lpo_html(o) or "Round off" in main._lpo_html(o):
        bad.append((o.ref, "printed page"))
ck("lpo_total, the Zoho list total, the edit snapshot and the printed page: all identical, no Round off line", not bad, bad[:5])
o0 = db.get(models.PurchaseOrder, next(iter(before_tot))) if before_tot else None
if o0:
    pdf = export_web.build_lpo_pdf(main._lpo_for_print(o0)).getvalue()
    ck(f"PDF of {o0.ref}: no Round off line", b"Round off" not in pdf)
db.close()


# ---- helpers for new LPOs ------------------------------------------------------------------------------------
def body(lines, ro=None, **k):
    b = {'supplier_name': 'Shafa Al Nahdah Ready Mixe Concrete Manufacturing LLC', 'order_date': '2026-10-30',
         'project_location': k.pop('site', '913'),
         'lines': [{'description': d, 'qty': q, 'unit': u, 'rate': r, 'tax_pct': 5, 'cost_to': t} for d, q, u, r, t in lines]}
    if ro is not None:
        b['round_off'] = ro
    b.update(k)
    return b


L130 = [('CONCRETE C40/20 400KG(64%OPC+36%GGBS) SUPER STRUCTURE', 34, 'M3', 353, ''), ('PUMP 42MTR', 1, 'NOS', 600, ''),
        ('TEMPERATURE CONTROL 32*', 34, 'M3', 12, '')]
L144 = [('CONCRETE C40/20 400KG(34%OPC+66%GGBS)SUPER STRUCTURE', 11, 'M3', 348, ''), ('PUMP 42MTR', 1, 'NOS', 600, '')]

# ---- create ----------------------------------------------------------------------------------------------------
r = c.post('/store/purchase/orders', json=body(L130, 0.5), headers=H)
ck('LPO like IC/LPO/20260130 raised with round off +0.50', r.status_code == 200, r.text)
a = r.json()
ck('saved: round off 0.50, total 13,661.00 (13,010.00 + VAT 650.50 + 0.50)', a['round_off'] == 0.5 and a['total'] == 13661.0, (a['round_off'], a['total']))
r = c.post('/store/purchase/orders', json=body(L144, -0.4), headers=H)
b = r.json()
ck('LPO like IC/LPO/20260144 raised with round off -0.40: total 4,649.00', r.status_code == 200 and b['total'] == 4649.0, r.text[:300])
r = c.post('/store/purchase/orders', json=body(L144), headers=H)
n0 = r.json()
ck('round off left out (blank): 0, total 4,649.40', n0['round_off'] == 0 and n0['total'] == 4649.4, (n0['round_off'], n0['total']))

# ---- limits ----------------------------------------------------------------------------------------------------
for v, ok_ in ((0.99, True), (-0.99, True), (0, True), (1, False), (-1, False), (1.5, False), (0.995, False), (0.505, False),
               (-12, False), ("abc", False), ("0.25", True)):
    r = c.post('/store/purchase/orders', json=body(L144, v), headers=H)
    good = r.status_code == 200
    ck(f"round off {v!r}: {'accepted' if ok_ else 'refused with a clear message'}",
       good == ok_ and (ok_ or (r.status_code in (400, 422) and ('Round off' in r.text or 'round_off' in r.text))), (r.status_code, r.text[:200]))
    if good:
        db = database.SessionLocal(); x = db.get(models.PurchaseOrder, r.json()['id'])
        db.delete(x); db.commit(); db.close()
r = c.post('/store/purchase/orders', json=body(L144, 1.0), headers=H)
ck('the message says the limit', 'less than 1.00' in r.text, r.text)

# ---- register, report, printed page, PDF, Excel ------------------------------------------------------------
reg = {x['id']: x for x in c.get('/store/purchase/orders', headers=H).json()}
ck('register: 13,661.00 and 4,649.00', reg[a['id']]['total'] == 13661.0 and reg[b['id']]['total'] == 4649.0,
   (reg[a['id']]['total'], reg[b['id']]['total']))
rep = c.get('/store/purchase/report?group_by=order&date_from=2026-10-30&date_to=2026-10-30', headers=H).json()
row = {x['id']: x for x in rep['rows']}
keys = [m['key'] for m in rep['measures']]
ck('report: a Round off column, just before Total', keys[-2:] == ['round_off', 'total'], keys)
ck('report rows: 13,010.00 + 650.50 + 0.50 = 13,661.00 and 4,428.00 + 221.40 - 0.40 = 4,649.00',
   (row[a['id']]['sub_total'], row[a['id']]['vat'], row[a['id']]['round_off'], row[a['id']]['total']) == (13010.0, 650.5, 0.5, 13661.0)
   and (row[b['id']]['round_off'], row[b['id']]['total']) == (-0.4, 4649.0), (row[a['id']], row[b['id']]))
ck('report totals: sub total + VAT + round off = total', abs(rep['totals']['sub_total'] + rep['totals']['vat'] + rep['totals']['round_off'] - rep['totals']['total']) < 0.005, rep['totals'])
rm = c.get('/store/purchase/report?group_by=material&date_from=2026-10-30&date_to=2026-10-30', headers=H).json()
ck('report by material: the round offs share out to the fils and add up to the orders\' (0.50 - 0.40 = 0.10)',
   abs(sum(x['round_off'] for x in rm['rows']) - 0.10) < 1e-9 and abs(rm['totals']['total'] - rep['totals']['total']) < 0.005, rm['totals'])
rs = c.get('/store/purchase/report?group_by=supplier&date_from=2026-10-30&date_to=2026-10-30', headers=H).json()
ck('report by supplier: same total', abs(rs['totals']['total'] - rep['totals']['total']) < 0.005, rs['totals'])
xl = c.get(f'/export/store/purchase-report?token={DL()}&format=excel&group_by=order&date_from=2026-10-30&date_to=2026-10-30')
from openpyxl import load_workbook  # noqa: E402
ws = load_workbook(io.BytesIO(xl.content)).active
hdr = [str(x.value) for r_ in ws.iter_rows() for x in r_ if x.value is not None]
ck('report Excel: a Round off (AED) column', 'Round off (AED)' in hdr, (xl.status_code, hdr[:30]))

db = database.SessionLocal()
oa, ob = db.get(models.PurchaseOrder, a['id']), db.get(models.PurchaseOrder, b['id'])
ha = main._lpo_html(oa)
i_vat, i_ro, i_tot = ha.find('Standard Rate'), ha.find('<td>Round off</td>'), ha.find('<td>Total</td>')
ck('printed page: a Round off line 0.50 between VAT and Total', 0 < i_vat < i_ro < i_tot and '<td class="r">0.50</td>' in ha and 'AED 13,661.00' in ha,
   (i_vat, i_ro, i_tot))
ck('printed page: VAT 650.50 (none on the round off)', '<td class="r">650.50</td>' in ha)
hb = main._lpo_html(ob)
ck('printed page -0.40: Round off -0.40, AED 4,649.00', '<td class="r">-0.40</td>' in hb and 'AED 4,649.00' in hb)
ck('printed page without a round off: no Round off line', 'Round off' not in main._lpo_html(db.get(models.PurchaseOrder, n0['id'])))
v = c.get(f"/export/purchase/{a['id']}/view?token={DL()}")
ck('the preview page (the one opened in the app) shows it too', 'Round off' in v.text and 'AED 13,661.00' in v.text)
with tempfile.TemporaryDirectory() as td:
    for o_, want_ro, want_tot in ((oa, '0.50', '13,661.00'), (ob, '-0.40', '4,649.00')):
        pdf = c.get(f"/export/purchase/{o_.id}?token={DL()}&format=pdf")
        fp = os.path.join(td, 'x.pdf'); open(fp, 'wb').write(pdf.content)
        txt = subprocess.run(['pdftotext', '-layout', fp, '-'], capture_output=True, text=True).stdout
        flat = ' '.join(txt.split())     # the terms print beside the money box, so the lines interleave
        i1, i2, i3 = flat.find('Standard Rate (5%)'), flat.find(f'Round off {want_ro}'), flat.find(f'Total AED {want_tot}')
        ck(f'PDF {o_.ref}: Round off {want_ro} after VAT, Total AED {want_tot}', 0 <= i1 < i2 < i3,
           (i1, i2, i3, flat[flat.find('Sub Total'):flat.find('Sub Total') + 400]))
    pdf = c.get(f"/export/purchase/{n0['id']}?token={DL()}&format=pdf")
    fp = os.path.join(td, 'n.pdf'); open(fp, 'wb').write(pdf.content)
    txt = subprocess.run(['pdftotext', '-layout', fp, '-'], capture_output=True, text=True).stdout
    ck('PDF without a round off: no Round off line, AED 4,649.40', 'Round off' not in txt and 'AED 4,649.40' in txt)
    # Excel: the Round off row is a value and the Total formula takes it in (recalculated by LibreOffice)
    xr = c.get(f"/export/purchase/{a['id']}?token={DL()}&format=excel")
    fx = os.path.join(td, 'a.xlsx'); open(fx, 'wb').write(xr.content)
    ws = load_workbook(fx).active
    cells = {(cl.row, cl.column): cl.value for r_ in ws.iter_rows() for cl in r_ if cl.value is not None}
    ro_row = next((rr for (rr, cc), vv in cells.items() if vv == 'Round off'), None)
    tot_row = next((rr for (rr, cc), vv in cells.items() if vv == 'Total' and cc == 5), None)
    ck('Excel: a Round off row (0.50) just above Total', ro_row and tot_row == ro_row + 1 and cells.get((ro_row, 7)) == 0.5, (ro_row, tot_row))
    f = str(cells.get((tot_row, 7)) or '')
    m = re.match(r'=SUM\(G(\d+):G(\d+)\)', f)
    ck('Excel: the Total formula adds the Round off row in', m and int(m.group(1)) < ro_row <= int(m.group(2)), f)
    try:
        subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx', '--outdir', os.path.join(td, 'calc'), fx],
                       capture_output=True, timeout=120)
        wv = load_workbook(os.path.join(td, 'calc', 'a.xlsx'), data_only=True).active
        tv = wv.cell(tot_row, 7).value
        ck(f'Excel recalculated: Total {tv} = 13,661.00', tv is not None and abs(float(tv) - 13661.0) < 0.005, tv)
    except Exception as e:
        print('NOTE Excel not recalculated (LibreOffice):', e)

# ---- Zoho: sent as the PO's adjustment (mocked - nothing reaches Zoho) ------------------------------------------
CALLS = []
Z.vat_tax = lambda db_: 'V5'; Z.zero_tax = lambda db_: 'V0'; Z.account_id = lambda db_, n: 'ACC'; Z.connected = lambda db_: True
pozoho._vendor = lambda db_, o_: 'VEN'; pozoho._vendor_treatment = lambda db_, v: ''


def fake_call(db_, method, path, params=None, body=None):
    CALLS.append((method, path, body))
    if method == 'GET' and path == 'purchaseorders':
        return {'purchaseorders': []}
    if method == 'POST' and path == 'purchaseorders':
        return {'purchaseorder': {'purchaseorder_id': 'ZPO1', 'total': 13661.0}}
    return {'purchaseorder': {}}


Z.call = fake_call
for o_ in (oa, ob, db.get(models.PurchaseOrder, n0['id'])):
    o_.zoho_account = 'Concrete Materials'
db.commit()
zb = pozoho._body(db, oa)
ck('Zoho PO body: adjustment 0.5, description "Round off"', zb.get('adjustment') == 0.5 and zb.get('adjustment_description') == 'Round off', zb)
ck('Zoho PO body: the lines carry no round off (VAT only on the lines: 13,010.00 at 5%)',
   sum(l['quantity'] * l['rate'] for l in zb['line_items']) == 13010 and all(l['tax_id'] == 'V5' for l in zb['line_items']))
ck('Zoho PO body for -0.40: adjustment -0.4', pozoho._body(db, ob).get('adjustment') == -0.4)
ck('Zoho PO body with no round off: no adjustment on a first send', 'adjustment' not in pozoho._body(db, db.get(models.PurchaseOrder, n0['id'])))
zl = {x['id']: x for x in c.get('/store/zoho/pos', headers=H).json()['rows']}
ck('Send to Zoho list: 13,661.00 and 4,649.00 (what Zoho will show)', zl[a['id']]['total'] == 13661.0 and zl[b['id']]['total'] == 4649.0,
   (zl.get(a['id'], {}).get('total'), zl.get(b['id'], {}).get('total')))
db.close()
r = c.post('/store/zoho/pos/send', json={'ids': [a['id']]}, headers=H)
posted = [x for x in CALLS if x[0] == 'POST' and x[1] == 'purchaseorders']
ck('sent (mocked): the POST carries adjustment 0.5 "Round off"', r.json().get('sent') == 1 and posted and posted[-1][2].get('adjustment') == 0.5
   and posted[-1][2].get('adjustment_description') == 'Round off', (r.text, posted[-1:] and posted[-1][2].get('adjustment')))

# ---- edit (an issued LPO, already in Zoho) ------------------------------------------------------------------------
CALLS.clear()
r = c.put(f"/store/purchase/orders/{a['id']}", json=body(L130, 0.3), headers=H)
ck('edit: round off 0.50 -> 0.30, total 13,660.80', r.status_code == 200 and r.json()['total'] == 13660.8, r.text[:300])
puts = [x for x in CALLS if x[0] == 'PUT']
ck('edit of an LPO in Zoho: the PO in Zoho updated with adjustment 0.3 (mocked)', puts and puts[-1][2].get('adjustment') == 0.3, puts[-1:])
db = database.SessionLocal()
log = db.query(models.AuditLog).filter(models.AuditLog.action == 'edit_lpo').order_by(models.AuditLog.id.desc()).first()
ck('activity log notes the round off change and the new total', log and 'Round off: +0.50 -> +0.30' in log.details and '13,660.80' in log.details,
   log and log.details)
nt = pozoho.note_of(db.get(models.PurchaseOrder, a['id']))
ck('accounts are told too', any('Round off: +0.50 -> +0.30' in x for x in nt['items'][-1]['changes']), nt)
db.close()
bd = body(L130); bd.pop('round_off', None)
r = c.put(f"/store/purchase/orders/{a['id']}", json=bd, headers=H)
ck('an edit from an older screen (round off not sent) keeps it: 0.30', r.json()['round_off'] == 0.3, r.json()['round_off'])
CALLS.clear()
r = c.put(f"/store/purchase/orders/{a['id']}", json=body(L130, 0), headers=H)
puts = [x for x in CALLS if x[0] == 'PUT']
ck('round off taken off (0): total back to 13,660.50, Zoho adjustment set to 0', r.json()['total'] == 13660.5 and puts and puts[-1][2].get('adjustment') == 0,
   (r.json()['total'], puts[-1:]))
db = database.SessionLocal()
log = db.query(models.AuditLog).filter(models.AuditLog.action == 'edit_lpo').order_by(models.AuditLog.id.desc()).first()
ck('activity log: Round off: +0.30 -> none', 'Round off: +0.30 -> none' in log.details, log.details)
ck('printed page: no Round off line any more', 'Round off' not in main._lpo_html(db.get(models.PurchaseOrder, a['id'])))
r = c.put(f"/store/purchase/orders/{a['id']}", json=body(L130, 2), headers=H)
ck('an edit with round off 2 is refused, nothing changed', r.status_code == 400 and db.get(models.PurchaseOrder, a['id']).round_off == 0, r.text)
c.put(f"/store/purchase/orders/{a['id']}", json=body(L130, 0.5), headers=H)
db.close()

# ---- Zoho check: the PO in Zoho compared with the LPO (round off included) ----------------------------------------
db = database.SessionLocal()
zlog = models.ZohoLog(source='lpo', item_id=a['id'], ref=a['ref'], action='create', kind='purchaseorder', zoho_id='ZPO1', ok=True, amount=13661.0)
db.add(zlog); db.commit(); lid = zlog.id; db.close()
zoholog._live = lambda db_, kind, zid: {'amount': 13661.0, 'from': '', 'to': '', 'project': '', 'supplier': '', 'number': '', 'status': ''}
r = c.get(f'/store/zoho/sent/{lid}/check', headers=H).json()
ck('Zoho check: PO 13,661.00 in Zoho matches the LPO with its round off - no difference', r.get('diffs') == [], r)
zoholog._live = lambda db_, kind, zid: {'amount': 13660.5, 'from': '', 'to': '', 'project': '', 'supplier': '', 'number': '', 'status': ''}
r = c.get(f'/store/zoho/sent/{lid}/check', headers=H).json()
ck('Zoho check: a PO in Zoho without the adjustment (13,660.50) is shown as different', any('13,660.50' in d and '13,661.00' in d for d in r.get('diffs', [])), r)

# ---- combined LPO: the round off shared over the sites, parts adding up exactly ------------------------------------
LC = [('SAFETY COVERALL', 20, 'pcs', 28, 'Store'), ('WIRE NAIL 2.5"', 25, 'kg', 6.5, '913'), ('SOFT BRUSH 12"', 6, 'pcs', 14, '910'),
      ('WHITE SILICONE 300ML', 24, 'pcs', 9.5, '910'), ('MMA WELDING MACHINE', 1, 'nos', 1450, 'STORE')]
r = c.post('/store/purchase/orders', json=body(LC, 0.99, combined=True, site='STORE', order_date='2026-10-28'), headers=H)
cb = r.json()
ck('combined LPO with round off 0.99: total 2,609.71 / .72 (2,484.50 + VAT 124.225 + 0.99)', r.status_code == 200 and abs(cb['total'] - 2609.715) < 0.006, r.text[:300])
db = database.SessionLocal(); oc = db.get(models.PurchaseOrder, cb['id'])
sp = main.lpo_split(oc, with_ro=True)
ck('split of 0.99 over Store 2,110.50 / 910 327.60 / 913 170.62: 0.80 + 0.12 + 0.06 in proportion, Store (largest) takes the odd fils: 0.81 / 0.12 / 0.06',
   [(k, round(n + v, 2), p) for k, n, v, p in sp] == [('STORE', 2110.5, 0.81), ('910', 327.6, 0.12), ('913', 170.62, 0.06)], sp)
ck('the parts add up to the round off exactly', round(sum(p for *_, p in sp), 2) == 0.99)
for ro_, ws_ in ((0.99, [1, 1, 1]), (-0.99, [3, 2, 1]), (0.01, [5, 5]), (-0.5, [0, 0]), (0.33, [1000, 1, 1])):
    pr = main.split_round_off(ro_, ws_)
    ck(f'split_round_off({ro_}, {ws_}) = {pr}: whole fils, adds up exactly', round(sum(pr), 2) == ro_ and all(abs(x * 100 - round(x * 100)) < 1e-9 for x in pr))
zl = {x['id']: x for x in c.get('/store/zoho/pos', headers=H).json()['rows']}
ck('Send to Zoho list: the split adds up to the total', zl[cb['id']]['split'] and abs(sum(s['amount'] for s in zl[cb['id']]['split']) - cb['total']) < 0.005, zl.get(cb['id']))
projexpense._CACHE.clear()
g = projexpense.gather(db, '2026-10-28', '2026-10-28')
mine = {code: [x for x in ls if x.get('cat') == 'lpo' and x.get('id') == cb['id']] for code, ls in g['lines'].items()}
a910 = sum(x['amount'] for x in mine.get('910', [])); a913 = sum(x['amount'] for x in mine.get('913', []))
ck('Project expense 910: 312.00 + its 0.12 of the round off = 312.12', abs(a910 - 312.12) < 0.001, a910)
ck('Project expense 913: 162.50 + 0.06 = 162.56', abs(a913 - 162.56) < 0.001, a913)
ck('the store part (stock, not a project cost): 2,010.00 + 0.81', abs(g['unlinked']['lpo_store'] - 2010.81) < 0.001, g['unlinked']['lpo_store'])
# an ordinary LPO: Project expense LPO cost = net + its round off (VAT left out as before)
g = projexpense.gather(db, '2026-10-30', '2026-10-30')
pa = sum(x['amount'] for x in g['lines'].get('913', []) if x.get('cat') == 'lpo' and x.get('id') == a['id'])
pb = sum(x['amount'] for x in g['lines'].get('913', []) if x.get('cat') == 'lpo' and x.get('id') == b['id'])
ck('Project expense 913: LPO cost net of VAT incl. round off - 13,010.50 and 4,427.60', abs(pa - 13010.5) < 0.001 and abs(pb - 4427.6) < 0.001, (pa, pb))
# the store's valuation of a material stays at the line rate after discount - a round off is not an item cost
book = main._price_book(db, {i.id: i for i in db.query(models.StoreItem).all()})
lines_a = [l for l in db.get(models.PurchaseOrder, a['id']).lines if l.item_id]
ck('store valuation per item: the LPO line rate (353.00), not touched by the round off',
   lines_a and book(lines_a[0].item_id, None) == 353.0, lines_a and book(lines_a[0].item_id, None))
db.close()

print(f"{len(FAIL)} FAILED" if FAIL else 'ALL PASS')
sys.exit(1 if FAIL else 0)
