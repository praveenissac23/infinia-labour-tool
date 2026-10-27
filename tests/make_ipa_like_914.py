"""A workbook laid out like 914's real IPA (the consultant's format): a cover with the bill no., date and total due,
a summary that collects each sheet's total, APP C (approved variations: Total % = Previous + Current, spare V02-V05
lines, a Sub-Total) and a Division sheet (Ref / description / unit / QTY x RATE = AMOUNT, Previous / Current / Total %,
a heading row, a spec paragraph, TOTAL CARRIED TO COLLECTION over the rows including blanks).
    python3 tests/make_ipa_like_914.py /tmp/claude-0/bills/914-like.xlsx"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side

out = sys.argv[1] if len(sys.argv) > 1 else "/tmp/claude-0/bills/914-like.xlsx"
wb = Workbook()
thin = Side(style="thin"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
B = Font(bold=True)

cv = wb.active; cv.title = "IPA"
cv["A1"] = "INTERIM PAYMENT APPLICATION"; cv["A1"].font = B
cv["A3"] = "Bill No."; cv["B3"] = 7
cv["A4"] = "Date"; cv["B4"] = "15.09.2026"
cv["A6"] = "Work done to date"; cv["B6"] = "=Summary!C6"
cv["A7"] = "Less previous"; cv["B7"] = 1000000
cv["A8"] = "Total amount due"; cv["B8"] = "=B6-B7"

sm = wb.create_sheet("Summary")
sm["A1"] = "SUMMARY"; sm["A1"].font = B
sm["A3"], sm["B3"], sm["C3"] = "Ref", "Section", "To date"
sm["A4"], sm["B4"], sm["C4"] = "C", "Approved variations", "='APP C'!I11"
sm["A5"], sm["B5"], sm["C5"] = "03", "Concrete", "='Division 03-Concrete Works'!M20"
sm["B6"], sm["C6"] = "Total", "=SUM(C4:C5)"

c = wb.create_sheet("APP C")
c["A1"] = "APPENDIX - C"; c["A2"] = "APPROVED VARIATIONS"
for col, t in zip("ABCDGHI", ["Ref.No", "Description", "Total Amount", "Percentage of Work Done", "Total Work Done", "", ""]):
    if t: c[f"{col}3"] = t
for col, t in zip("DEFGHI", ["Previous %", "Current %", "Total %", "Previous", "Current", "Total"]):
    c[f"{col}4"] = t
c["A5"], c["B5"], c["C5"], c["D5"], c["E5"], c["F5"] = "V01", "Design Revision for Lift Pit", "=52826.96/2", 1, "=F5-D5", 1
c["G5"], c["H5"], c["I5"] = "=D5*C5", "=E5*C5", "=C5*F5"
for r in range(6, 10):
    c[f"A{r}"] = f"V0{r - 4}"; c[f"D{r}"] = 0; c[f"E{r}"] = 0; c[f"F{r}"] = f"=SUM(D{r}:E{r})"
    c[f"G{r}"], c[f"H{r}"], c[f"I{r}"] = f"=D{r}*C{r}", f"=E{r}*C{r}", f"=C{r}*F{r}"
c["C10"], c["D10"], c["E10"], c["F10"], c["G10"], c["H10"], c["I10"] = "=0", 0, "=F10-D10", 0, "=C10*D10", 0, "=C10*F10"
c["B11"] = "Sub - Total"
c["C11"], c["D11"], c["E11"], c["F11"] = "=SUM(C5:C10)", "=G11/C11", "=H11/C11", "=I11/$C$11"
c["G11"], c["H11"], c["I11"] = "=SUM(G5:G10)", "=SUM(H5:H10)", "=SUM(I5:I10)"
for r in range(3, 12):
    for col in "ABCDEFGHI":
        c[f"{col}{r}"].border = box

d = wb.create_sheet("Division 03-Concrete Works")
for col, t in zip("ABCDEFGHK", ["Ref.No.CSI", "Ref.No.BOQ", "WORK DESCRIPTION", "UNIT", "QTY", "RATE (AED)", "AMOUNT (AED)", "Percentage of Work Done", "Total Work Done"]):
    d[f"{col}4"] = t
for col, t in zip("HIJKLM", ["Previous\n%", "Current\n%", "Total\n%", "Previous", "Current", "Total"]):
    d[f"{col}5"] = t
d["B7"], d["C7"] = "3.1.1", "Concrete (PCC/Blinding C20)"; d["C7"].font = B
d["C8"] = "Execution of all works beneath the foundations of the villa as per drawings and specification."
rows = [(9, "3.1.1.1", "for villa", "M3", 18.97, 492, 1, 1), (10, "3.1.1.2", "Boundary wall", "M3", 3.17, "=+F9", 0.9, 0.9),
        (12, "3.2.2.1", "Villa Parapet & capping Beam", "M3", 9, 2428, 0, 0)]
d["B11"], d["C11"] = "3.2.2", "Superstructure"; d["C11"].font = B
for r, ref, desc, u, q, rate, pv, tv in rows:
    d[f"B{r}"], d[f"C{r}"], d[f"D{r}"], d[f"E{r}"], d[f"F{r}"] = ref, desc, u, q, rate
    d[f"G{r}"] = f"=+F{r}*E{r}"; d[f"H{r}"] = pv; d[f"I{r}"] = f"=J{r}-H{r}"; d[f"J{r}"] = tv
    d[f"K{r}"], d[f"L{r}"], d[f"M{r}"] = f"=H{r}*G{r}", f"=I{r}*G{r}", f"=J{r}*G{r}"
    d[f"N{r}"] = "=30.02+0.745"                       # a side note outside the table
d["C20"], d["E20"] = "TOTAL CARRIED TO COLLECTION", "DIV03/4"
d["G20"], d["K20"], d["L20"], d["M20"] = "=SUM(G9:G19)", "=SUM(K9:K19)", "=SUM(L9:L19)", "=SUM(M9:M19)"
d["H20"], d["J20"] = "=K20/G20", "=M20/G20"
for r in range(4, 21):
    for col in "ABCDEFGHIJKLM":
        d[f"{col}{r}"].border = box
wb.save(out)
print(out)
