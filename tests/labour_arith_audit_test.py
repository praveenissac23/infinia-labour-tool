"""Labour payroll arithmetic, worked out again by hand and compared everywhere the figure appears (31 Oct 2026).

    cd app && DATABASE_URL=sqlite:////tmp/la.db INFINIA_NO_SCHED=1 INFINIA_NO_PUSH=1 python3 ../tests/labour_arith_audit_test.py

On an empty database it puts in its own awkward workers (AUD-xx); on a copy of the live database
(cp full.db /tmp/la.db first) it also audits every real card of every cycle on file.

The figures are worked out here with exact fractions, from the raw attendance and the rules - nothing is
imported from the app's own calculation:
  a day is 1/30 of the monthly rate (the rate in force at the cycle end, from the rate history);
  Present, Sick, Medical, Friday, Sunday, Holiday are paid days; a half day is half;
  daily pay: Payable = paid days x day, Deduction = absent days x day (an absent day is not earned AND costs a day);
  fixed pay: Payable = min((paid + absent + leave) x day, rate), Deduction = min((absent + leave) x day, Payable);
  OT and BH = hours x day / 8; Final = Payable + OT + BH - Deduction, rounded to the dirham, half away from zero;
  a day after the leaving date counts as Terminated; additions / deductions are added to Final afterwards.
Then it checks the same worker's figure is the same on: the stored card (DB), the salary cards screen
(/summaries), the live card, the Paid status card (Labour payroll), the report table (rows, Excel),
the salary card Excel / PDF / preview, the report builder, the site cost report, and a payment marked
paid with its payslip - and that every total equals the sum of the rows it is the total of.
"""
import sys, io, os, re, json
sys.path.insert(0, '.')
from datetime import date, timedelta, datetime
from fractions import Fraction as F
from collections import defaultdict
from fastapi.testclient import TestClient
import database, models, auth
models.Base.metadata.create_all(database.engine)
import main, services
import payroll_cycle as pcyc

FAIL = []
CHECKS = defaultdict(int)


def ck(label, ok, info=''):
    CHECKS[label.split(':')[0]] += 1
    if not ok:
        FAIL.append(label)
        print(f'FAIL {label}  [{str(info)[:400]}]')


def say(label):
    print('PASS ' + label)


# ---- the independent working ----------------------------------------------------------------------------------------
PAID = ('present', 'sick', 'medical', 'friday', 'sunday', 'holiday')


def bounds(label):
    return pcyc.cycle_bounds_for(datetime.strptime('25 ' + label, '%d %B %Y').date())[:2]


def half_up(x):
    """Round to the whole dirham, half away from zero (Excel ROUND), on an exact fraction."""
    n = (abs(x) + F(1, 2)).__floor__()
    return n if x >= 0 else -n


def rate_for(db, e, ce):
    hist = (db.query(models.SalaryChange).filter(models.SalaryChange.employee_id == e.id,
                                                 models.SalaryChange.kind.in_(('rate', 'opening')))
              .order_by(models.SalaryChange.effective_on, models.SalaryChange.id).all())
    if e.staff or not hist:
        return F(str(e.total_salary or 0))
    upto = [h for h in hist if h.effective_on <= ce]
    h = (upto or hist[:1])[-1]
    return F(str(h.basic or 0)) + F(str(h.allowance or 0))


def expected(db, e, s):
    cs, ce = bounds(s.month_year)
    rows = db.query(models.DailyRow).filter(models.DailyRow.emp_no == e.emp_no, models.DailyRow.month_year == s.month_year).all()
    cnt = defaultdict(F)
    ot = bh = F(0)
    for r in rows:
        after = e.terminated_on and r.full_date > e.terminated_on
        for st in (r.am, r.pm):
            cnt['terminated' if after else (st or '').strip().lower()] += F(1, 2)
        if not after:
            ot += F(str(r.ot or 0)); bh += F(str(r.bh or 0))
    rate = rate_for(db, e, ce)
    day = rate / 30
    paid = sum(cnt[k] for k in PAID)
    if (e.pay_type or 'daily') == 'fixed':
        off = cnt['absent'] + cnt['leave']
        comp = min((paid + off) * day, rate)
        ded = min(off * day, comp)
    else:
        comp = paid * day
        ded = cnt['absent'] * day
    ota, bha = day / 8 * ot, day / 8 * bh
    allow, oded = F(str(s.allowances or 0)), F(str(s.other_deduction or 0))
    final = half_up(comp + ota + bha + allow - ded - oded)
    adj = sum((-F(str(a.amount)) if a.is_deduction else F(str(a.amount))) for a in s.adjustments)
    return {'rate': rate, 'present': cnt['present'], 'absent': cnt['absent'], 'sick': cnt['sick'], 'medical': cnt['medical'],
            'sunday': cnt['sunday'], 'holiday': cnt['holiday'], 'leave': cnt['leave'], 'terminated': cnt['terminated'],
            'paid_days': paid, 'ot_h': ot, 'bh_h': bh, 'comp': comp, 'ded': ded, 'ota': ota, 'bha': bha,
            'final': final, 'adj': adj, 'net': final + adj}


def fils(x):
    """An exact amount to the fils, half away from zero."""
    return float(half_up(F(x) * 100) / 100)


def near(a, b, tol=0.005):
    return abs(float(a) - float(b)) <= tol


def money(t):
    if isinstance(t, (int, float)):
        return float(t)
    return float(str(t).replace('AED', '').replace(',', '').replace('+', '').replace(' ', '').strip() or 0)


# ---- awkward workers of our own -------------------------------------------------------------------------------------
c = TestClient(main.app)
c.__enter__()                     # the app's start-up: brings an older database's tables up to date, as the server does
db = database.SessionLocal()
if not db.query(models.User).filter(models.User.username == 'audadmin').first():
    db.add(models.User(username='audadmin', hashed_password=auth.hash_password('pw123456'), full_name='Audit', role='admin'))
    db.commit()
T = main._dubai_today()
cur_s, cur_e, CUR = pcyc.cycle_bounds_for(T)
p_s, p_e, PREV = pcyc.cycle_bounds_for(cur_s - timedelta(days=1))
# the most recent closed cycle that runs 31 days - where a fixed salary must still stop at the month's figure
d = p_s - timedelta(days=1)
while True:
    a, b, L31 = pcyc.cycle_bounds_for(d)
    if (b - a).days + 1 == 31:
        break
    d = a - timedelta(days=1)


def worker(no, name, rate, pay_type='daily', joined=None, left=None, basic=None):
    e = db.query(models.Employee).filter(models.Employee.emp_no == no).first()
    if e:
        return e
    e = models.Employee(emp_no=no, name=name, trade='HELPER', total_salary=rate, basic_salary=basic or round(rate * 2 / 3, 2),
                        active=left is None, company='Infinia', pay_type=pay_type, staff=False, joined_on=joined,
                        terminated_on=left, pay_route='wps')
    db.add(e); db.flush()
    return e


def mark(e, label, pattern, ot=None, bh=None, start=None, end=None):
    """pattern(day_index, date) -> (am, pm); ot/bh(day_index, date) -> hours."""
    cs, ce = bounds(label)
    d, i = max(cs, start or cs), 0
    while d <= min(ce, end or ce, T - timedelta(days=1)):
        am, pm = pattern(i, d)
        if am or pm:
            db.add(models.DailyRow(employee_id=e.id, emp_no=e.emp_no, emp_name=e.name, trade=e.trade, month_year=label,
                                   full_date=d, day=d.day, am=am, pm=pm, ot=(ot(i, d) if ot else 0), bh=(bh(i, d) if bh else 0),
                                   site=('911' if i % 3 else '704'), engineer='ENG', comments=''))
        d += timedelta(days=1); i += 1
    db.flush()


def normal(i, d):
    return ('Sunday', 'Sunday') if d.weekday() == 6 else ('Present', 'Present')


SYN = {}
if not db.query(models.Employee).filter(models.Employee.emp_no == 'AUD-01').first():
    # 1. daily 1,400, a bit of everything, OT that lands exactly on a half dirham
    e = worker('AUD-01', 'AUDIT MIXED', 1400)
    pat = lambda i, d: (('Absent', 'Absent') if i == 3 else ('Present', 'Absent') if i == 5 else ('Sick', 'Sick') if i == 8
                        else ('Medical', 'Medical') if i == 9 else ('Holiday', 'Holiday') if i == 12 else ('Leave', 'Leave') if i == 14
                        else normal(i, d))
    for lab in (PREV, CUR, L31):
        mark(e, lab, pat, ot=lambda i, d: 3 if i % 4 == 0 and d.weekday() != 6 and i not in (3, 8, 9, 12, 14) else 0,
             bh=lambda i, d: 1 if i == 20 else 0)
    # 2. fixed 3,100 in a 31-day cycle: full month = 3,100, one absence costs exactly 1/30
    e = worker('AUD-02', 'AUDIT FIXED', 3100, 'fixed', basic=2000)
    mark(e, L31, normal)
    e = worker('AUD-03', 'AUDIT FIXED ABSENT', 3100, 'fixed', basic=2000)
    mark(e, L31, lambda i, d: ('Absent', 'Absent') if i == 10 else normal(i, d))
    mark(e, PREV, lambda i, d: ('Leave', 'Leave') if i in (4, 5) else ('Present', 'Absent') if i == 6 else normal(i, d))
    # 3. a leaver: days after his leaving date still read Present on the sheet - they must count as Terminated
    e = worker('AUD-04', 'AUDIT LEAVER', 1500, left=p_s + timedelta(days=11))
    mark(e, PREV, normal, ot=lambda i, d: 2.5 if i in (1, 2) else 0)
    # ... and one whose days after leaving were saved as Terminated, as the attendance screen does
    lt = p_s + timedelta(days=9)
    e = worker('AUD-08', 'AUDIT LEAVER SAVED', 1400, left=lt)
    mark(e, PREV, lambda i, d: ('Terminated', 'Terminated') if d > lt else normal(i, d))
    # 4. a joiner part way through the cycle
    e = worker('AUD-05', 'AUDIT JOINER', 1650, joined=p_s + timedelta(days=17))
    mark(e, PREV, normal, start=p_s + timedelta(days=17), ot=lambda i, d: 1.5 if d.weekday() != 6 else 0)
    # 5. a rate change in the middle of the cycle: the cycle is paid at the rate in force at its end
    e = worker('AUD-06', 'AUDIT RATE CHANGE', 1800)
    db.add(models.SalaryChange(employee_id=e.id, effective_on=L31 and bounds(L31)[0] - timedelta(days=60), kind='opening', basic=1000, allowance=600))
    db.add(models.SalaryChange(employee_id=e.id, effective_on=p_s + timedelta(days=10), kind='rate', basic=1100, allowance=700))
    mark(e, PREV, normal, ot=lambda i, d: 0.5 if i % 2 else 0)
    mark(e, L31, normal, ot=lambda i, d: 1 if i % 5 == 0 else 0)
    # 6. mostly absent: the card goes below zero and rounds away from zero
    e = worker('AUD-07', 'AUDIT NEGATIVE', 1300)
    mark(e, PREV, lambda i, d: ('Present', 'Present') if i in (0,) else ('Present', 'Absent') if i == 1 else ('Absent', 'Absent'),
         end=p_s + timedelta(days=6), ot=lambda i, d: 1.5 if i == 0 else 0)
    db.commit()
    # additions and deductions, with fils, through the app
    c0 = c
    H0 = {'Authorization': 'Bearer ' + c0.post('/auth/login', data={'username': 'audadmin', 'password': 'pw123456'}).json()['access_token']}
    for e in db.query(models.Employee).filter(models.Employee.emp_no.like('AUD-%')).all():
        for lab in (PREV, CUR, L31):
            services.recalculate_summary(db, e, lab)
    for no, lab, desc, amt, ded in (('AUD-01', PREV, 'Site bonus', 300, False), ('AUD-01', PREV, 'Parking fine', 170.35, True),
                                    ('AUD-03', PREV, 'Food', 0.1, False), ('AUD-03', PREV, 'Food 2', 0.2, False),
                                    ('AUD-06', PREV, 'Advance', 250.5, True), ('AUD-01', L31, 'Tools', 45.25, True)):
        sm = db.query(models.EmployeeSummary).filter_by(emp_no=no, month_year=lab).first()
        r = c0.post(f'/summaries/{sm.id}/adjustments', headers=H0, json={'description': desc, 'amount': amt, 'is_deduction': ded})
        ck(f'adjustment saved {no} {desc}', r.status_code == 200, r.text)
    db.expire_all()

# ---- 1. every card: stored figures against the independent working -------------------------------------------------
H = {'Authorization': 'Bearer ' + c.post('/auth/login', data={'username': 'audadmin', 'password': 'pw123456'}).json()['access_token']}
def dl():
    return c.post('/auth/download-token', headers=H).json()['token']

emps = {e.emp_no: e for e in db.query(models.Employee).all()}
cycles = sorted({s.month_year for s in db.query(models.EmployeeSummary.month_year).all()}, key=lambda l: bounds(l)[0])
early = {s.key for s in db.query(models.Setting).filter(models.Setting.key.like('labour_full_cycle:%')).all() if s.value}
EXP = {}
n_cards = 0
for lab in cycles:
    # the screen first (it clears a leaver's empty card the same way the app does when the page is opened)
    scr = c.get(f'/summaries/{lab}', headers=H)
    ck(f'salary cards screen opens: {lab}', scr.status_code == 200, scr.text[:200])
    scr = {x['emp_no']: x for x in scr.json()}
    db.expire_all()
    for s in db.query(models.EmployeeSummary).filter(models.EmployeeSummary.month_year == lab).all():
        e = emps.get(s.emp_no)
        if e is None or e.staff:
            continue
        if f'labour_full_cycle:{lab}' in early or f'labour_full_cycle:{lab}:{s.emp_no}' in early:
            print(f'NOTE {s.emp_no} {lab}: paid early for the whole cycle - not worked out again here (pay_full_early_test covers it)')
            continue
        x = expected(db, e, s)
        EXP[(s.emp_no, lab)] = x
        n_cards += 1
        who = f'{s.emp_no} {lab}'
        ck(f'rate on the card: {who}', near(s.total_salary, x['rate']), (s.total_salary, float(x['rate'])))
        for fld, k in (('present_days', 'present'), ('absent_days', 'absent'), ('sick_days', 'sick'), ('medical_days', 'medical'),
                       ('sunday_days', 'sunday'), ('holiday_days', 'holiday'), ('leave_days', 'leave'),
                       ('terminated_days', 'terminated'), ('ot_hours', 'ot_h'), ('bh_hours', 'bh_h')):
            ck(f'day / hour count: {who} {fld}', near(getattr(s, fld) or 0, x[k]), (getattr(s, fld), float(x[k])))
        for fld, k in (('total_salary_component', 'comp'), ('deduction', 'ded'), ('ot_amount', 'ota'), ('bh_amount', 'bha')):
            ck(f'money line: {who} {fld}', near(getattr(s, fld) or 0, x[k]), (getattr(s, fld), float(x[k])))
        ck(f'final salary: {who}', s.final_salary == x['final'], (s.final_salary, x['final'], float(x['comp'] + x['ota'] + x['bha'] - x['ded'])))
        ck(f'final + additions / deductions: {who}', near(services.adjusted_final_salary(s), x['net']), (services.adjusted_final_salary(s), float(x['net'])))
        # the salary cards screen shows the stored card
        y = scr.get(s.emp_no)
        ck(f'screen = stored card: {who}', y is not None and y['final_salary'] == s.final_salary
           and near(sum((-a['amount'] if a['is_deduction'] else a['amount']) for a in y['adjustments']), x['adj']), (y and y['final_salary'], s.final_salary))
say(f'{n_cards} labour cards in {len(cycles)} cycles ({", ".join(cycles)}) worked out again from the attendance')

# leaver cards that pay nothing are taken off the screen - nothing with money on it may disappear
for lab in cycles:
    on_screen = {x['emp_no'] for x in c.get(f'/summaries/{lab}', headers=H).json()}
    for (no, l2), x in EXP.items():
        if l2 == lab and (x['net'] != 0 or x['paid_days']):
            ck(f'card with pay on it is on the screen: {no} {lab}', no in on_screen)

# ---- 2. the live card -----------------------------------------------------------------------------------------------
lc_n = 0
for (no, lab), x in EXP.items():
    if not (no.startswith('AUD-') or lc_n < 40):
        continue
    r = c.get(f'/live-card/{no}/{lab}', headers=H)
    if r.status_code != 200:
        ck(f'live card opens: {no} {lab}', False, r.text[:200]); continue
    j = r.json(); sm = j['summary'] or {}
    ck(f'live card = card: {no} {lab}', sm.get('final_salary') == x['final'], (sm.get('final_salary'), x['final']))
    lc_n += 1
say(f'live card checked for {lc_n} cards')

# ---- 3. Paid status card (Labour payroll > Salary cards) ------------------------------------------------------------
for lab in cycles:
    r = c.get('/employees/pay/labour', params={'month_year': lab}, headers=H)
    ck(f'paid status card opens: {lab}', r.status_code == 200, r.text[:200])
    if r.status_code != 200:
        continue
    for g in r.json()['groups']:
        ck(f'paid status company total = its rows: {lab} {g["company"]}', near(g['net'], sum(x['net'] for x in g['rows']), 0.01),
           (g['net'], sum(x['net'] for x in g['rows'])))
        for row in g['rows']:
            x = EXP.get((row['emp_no'], lab))
            if not x:
                continue
            ck(f'paid status card salary = card: {row["emp_no"]} {lab}', row['salary'] == x['final'], (row['salary'], x['final']))
            ck(f'paid status card net = card + adjustments: {row["emp_no"]} {lab}', near(row['net'], x['net']), (row['net'], float(x['net'])))
            ck(f'paid status card days = paid days: {row["emp_no"]} {lab}', near(row['days'], x['paid_days']), (row['days'], float(x['paid_days'])))
            if not row['pay']['payments']:
                ck(f'nothing paid -> due is the card: {row["emp_no"]} {lab}', near(row['pay']['due'], x['net']), (row['pay'], float(x['net'])))
say('paid status card rows and company totals checked')

# ---- 4. the report table (screen rows, Excel) -----------------------------------------------------------------------
import openpyxl
COLS = 'emp_no,emp_name,total_salary,present_days,absent_days,ot_hours,deduction,ot_amount,bh_amount,final_salary,adjusted_final_salary'
for lab in cycles:
    tk = dl()
    r = c.get(f'/export/{lab}/report-table', params={'token': tk, 'columns': COLS, 'format': 'rows'})
    ck(f'report table rows: {lab}', r.status_code == 200, r.text[:200])
    rows = r.json()['rows']
    for row in rows:
        x = EXP.get((row['Emp No'], lab))
        if not x:
            continue
        ck(f'report row final: {row["Emp No"]} {lab}', row['Final Salary'] == x['final'], (row['Final Salary'], x['final']))
        ck(f'report row adjusted final: {row["Emp No"]} {lab}', near(row['Adjusted Final Salary'], x['net']), (row['Adjusted Final Salary'], float(x['net'])))
        ck(f'report row OT amount (to the fils): {row["Emp No"]} {lab}', near(row['OT Amount'], fils(x['ota']), 0.0001), (row['OT Amount'], float(x['ota'])))
    xr = c.get(f'/export/{lab}/report-table', params={'token': dl(), 'columns': COLS, 'format': 'excel'})
    ck(f'report table Excel: {lab}', xr.status_code == 200, xr.text[:200])
    ws = openpyxl.load_workbook(io.BytesIO(xr.content), data_only=True).active
    grid = [[c_.value for c_ in r_] for r_ in ws.iter_rows()]
    head_i = next(i for i, r_ in enumerate(grid) if 'Emp No' in r_)
    head = grid[head_i]
    body = [r_ for r_ in grid[head_i + 1:] if r_[head.index('Emp No')] not in (None, '')]
    tot = next((r_ for r_ in body if str(r_[0]).strip().upper().startswith('TOTAL')), None)
    data = [r_ for r_ in body if r_ is not tot and str(r_[head.index('Emp No')]).strip().upper() not in ('TOTAL', '')]
    ck(f'report Excel has every row: {lab}', len(data) == len(rows), (len(data), len(rows)))
    fmt = {}
    for r_ in ws.iter_rows(min_row=head_i + 2, max_row=head_i + 2):
        fmt = {head[k]: c_.number_format for k, c_ in enumerate(r_)}
    for col in ('Final Salary', 'Adjusted Final Salary', 'OT Amount', 'Absence/Leave Deduction', 'BH Amount'):
        ci = head.index(col)
        # what a person sees in the cell: the value at its number format (two decimals for money)
        dp = 2 if '0.00' in (fmt.get(col) or '') or isinstance(data[0][ci] if data else '', str) else None
        ck(f'report Excel money column shows two decimals: {lab} {col}', dp == 2, fmt.get(col))
        shown = [round(money(r_[ci]) + 1e-9, 2) for r_ in data]
        want = [round(float(rw[col]) + 1e-9, 2) for rw in rows]
        ck(f'report Excel column = screen rows: {lab} {col}', all(near(a, b) for a, b in zip(shown, want)), list(zip(shown, want))[:5])
        if tot is not None and tot[ci] not in (None, ''):
            ck(f'report Excel total = sum of the rows it shows: {lab} {col}', near(money(tot[ci]), round(sum(shown), 2), 0.005),
               (tot[ci], round(sum(shown), 2)))
say('report table rows and Excel checked, totals = sum of shown rows')

# ---- 5. salary card Excel / PDF / preview: the final figure on every card -------------------------------------------
from pypdf import PdfReader
for lab in cycles:
    want = {no: x for (no, l2), x in EXP.items() if l2 == lab}
    xr = c.get(f'/export/{lab}/excel', params={'token': dl()})
    ck(f'salary cards Excel: {lab}', xr.status_code == 200, xr.text[:200])
    wb = openpyxl.load_workbook(io.BytesIO(xr.content), data_only=True)
    seen = {}
    for ws in wb.worksheets:
        cells = {(c_.row, c_.column): c_.value for r_ in ws.iter_rows() for c_ in r_ if c_.value not in (None, '')}
        nos = sorted((r_, v) for (r_, col), v in cells.items() if col == 4 and cells.get((r_, 3)) == 'Employee No:')
        finals = sorted(r_ for (r_, col), v in cells.items() if v == 'FINAL SALARY TO PROCESS')
        lines = defaultdict(dict)
        for (r_, col), v in cells.items():
            if col == 4 and v in ('Payable Salary', 'Absence or Leave', 'OT Amount', 'BH Amount'):
                lines[r_][v] = cells.get((r_, 5))
        for (rn, no) in nos:
            fr = next((f for f in finals if f > rn), None)
            seen[str(no)] = (money(cells.get((fr + 1, 7))) if fr else None,
                             {v: money(t) for r_, d_ in lines.items() if fr and fr <= r_ < fr + 6 for v, t in d_.items()})
    for no, x in want.items():
        got = seen.get(no)
        ck(f'card Excel final = card: {no} {lab}', got is not None and near(got[0], x['net']), (got and got[0], float(x['net'])))
        if got:
            ck(f'card Excel money lines = card: {no} {lab}',
               near(got[1].get('Payable Salary', 0), round(float(x['comp']), 2)) and near(got[1].get('OT Amount', 0), round(float(x['ota']), 2)),
               (got[1], float(x['comp']), float(x['ota'])))
    pr = c.get(f'/export/{lab}/pdf', params={'token': dl()})
    ck(f'salary cards PDF: {lab}', pr.status_code == 200, pr.text[:200])
    pdf_finals = {}
    pages = PdfReader(io.BytesIO(pr.content)).pages
    heads = sum(1 for pg in pages if 'Employee No:' in (pg.extract_text() or ''))
    ck(f'salary cards PDF: one page per card (a leaver\'s Terminated days used to push his card onto two): {lab}',
       len(pages) == heads, (len(pages), heads))
    for pg in pages:
        t = pg.extract_text() or ''
        m_no = re.search(r'Employee No:\s*([A-Z0-9-]+)', t)
        m_f = re.search(r'FINAL SALARY TO\s+PROCESS\s*AED\s*(-?[\d,]+\.\d\d)', t)
        if m_no and m_f:
            pdf_finals[m_no.group(1)] = money(m_f.group(1))
    for no, x in want.items():
        ck(f'card PDF final = card: {no} {lab}', no in pdf_finals and near(pdf_finals[no], x['net']), (pdf_finals.get(no), float(x['net'])))
    vr = c.get(f'/export/{lab}/cards/view', params={'token': dl()})
    ck(f'salary cards preview: {lab}', vr.status_code == 200, vr.text[:200])
    html = vr.text
    view = {}
    for blk in html.split('<div class="card">')[1:]:
        m_no = re.search(r'<th>Employee No:</th><td>([^<]+)</td>', blk)
        m_f = re.search(r'FINAL SALARY TO PROCESS</th></tr>\s*<tr><td[^>]*>AED (-?[\d,]+\.\d\d)', blk)
        if m_no and m_f:
            view[m_no.group(1)] = money(m_f.group(1))
            # the days shown in the grid add up to the day totals on the same card
            sts = re.findall(r'<td(?: style="[^"]*")?>(Present|Absent|Sick|Medical|Sunday|Holiday|Leave|Friday|Terminated)</td>', blk.split('class="days"')[0])
            x = want.get(m_no.group(1))
            if x:
                ck(f'card preview grid days = card paid days: {m_no.group(1)} {lab}',
                   near(sum(F(1, 2) for s_ in sts if s_.lower() in PAID), x['paid_days']), (sum(F(1, 2) for s_ in sts if s_.lower() in PAID), float(x['paid_days'])))
    for no, x in want.items():
        ck(f'card preview final = card: {no} {lab}', no in view and near(view[no], x['net']), (view.get(no), float(x['net'])))
say('salary card Excel, PDF and preview: the same final on every card')

# ---- 6. report builder and site cost: totals add up -----------------------------------------------------------------
for lab in cycles:
    r = c.get('/reports/custom', params={'month_year': lab, 'data_source': 'summary', 'dimensions': 'emp_no',
                                         'measures': 'final_salary,adjusted_final_salary,ot_amount'}, headers=H)
    ck(f'report builder (cards): {lab}', r.status_code == 200, r.text[:200])
    j = r.json()
    rows = [x for x in j['rows'] if (x.get('emp_no'), lab) in EXP]
    for x in rows:
        e_ = EXP[(x['emp_no'], lab)]
        ck(f'report builder final: {x["emp_no"]} {lab}', near(x.get('final_salary', 0), e_['final']), (x.get('final_salary'), e_['final']))
        ck(f'report builder adjusted final: {x["emp_no"]} {lab}', near(x.get('adjusted_final_salary', 0), e_['net']), (x.get('adjusted_final_salary'), float(e_['net'])))
    if j.get('totals') and 'adjusted_final_salary' in j['totals']:
        ck(f'report builder total = sum of rows: {lab}', near(j['totals']['adjusted_final_salary'], sum(x.get('adjusted_final_salary', 0) for x in j['rows']), 0.01),
           (j['totals']['adjusted_final_salary'], sum(x.get('adjusted_final_salary', 0) for x in j['rows'])))
    # spread across sites by the day: every dirham of the cards lands on some site
    r = c.get('/reports/custom', params={'month_year': lab, 'data_source': 'daily', 'dimensions': 'site', 'measures': 'final_salary_cost,ot_amount'}, headers=H)
    ck(f'report builder (by site): {lab}', r.status_code == 200, r.text[:200])
    j = r.json()
    tot_cost = sum(float(x.get('final_salary_cost') or 0) for x in j['rows'])
    want_cost = sum(float(x['net']) for (no, l2), x in EXP.items() if l2 == lab)
    ck(f'cost by site adds up to the cards (final + adjustments): {lab}', near(tot_cost, want_cost, 0.05), (round(tot_cost, 2), round(want_cost, 2)))
    tot_ot = sum(float(x.get('ot_amount') or 0) for x in j['rows'])
    ck(f'OT by site adds up to the cards: {lab}', near(tot_ot, sum(float(x['ota']) for (no, l2), x in EXP.items() if l2 == lab), 0.05), tot_ot)
    r = c.get(f'/summaries/{lab}/by-site', headers=H)
    ck(f'site cost report: {lab}', r.status_code == 200, r.text[:200])
    j = r.json()
    ck(f'site cost total = sum of the sites: {lab}', near(j['totals'].get('total_cost', 0), sum(x['total_cost'] for x in j['rows']), 0.01),
       (j['totals'], sum(x['total_cost'] for x in j['rows'])))
say('report builder and site cost totals checked')

# ---- 7. marking paid: the amount, the payment record and the payslip agree with the card ----------------------------
sm = db.query(models.EmployeeSummary).filter_by(emp_no='AUD-01', month_year=PREV).first()
if sm:
    x = EXP[('AUD-01', PREV)]
    pv = c.post('/employees/pay/preview', headers=H, json={'kind': 'labour', 'source_id': sm.id, 'type': 'regular'}).json()
    ck('pay preview = card net', near(pv.get('amount', 0), x['net']) and near(pv.get('due', 0), x['net']), (pv, float(x['net'])))
    mk = c.post('/employees/pay/mark', headers=H, json={'kind': 'labour', 'source_id': sm.id, 'type': 'regular', 'mode': 'wps', 'payslip': True})
    ck('mark paid', mk.status_code == 200, mk.text[:300])
    if mk.status_code == 200:
        p = db.get(models.PayPayment, mk.json()['payment']['id'])
        ck('payment amount = card net', near(p.amount, x['net']), (p.amount, float(x['net'])))
        figs = json.loads(p.figures)
        ck('payslip net = payment', near(figs['net'], p.amount), (figs['net'], p.amount))
        e_sum = sum(v for _, v in figs['earnings']) - sum(v for _, v in figs['deductions'])
        ck('payslip lines add up to its net (to the rounding of the card)', abs(e_sum - figs['net']) <= 0.5 + 1e-6, (e_sum, figs['net'], figs['earnings'], figs['deductions']))
        row = next(rw for g in c.get('/employees/pay/labour', params={'month_year': PREV}, headers=H).json()['groups'] for rw in g['rows'] if rw['emp_no'] == 'AUD-01')
        ck('paid status card shows it paid in full, no balance', row['pay']['state'] in ('paid', 'todate') and near(row['pay']['balance'], 0), row['pay'])
        # a change after payment is carried to the next cycle, to the fils
        c.post(f'/summaries/{sm.id}/adjustments', headers=H, json={'description': 'Late bonus', 'amount': 12.5, 'is_deduction': False})
        db.expire_all()
        carry = db.query(models.PayCarry).filter(models.PayCarry.payment_id == p.id).first()
        if row['pay']['state'] == 'paid':
            ck('a change after payment is carried to next cycle (12.50)', carry is not None and near(carry.amount, 12.5), carry and carry.amount)
say('mark paid / payslip / carry checked')

# ---- 8. paying a running cycle early, for the whole cycle: preview = payment = card -----------------------------------
sm = db.query(models.EmployeeSummary).filter_by(emp_no='AUD-01', month_year=CUR).first()
if sm and T < bounds(CUR)[1]:
    pv = c.post('/employees/pay/preview', headers=H, json={'kind': 'labour', 'source_id': sm.id, 'type': 'early_full'}).json()
    e = emps['AUD-01']
    # by hand: the days still to come (from today) count as worked, a Sunday as Sunday
    cs, ce = bounds(CUR)
    rows = {r.full_date: r for r in db.query(models.DailyRow).filter_by(emp_no='AUD-01', month_year=CUR).all()}
    cnt, ot = defaultdict(F), F(0)
    d = cs
    while d <= ce:
        r = rows.get(d)
        if r is not None:
            for st in (r.am, r.pm):
                cnt[st.lower()] += F(1, 2)
            ot += F(str(r.ot or 0)) + F(str(r.bh or 0))
        elif d >= T:
            cnt['sunday' if d.weekday() == 6 else 'present'] += 1
        d += timedelta(days=1)
    day = rate_for(db, e, ce) / 30
    full = half_up(sum(cnt[k] for k in PAID) * day + day / 8 * ot - cnt['absent'] * day)
    adj = sum((-F(str(a.amount)) if a.is_deduction else F(str(a.amount))) for a in sm.adjustments)
    ck('early full cycle: preview = worked out by hand', near(pv.get('amount', 0), full + adj), (pv, full + adj))
    mk = c.post('/employees/pay/mark', headers=H, json={'kind': 'labour', 'source_id': sm.id, 'type': 'early_full', 'mode': 'cash'})
    ck('early full cycle: marked paid', mk.status_code == 200, mk.text[:300])
    if mk.status_code == 200:
        db.expire_all()
        sm = db.query(models.EmployeeSummary).filter_by(emp_no='AUD-01', month_year=CUR).first()
        ck('early full cycle: paid = preview = the card afterwards', near(mk.json()['payment']['amount'], pv.get('amount', 0))
           and near(services.adjusted_final_salary(sm), pv.get('amount', 0)), (mk.json()['payment']['amount'], pv.get('amount'), services.adjusted_final_salary(sm)))
    say('early full cycle checked')

db.close()
print(f'\nchecks run: {sum(CHECKS.values())} over {n_cards} cards')
print(f'{len(FAIL)} FAILED' if FAIL else 'LABOUR ARITHMETIC HOLDS UP - SCREEN, PREVIEW, PDF, EXCEL, DB AND PAYMENT AGREE')
sys.exit(1 if FAIL else 0)
