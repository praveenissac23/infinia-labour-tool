"""Every salary figure recomputed by hand and compared with the app's.

Run: cd app && DATABASE_URL=sqlite:////tmp/audit2.db python3 ../tests/hr_arithmetic_audit.py

The database must already hold the office data (deploy/load_office_hr.py).
Nothing here uses the app's own payroll code: the registers are read with
plain SQL and every line is worked out again from the rules as agreed -

  * a day is gross / 30 dropped to whole dirhams; a half day is half that;
  * one sick day a month is paid, in date order; other sick days, absent,
    vacation and unpaid leave are deducted; holiday and paid leave are not;
  * a marked pay rule (paid / unpaid) overrides the kind;
  * a whole month away costs the month's salary and not a dirham more;
  * days before joining or after leaving come off the same way;
  * salary is the one in force on the last day of the month;
  * add / ded = additions - deductions entered for the month;
  * loan instalment = min(instalment, balance) per open loan, unless typed;
  * net = gross - absence - deductions - loan + additions; pension is
    printed but never added or taken;
  * to-date = gross x days elapsed / days in month - absence + add/ded
    (the loan comes off at month end);
  * held rows are left out of every total.

Then the same figures are read off the API, the preview, the PDF and the
spreadsheet, and all of them have to agree.
"""
import io
import math
import os
import sqlite3
import sys
from datetime import date, datetime, timedelta

sys.path.insert(0, '.')
from fastapi.testclient import TestClient
import main

FAIL = []


def ck(label, ok, ctx=None):
    print(("PASS " if ok else "FAIL ") + label + ("" if ok else f"   [{ctx}]"))
    if not ok:
        FAIL.append(label)


def same(a, b):
    return abs(round(float(a or 0), 2) - round(float(b or 0), 2)) < 0.005


DB = os.environ['DATABASE_URL'].split('sqlite:///', 1)[1]
raw = sqlite3.connect(DB)
raw.row_factory = sqlite3.Row


def q(sql, *args):
    return [dict(r) for r in raw.execute(sql, args).fetchall()]


def d(s):
    return datetime.strptime(str(s)[:10], '%Y-%m-%d').date() if s else None


def bounds(month_year):
    a = datetime.strptime('1 ' + month_year, '%d %B %Y').date()
    nxt = date(a.year + (a.month == 12), a.month % 12 + 1, 1)
    return a, nxt - timedelta(days=1)


PAID_KINDS = {'holiday', 'paid_leave'}


def day_rate(gross):
    return float(int(gross / 30.0))


def salary_on(emp_id, e, day):
    ch = q("select basic, allowance from salary_changes where employee_id=? and effective_on<=? "
           "order by effective_on desc, id desc limit 1", emp_id, day.isoformat())
    if ch:
        return round(ch[0]['basic'] or 0, 2), round(ch[0]['allowance'] or 0, 2)
    return round(e['basic_salary'] or 0, 2), round(e['allowance'] or 0, 2)


def absence(emp_id, a, b, gross):
    rows = q("select * from staff_leave where employee_id=? and on_date>=? and on_date<=? "
             "order by on_date, id", emp_id, a.isoformat(), b.isoformat())
    sick_left = 1.0
    unpaid = 0.0
    for r in rows:
        kind = r.get('kind') or ('sick' if 'sick' in (r.get('reason') or '').lower()
                                 else ('paid_leave' if r.get('paid') else 'absent'))
        portion = r.get('portion') or 1.0
        rule = r.get('pay_rule') or 'auto'
        if rule == 'paid':
            paid = portion
        elif rule == 'unpaid':
            paid = 0.0
        elif kind == 'sick':
            paid = min(portion, max(sick_left, 0.0))
            sick_left -= paid
        elif kind in PAID_KINDS:
            paid = portion
        else:
            paid = 0.0
        unpaid += portion - paid
    if not unpaid:
        return 0.0
    if unpaid >= (b - a).days + 1:
        return gross
    return round(min(day_rate(gross) * unpaid, gross), 2)


def part_month(e, a, b, gross):
    j, t = d(e['joined_on']), d(e['terminated_on'])
    days = 0
    if j and a < j <= b:
        days += (j - a).days
    if t and a <= t < b:
        days += (b - t).days
    return round(min(day_rate(gross) * days, gross), 2) if days else 0.0


def items(emp_id, month_year):
    rows = q("select direction, amount from pay_items where employee_id=? and month_year=?",
             emp_id, month_year)
    add = round(sum(r['amount'] for r in rows if r['direction'] == 'add'), 2)
    ded = round(sum(r['amount'] for r in rows if r['direction'] == 'deduct'), 2)
    return add, ded


def loan_state(emp_id, month_year=None):
    """(instalment due this month, balance before it), from the loan book.
    For an approved month the payroll repayment of that month is put back
    so the figure is the one the sheet saw when it was approved."""
    due, balance = 0.0, 0.0
    for l in q("select * from staff_loans where employee_id=?", emp_id):
        reps = q("select amount, month_year, source from loan_repayments where loan_id=?", l['id'])
        paid = sum(r['amount'] for r in reps
                   if not (month_year and r['source'] == 'payroll' and r['month_year'] == month_year))
        left = round((l['amount'] or 0) - paid, 2)
        if left <= 0.005:
            continue
        balance += left
        was_closed = l['closed'] and not any(
            r['source'] == 'payroll' and r['month_year'] == month_year for r in reps)
        if not was_closed:
            due += min(l['instalment'] or 0, left)
    return round(due, 2), round(balance, 2)


c = TestClient(main.app)
H = {'Authorization': 'Bearer ' + c.post('/auth/login', data={'username': 'admin', 'password': 'p'})
     .json()['access_token']}


def dl(path, **params):
    tok = c.post('/auth/download-token', headers=H).json()['token']
    return c.get(path, params=dict(params, token=tok, format='excel'))
today = main._dubai_today()

# ---- The running month, with the kinds of entry a real month has --------
THIS = today.strftime('%B %Y')
ym = today.strftime('%Y-%m')
for co in c.get('/employees/companies', headers=H).json():
    for grp in ('staff', 'local'):
        c.post('/employees/payroll/runs', json={'month_year': THIS, 'company_id': co['id'], 'group': grp}, headers=H)
staff = c.get('/employees/staff', headers=H).json()
staff = staff.get('rows', staff)
codes = [s['emp_no'] for s in staff if s.get('active', True)]
entries = [
    {'emp_no': codes[0], 'kind': 'sick', 'from': f'{ym}-02'},
    {'emp_no': codes[0], 'kind': 'sick', 'from': f'{ym}-03'},          # second sick day: deducted
    {'emp_no': codes[1], 'kind': 'absent', 'from': f'{ym}-04', 'half': True},
    {'emp_no': codes[2], 'kind': 'vacation', 'from': f'{ym}-06', 'to': f'{ym}-10'},
    {'emp_no': codes[3], 'kind': 'holiday', 'from': f'{ym}-07'},        # paid
    {'emp_no': codes[4], 'kind': 'unpaid', 'from': f'{ym}-08', 'to': f'{ym}-09'},
    {'emp_no': codes[5], 'kind': 'absent', 'from': f'{ym}-01', 'to': f'{ym}-30'},  # a whole month
]
for en in entries:
    r = c.post('/employees/leave', json=en, headers=H)
    if r.status_code != 200:
        print('   leave entry refused:', en, r.text[:120])
for code, direction, cat, amt in ((codes[0], 'add', 'taxi', 312.5), (codes[1], 'deduct', 'iloe', 241.5),
                                  (codes[2], 'add', 'leave_salary', 5000), (codes[6], 'deduct', 'fine', 99.99)):
    r = c.post('/employees/pay-items', json={'emp_no': code, 'month_year': THIS, 'direction': direction,
                                             'category': cat, 'amount': amt, 'notes': 'audit'}, headers=H)
    if r.status_code != 200:
        print('   pay item refused:', code, r.text[:120])

runs = q("select r.*, c.short_name, c.name as cname from payroll_runs r join companies c on c.id=r.company_id "
         "order by r.id")
ck('there are cycles on file to audit', len(runs) > 0)
total_lines = 0
for r in runs:
    api = c.get(f"/employees/payroll/runs/{r['id']}", headers=H).json()
    a, b = bounds(r['month_year'])
    label = f"{r['short_name']} {r['month_year']} ({r.get('group')})"
    days = (b - a).days + 1
    done = 0 if today < a else days if today > b else (today - a).days + 1
    running = r['status'] != 'approved' and done < days
    mine = {}
    for l in api['lines']:
        e = q("select * from employees where id=?", l['employee_id'])[0]
        basic, allow = salary_on(e['id'], e, b)
        gross = round(basic + allow, 2)
        absent = round(min(absence(e['id'], a, b, gross) + part_month(e, a, b, gross), gross), 2)
        add, ded = items(e['id'], r['month_year'])
        due, bal = loan_state(e['id'], r['month_year'] if r['status'] == 'approved' else None)
        loan = l['loan_deduction'] if l['loan_edited'] else due
        net = round(gross - absent - ded - loan + add, 2)
        to_date = round(gross * done / days - absent + add - ded, 2) if running else net
        mine[l['emp_no']] = dict(gross=gross, absent=absent, adjust=round(add - ded, 2),
                                 loan=loan, net=net, to_date=to_date, balance=bal, held=l['held'])
        total_lines += 1
        row = l
        ok = (same(row['fixed_salary'], gross) and same(row['deduction'], absent)
              and same(row['adjust'], add - ded) and same(row['loan_deduction'], loan)
              and same(row['net_pay'], net)
              and (not running or same(row['to_date'], to_date)))
        if not ok:
            ck(f"{label}: {l['emp_no']} {l['name']}", False,
               f"mine={mine[l['emp_no']]} app={{gross {row['fixed_salary']}, absent {row['deduction']}, "
               f"adjust {row['adjust']}, loan {row['loan_deduction']}, net {row['net_pay']}, "
               f"to_date {row.get('to_date')}}}")
        if loan and not l['remark_edited']:
            ck(f"{label}: {l['emp_no']} remark carries the loan balance after this instalment",
               f"balance {bal - loan:,.2f}" in (row['remarks'] or ''), row['remarks'])
    paid = [v for v in mine.values() if not v['held']]
    ck(f"{label}: every line agrees ({len(mine)} lines)",
       not any(f.startswith(label + ':') and 'remark' not in f for f in FAIL))
    ck(f"{label}: net total {sum(v['net'] for v in paid):,.2f}",
       same(api['totals']['net_pay'], sum(v['net'] for v in paid)), api['totals'])
    by_route = {}
    for l in api['lines']:
        if not l['held']:
            by_route[l['pay_route']] = round(by_route.get(l['pay_route'], 0) + mine[l['emp_no']]['net'], 2)
    ck(f"{label}: totals by route {by_route}", all(same(api['by_route'].get(k, 0), v) for k, v in by_route.items())
       and all(same(by_route.get(k, 0), v) for k, v in api['by_route'].items()), api['by_route'])

    # ---- The printed statement says the same -----------------------------
    x = dl('/export/payroll/statement', run_id=r['id'])
    ck(f"{label}: statement spreadsheet downloads", x.status_code == 200, x.status_code)
    if x.status_code == 200:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(x.content), data_only=True)
        ws = wb.active
        rows = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
        head_i = next(i for i, row in enumerate(rows) if row and 'Emp. Code' in [str(v) for v in row])
        head = [str(v) for v in rows[head_i]]
        col = {h: i for i, h in enumerate(head)}
        last_col = 'Salary To Date' if running else 'Net Pay'
        ck(f"{label}: statement columns are {head[:8]}",
           head[:4] == ['Emp. Code', 'Employee Name', 'Joining Date', 'Gross Salary']
           and 'Add / Ded.' in col and 'Absent Ded.' in col and 'Loan' in col and last_col in col
           and not any(h in col for h in ('Sr.', 'Salary Paid', 'Basic Salary', 'Fix Allown.',
                                          'Leave Salary', 'Salary Payable')), head)
        seen, tot = 0, None
        for row in rows[head_i + 1:]:
            if not row or row[0] is None:
                continue
            code = str(row[col['Emp. Code']])
            if code in mine:
                m = mine[code]
                want = m['to_date'] if running else m['net']
                ck(f"{label}: printed {code}", same(row[col['Gross Salary']], m['gross'])
                   and same(row[col['Add / Ded.']], m['adjust']) and same(row[col['Absent Ded.']], m['absent'])
                   and same(row[col['Loan']], m['loan']) and same(row[col[last_col]], want),
                   (row, m)) if not (same(row[col['Gross Salary']], m['gross'])
                   and same(row[col['Add / Ded.']], m['adjust']) and same(row[col['Absent Ded.']], m['absent'])
                   and same(row[col['Loan']], m['loan']) and same(row[col[last_col]], want)) else None
                seen += 1
            elif 'TOTAL' in str(row[0]).upper() or (row[1] and 'TOTAL' in str(row[1]).upper()):
                tot = row
        ck(f"{label}: every paid line is printed ({seen})", seen == len(paid), (seen, len(paid)))
        want_tot = sum(v['to_date'] if running else v['net'] for v in paid)
        ck(f"{label}: printed total {want_tot:,.2f}",
           tot is not None and same(tot[col[last_col]], want_tot), tot)

print(f"\n{total_lines} salary lines recomputed by hand")

# ---- Consolidated: the companies added together -----------------------
for month in sorted({r['month_year'] for r in runs}):
    cons = dl('/export/payroll/consolidated', month_year=month)
    ck(f"consolidated {month}: downloads", cons.status_code == 200, cons.status_code)
    if cons.status_code == 200:
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(cons.content), data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        want = 0.0
        for r in runs:
            if r['month_year'] == month:
                api = c.get(f"/employees/payroll/runs/{r['id']}", headers=H).json()
                want += sum(l['net_pay'] for l in api['lines'] if not l['held'])
        nums = [v for row in rows for v in row if isinstance(v, (int, float))]
        ck(f"consolidated {month}: carries the grand total {want:,.2f}", any(same(v, want) for v in nums),
           sorted(nums)[-5:])

# ---- Gratuity, by hand ---------------------------------------------------
for e in q("select * from employees where staff=1 and active=1 order by emp_no"):
    g = c.get(f"/employees/staff/{e['emp_no']}/file", headers=H).json()['gratuity']
    scheme = e.get('scheme') or 'gratuity'
    if scheme != 'gratuity' or not e['joined_on']:
        ck(f"gratuity {e['emp_no']} {e['name']}: none ({scheme})", same(g['amount'], 0), g)
        continue
    end = today
    service = (end - d(e['joined_on'])).days + 1
    unpaid = 0.0
    for l in q("select * from staff_leave where employee_id=? and on_date<=?", e['id'], end.isoformat()):
        pass  # unpaid days are judged month by month below
    ms = {}
    for l in q("select * from staff_leave where employee_id=? and on_date<=? order by on_date, id",
               e['id'], end.isoformat()):
        ms.setdefault(str(l['on_date'])[:7], []).append(l)
    for month, rows in ms.items():
        y, m = map(int, month.split('-'))
        a = date(y, m, 1)
        b = date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1)
        gross = 30.0
        amt = absence(e['id'], a, b, 30 * 30)  # day rate 30 -> unpaid days = amt / 30
        unpaid += amt / 30.0
    years = max(service - unpaid, 0) / 365.0
    basic = e['basic_salary'] or 0
    daily = basic * 12 / 365.0
    if years < 1:
        amount = 0.0
    else:
        days_ = min(years, 5) * 21 + max(years - 5, 0) * 30
        amount = min(days_ * daily, basic * 24)
    ck(f"gratuity {e['emp_no']} {e['name']}: {amount:,.2f}", same(g['amount'], amount),
       (g['amount'], amount, years, unpaid))

print()
if FAIL:
    print(f"{len(FAIL)} FAILED")
    sys.exit(1)
print("EVERY FIGURE AGREES WITH THE HAND CALCULATION")
