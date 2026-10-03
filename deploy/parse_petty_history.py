"""Turn the petty cash files kept outside the app into the lines it keeps.

    python3 deploy/parse_petty_history.py SITE.xlsx COMPANY_2026.xlsx TALLY_2025.xlsx JOMON.xlsx

Writes deploy/petty_books_history.json for deploy/load_petty_history.py.

  site    <- site petty cash (Zoho "Account Transactions", Jun - 7 Sep 2026)
  office  <- company petty cash: Tally cash book 2025 + Zoho 2026
  pro     <- Jomon's petty cash (Zoho, Jan 2025 - Aug 2026)

Zoho exports: debit = cash received, credit = paid; "Vendor Name : X"
inside the details becomes the supplier.

The Tally 2025 book was exported from Tally's Educational mode, which
prints a date only on the 1st, 2nd and 31st of a month ("Educational"
everywhere else). Rows are in date order, so each takes the last date
printed above it - the right month where Tally gives one, the exact day
unknown. The voucher number goes in the supplier / reference column.
The file also carries October-November twice; the copy that ends at
Tally's own closing balance (2,703.36 - the 2026 book's opening) is
used. The 2026 book's "op" line is then not loaded again.

Every book's lines must come to the files' own totals before anything
is written.
"""
import datetime
import json
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "petty_books_history.json")


def _r2(v):
    return round(float(v or 0), 2)


def zoho(path, skip_op=False):
    ws = openpyxl.load_workbook(path, data_only=True).active
    out = []
    for r in range(3, ws.max_row + 1):
        d, det, dr, cr = (ws.cell(r, c).value for c in range(1, 5))
        if isinstance(d, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", d.strip()):
            d = datetime.datetime.strptime(d.strip(), "%Y-%m-%d")      # some exports keep dates as text
        if not isinstance(d, datetime.datetime):
            continue                                   # header / totals
        dr = dr if isinstance(dr, (int, float)) else 0
        cr = cr if isinstance(cr, (int, float)) else 0
        det = str(det or "").strip()
        sup = ""
        m = re.search(r"\s*Vendor Name\s*:\s*(.+)$", det, re.S)
        if m:
            sup = " ".join(m.group(1).split())
            det = det[:m.start()].strip()
        det = re.sub(r",(?=\S)", ", ", " ".join(det.split()))      # "A,B" -> "A, B"
        if skip_op and det.lower() in ("op", "opening balance"):
            continue
        if det.lower() in ("op", "opening balance"):
            det = "Opening balance"
        out.append({"date": d.date().isoformat(), "description": det or sup or "-", "supplier": sup, "site": "",
                    "received": _r2(dr), "paid": _r2(cr)})
    return out


def tally(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = list(range(11, ws.max_row + 1))
    # The second copy of Oct-Nov starts where the dates go back to 1-Oct
    # after November; the first copy (the stale one) is dropped.
    marks, prev, cut_a, cut_b = [], None, None, None
    for r in rows:
        v = ws.cell(r, 1).value
        if isinstance(v, datetime.datetime):
            d = v.date()
            if prev and d < prev and cut_b is None:
                cut_b = r
                cut_a = next(x for x, dd in marks if dd == d)
            marks.append((r, d))
            prev = d
    keep = [r for r in rows if not (cut_a and cut_a <= r < cut_b)]
    out, last = [], None
    i = 0
    while i < len(keep):
        r = keep[i]
        a, _, led = (ws.cell(r, c).value for c in (1, 2, 3))
        if isinstance(a, datetime.datetime):
            last = a.date()
        vt, vn = ws.cell(r, 6).value, ws.cell(r, 7).value
        dr, cr = ws.cell(r, 8).value, ws.cell(r, 9).value
        led = " ".join(str(led or "").split())
        if led == "Opening Balance":
            out.append({"date": last.isoformat(), "description": "Opening balance", "supplier": "", "site": "",
                        "received": _r2(dr), "paid": _r2(cr)})
            i += 1
            continue
        if not vt:
            i += 1
            continue
        if led.startswith("Closing Balance"):
            break
        desc = led
        if led == "(as per details)":
            # The breakdown lines below say what the payment was for.
            subs, j = [], i + 1
            while j < len(keep) and not ws.cell(keep[j], 6).value and ws.cell(keep[j], 3).value \
                    and not str(ws.cell(keep[j], 3).value).startswith("Closing"):
                s = " ".join(str(ws.cell(keep[j], 3).value).split())
                if not s.lower().startswith("vat@"):
                    subs.append(s)
                j += 1
            desc = ", ".join(dict.fromkeys(subs)) or "As per details"
            if len(desc) > 160:
                desc = desc[:157].rsplit(",", 1)[0] + ", ..."
        out.append({"date": last.isoformat(), "description": desc, "supplier": str(vn or ""), "site": "",
                    "received": _r2(dr), "paid": _r2(cr)})
        i += 1
    return out


def bal(lines):
    return round(sum(x["received"] - x["paid"] for x in lines), 2)


def main(site, comp26, tally25, jomon):
    books = {}
    books["site"] = zoho(site)
    t = tally(tally25)
    assert bal(t) == 2703.36, ("Tally 2025 does not close at 2,703.36", bal(t))
    books["office"] = t + zoho(comp26, skip_op=True)
    books["pro"] = zoho(jomon)
    want = {"site": 1199.95, "office": 1856.64, "pro": 20898.20}
    for b, lines in books.items():
        rec = round(sum(x["received"] for x in lines), 2)
        paid = round(sum(x["paid"] for x in lines), 2)
        print(f"{b:<7} {len(lines):>5} lines  {lines[0]['date']} to {max(x['date'] for x in lines)}"
              f"  received {rec:>13,.2f}  paid {paid:>13,.2f}  balance {bal(lines):>11,.2f}  (file {want[b]:,.2f})")
        assert bal(lines) == want[b], b
        for x in lines:
            assert bool(x["received"]) != bool(x["paid"]), x
    json.dump(books, open(OUT, "w"), indent=0)
    print("written", OUT)


if __name__ == "__main__":
    main(*sys.argv[1:5])
