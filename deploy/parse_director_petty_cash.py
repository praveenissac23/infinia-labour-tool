"""Turn the directors' petty cash sheets into the lines the app keeps.

    python3 deploy/parse_director_petty_cash.py NAVEEN.xlsx PRAVEEN.xlsx

Writes deploy/director_petty_cash.json, which
deploy/load_director_petty_cash.py puts into the app.

How the sheets read: DR is money the director put in or paid for the
company (what the company owes him goes up), CR is money paid back to
him (it goes down), BAL is what the company owes him. Each block opens
with the balance brought forward and ends with "Balance as on ...".

The sheets are kept by hand, so the columns are not always used the
same way: the first Praveen blocks put his payments under CR, and the
very first Naveen block (2022) counts down cash he held for Infinia.
Each line is therefore read the way the sheet's own balance formula
reads it. After every balance the sheet states, the running total here
is compared with the sheet's figure; where they differ (a formula on
the sheet skipped a line, a balance was typed in) a correction line is
added so the app shows exactly the balance the sheet shows, and the
difference is printed.
"""
import json
import os
import re
import sys
from datetime import date, datetime

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "director_petty_cash.json")


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.fullmatch(r"\s*(\d{1,2})[./](\d{1,2})[./](\d{2,4})\s*", str(v or ""))
    if not m:
        return None
    d, mo, y = (int(x) for x in m.groups())
    if y < 100:
        y += 2000
    try:
        return date(y, mo, d)
    except ValueError:
        return None


def _num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def _project(v):
    if isinstance(v, (int, float)):
        return str(int(v))
    s = str(v or "").strip()
    return s if re.fullmatch(r"\d{3}", s) else ""


def _sign(formula, col, row):
    """+1 / -1 if the balance formula adds / takes away this cell."""
    m = re.search(r"([+-])\s*" + col + str(row) + r"(?!\d)", formula or "")
    return (1 if m.group(1) == "+" else -1) if m else 0


# What each difference on these two sheets is, found by reading them.
WHY = {
    ("naveen", 27): "transfer of 54,202.00 settled 54,202.94",
    ("naveen", 240): "its formula left out Permit & Local fees 146.00 of 15-Jan-25",
    ("praveen", 29): "its formula left out JVT villa cleaning 89.25 of 25-Oct-23",
    ("praveen", 126): "its formulas left out Design Revision fee 1,020.00, Laptop VAT 381.86 and Sand shifting 2,645.00 of Jan-Feb 25",
}


def parse(path, who):
    wf = openpyxl.load_workbook(path).active
    wv = openpyxl.load_workbook(path, data_only=True).active
    formulas = {r: wf.cell(r, 6).value for r in range(1, wf.max_row + 1)}
    referenced = set()
    for f in formulas.values():
        if isinstance(f, str) and f.startswith("="):
            referenced.update(int(x) for x in re.findall(r"F(\d+)", f))

    lines, report = [], []
    bal = 0.0               # what the company owes the director
    invert = False          # the 2022 block: cash HE held for Infinia
    last_date = None
    pending_adj = []        # corrections before any dated line

    def add(d, desc, project, delta, note=""):
        delta = round(delta, 2)
        if not delta:
            return
        lines.append({"date": d.isoformat() if d else None, "description": desc, "site": project,
                      "paid": delta if delta > 0 else 0.0, "received": -delta if delta < 0 else 0.0, "note": note})

    def check(r, stated, label):
        nonlocal bal
        want = round(-stated if invert else stated, 2)
        diff = round(want - bal, 2)
        if abs(diff) >= 0.01:
            first = not lines and not report
            desc = (f"Opening balance - cash held by {who.title()} for Infinia (as per sheet)" if first and diff < 0 else
                    f"Opening balance (as per sheet)" if first else
                    f"Correction to the sheet's balance of {want:,.2f}: {WHY.get((who, r), f'sheet row {r}')}")
            report.append(f"  row {r}: {label!r} sheet says {want:,.2f}, lines add up to {bal:,.2f} -> correction {diff:+,.2f}")
            if last_date:
                add(last_date, desc, "", diff, "correction")
            else:
                pending_adj.append((desc, diff))
            bal = want

    for r in range(1, wv.max_row + 1):
        a, b, c, dv, ev, fv = (wv.cell(r, i).value for i in range(1, 7))
        f = formulas.get(r)
        text = " ".join(str(x) for x in (a, b) if isinstance(x, str)).strip()
        low = text.lower()
        if "out to infinia" in low:
            invert = True
            continue
        is_fx = isinstance(f, str) and f.startswith("=")
        amount_d, amount_e = _num(dv), _num(ev)
        if low.endswith("petty cash") and not amount_d and not amount_e and not _date(a):
            invert = False      # a new block's heading
            continue
        label_row = any(k in low for k in ("balance", "total")) or low.startswith("paid off")
        if (amount_d or amount_e) and not label_row:
            d = _date(a) or last_date
            desc = re.sub(r"\s+", " ", str(b or "").strip())
            if not desc:
                raise SystemExit(f"{who}: row {r} has an amount but no description")
            sd, se = (_sign(f, "D", r), _sign(f, "E", r)) if is_fx else (0, 0)
            if not is_fx and isinstance(fv, (int, float)):
                # A balance typed in by hand: read the line whichever way
                # lands nearer to it.
                amt = amount_d or amount_e
                expect = -fv if invert else fv
                delta = amt if abs(bal + amt - expect) <= abs(bal - amt - expect) else -amt
            else:
                if not sd and amount_d:
                    sd = -1 if invert else 1
                if not se and amount_e:
                    se = 1 if invert else -1
                delta = sd * amount_d + se * amount_e
                if invert:
                    delta = -delta
            if pending_adj and d:
                for desc0, diff0 in pending_adj:
                    add(d, desc0, "", diff0, "correction")
                pending_adj.clear()
            add(d, desc, _project(c), delta)
            bal = round(bal + delta, 2)
            last_date = d
            if not is_fx and isinstance(fv, (int, float)):
                check(r, fv, desc)
            continue
        # A balance the sheet states and carries on from.
        if isinstance(fv, (int, float)) and r in referenced and (label_row or not is_fx):
            check(r, fv, text or "(balance)")
    return lines, round(bal, 2), report


def main(naveen, praveen):
    out = {}
    for who, path in (("naveen", naveen), ("praveen", praveen)):
        lines, bal, report = parse(path, who)
        paid = round(sum(x["paid"] for x in lines), 2)
        rec = round(sum(x["received"] for x in lines), 2)
        assert round(paid - rec, 2) == bal, (who, paid, rec, bal)
        print(f"{who}: {len(lines)} lines, {lines[0]['date']} to {max(x['date'] for x in lines)}, "
              f"DR {paid:,.2f}  CR {rec:,.2f}  balance due {bal:,.2f}")
        for x in report:
            print(x)
        out[who] = {"lines": lines, "balance": bal}
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w") as fh:
        json.dump(out, fh, indent=1)
    print("written", OUT)


if __name__ == "__main__":
    main(*sys.argv[1:3])
