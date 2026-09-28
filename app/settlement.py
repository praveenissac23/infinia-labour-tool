"""Final settlement - what a member of staff is paid when he leaves.

One working, drawn three ways: on the Gratuity screen, as a preview
page, and as a PDF or Excel sheet to hand over and sign. Nothing is
stored: the figures come from the record (joining date, basic, gross,
leave taken, loans) and from what is typed on the screen (the last
working day, why he is leaving, notice served, a ticket, anything
else), so the paper can never disagree with the screen.

The law (UAE Federal Decree-Law 33 of 2021) and the house rule of
paying only what the law requires:
  * Gratuity (art. 51): none under a year's service; 21 days' BASIC for
    each of the first five years, 30 after, pro rata, capped at two
    years' basic. Unpaid absence is not service. A day's basic is taken
    as basic x 12 / 365 (see main.gratuity_detail).
  * Resigning no longer reduces gratuity - the old one-third / two-
    thirds cut went with the 1980 law - so the reason changes only the
    notice, never the gratuity.
  * Unused leave (art. 29) is paid on the BASIC wage, not the gross.
    Under a year: nothing for the first six months, then 2 days a month.
  * Notice (art. 43): whoever does not give the notice owes the wage for
    the part not served. When he resigns and leaves early, that comes
    off; when the company ends it without notice, it is added. Dismissal
    under article 44 needs no notice either way.
  * Loans and advances still owing are deducted.
"""
import calendar
import math
import io
from datetime import date, timedelta
from html import escape
from urllib.parse import quote, urlencode

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse
from sqlalchemy.orm import Session

import main as M
import models, auth, people, export_web
from database import get_db

router = APIRouter()

REASONS = {
    "resignation": "Resignation",
    "termination": "Termination by the company",
    "end_of_contract": "End of contract",
    "art44": "Dismissal under Article 44",
    "death": "Death",
    "retirement": "Retirement",
}
CONTRACTS = {"limited": "Limited", "unlimited": "Unlimited", "part-time": "Part-time"}


def _f(v, default=0.0):
    try:
        return float(str(v).replace(",", "")) if v not in (None, "") else default
    except ValueError:
        return default


def _money(v):
    return f"{v:,.2f}"


def _dmy(d):
    return d.strftime("%d %b %Y") if d else "-"


def _ymd_text(a, b):
    """Service as years, months and days, both ends counted."""
    if not a or not b or b < a:
        return "-"
    end = b + timedelta(days=1)
    y = end.year - a.year
    m = end.month - a.month
    d = end.day - a.day
    if d < 0:
        m -= 1
        prev = (end.replace(day=1) - timedelta(days=1))
        d += calendar.monthrange(prev.year, prev.month)[1]
    if m < 0:
        y -= 1
        m += 12
    return f"{y} yr {m} mo {d} d"


def _leave_owed(db, e, last):
    """Unused leave on the last day, at the legal minimum."""
    p = db.query(models.PeopleProfile).filter(models.PeopleProfile.employee_id == e.id).first()
    lv = people.leave_state(db, e, p, "office", today=last)
    bal = lv.get("balance") or 0.0
    note = f"{lv.get('accrued', 0):g} earned - {lv.get('taken', 0):g} taken"
    if e.joined_on:
        months = people._months_between(e.joined_on, last)
        if months < 12:
            # Art. 29: nothing before six months, 2 days a month after.
            earned = 0.0 if months < 6 else 2.0 * months
            earned += (p.leave_opening or 0.0) if p else 0.0
            bal = min(bal, earned - (lv.get("taken") or 0.0))
            note = f"under a year: {earned:g} earned (2 days a month after 6 months) - {lv.get('taken', 0):g} taken"
    return max(round(bal, 1), 0.0), note, p


def settlement(db, emp_no, q):
    """The whole working. q holds what was typed on the screen."""
    e = M._staff_by_code(db, emp_no)
    today = M._dubai_today()
    last = M._as_date(q.get("last_day")) or e.terminated_on or today
    if e.joined_on and last < e.joined_on:
        raise HTTPException(status_code=400, detail="The last working day is before he joined.")
    reason = q.get("reason") if q.get("reason") in REASONS else "resignation"
    basic = round(e.basic_salary or 0, 2)
    gross = M._gross(e)
    day_basic = round(basic * 12 / 365.0, 2)       # gratuity and leave
    day_gross = round(gross / 30.0, 2)             # notice: a month's wage for 30 days

    g = M.gratuity_detail(e, db, last)
    grat = g["amount"] if g["entitled"] else 0.0
    # The two lines of the table, from the exact figures, so they add up
    # to the total printed under them.
    yrs = max(g["service_days"] - g["unpaid_days"], 0) / 365.0
    f5 = a5 = 0.0
    if grat:
        f5 = round(min(yrs, 5) * 21 * basic * 12 / 365.0, 2)
        a5 = round(max(yrs - 5, 0) * 30 * basic * 12 / 365.0, 2)
        if not g.get("capped"):
            grat = round(f5 + a5, 2)

    # Unused leave, on the basic.
    auto_leave, leave_note, prof = _leave_owed(db, e, last)
    leave_days = auto_leave if q.get("leave_days") in (None, "") else max(_f(q.get("leave_days")), 0.0)
    leave_pay = round(leave_days * day_basic, 2)

    # Salary for the last month's days not yet paid.
    mdays = calendar.monthrange(last.year, last.month)[1]
    auto_sal = last.day if not (e.joined_on and e.joined_on > last.replace(day=1)) else (last - e.joined_on).days + 1
    sal_days = auto_sal if q.get("salary_days") in (None, "") else max(_f(q.get("salary_days")), 0.0)
    sal_days = min(sal_days, mdays)
    salary = round(gross / mdays * sal_days, 2)

    # Notice.
    notice_req = int(_f(q.get("notice_required"), e.notice_days or 30))
    notice_srv = max(min(_f(q.get("notice_served"), notice_req), notice_req), 0)
    short = notice_req - notice_srv
    notice_ded = notice_add = 0.0
    if reason == "resignation" and short > 0:
        notice_ded = round(short * day_gross, 2)
    if reason == "termination" and short > 0:
        notice_add = round(short * day_gross, 2)

    ticket = max(_f(q.get("ticket")), 0.0)
    other_add = max(_f(q.get("other_add")), 0.0)
    loans = [l for l in M.list_loans(open_only=False, db=db, user=None)["rows"] if l["emp_no"] == e.emp_no and not l["closed"]]
    auto_loan = round(sum(l["balance"] for l in loans), 2)
    loan = auto_loan if q.get("loan") in (None, "") else max(_f(q.get("loan")), 0.0)
    other_ded = max(_f(q.get("other_ded")), 0.0)

    payables = [("End of service gratuity", grat),
                ("Leave encashment", leave_pay),
                (f"Salary for {sal_days:g} day(s) of {last:%b %Y}", salary)]
    if notice_add:
        payables.append((f"Notice pay ({short:g} day(s) not given)", notice_add))
    payables += [("Air ticket", ticket),
                 (q.get("other_add_note") or "Other additions", other_add)]
    deductions = [("Loan / advance balance", loan)]
    if notice_ded:
        deductions.append((f"Notice not served ({short:g} of {notice_req} days)", notice_ded))
    deductions.append((q.get("other_ded_note") or "Other deductions / recoveries", other_ded))
    # Optional lines only when there is something on them.
    payables = payables[:3] + [x for x in payables[3:] if x[1]]
    deductions = [x for x in deductions if x[1]] or [("No deductions", 0.0)]
    total_pay = round(sum(v for _, v in payables), 2)
    total_ded = round(sum(v for _, v in deductions), 2)
    net = float(math.floor(total_pay - total_ded + 0.5))   # whole dirhams, as on the paper

    docs = {d.kind: d for d in db.query(models.EmployeeDocument)
            .filter(models.EmployeeDocument.employee_id == e.id).all()}
    passport = docs.get("passport")
    visa = docs.get("visa")
    ctype = q.get("contract") or (prof.employment_type if prof else "") or ""

    notes = []
    if g["entitled"] and not grat:
        notes.append("No gratuity: under one year's continuous service (art. 51).")
    if not g["entitled"]:
        notes.append(f"No gratuity: {g['why']}.")
    if g.get("capped"):
        notes.append("Gratuity capped at two years' basic wage (art. 51).")
    if reason == "resignation":
        notes.append("Resignation does not reduce gratuity under Decree-Law 33 of 2021; "
                     "only notice not served is deducted (art. 43).")
    if reason == "art44":
        notes.append("Dismissal under article 44: no notice either way; gratuity is still due.")
    notes.append("Gratuity and leave are paid on the basic wage only; unpaid absence is not service.")

    return {
        "emp": {"emp_no": e.emp_no, "name": e.name, "designation": e.designation or e.trade or "",
                "company": e.company or "", "nationality": (prof.nationality if prof else "") or "",
                "mobile": (prof.mobile if prof else "") or "",
                "passport": (passport.number if passport else "") or "",
                "visa_expiry": visa.expires_on.isoformat() if visa and visa.expires_on else "",
                "contract": ctype, "joined_on": e.joined_on.isoformat() if e.joined_on else "",
                "iban": e.iban or "", "pay_route": e.pay_route or "wps"},
        "input": {"last_day": last.isoformat(), "reason": reason, "reason_label": REASONS[reason],
                  "leave_days": leave_days, "leave_days_auto": auto_leave, "leave_note": leave_note,
                  "salary_days": sal_days, "salary_days_auto": auto_sal, "month_days": mdays,
                  "notice_required": notice_req, "notice_served": notice_srv,
                  "loan": loan, "loan_auto": auto_loan, "ticket": ticket, "other_add": other_add,
                  "other_ded": other_ded, "other_add_note": q.get("other_add_note") or "",
                  "other_ded_note": q.get("other_ded_note") or "", "contract": ctype,
                  "settled_on": (M._as_date(q.get("settled_on")) or today).isoformat()},
        "salary": {"basic": basic, "gross": gross, "day_basic": day_basic, "day_gross": day_gross},
        "service": {"days": g["service_days"], "unpaid": g["unpaid_days"],
                    "counted": max(g["service_days"] - g["unpaid_days"], 0),
                    "years": g["years"], "text": _ymd_text(e.joined_on, last)},
        "gratuity": {"amount": grat, "entitled": g["entitled"], "capped": g.get("capped", False),
                     "first5": round(min(yrs, 5), 3), "after5": round(max(yrs - 5, 0), 3),
                     "days": g["days"], "why": g["why"], "first5_amt": f5, "after5_amt": a5,
                     "cap_cut": round(f5 + a5 - grat, 2) if grat and g.get("capped") else 0.0},
        "payables": [{"label": k, "amount": round(v, 2)} for k, v in payables],
        "deductions": [{"label": k, "amount": round(v, 2)} for k, v in deductions],
        "total_pay": total_pay, "total_ded": total_ded, "net": net,
        "notes": notes,
    }


HR = Depends(M.require_screen("hrpayroll"))
FIELDS = ("last_day", "reason", "leave_days", "salary_days", "notice_required", "notice_served",
          "loan", "ticket", "other_add", "other_ded", "other_add_note", "other_ded_note",
          "contract", "settled_on")


def _q(**kw):
    return {k: v for k, v in kw.items() if v not in (None, "")}


@router.get("/employees/staff/{emp_no}/settlement")
def get_settlement(emp_no: str, last_day: str = "", reason: str = "", leave_days: str = "",
                   salary_days: str = "", notice_required: str = "", notice_served: str = "",
                   loan: str = "", ticket: str = "", other_add: str = "", other_ded: str = "",
                   other_add_note: str = "", other_ded_note: str = "", contract: str = "",
                   settled_on: str = "", db: Session = Depends(get_db), user: models.User = HR):
    return settlement(db, emp_no, _q(**{k: v for k, v in locals().items() if k in FIELDS}))


# ---- The paper ---------------------------------------------------------

def _details(s):
    e, i = s["emp"], s["input"]
    return [("Employee name", e["name"]), ("Employee code", e["emp_no"]),
            ("Designation", e["designation"] or "-"), ("Company", e["company"] or "-"),
            ("Passport no.", e["passport"] or "-"), ("Nationality", e["nationality"] or "-"),
            ("Contact no.", e["mobile"] or "-"), ("Contract type", CONTRACTS.get(i["contract"], i["contract"] or "-")),
            ("Reason for leaving", i["reason_label"]), ("Visa expiry", _dmy(M._as_date(e["visa_expiry"]))),
            ("Date of joining", _dmy(M._as_date(e["joined_on"]))), ("Last working day", _dmy(M._as_date(i["last_day"]))),
            ("Total service", s["service"]["text"]), ("Service counted", f"{s['service']['counted']:,.0f} days ({s['service']['years']:.3f} yrs)")]


def _basis(s):
    i, sal, sv = s["input"], s["salary"], s["service"]
    return [("Basic salary", _money(sal["basic"])), ("Gross salary", _money(sal["gross"])),
            ("Daily basic (basic x 12 / 365)", _money(sal["day_basic"])),
            ("Service days", f"{sv['days']:,}"),
            ("Unpaid days (not service)", f"{sv['unpaid']:g}"),
            ("Unused leave days", f"{i['leave_days']:g}"),
            ("Notice required / served", f"{i['notice_required']} / {i['notice_served']:g} days")]


def _grat_rows(s):
    g, d = s["gratuity"], s["salary"]["day_basic"]
    rows = [("First 5 years", f"{g['first5']:.3f}", "21", _money(d), _money(g["first5_amt"])),
            ("Above 5 years", f"{g['after5']:.3f}", "30", _money(d), _money(g["after5_amt"]))]
    if g.get("cap_cut"):
        rows.append(("Less: limit of two years' basic", "", "", "", "-" + _money(g["cap_cut"])))
    return rows


DECLARATION = ("I, the undersigned, confirm that I have received my original passport, my ticket "
               "entitlement and all my final dues and settlement up to the date above from the company, "
               "and that I have no further claim against it.")


def _html(s, pdf_url, excel_url):
    logo = export_web.logo_data_uri()
    e, i = s["emp"], s["input"]
    kv = lambda rows: "".join(f"<div class='kv'><span>{escape(k)}</span><b>{escape(str(v))}</b></div>" for k, v in rows)
    money_rows = lambda rows, tot, label: ("".join(
        f"<tr><td>{escape(r['label'])}</td><td class='n'>{_money(r['amount'])}</td></tr>" for r in rows)
        + f"<tr class='t'><td>{label}</td><td class='n'>{_money(tot)}</td></tr>")
    g = s["gratuity"]
    grat = "".join(f"<tr><td>{a}</td><td class='n'>{b}</td><td class='n'>{c}</td><td class='n'>{d}</td><td class='n'>{x}</td></tr>"
                   for a, b, c, d, x in _grat_rows(s))
    net = s["net"]
    return f"""<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Final Settlement - {escape(e['name'])}</title><style>
:root{{--ink:#1d2433;--mute:#6b7280;--line:#dfe3ea;--head:#1f2d45;--soft:#f5f7fa;--bg:#e9ecf1}}
*{{box-sizing:border-box}} body{{margin:0;background:var(--bg);font:13px/1.45 Arial,Helvetica,sans-serif;color:var(--ink)}}
.bar{{position:sticky;top:0;background:#fff;border-bottom:1px solid var(--line);padding:10px 16px;display:flex;gap:8px;justify-content:flex-end;z-index:2}}
.bar a,.bar button{{font:600 13px Arial;padding:8px 14px;border-radius:6px;border:1px solid var(--line);background:#fff;color:var(--ink);text-decoration:none;cursor:pointer}}
.bar .p{{background:var(--head);color:#fff;border-color:var(--head)}}
.page{{max-width:800px;margin:18px auto;background:#fff;padding:34px 40px;box-shadow:0 2px 10px rgba(0,0,0,.08)}}
.top{{display:flex;justify-content:space-between;align-items:flex-end;border-bottom:2px solid var(--head);padding-bottom:12px}}
.top img{{height:30px}} .top h1{{margin:0;font-size:18px;letter-spacing:.06em;color:var(--head)}} .top .d{{color:var(--mute);font-size:12px;text-align:right}}
h2{{font-size:11px;letter-spacing:.08em;text-transform:uppercase;color:var(--head);margin:20px 0 8px;border-bottom:1px solid var(--line);padding-bottom:4px}}
.grid{{display:grid;grid-template-columns:1fr 1fr;column-gap:28px}}
.kv{{display:flex;justify-content:space-between;gap:12px;padding:4px 0;border-bottom:1px dotted var(--line)}} .kv span{{color:var(--mute)}} .kv b{{font-weight:600;text-align:right}}
table{{width:100%;border-collapse:collapse}} th{{background:var(--head);color:#fff;font-size:11px;text-align:left;padding:6px 8px}} th.n,td.n{{text-align:right}}
td{{padding:6px 8px;border-bottom:1px solid var(--line)}} tr.t td{{font-weight:700;background:var(--soft);border-top:1.5px solid var(--ink)}}
.two{{display:grid;grid-template-columns:1fr 1fr;gap:20px;align-items:start}}
.net{{margin-top:18px;display:flex;justify-content:space-between;align-items:center;background:var(--head);color:#fff;padding:12px 16px;border-radius:4px}}
.net span{{font-size:12px;letter-spacing:.08em;text-transform:uppercase}} .net b{{font-size:20px}}
.notes{{margin:10px 0 0;padding-left:18px;color:var(--mute);font-size:11px}}
.decl{{margin-top:22px;font-size:12px;border:1px solid var(--line);padding:10px 12px;background:var(--soft)}}
.sign{{display:grid;grid-template-columns:repeat(3,1fr);gap:26px;margin-top:46px}} .sign div{{border-top:1px solid var(--ink);padding-top:6px;font-size:12px;text-align:center}}
@media(max-width:640px){{.page{{padding:20px 16px;margin:0}} .grid,.two{{grid-template-columns:1fr}}}}
@media print{{.bar{{display:none}} body{{background:#fff}} .page{{box-shadow:none;margin:0;max-width:none;padding:0}}}}
</style></head><body>
<div class="bar"><button onclick="print()">Print</button><a href="{escape(excel_url)}">Excel</a><a class="p" href="{escape(pdf_url)}">Download PDF</a></div>
<div class="page">
 <div class="top"><div>{f'<img src="{logo}" alt="">' if logo else '<b>INFINIA</b>'}</div>
  <div><h1>FINAL SETTLEMENT</h1><div class="d">Settlement date: {_dmy(M._as_date(i['settled_on']))}</div></div></div>
 <h2>Employee details</h2><div class="grid">{kv(_details(s))}</div>
 <h2>Salary and service</h2><div class="grid">{kv(_basis(s))}</div>
 <h2>End of service gratuity (UAE Decree-Law 33/2021, art. 51)</h2>
 <table><thead><tr><th>Period</th><th class="n">Years</th><th class="n">Days / year</th><th class="n">Daily basic</th><th class="n">Amount (AED)</th></tr></thead>
 <tbody>{grat}<tr class="t"><td colspan="4">Total gratuity{' (capped at 2 years basic)' if g['capped'] else ''}</td><td class="n">{_money(g['amount'])}</td></tr></tbody></table>
 <div class="two">
  <div><h2>Payable to employee</h2><table><tbody>{money_rows(s['payables'], s['total_pay'], 'Gross payable')}</tbody></table></div>
  <div><h2>Deductions</h2><table><tbody>{money_rows(s['deductions'], s['total_ded'], 'Total deductions')}</tbody></table></div>
 </div>
 <div class="net"><span>{'Final settlement payable' if net >= 0 else 'Balance due from employee'}</span><b>AED {_money(abs(net))}</b></div>
 <ul class="notes">{''.join(f'<li>{escape(n)}</li>' for n in s['notes'])}</ul>
 <div class="decl">{escape(DECLARATION)}</div>
 <div class="sign"><div>Prepared by</div><div>Approved by</div><div>Employee signature &amp; date</div></div>
</div></body></html>"""


def _pdf(s):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, KeepTogether
    HEAD, LINE, SOFT, MUTE = colors.HexColor("#1f2d45"), colors.HexColor("#dfe3ea"), colors.HexColor("#f5f7fa"), colors.HexColor("#6b7280")
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=8.5, leading=11)
    bold = ParagraphStyle("bb", parent=base, fontName="Helvetica-Bold")
    mute = ParagraphStyle("m", parent=base, textColor=MUTE)
    right = ParagraphStyle("r", parent=base, alignment=TA_RIGHT)
    rightb = ParagraphStyle("rb", parent=bold, alignment=TA_RIGHT)
    white = ParagraphStyle("w", parent=bold, textColor=colors.white)
    whiter = ParagraphStyle("wr", parent=white, alignment=TA_RIGHT)
    sec = ParagraphStyle("s", parent=bold, fontSize=8, textColor=HEAD, spaceBefore=8, spaceAfter=3)
    P = lambda t, st=base: Paragraph(escape(str(t)), st)
    W = A4[0] - 30 * mm
    i = s["input"]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=f"Final Settlement - {s['emp']['name']}")
    logo = export_web._logo_image(40)
    top = Table([[logo or P("INFINIA", bold),
                  [Paragraph("FINAL SETTLEMENT", ParagraphStyle("t", parent=bold, fontSize=14, leading=18, textColor=HEAD, alignment=TA_RIGHT)),
                   Paragraph(f"Settlement date: {_dmy(M._as_date(i['settled_on']))}", ParagraphStyle("d", parent=mute, alignment=TA_RIGHT))]]],
                colWidths=[W / 2, W / 2])
    top.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("LINEBELOW", (0, 0), (-1, 0), 1.5, HEAD),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))

    def kv_grid(rows):
        half = (len(rows) + 1) // 2
        L, R = rows[:half], rows[half:]
        data = []
        for n in range(half):
            a = L[n]; b = R[n] if n < len(R) else ("", "")
            data.append([P(a[0], mute), P(a[1], rightb), "", P(b[0], mute), P(b[1], rightb)])
        t = Table(data, colWidths=[W * .22, W * .26, W * .04, W * .22, W * .26])
        t.setStyle(TableStyle([("LINEBELOW", (0, 0), (1, -1), .4, LINE), ("LINEBELOW", (3, 0), (4, -1), .4, LINE),
                               ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                               ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
        return t

    g = s["gratuity"]
    gd = [[P("Period", white), P("Years", whiter), P("Days / year", whiter), P("Daily basic", whiter), P("Amount (AED)", whiter)]]
    gd += [[P(a), P(b, right), P(c, right), P(d, right), P(x, right)] for a, b, c, d, x in _grat_rows(s)]
    gd.append([P("Total gratuity" + (" (capped at 2 years basic)" if g["capped"] else ""), bold), "", "", "", P(_money(g["amount"]), rightb)])
    gt = Table(gd, colWidths=[W * .32, W * .15, W * .15, W * .18, W * .20])
    gt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), HEAD), ("LINEBELOW", (0, 1), (-1, -2), .4, LINE),
                            ("SPAN", (0, -1), (3, -1)), ("BACKGROUND", (0, -1), (-1, -1), SOFT), ("LINEABOVE", (0, -1), (-1, -1), 1, colors.black)]))

    def money_tbl(title, rows, tot, label):
        d = [[P(title, white), P("AED", whiter)]] + [[P(r["label"]), P(_money(r["amount"]), right)] for r in rows]
        d.append([P(label, bold), P(_money(tot), rightb)])
        t = Table(d, colWidths=[W * .34, W * .14])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), HEAD), ("LINEBELOW", (0, 1), (-1, -2), .4, LINE),
                               ("BACKGROUND", (0, -1), (-1, -1), SOFT), ("LINEABOVE", (0, -1), (-1, -1), 1, colors.black),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        return t

    two = Table([[money_tbl("Payable to employee", s["payables"], s["total_pay"], "Gross payable"), "",
                  money_tbl("Deductions", s["deductions"], s["total_ded"], "Total deductions")]],
                colWidths=[W * .48, W * .04, W * .48])
    two.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    net = s["net"]
    nt = Table([[P("FINAL SETTLEMENT PAYABLE" if net >= 0 else "BALANCE DUE FROM EMPLOYEE", white),
                 Paragraph(f"AED {_money(abs(net))}", ParagraphStyle("n", parent=whiter, fontSize=13, leading=16))]],
               colWidths=[W * .6, W * .4])
    nt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), HEAD), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                            ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
    notes = [Paragraph("&bull; " + escape(n), ParagraphStyle("nn", parent=mute, fontSize=7.5, leading=9.5)) for n in s["notes"]]
    decl = Table([[P(DECLARATION)]], colWidths=[W])
    decl.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .5, LINE), ("BACKGROUND", (0, 0), (-1, -1), SOFT),
                              ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    sign = Table([[P("Prepared by", ParagraphStyle("c", parent=base, alignment=1)), "", P("Approved by", ParagraphStyle("c2", parent=base, alignment=1)), "",
                   P("Employee signature & date", ParagraphStyle("c3", parent=base, alignment=1))]],
                 colWidths=[W * .28, W * .08, W * .28, W * .08, W * .28])
    sign.setStyle(TableStyle([("LINEABOVE", (0, 0), (0, 0), .8, colors.black), ("LINEABOVE", (2, 0), (2, 0), .8, colors.black),
                              ("LINEABOVE", (4, 0), (4, 0), .8, colors.black)]))
    story = [top, P("EMPLOYEE DETAILS", sec), kv_grid(_details(s)), P("SALARY AND SERVICE", sec), kv_grid(_basis(s)),
             P("END OF SERVICE GRATUITY (UAE DECREE-LAW 33/2021, ART. 51)", sec), gt, Spacer(1, 8), two, Spacer(1, 10), nt,
             Spacer(1, 6), *notes, Spacer(1, 10), KeepTogether([decl, Spacer(1, 34), sign])]
    doc.build(story)
    buf.seek(0)
    return buf


def _excel(s):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook(); ws = wb.active; ws.title = "Final Settlement"
    HEAD = PatternFill("solid", fgColor="1F2D45"); SOFT = PatternFill("solid", fgColor="F5F7FA")
    wf, bf, mf = Font(bold=True, color="FFFFFF"), Font(bold=True), Font(color="6B7280")
    thin = Side(style="thin", color="DFE3EA"); top = Border(top=Side(style="medium", color="000000"))
    for col, w in zip("ABCDE", (34, 16, 14, 16, 18)):
        ws.column_dimensions[col].width = w
    M_ = '#,##0.00'
    r = 1
    ws.cell(r, 1, "FINAL SETTLEMENT").font = Font(bold=True, size=15, color="1F2D45")
    ws.cell(r, 5, f"Settlement date: {_dmy(M._as_date(s['input']['settled_on']))}").alignment = Alignment(horizontal="right")
    r += 2

    def section(t):
        nonlocal r
        c = ws.cell(r, 1, t.upper()); c.font = Font(bold=True, color="1F2D45")
        for k in range(1, 6):
            ws.cell(r, k).border = Border(bottom=thin)
        r += 1

    def kvs(rows):
        nonlocal r
        for k, v in rows:
            ws.cell(r, 1, k).font = mf
            c = ws.cell(r, 2, v); c.font = bf
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            r += 1
        r += 1

    section("Employee details"); kvs(_details(s))
    section("Salary and service"); kvs(_basis(s))
    section("End of service gratuity (UAE Decree-Law 33/2021, art. 51)")
    for k, h in enumerate(("Period", "Years", "Days / year", "Daily basic", "Amount (AED)"), 1):
        c = ws.cell(r, k, h); c.font = wf; c.fill = HEAD
        c.alignment = Alignment(horizontal="left" if k == 1 else "right")
    r += 1
    g, d = s["gratuity"], s["salary"]["day_basic"]
    for label, yrs, per, amt in (("First 5 years", g["first5"], 21, g["first5_amt"]), ("Above 5 years", g["after5"], 30, g["after5_amt"])):
        ws.cell(r, 1, label); ws.cell(r, 2, yrs).number_format = "0.000"; ws.cell(r, 3, per)
        ws.cell(r, 4, d).number_format = M_; ws.cell(r, 5, amt).number_format = M_
        r += 1
    if g.get("cap_cut"):
        ws.cell(r, 1, "Less: limit of two years' basic"); ws.cell(r, 5, -g["cap_cut"]).number_format = M_
        r += 1
    ws.cell(r, 1, "Total gratuity" + (" (capped at 2 years basic)" if g["capped"] else "")).font = bf
    c = ws.cell(r, 5, g["amount"]); c.number_format = M_; c.font = bf
    for k in range(1, 6):
        ws.cell(r, k).fill = SOFT; ws.cell(r, k).border = top
    r += 2

    def money(title, rows, tot, label):
        nonlocal r
        for k, h in ((1, title), (5, "AED")):
            c = ws.cell(r, k, h); c.font = wf
        for k in range(1, 6):
            ws.cell(r, k).fill = HEAD
        ws.cell(r, 5).alignment = Alignment(horizontal="right")
        r += 1
        for x in rows:
            ws.cell(r, 1, x["label"]); ws.cell(r, 5, x["amount"]).number_format = M_
            r += 1
        ws.cell(r, 1, label).font = bf
        c = ws.cell(r, 5, tot); c.font = bf; c.number_format = M_
        for k in range(1, 6):
            ws.cell(r, k).fill = SOFT; ws.cell(r, k).border = top
        r += 2

    money("Payable to employee", s["payables"], s["total_pay"], "Gross payable")
    money("Deductions", s["deductions"], s["total_ded"], "Total deductions")
    net = s["net"]
    ws.cell(r, 1, "FINAL SETTLEMENT PAYABLE" if net >= 0 else "BALANCE DUE FROM EMPLOYEE").font = Font(bold=True, color="FFFFFF", size=12)
    c = ws.cell(r, 5, abs(net)); c.number_format = '"AED "#,##0.00'; c.font = Font(bold=True, color="FFFFFF", size=12)
    for k in range(1, 6):
        ws.cell(r, k).fill = HEAD
    r += 2
    for n in s["notes"]:
        ws.cell(r, 1, "- " + n).font = Font(color="6B7280", size=9)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        r += 1
    r += 1
    c = ws.cell(r, 1, DECLARATION); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5); ws.row_dimensions[r].height = 44
    r += 4
    for col, t in ((1, "Prepared by"), (3, "Approved by"), (5, "Employee signature & date")):
        c = ws.cell(r, col, t); c.border = Border(top=Side(style="thin", color="000000")); c.alignment = Alignment(horizontal="center")
    ws.page_setup.orientation = "portrait"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    try:
        export_web._excel_logo_header(ws, 3)
    except Exception:
        pass
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


def _stem(s):
    return "Final_Settlement_" + "".join(ch if ch.isalnum() else "_" for ch in s["emp"]["name"]).strip("_")[:40]


@router.get("/export/payroll/settlement")
def export_settlement(token: str, emp_no: str, format: str = "pdf", last_day: str = "", reason: str = "",
                      leave_days: str = "", salary_days: str = "", notice_required: str = "", notice_served: str = "",
                      loan: str = "", ticket: str = "", other_add: str = "", other_ded: str = "",
                      other_add_note: str = "", other_ded_note: str = "", contract: str = "", settled_on: str = "",
                      db: Session = Depends(get_db)):
    M._require_hr_reader(auth.get_download_user_from_token(token, db))
    s = settlement(db, emp_no, _q(**{k: v for k, v in locals().items() if k in FIELDS}))
    if format == "excel":
        return StreamingResponse(_excel(s), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f"attachment; filename={_stem(s)}.xlsx"})
    return StreamingResponse(_pdf(s), media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename={_stem(s)}.pdf"})


@router.get("/export/payroll/settlement/view", response_class=HTMLResponse)
def view_settlement(token: str, emp_no: str, last_day: str = "", reason: str = "",
                    leave_days: str = "", salary_days: str = "", notice_required: str = "", notice_served: str = "",
                    loan: str = "", ticket: str = "", other_add: str = "", other_ded: str = "",
                    other_add_note: str = "", other_ded_note: str = "", contract: str = "", settled_on: str = "",
                    db: Session = Depends(get_db)):
    user = auth.get_download_user_from_token(token, db)
    M._require_hr_reader(user)
    q = _q(**{k: v for k, v in locals().items() if k in FIELDS})
    s = settlement(db, emp_no, q)
    t = auth.create_view_token(user.username)
    base = "/export/payroll/settlement?" + urlencode({"emp_no": emp_no, **q, "token": t})
    return HTMLResponse(_html(s, base + "&format=pdf", base + "&format=excel"))
