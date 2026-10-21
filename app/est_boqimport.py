"""Read Infinia's internal BOQ workbook (the estimation team's priced take-off) into BOQ lines.

The file itself is never kept: it is read in memory and only the lines go into the database.

Layout it understands (T247, T252, 922 ...):
  * one sheet per trade, named "N. NAME" (1.PRELIM, 2.SUBSTR ... 20. SPECIALIZED ITEMS) - any
    extra numbered sheet added later is read the same way;
  * row 1 the bill title ("2. SUBSTRUCTURE"), then sections ("A. EARTHWORKS") each with a
    header row "S/NO. | ITEMS | CONS. QTY | IC QTY | UNIT | RATE | AMOUNT" and its items;
  * TOTAL rows, the sheet's own SUMMARY block and price comparisons (SL.NO tables) are not items;
  * an amount written as words ("Phase 02", "Excluded", "Included above", "By Client") is kept
    on the line as its remark, with no money;
  * the SUMMARY / TOTAL SUMMARY sheet gives each bill's total, used to check the read.
Lines with nothing priced (blank / zero amount and no remark) are left out - they are the
template's unused rows.
"""
import io
import re

import openpyxl

TITLE = re.compile(r"^\s*(\d{1,2})\s*\.\s*(.+)$")
SECTION = re.compile(r"^\s*([A-Z])\s*[.)]\s*(.+)$")


def _s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _n(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).replace(",", "").replace("AED", "").strip()
    try:
        return float(t)
    except ValueError:
        return None


SMALL = {"And": "and", "Of": "of", "For": "for", "To": "to", "In": "in", "Ac": "AC", "Mep": "MEP", "Rcc": "RCC", "Ss": "SS", "Pt": "PT", "Ff": "FF", "Fa": "FA"}


def _bill_name(num, text):
    words = [SMALL.get(w, w) for w in text.strip().title().split()]
    return f"{int(num):02d}. " + " ".join(words)


def _header(row):
    """Column positions from an S/NO. header row, or None."""
    first = _s(row[0]).upper() if row else ""
    if not first.startswith("S/NO") and not first.startswith("S.NO") and first not in ("SNO", "S NO"):
        return None
    cols = {}
    for i, v in enumerate(row):
        t = _s(v).upper().replace(".", "").strip()
        if not t:
            continue
        if t in ("ITEM", "ITEMS", "DESCRIPTION", "ITEM DESCRIPTION") and "desc" not in cols:
            cols["desc"] = i
        elif t in ("IC QTY", "QTY", "QUANTITY") and "qty" not in cols:
            cols["qty"] = i
        elif t == "CONS QTY":
            cols["cons"] = i
        elif t == "UNIT" and "unit" not in cols:
            cols["unit"] = i
        elif t == "RATE" and "rate" not in cols:
            cols["rate"] = i
        elif t in ("AMOUNT", "LUMPSUM AMOUNT", "LUMP SUM AMOUNT", "TOTAL AMOUNT") and "amount" not in cols:
            cols["amount"] = i
        elif t == "SIZE":
            cols["size"] = i
        elif t == "TYPE":
            cols["type"] = i
    if "desc" not in cols or "amount" not in cols:
        return None
    if "qty" not in cols and "cons" in cols:
        cols["qty"] = cols["cons"]
    return cols


def _summary(wb):
    """The workbook's own summary sheet (the one listing the most bills):
    {bill number: (amount or None, words, note)}, its total, its name."""
    best = ({}, None, "")
    for name in wb.sheetnames:
        if "SUMMARY" not in name.upper():
            continue
        out, total = {}, None
        for row in wb[name].iter_rows(max_col=5, values_only=True):
            row = list(row) + [None] * 5
            m = re.match(r"^\s*(\d{1,2})\s*\.?\s*$", _s(row[0]))
            if m and _s(row[1]):
                v = row[2]
                out.setdefault(int(m.group(1)), (_n(v), "" if _n(v) is not None else _s(v), _s(row[3])))
            elif out and total is None and "TOTAL" in _s(row[1]).upper() and _n(row[2]) is not None:
                total = _n(row[2])
        if len(out) > len(best[0]):
            best = (out, total, name)
    return best


def _sheet_summary(rows):
    """A trade sheet's own SUMMARY block: [(letter, amount or None, words)] in order."""
    out, on = [], False
    for r in rows:
        first = _s(r[0]) if r else ""
        if first.upper().startswith("SUMMARY"):
            on = True
            continue
        if not on:
            continue
        m = re.match(r"^([A-Z])\s*[.)]?$", first)
        if m:
            val = next((v for v in list(r)[2:12] if _s(v) != ""), None)
            out.append((m.group(1), _n(val), "" if _n(val) is not None else _s(val)))
    return out


OUT_WORDS = re.compile(r"(exclud|phase|n/?a$|not in|by client|by others|omit)", re.I)


def _scope_word(words, note=""):
    w = (words or note or "").strip()
    if not w or w in ("-", "0"):
        return "Not in summary"
    if re.match(r"^exclud", w, re.I) or w.lower() == "exclude":
        return "Excluded"
    if re.match(r"^phase\s*0?2", w, re.I):
        return "Phase 02"
    if re.match(r"^n/?a$", w, re.I):
        return "N/A"
    return w[:40]


def read_internal_boq(data: bytes):
    """Lines, bills and checks from an internal BOQ workbook. A line with scope "" is in the
    awarded scope and counts; anything else (Excluded, Phase 02 ...) is kept for reference
    but not counted, exactly as the workbook's summary does."""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    lines, bills, notes = [], [], []
    summary, summary_total, summary_sheet = _summary(wb)
    for ws in wb.worksheets:
        m = TITLE.match(ws.title)
        if not m:
            continue
        num = int(m.group(1))
        rows = [list(r) + [None] * 4 for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        title = next((_s(r[0]) for r in rows[:3] if TITLE.match(_s(r[0]))), "")
        tm = TITLE.match(title) if title else None
        bill = _bill_name(num, tm.group(2) if tm else m.group(2))
        # the sheet's own summary says which of its sections are left out
        ssum, sec_scope, sec_i = _sheet_summary(rows), {}, -1
        section, cols, parent, mine = "", None, "", []
        for r in rows:
            first = _s(r[0])
            h = _header(r)
            if h:
                cols, parent = h, ""
                continue
            up = first.upper()
            if up.startswith("SUMMARY"):
                break
            if up.startswith("PRICE COMPARISON") or up.startswith("SL.NO") or up.startswith("SL NO"):
                cols = None
                continue
            sm = SECTION.match(first)
            if sm and not any(_s(x) for x in r[1:8]):
                section, cols, parent = f"{sm.group(1)}. {sm.group(2).strip()}", None, ""
                sec_i += 1
                if sec_i < len(ssum):
                    _, a, w = ssum[sec_i]
                    sec_scope[section] = "" if a is not None else _scope_word(w)
                continue
            if cols is None:
                continue
            desc = _s(r[cols["desc"]])
            if not desc or re.match(r"^(sub ?total|total)\b", desc, re.I) or re.search(r"\btotal$", desc, re.I) or desc.upper().startswith("TOTAL"):
                if desc:
                    parent = ""
                continue
            qty = _n(r[cols["qty"]]) if "qty" in cols else None
            unit = _s(r[cols["unit"]]) if "unit" in cols else ""
            rate = _n(r[cols["rate"]]) if "rate" in cols else None
            raw = r[cols["amount"]]
            amt = _n(raw)
            words = "" if amt is not None else _s(raw)
            if words in ("-", "--", "0"):
                words = ""
            if not amt and not words:
                if not (qty or rate) and len(desc) > 6:
                    parent = desc
                continue
            if re.fullmatch(r"(others?|ohp)", desc, re.I) and not amt:
                continue
            if _n(unit) is not None:
                unit = ""
            if (not first or len(desc) <= 6) and parent:
                desc = f"{parent} - {desc}"
            extra = []
            if "size" in cols:
                a, b = _n(r[cols["size"]]), _n(r[cols["size"] + 1])
                if a and b:
                    extra.append(f"{a:g} x {b:g} m")
            if "type" in cols and _s(r[cols["type"]]):
                extra.append(_s(r[cols["type"]]))
            if extra:
                desc += " (" + ", ".join(extra) + ")"
            if "qty" not in cols and "rate" not in cols:       # preliminaries: a lump sum
                qty, rate, unit = (1.0, amt, "LS") if amt else (0.0, 0.0, "LS")
            mine.append({"bill": bill, "section": section, "code": first[:20], "description": desc[:500], "unit": unit[:20],
                         "qty": round(qty or 0.0, 4), "rate": round(rate or 0.0, 4), "amount": round(amt, 2) if amt else 0.0,
                         "remark": words[:120], "scope": sec_scope.get(section, "")})
        if not mine:
            continue
        s_amt, s_words, s_note = summary.get(num, (None, "", ""))
        bill_scope = "" if (s_amt is not None or not summary) else _scope_word(s_words, s_note)
        for x in mine:
            if bill_scope:
                x["scope"] = bill_scope
        counted = round(sum(x["amount"] for x in mine if not x["scope"]), 2)
        adj = 0.0
        if not bill_scope and s_amt is not None and abs(counted - round(s_amt, 2)) > 0.005:
            adj = round(s_amt - counted, 2)
            note = f" ({s_note})" if s_note else ""
            mine.append({"bill": bill, "section": "Summary", "code": "", "description": f"Difference to the summary figure{note}",
                         "unit": "", "qty": 0.0, "rate": 0.0, "amount": adj, "remark": "as per summary", "scope": ""})
            notes.append(f"{bill}: lines {counted:,.2f}, summary {s_amt:,.2f}{note} - difference {adj:,.2f} added as a line")
        lines += mine
        out = sorted({x["scope"] for x in mine if x["scope"]})
        bills.append({"bill": bill, "lines": len(mine), "counted": round(counted + adj, 2), "scope": bill_scope,
                      "left_out": round(sum(x["amount"] for x in mine if x["scope"]), 2),
                      "sections_out": [sct for sct, v in sec_scope.items() if v] if not bill_scope else [],
                      "summary": s_amt, "summary_words": s_words or s_note})
    if not lines:
        raise ValueError("No BOQ lines found. Upload the internal BOQ workbook (sheets 1.PRELIM, 2.SUBSTR ...).")
    total = round(sum(x["amount"] for x in lines if not x["scope"]), 2)
    if summary_total is not None and abs(total - round(summary_total, 2)) > 0.005:
        d = round(summary_total - total, 2)
        lines.append({"bill": "99. Summary", "section": "", "code": "", "description": "Discount / rounding as per summary total",
                      "unit": "", "qty": 0.0, "rate": 0.0, "amount": d, "remark": "as per summary", "scope": ""})
        bills.append({"bill": "99. Summary", "lines": 1, "counted": d, "scope": "", "left_out": 0.0, "sections_out": [],
                      "summary": None, "summary_words": ""})
        notes.append(f"Bills add up to {total:,.2f}, the summary total is {summary_total:,.2f} - {d:,.2f} added as a discount / rounding line")
        total = round(total + d, 2)
    return {"lines": lines, "bills": bills, "notes": notes, "summary_sheet": summary_sheet, "summary_total": summary_total,
            "total": total, "left_out": round(sum(x["amount"] for x in lines if x["scope"]), 2)}
