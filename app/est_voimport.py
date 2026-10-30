"""Reads a project's variation log out of the QS Excel file (the layout the team
keeps per project: Ref.No, Description, Submitted / Revised / Accepted /
Absorbed amount, dates, Claimed in PC, Remarks), plus the "OHP" sheet
(coordination, builders' work and attendance) and the breakdown sheets behind
absorbed items. Returns plain rows for est_variations; nothing is saved here."""
import datetime as _dt
import io
import re

import openpyxl

HEAD = {  # header text (lower case, spaces squeezed) -> our key
    "ref.no": "ref", "ref no": "ref", "ref": "ref", "phase": "phase", "description": "title", "item": "title",
    "submitted amount": "submitted_amount", "revised amount": "revised_amount", "accepted amount": "approved_amount",
    "approved amount": "approved_amount", "absorbed amount": "absorbed_amount", "date of submission": "submitted",
    "date of acceptance": "approved_on", "date of approval": "approved_on", "claimed in pc": "pc_ref", "remarks": "remarks",
    "status": "ohp_status", "claimed %": "claimed_pct",
}
MONEY = ("submitted_amount", "revised_amount", "approved_amount", "absorbed_amount")


def _txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"[ \t]+", " ", str(v)).strip()


def _num(v):
    """A money cell: a number, or text such as '*106,000.00'. Text that is not a
    plain figure ('41,900+Per day items', 'Included') gives None."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = str(v).replace(",", "").replace("AED", "").replace("*", "").strip()
    if s in ("", "-"):
        return None
    neg = s.startswith("(") and s.endswith(")")          # (5,000.00) - a minus figure
    try:
        x = round(float(s.strip("()").strip()), 2)
    except ValueError:
        return None
    return -x if neg else x


def _date(v):
    if isinstance(v, _dt.datetime):
        return v.date()
    if isinstance(v, _dt.date):
        return v
    s = _txt(v)
    m = re.match(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})$", s)
    if m:
        d, mo, y = (int(x) for x in m.groups())
        y += 2000 if y < 100 else 0
        try:
            return _dt.date(y, mo, d)
        except ValueError:
            return None
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)                         # 2026-10-21
    if m:
        try:
            return _dt.date(*(int(x) for x in m.groups()))
        except ValueError:
            return None
    for f in ("%d-%b-%Y", "%d-%b-%y", "%d %b %Y", "%d %B %Y", "%d-%B-%Y", "%d %b %y"):   # 21-Oct-2026, 21 October 2026
        try:
            return _dt.datetime.strptime(s, f).date()
        except ValueError:
            continue
    return None


def _fill_merged(ws):
    for rng in list(ws.merged_cells.ranges):
        top = ws.cell(rng.min_row, rng.min_col).value
        ws.unmerge_cells(str(rng))
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                ws.cell(r, c).value = top


def _header(ws):
    for r in range(1, min(ws.max_row, 15) + 1):
        cols = {}
        for c in range(1, ws.max_column + 1):
            k = HEAD.get(re.sub(r"\s+", " ", _txt(ws.cell(r, c).value).lower()))
            if k and k not in cols:
                cols[k] = c
        if "ref" in cols and "title" in cols:
            return r, cols
    return None, None


def _status(row):
    """The app's status from the figures and the remarks."""
    rem = (row.get("remarks") or "").lower()
    acc, absd = row.get("approved_amount") or 0, row.get("absorbed_amount") or 0
    if "reject" in rem or "not accepted" in rem:
        return "Rejected"
    if absd and not acc:
        return "Absorbed"
    if "absorbed" in rem and not acc:
        return "Absorbed"
    if acc:
        return "Approved"
    if re.search(r"accept|approved", rem) and row.get("approved_on"):
        return "Approved"
    if row.get("submitted_amount") or row.get("submitted") or row.get("revised_amount"):
        return "Submitted"
    return "Identified"


def _breakdowns(wb, skip):
    """Blocks on the side sheets: a title line, item lines, then 'Total'. Gives
    {title (lower case): 'item 11,300; item 6,800'}."""
    out = {}
    for ws in wb.worksheets:
        if ws.title in skip:
            continue
        title, items = None, []
        for r in range(1, ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column, 8) + 1)]
            a = _txt(vals[0])
            nums = [_num(v) for v in vals[1:] if _num(v) is not None]
            low = a.lower()
            if low.startswith("total") or low.startswith("grand total"):
                if title and items:
                    out.setdefault(title.lower(), "; ".join(items))
                title, items = None, []
            elif a and not nums and not any(_txt(v) for v in vals[1:]):
                if title and items:
                    items.append(a)                  # a line of the same breakdown with no figure of its own
                else:
                    title, items = a, []
            elif title and a and nums:
                items.append(f"{a} {nums[-1]:,.2f}")
    return out


def read_log(data: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    rows, used = [], set()
    main = None
    for ws in wb.worksheets:
        hr, cols = _header(ws)
        if not hr:
            continue
        is_ohp = "ohp_status" in cols or "coordination" in " ".join(_txt(ws.cell(r, 1).value).lower() for r in range(1, hr))
        if main is None and not is_ohp:
            main = ws.title
        used.add(ws.title)
        _fill_merged(ws)
        rows += (_read_ohp if is_ohp else _read_var)(ws, hr, cols)
    if not rows:
        raise ValueError("No variation log found - the sheet needs a header row with 'Ref.No' and 'Description'.")
    notes = _breakdowns(wb, used)
    for x in rows:
        b = notes.get((x["title"] or "").lower())
        if b:
            x["notes"] = (x["notes"] + "\n" if x["notes"] else "") + "Breakdown: " + b
    for i, x in enumerate(rows):
        x["seq"] = i + 1
    return rows


def _cells(ws, r, cols):
    g = lambda k: ws.cell(r, cols[k]).value if k in cols else None
    known = set(cols.values())
    extra = [_txt(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)
             if c not in known and isinstance(ws.cell(r, c).value, str) and _txt(ws.cell(r, c).value)]
    return g, extra


def _read_var(ws, hr, cols):
    out, parent, section = [], None, ""
    for r in range(hr + 1, ws.max_row + 1):
        g, extra = _cells(ws, r, cols)
        ref, title = _txt(g("ref")), _txt(g("title"))
        if not title:
            continue
        if title.lower().startswith(("previous", "total")) or title.lower() == "description" or ref.lower().startswith("ref"):
            continue                                # the header's second line, or the totals line
        raw = {k: g(k) for k in MONEY}
        money = {k: _num(v) for k, v in raw.items()}
        text_money = {k: _txt(v) for k, v in raw.items() if v is not None and _num(v) is None and _txt(v) not in ("", "-")}
        sub, acc = _date(g("submitted")), _date(g("approved_on"))
        remarks = "\n".join(x for x in [_txt(g("remarks"))] + extra if x and x != "0")
        has = any(money.values()) or text_money or sub or acc
        if not has and not ref and re.search(r"absorbed", title, re.I):
            section, parent = "absorbed", None     # "Absorbed Variations" heading: the lines under it are absorbed items
            continue
        if not has:
            nxt_ref = _txt(ws.cell(r + 1, cols["ref"]).value) if r < ws.max_row else ""
            nxt_title = _txt(ws.cell(r + 1, cols["title"]).value) if r < ws.max_row else ""
            if ref and not nxt_ref and nxt_title and section != "absorbed":
                parent = (ref, title)               # "V05 Design revision of stair case" with options under it
                continue
            if out and ref and out[-1]["ref"] == ref and not out[-1]["_has"]:
                out[-1]["notes"] = (out[-1]["notes"] + "\n" if out[-1]["notes"] else "") + title   # same VO, second line
                continue
        if ref:
            if parent and parent[0] != ref:
                parent = None
        elif parent:
            ref, title = parent[0], f"{parent[1]} - {title}"
        phase = _txt(g("phase"))
        notes = []
        if phase and phase != "-":
            notes.append(f"Phase {phase}")
        for k, label in (("submitted_amount", "Submitted"), ("revised_amount", "Revised"), ("approved_amount", "Accepted"),
                         ("absorbed_amount", "Absorbed")):
            if k in text_money:
                notes.append(f"{label}: {text_money[k]}")
        if remarks:
            notes.append(remarks)
        row = {"ref": (ref if ref != "-" else "")[:40], "title": title[:1000], "submitted": sub, "approved_on": acc,
               "pc_ref": _pcref(_txt(g("pc_ref"))), "remarks": remarks, "_has": bool(has),
               "kind": "Addition", "source": "Design change" if re.search(r"design", title, re.I) else "Client instruction"}
        for k in MONEY:
            row[k] = money[k] or 0.0
        if section == "absorbed":
            row["status"] = "Absorbed"
        else:
            row["status"] = _status(row)
        neg = [row[k] for k in MONEY if row[k]]
        if neg and all(v < 0 for v in neg):
            row["kind"] = "Omission"
        row["notes"] = "\n".join(notes)
        out.append(row)
    for x in out:
        x.pop("_has", None)
        x.pop("remarks", None)
    return out


def _pcref(s):
    return "" if s in ("", "-", "0") else s[:60]


def _read_ohp(ws, hr, cols):
    out = []
    for r in range(hr + 1, ws.max_row + 1):
        g, extra = _cells(ws, r, cols)
        title = _txt(g("title"))
        if not title or title.lower().startswith("total") or title.lower() == "description":
            continue
        amt_raw, amt = g("submitted_amount"), _num(g("submitted_amount"))
        pct = _num(g("claimed_pct"))
        notes = []
        if amt is None and _txt(amt_raw):
            notes.append(f"Amount: {_txt(amt_raw)}")
        st = _txt(g("ohp_status"))
        if pct is not None:
            notes.append(f"{st or 'Claimed'} {round(pct * 100) if pct <= 1 else round(pct)}%")
        rem = "\n".join(x for x in [_txt(g("remarks"))] + extra if x)
        if rem:
            notes.append(rem)
        phase = _txt(g("phase"))
        if phase and phase != "-":
            notes.insert(0, f"Phase {phase}")
        out.append({"ref": f"OHP-{_txt(g('ref'))}" if _txt(g("ref")) else "OHP", "title": title[:1000], "kind": "Coordination",
                    "source": "Coordination & attendance", "submitted": None, "approved_on": None, "pc_ref": "",
                    "submitted_amount": amt or 0.0, "revised_amount": 0.0, "approved_amount": 0.0,
                    "absorbed_amount": _num(g("absorbed_amount")) or 0.0,
                    "status": "Claimed" if st.lower().startswith("claim") else ("Identified" if not st else st[:20]),
                    "notes": "\n".join(notes)})
    return out
