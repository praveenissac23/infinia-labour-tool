"""The camp allocation as it looks on screen: room cards, six beds each,
on vacation / booked / free marked in colour - on one A4 sheet (Preview
and PDF) and as Jomon's room-by-room sheet in Excel."""
import io
from datetime import date
from html import escape

from fastapi.responses import HTMLResponse
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

import export_web

RED = "#B03A2E"
INK, MUTED, LINE, SOFT = "#2B2622", "#8A8079", "#E6DED6", "#FCFAF7"
TAG = {"vac": ("#FFF4E0", "#9A5B00"), "book": ("#E8F0FB", "#1F5FA8"), "soon": ("#F4EFEB", "#6B625A"),
       "guest": ("#FDECEA", RED), "free": ("#E8F5EA", "#2E7D32")}


def _d(iso, year=False):
    if not iso:
        return ""
    x = date.fromisoformat(iso)
    return x.strftime("%d/%m/%Y")


def bed_view(b):
    """(name, sub line, tag text, tag kind) for one bed."""
    k = b["kind"]
    if k == "free":
        return "Free", "", "", "free"
    if k == "vac":
        return b["name"], b.get("no", ""), "On vacation" + (f" · back {_d(b.get('back'))}" if b.get("back") else ""), "vac"
    if k == "book":
        return b["name"], b.get("trade") or "new joiner", ("Was due " if b.get("late") else "Booked · ") + (_d(b.get("due")) or "date not set"), "book"
    if k == "guest":
        return b["name"], "no ID", "No ID - check", "guest"
    return b["name"], b.get("no", ""), (f"leave {_d(b['soon'])}" if b.get("soon") else ""), "soon" if b.get("soon") else ""


def _title_case(s):
    s = s or ""
    return s.title() if s.isupper() else s


# ---- Preview (HTML, the sheet as it prints) ---------------------------------

def preview(d, pdf_url, excel_url):
    t = d["totals"]
    logo = export_web.logo_data_uri()
    tiles = "".join(f'<div class="tile {c}"><span>{lab}</span><b>{n}</b></div>' for lab, n, c in (
        ("All beds", t["beds"], ""), ("In camp", t["in"], ""), ("On vacation", t["vac"], "vac"),
        ("Booked", t["book"], "book"), ("Free beds", t["free"], "free")))
    rooms = []
    for r in d["rooms"]:
        rows = []
        for b in r["beds"]:
            name, sub, tag, kind = bed_view(b)
            rows.append(f'<div class="bed {"free" if kind == "free" else ""}"><span class="n">{b["bed"]}</span>'
                        f'<span class="who"><b>{escape(_title_case(name))}</b>'
                        f'{f"<small>{escape(sub)}</small>" if sub else ""}'
                        f'{f"<i class=t-{kind}>{escape(tag)}</i>" if tag else ""}</span></div>')
        rooms.append(f'<div class="room"><div class="rh"><b>Room {escape(r["name"])}</b>'
                     f'<span class="{"has" if r["free"] else ""}">{str(r["free"]) + " free" if r["free"] else "Full"}</span></div>{"".join(rows)}</div>')
    nb = d.get("nobed") or []
    nobed = (f'<div class="nobed"><b>No bed here ({len(nb)}):</b> ' + ", ".join(escape(f"{p['no']} {_title_case(p['name'])}") for p in nb) + "</div>") if nb else ""
    css = f"""
  body{{margin:0;background:#F1EFEA;color:{INK};font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,"Helvetica Neue",Arial,sans-serif}}
  .bar{{display:flex;align-items:center;gap:10px;flex-wrap:wrap;position:sticky;top:0;padding:10px 16px;background:#fff;border-bottom:1px solid #E2E0DC;z-index:5}}
  .bar h1{{font-size:16px;margin:0 6px 0 0}}
  a.btn{{display:inline-block;text-decoration:none;font-size:13px;font-weight:600;padding:7px 14px;border-radius:6px;border:1px solid #D9B8B3;background:#FDF4F3;color:#8C2F26}}
  a.btn.dark{{background:#2E3238;border-color:#2E3238;color:#fff}} a.btn.wa{{background:#25D366;border-color:#25D366;color:#fff}}
  .page{{width:210mm;max-width:calc(100% - 24px);margin:16px auto;background:#fff;padding:9mm 9mm 7mm;box-sizing:border-box;box-shadow:0 1px 6px rgba(0,0,0,.14)}}
  .lh{{display:flex;align-items:center;justify-content:space-between;border-bottom:2px solid {RED};padding-bottom:4mm;margin-bottom:4mm}}
  .lh img{{width:40mm}} .lh .r{{text-align:right}} .lh .co{{font-weight:700;font-size:12pt}} .lh .ti{{font-size:10pt;color:{INK};font-weight:700;margin-top:2px}} .lh .dt{{font-size:8.5pt;color:{MUTED};margin-top:2px}}
  .tiles{{display:grid;grid-template-columns:repeat(5,1fr);gap:2.5mm;margin-bottom:4mm}}
  .tile{{border:1px solid {LINE};border-radius:3mm;padding:2mm 3mm}} .tile span{{display:block;font-size:6.8pt;letter-spacing:.06em;text-transform:uppercase;color:{MUTED}}} .tile b{{font-size:15pt}}
  .tile.vac b{{color:#B26A00}} .tile.book b{{color:#1F5FA8}} .tile.free b{{color:#2E7D32}}
  .grid{{display:grid;grid-template-columns:repeat(3,1fr);gap:2.5mm}}
  .room{{border:1px solid {LINE};border-radius:3mm;overflow:hidden;break-inside:avoid}}
  .rh{{display:flex;justify-content:space-between;align-items:center;padding:1.6mm 3mm;background:{SOFT};border-bottom:1px solid {LINE}}}
  .rh b{{font-size:9.5pt}} .rh span{{font-size:7.5pt;color:{MUTED}}} .rh span.has{{color:#2E7D32;font-weight:700}}
  .bed{{display:flex;align-items:center;gap:2mm;padding:.9mm 3mm;min-height:8.1mm;box-sizing:border-box;border-bottom:1px solid #F3EEE9}} .bed:last-child{{border-bottom:0}}
  .n{{width:3.8mm;height:3.8mm;border-radius:50%;background:#F1ECE6;color:{MUTED};font-size:6.2pt;display:flex;align-items:center;justify-content:center;flex:none}}
  .bed.free .n{{background:#E8F5EA;color:#2E7D32}} .bed.free b{{color:#2E7D32;font-weight:500}}
  .who{{min-width:0;flex:1;line-height:1.25;display:block;font-size:6.4pt}} .who b{{display:block;font-size:7.6pt;line-height:1.25;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}}
  .who small{{font-size:6.4pt;color:{MUTED}}} .who i{{font-style:normal;font-size:6pt;font-weight:700;border-radius:2mm;padding:0 1.4mm;margin-left:1.4mm}}
  {"".join(f".t-{k}{{background:{bg};color:{fg}}}" for k, (bg, fg) in TAG.items())}
  .legend{{display:flex;gap:5mm;font-size:7pt;color:{MUTED};margin-top:3mm;flex-wrap:wrap}} .legend i{{display:inline-block;width:2.4mm;height:2.4mm;border-radius:.6mm;margin-right:1.2mm;vertical-align:-.3mm}}
  .nobed{{font-size:7pt;color:{MUTED};margin-top:2mm}}
  @media print{{.bar{{display:none}} body{{background:#fff}} .page{{box-shadow:none;margin:0;max-width:none}}}} @page{{size:A4 portrait;margin:0}}
"""
    html = f"""<!doctype html><html><head><meta charset="utf-8"><title>Labour Accommodation Allocation</title>
<meta name="viewport" content="width=820, minimum-scale=0.3, maximum-scale=4, user-scalable=yes"><style>{css}</style></head><body>
<div class="bar"><h1>Labour Accommodation Allocation</h1>
<a class="btn dark" href="{escape(pdf_url)}">Download PDF</a><a class="btn" href="{escape(excel_url)}">Download Excel</a>
{export_web.share_button(pdf_url, "Labour Accommodation Allocation")}</div>
<div class="page"><div class="lh">{f'<img src="{logo}" alt="Infinia">' if logo else '<b>INFINIA</b>'}
<div class="r"><div class="co">INFINIA CONTRACTING LLC</div><div class="ti">Labour Accommodation Allocation</div><div class="dt">Camp · {len(d["rooms"])} rooms · as on {_d(d["today"], True)}</div></div></div>
<div class="tiles">{tiles}</div><div class="grid">{"".join(rooms)}</div>
<div class="legend"><span><i style="background:{INK}"></i>In camp</span><span><i style="background:#F0B44C"></i>On vacation - bed kept</span><span><i style="background:#5B8FD6"></i>Booked for a new joiner</span><span><i style="background:#7CC08A"></i>Free</span><span><i style="background:{RED}"></i>No ID - check</span></div>
{nobed}</div>{export_web.SHARE_JS}</body></html>"""
    return HTMLResponse(html)


# ---- PDF (same sheet) ---------------------------------------------------------

def pdf(d):
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setTitle("Labour Accommodation Allocation"); c.setAuthor("Infinia Contracting LLC")
    W, H = A4
    M = 10 * mm
    CW = W - 2 * M
    col = lambda h: HexColor(h)

    def head():
        y = H - M
        try:
            import os
            logo = ImageReader(os.path.join(os.path.dirname(os.path.abspath(__file__)), "logo.png"))
            iw, ih = logo.getSize(); h = 10 * mm; w = h * iw / ih
            c.drawImage(logo, M, y - h, w, h, mask="auto")
        except Exception:
            c.setFont("Helvetica-Bold", 14); c.drawString(M, y - 8 * mm, "INFINIA")
        c.setFillColor(col(INK)); c.setFont("Helvetica-Bold", 11.5); c.drawRightString(W - M, y - 4 * mm, "INFINIA CONTRACTING LLC")
        c.setFillColor(col(RED)); c.setFont("Helvetica-Bold", 10); c.drawRightString(W - M, y - 8.5 * mm, "Labour Accommodation Allocation")
        c.setFillColor(col(MUTED)); c.setFont("Helvetica", 7.8)
        c.drawRightString(W - M, y - 12.3 * mm, f"Camp · {len(d['rooms'])} rooms · as on {_d(d['today'], True)}")
        c.setStrokeColor(col(RED)); c.setLineWidth(1.4); c.line(M, y - 14.5 * mm, W - M, y - 14.5 * mm)
        return y - 18 * mm

    def foot(page):
        c.setFillColor(col(MUTED)); c.setFont("Helvetica", 7)
        c.drawString(M, 6 * mm, "Infinia Contracting LLC  |  Labour accommodation")
        c.drawRightString(W - M, 6 * mm, f"Page {page}")

    y = head()
    t = d["totals"]
    # tiles
    gap = 2.5 * mm; tw = (CW - 4 * gap) / 5; th = 12 * mm
    for i, (lab, n, fg) in enumerate((("ALL BEDS", t["beds"], INK), ("IN CAMP", t["in"], INK), ("ON VACATION", t["vac"], "#B26A00"),
                                       ("BOOKED", t["book"], "#1F5FA8"), ("FREE BEDS", t["free"], "#2E7D32"))):
        x = M + i * (tw + gap)
        c.setStrokeColor(col(LINE)); c.setFillColor(white); c.setLineWidth(.7); c.roundRect(x, y - th, tw, th, 2.5 * mm, fill=1, stroke=1)
        c.setFillColor(col(MUTED)); c.setFont("Helvetica", 6.3); c.drawString(x + 3 * mm, y - 4 * mm, lab)
        c.setFillColor(col(fg)); c.setFont("Helvetica-Bold", 14); c.drawString(x + 3 * mm, y - 10 * mm, str(n))
    y -= th + 4 * mm
    # room cards, 3 across
    cols = 3; cg = 2.5 * mm; cw = (CW - (cols - 1) * cg) / cols
    rh_h, bed_h = 6.2 * mm, 6.05 * mm
    page = 1
    rooms = d["rooms"]
    for i in range(0, len(rooms), cols):
        row = rooms[i:i + cols]
        card_h = rh_h + bed_h * max(len(r["beds"]) for r in row)
        if y - card_h < 22 * mm:
            foot(page); c.showPage(); page += 1; y = head()
        for j, r in enumerate(row):
            x = M + j * (cw + cg); top = y
            h = rh_h + bed_h * len(r["beds"])
            c.setStrokeColor(col(LINE)); c.setFillColor(white); c.setLineWidth(.7)
            c.roundRect(x, top - h, cw, h, 2.5 * mm, fill=1, stroke=1)
            # header band
            p = c.beginPath(); rr = 2.5 * mm
            p.moveTo(x, top - rh_h); p.lineTo(x, top - rr); p.arcTo(x, top - 2 * rr, x + 2 * rr, top, 180, -90)
            p.lineTo(x + cw - rr, top); p.arcTo(x + cw - 2 * rr, top - 2 * rr, x + cw, top, 90, -90); p.lineTo(x + cw, top - rh_h); p.close()
            c.setFillColor(col(SOFT)); c.drawPath(p, fill=1, stroke=0)
            c.setStrokeColor(col(LINE)); c.line(x, top - rh_h, x + cw, top - rh_h)
            c.setFillColor(col(INK)); c.setFont("Helvetica-Bold", 9); c.drawString(x + 3 * mm, top - 4.4 * mm, f"Room {r['name']}")
            if r["free"]:
                c.setFillColor(col("#2E7D32")); c.setFont("Helvetica-Bold", 7.2); c.drawRightString(x + cw - 3 * mm, top - 4.3 * mm, f"{r['free']} free")
            else:
                c.setFillColor(col(MUTED)); c.setFont("Helvetica", 7.2); c.drawRightString(x + cw - 3 * mm, top - 4.3 * mm, "Full")
            by = top - rh_h
            for k, b in enumerate(r["beds"]):
                name, sub, tag, kind = bed_view(b)
                cy = by - bed_h / 2
                if k:
                    c.setStrokeColor(col("#F1ECE6")); c.line(x + 2 * mm, by, x + cw - 2 * mm, by)
                bg, fg = TAG.get("free" if kind == "free" else "x", ("#F1ECE6", MUTED))
                c.setFillColor(col(bg)); c.circle(x + 4.6 * mm, cy, 1.85 * mm, fill=1, stroke=0)
                c.setFillColor(col(fg)); c.setFont("Helvetica", 6); c.drawCentredString(x + 4.6 * mm, cy - 2.1, str(b["bed"]))
                tx = x + 8 * mm
                if kind == "free":
                    c.setFillColor(col("#2E7D32")); c.setFont("Helvetica", 7.6); c.drawString(tx, cy - 2.6, "Free")
                else:
                    nm = _title_case(name)
                    c.setFillColor(col(MUTED if kind == "vac" else INK)); c.setFont("Helvetica-Bold", 7.4)
                    while c.stringWidth(nm, "Helvetica-Bold", 7.4) > cw - 11 * mm and len(nm) > 3:
                        nm = nm[:-2]
                    if nm != _title_case(name):
                        nm = nm.rstrip() + "…"
                    c.drawString(tx, cy + 0.6 * mm, nm)
                    c.setFillColor(col(MUTED)); c.setFont("Helvetica", 6.2); c.drawString(tx, cy - 2.2 * mm, sub)
                    if tag:
                        sx = tx + c.stringWidth(sub, "Helvetica", 6.2) + 1.6 * mm
                        tbg, tfg = TAG[kind]
                        w = c.stringWidth(tag, "Helvetica-Bold", 5.8) + 2.6 * mm
                        c.setFillColor(col(tbg)); c.roundRect(sx, cy - 2.85 * mm, w, 3 * mm, 1.4 * mm, fill=1, stroke=0)
                        c.setFillColor(col(tfg)); c.setFont("Helvetica-Bold", 5.8); c.drawString(sx + 1.3 * mm, cy - 2.05 * mm, tag)
                by -= bed_h
        y -= card_h + cg
    # legend + no bed
    y -= 1 * mm
    x = M
    c.setFont("Helvetica", 7)
    for lab, colr in (("In camp", INK), ("On vacation - bed kept", "#F0B44C"), ("Booked for a new joiner", "#5B8FD6"), ("Free", "#7CC08A"), ("No ID - check", RED)):
        c.setFillColor(col(colr)); c.roundRect(x, y - 0.4 * mm, 2.4 * mm, 2.4 * mm, .6 * mm, fill=1, stroke=0)
        c.setFillColor(col(MUTED)); c.drawString(x + 3.4 * mm, y, lab); x += c.stringWidth(lab, "Helvetica", 7) + 9 * mm
    nb = d.get("nobed") or []
    if nb:
        y -= 5 * mm
        from reportlab.lib.utils import simpleSplit
        txt = f"No bed here ({len(nb)}): " + ", ".join(f"{p['no']} {_title_case(p['name'])}" for p in nb)
        for line in simpleSplit(txt, "Helvetica", 7, CW)[:3]:
            c.setFillColor(col(MUTED)); c.drawString(M, y, line); y -= 3.6 * mm
    foot(page)
    c.save()
    buf.seek(0)
    return buf


# ---- Excel (Jomon's room-by-room sheet) ------------------------------------

def excel(d):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook(); ws = wb.active; ws.title = "Accommodation"
    thin = Side(style="thin", color="BFB6AE"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.merge_cells("A1:E1"); ws["A1"] = "Infinia Contracting LLC  -  Labour Accommodation Allocation"
    ws["A1"].font = Font(bold=True, size=13, color="B03A2E")
    t = d["totals"]
    ws.merge_cells("A2:E2")
    ws["A2"] = (f"As on {_d(d['today'], True)}  |  {t['beds']} beds  |  in camp {t['in']}  |  on vacation {t['vac']}"
                f"  |  booked {t['book']}  |  free {t['free']}")
    ws["A2"].font = Font(size=9, color="6B625A")
    hdr = ["Room no", "Bed", "Name", "Employee ID", "Status"]
    for i, h in enumerate(hdr, 1):
        cl = ws.cell(row=4, column=i, value=h)
        cl.font = Font(bold=True, color="FFFFFF"); cl.fill = PatternFill("solid", fgColor="B03A2E")
        cl.alignment = Alignment(horizontal="center", vertical="center"); cl.border = box
    fills = {"vac": "FFF4E0", "book": "E8F0FB", "guest": "FDECEA", "free": "E8F5EA"}
    row = 5
    for r in d["rooms"]:
        start = row
        for b in r["beds"]:
            name, sub, tag, kind = bed_view(b)
            vals = [r["name"], b["bed"], "FREE" if kind == "free" else _title_case(name),
                    b.get("no", "") if kind not in ("book", "guest") else "", tag or ("In camp" if kind not in ("free",) else "")]
            for ci, v in enumerate(vals, 1):
                cl = ws.cell(row=row, column=ci, value=v); cl.border = box
                cl.alignment = Alignment(horizontal="center" if ci in (1, 2, 4) else "left", vertical="center")
                if kind in fills and ci > 1:
                    cl.fill = PatternFill("solid", fgColor=fills[kind])
            row += 1
        ws.merge_cells(start_row=start, start_column=1, end_row=row - 1, end_column=1)
        ws.cell(row=start, column=1).font = Font(bold=True, size=12)
    ws.cell(row=row + 1, column=3, value=f"Total vacant beds: {t['free']}").font = Font(bold=True)
    for col, w in zip("ABCDE", (10, 6, 30, 14, 26)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "portrait"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4; ws.print_title_rows = "4:4"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf
