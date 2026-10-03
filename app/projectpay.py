"""Project payment tracker - Accounts > Project payments.

Each project is split into the scopes it was contracted for (Design &
Build, Mobilization, Core & Shell, MEP ...). A scope holds its contract
value and every payment received against it, with the date. What has
been received, what remains and the status are worked out from those -
never typed - so they cannot disagree with the payments.

Two papers, from the same data:
  Summary            one line per scope: contract, received, remaining,
                     how much is collected, the last payment, status.
  With payment dates the same, with every payment and its date.
Both group the scopes under their project, with a project total where a
project has more than one scope, and open with the four figures a
manager asks first: contract value, received, remaining, % collected.

Admin and the chief accountant only (right: accounts_projects).
"""
import io
import re
from datetime import date
from html import escape
from urllib.parse import urlencode

from fastapi import APIRouter, Body, Depends, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse
from sqlalchemy import Column, Date, DateTime, Float, ForeignKey, Integer, String, Text, func
from sqlalchemy.orm import Session, relationship

import main as M
import auth, export_web, models
from database import Base, get_db

router = APIRouter()
RIGHT = "accounts_projects"
PP = Depends(M.require_screen(RIGHT))
COMPANY = "INFINIA CONTRACTING L.L.C."
RED, DARK, MUTED = "#B7322A", "#22262A", "#6B7178"
STATUS_COLOURS = {"Completed": ("#E7F4EA", "#1E7B34"), "Partial": ("#FFF3DC", "#9A5B00"), "Pending": ("#FDE8E7", "#B42318")}


class ProjectScope(Base):
    """One contracted scope of a project, e.g. 914 - CORE & SHELL."""
    __tablename__ = "project_scopes"
    id = Column(Integer, primary_key=True)
    project = Column(String, nullable=False, index=True)       # 914
    scope = Column(String, nullable=False)                      # CORE & SHELL
    contract = Column(Float, default=0.0)                       # AED
    remarks = Column(Text, default="")
    sort = Column(Integer, default=0)
    created_by = Column(Integer, ForeignKey("users.id"), nullable=True)
    updated_by = Column(Integer, ForeignKey("users.id"), nullable=True)
    created_at = Column(DateTime(timezone=True), server_default=func.now())
    updated_at = Column(DateTime(timezone=True), onupdate=func.now())
    payments = relationship("ProjectPayment", cascade="all, delete-orphan", order_by="ProjectPayment.on_date, ProjectPayment.id")


class ProjectPayment(Base):
    """A payment received against a scope."""
    __tablename__ = "project_payments"
    id = Column(Integer, primary_key=True)
    scope_id = Column(Integer, ForeignKey("project_scopes.id"), nullable=False, index=True)
    on_date = Column(Date, nullable=False)
    amount = Column(Float, default=0.0)
    note = Column(String, default="")


def create_tables(engine):
    for t in (ProjectScope.__table__, ProjectPayment.__table__):
        t.create(bind=engine, checkfirst=True)


def seed_once(SessionLocal):
    """The tracker as it stood on 02-Oct-2026 (PROJECT_PAYMENT_TRACKER_
    02.10.26), put in once on an empty table - checked line by line
    against that sheet's Total Received column. Never again after that,
    so nothing typed in the app is ever overwritten."""
    import json, os
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "deploy", "project_payments.json")
    db = SessionLocal()
    try:
        if db.query(models.Setting).filter(models.Setting.key == "project_payments_seeded").first():
            return
        if not db.query(ProjectScope).first() and os.path.exists(path):
            for x in json.load(open(path)):
                s = ProjectScope(project=x["project"], scope=x["scope"], contract=x["contract"], remarks=x["remarks"], sort=x["sort"])
                s.payments = [ProjectPayment(on_date=date.fromisoformat(p["date"]), amount=p["amount"]) for p in x["payments"]]
                db.add(s)
            print("Project payment tracker: loaded the 02-Oct-2026 sheet")
        db.add(models.Setting(key="project_payments_seeded", value="1"))
        db.commit()
    except Exception as e:                       # a failed seed must never stop the app
        db.rollback()
        print(f"Project payment tracker seed skipped: {e}")
    finally:
        db.close()


# ---- the figures --------------------------------------------------------------

def _status(contract, received):
    if received <= 0.005:
        return "Pending"
    return "Completed" if received >= contract - 0.005 else "Partial"


def _scope_out(s, start=None, end=None):
    """A scope's figures. With a period: the payments inside it, and the
    position (received, remaining, status) as at the period's end."""
    pays = [{"id": p.id, "date": p.on_date.isoformat(), "amount": round(p.amount or 0, 2), "note": p.note or ""} for p in s.payments]
    upto = [p for p in pays if not end or p["date"] <= end.isoformat()]
    inside = [p for p in upto if not start or p["date"] >= start.isoformat()]
    rec = round(sum(p["amount"] for p in upto), 2)
    con = round(s.contract or 0, 2)
    last = upto[-1] if upto else None
    return {"id": s.id, "project": s.project, "scope": s.scope, "contract": con, "remarks": s.remarks or "", "sort": s.sort or 0,
            "payments": pays, "period_payments": inside, "in_period": round(sum(p["amount"] for p in inside), 2),
            "received": rec, "remaining": round(con - rec, 2),
            "pct": round(100 * rec / con, 1) if con else 0.0, "status": _status(con, rec),
            "last_date": last["date"] if last else "", "last_amount": last["amount"] if last else 0.0}


def period_label(start, end):
    """'September 2026', 'Year 2026', '01-Sep-26 to 15-Sep-26' ..."""
    if not start and not end:
        return ""
    if start and end:
        nxt = date(end.year + (end.month == 12), end.month % 12 + 1, 1)
        month_end = end == date.fromordinal(nxt.toordinal() - 1)
        if start.day == 1 and month_end and start.year == end.year and start.month == end.month:
            return f"{start:%B %Y}"
        if (start.month, start.day) == (1, 1) and (end.month, end.day) == (12, 31) and start.year == end.year:
            return f"Year {start.year}"
        return f"{start:%d-%b-%y} to {end:%d-%b-%y}"
    return f"From {start:%d-%b-%y}" if start else f"Up to {end:%d-%b-%y}"


def tracker(db, hide_completed=False, q="", start=None, end=None):
    if start and end and start > end:
        raise HTTPException(status_code=400, detail="The From date is after the To date.")
    period = bool(start or end)
    scopes = db.query(ProjectScope).all()
    rows = [_scope_out(s, start, end) for s in scopes]
    # Projects in the order they were first listed; scopes in theirs.
    first = {}
    for r in sorted(rows, key=lambda r: (r["sort"], r["id"])):
        first.setdefault(r["project"], (r["sort"], r["id"]))
    rows.sort(key=lambda r: (first[r["project"]], r["sort"], r["id"]))
    totals = _totals(rows)                     # the whole company, whatever is shown
    shown = rows
    if period:
        # A month / year / range lists the scopes paid in it.
        shown = [r for r in shown if r["period_payments"]]
    if hide_completed:
        shown = [r for r in shown if r["status"] != "Completed"]
    q = (q or "").strip().lower()
    if q:
        shown = [r for r in shown if q in f"{r['project']} {r['scope']} {r['remarks']}".lower()]
    groups = []
    for r in shown:
        if not groups or groups[-1]["project"] != r["project"]:
            groups.append({"project": r["project"], "rows": []})
        groups[-1]["rows"].append(r)
    for g in groups:
        g.update(_totals(g["rows"]))
    as_on = end or M._dubai_today()
    return {"groups": groups, "rows": shown, "totals": totals, "shown": _totals(shown),
            "as_on": min(as_on, M._dubai_today()).isoformat(), "hide_completed": bool(hide_completed), "q": q,
            "period": period, "period_label": period_label(start, end),
            "from": start.isoformat() if start else "", "to": end.isoformat() if end else ""}


def _totals(rows):
    c = round(sum(r["contract"] for r in rows), 2)
    rec = round(sum(r["received"] for r in rows), 2)
    return {"contract": c, "received": rec, "remaining": round(c - rec, 2), "pct": round(100 * rec / c, 1) if c else 0.0,
            "in_period": round(sum(r["in_period"] for r in rows), 2),
            "payments": sum(len(r["period_payments"]) for r in rows),
            "count": len(rows), "projects": len({r["project"] for r in rows}),
            "completed": sum(1 for r in rows if r["status"] == "Completed"),
            "partial": sum(1 for r in rows if r["status"] == "Partial"),
            "pending": sum(1 for r in rows if r["status"] == "Pending")}


@router.get("/employees/accounts/projects")
def get_tracker(hide_completed: int = 0, q: str = "", start: str = "", end: str = "", db: Session = Depends(get_db), user: models.User = PP):
    return tracker(db, bool(hide_completed), q, M._as_date(start), M._as_date(end))


def _clean(p):
    project = " ".join(str(p.get("project") or "").split())[:20]
    scope = " ".join(str(p.get("scope") or "").split())[:80]
    if not project:
        raise HTTPException(status_code=400, detail="Which project? Type the project number.")
    if not scope:
        raise HTTPException(status_code=400, detail="Type the scope (e.g. CORE & SHELL).")
    try:
        contract = round(float(str(p.get("contract") or 0).replace(",", "")), 2)
    except ValueError:
        raise HTTPException(status_code=400, detail="The contract value must be a number.")
    if contract <= 0:
        raise HTTPException(status_code=400, detail="Type the contract value.")
    pays = []
    today = M._dubai_today()
    for i, x in enumerate(p.get("payments") or [], 1):
        d = M._as_date(x.get("date"))
        raw = str(x.get("amount") or "").replace(",", "").strip()
        if not d and not raw:
            continue                                   # an empty row left in the form
        if not d:
            raise HTTPException(status_code=400, detail=f"Payment {i}: put the date.")
        if d > today:
            raise HTTPException(status_code=400, detail=f"Payment {i}: {d:%d-%b-%Y} is in the future.")
        try:
            amt = round(float(raw), 2)
        except ValueError:
            raise HTTPException(status_code=400, detail=f"Payment {i}: the amount must be a number.")
        if amt <= 0:
            raise HTTPException(status_code=400, detail=f"Payment {i}: type the amount.")
        pays.append((d, amt, " ".join(str(x.get("note") or "").split())[:120]))
    if round(sum(a for _, a, _ in pays), 2) > contract + 0.005:
        raise HTTPException(status_code=400, detail="The payments come to more than the contract value - check the amounts or the contract value.")
    return {"project": project, "scope": scope, "contract": contract,
            "remarks": " ".join(str(p.get("remarks") or "").split())[:300]}, sorted(pays, key=lambda t: t[0])


def _save(db, s, payload, user):
    fields, pays = _clean(payload)
    for k, v in fields.items():
        setattr(s, k, v)
    s.payments = [ProjectPayment(on_date=d, amount=a, note=n) for d, a, n in pays]
    s.updated_by = user.id


@router.post("/employees/accounts/projects")
def add_scope(payload: dict = Body(...), db: Session = Depends(get_db), user: models.User = PP):
    s = ProjectScope(created_by=user.id)
    _save(db, s, payload, user)
    # A new scope of an existing project joins it; a new project goes last.
    same = db.query(func.max(ProjectScope.sort)).filter(ProjectScope.project == s.project).scalar()
    s.sort = (same if same is not None else (db.query(func.max(ProjectScope.sort)).scalar() or 0)) + 1
    if same is not None:
        for o in db.query(ProjectScope).filter(ProjectScope.sort >= s.sort).all():
            o.sort += 1
    db.add(s); db.commit()
    M.log_action(db, user.id, "project_payment", f"{s.project} {s.scope} added")
    return {"ok": True, "id": s.id}


@router.put("/employees/accounts/projects/{sid}")
def edit_scope(sid: int, payload: dict = Body(...), db: Session = Depends(get_db), user: models.User = PP):
    s = db.get(ProjectScope, sid)
    if not s:
        raise HTTPException(status_code=404, detail="That scope is no longer there.")
    _save(db, s, payload, user)
    db.commit()
    M.log_action(db, user.id, "project_payment", f"{s.project} {s.scope} changed")
    return {"ok": True}


@router.delete("/employees/accounts/projects/{sid}")
def delete_scope(sid: int, db: Session = Depends(get_db), user: models.User = PP):
    s = db.get(ProjectScope, sid)
    if not s:
        raise HTTPException(status_code=404, detail="That scope is no longer there.")
    M.log_action(db, user.id, "project_payment", f"{s.project} {s.scope} removed")
    db.delete(s); db.commit()
    return {"ok": True}


# ---- the papers -----------------------------------------------------------------

def _m(v):
    return f"{v:,.2f}"


def _dmy(iso):
    d = M._as_date(iso) if iso else None
    return f"{d:%d-%b-%y}" if d else ""


def _title(detail):
    return "With payment dates" if detail else "Summary"


def _sub(r, detail):
    bits = [_title(detail), r["period_label"] or "Full", f"As on {_dmy(r['as_on'])}"]
    if r["hide_completed"]:
        bits.append("completed scopes not shown")
    if r["q"]:
        bits.append(f"filtered: {r['q']}")
    return " · ".join(bits)


def _cols(r, detail):
    """The columns of the paper: (key, heading, pdf width mm or 0 = the rest).
    Full: received to date and the last payment. A period: what came in
    during it, then where each scope stands at its end."""
    per = r["period"]
    cols = [("proj", "Project", 15), ("scope", "Scope", 30 if detail else 36), ("contract", "Contract (AED)", 23)]
    if detail:
        cols.append(("pays", "Payments in period (date · AED)" if per else "Payments received (date · AED)", 62 if per else 70))
    if per:
        cols.append(("inper", "Received in period (AED)", 24))
    cols += [("received", "Received to date (AED)" if per else "Received (AED)", 24), ("remaining", "Remaining (AED)", 24),
             ("pct", "Collected", 23)]
    if not detail and not per:
        cols.append(("last", "Last payment", 36))
    cols += [("status", "Status", 19), ("remarks", "Remarks", 0)]
    return cols


NUMS = ("contract", "inper", "received", "remaining", "last")


def _kpis(r):
    """The four figures on top: (label, value, note, colour, bar%)."""
    t = r["shown"]
    if r["period"]:
        return [(f"Received - {r['period_label']}", f"AED {_m(t['in_period'])}", f"{t['payments']} payments · {t['projects']} projects", "#1E7B34", None),
                ("Contract value (these scopes)", f"AED {_m(t['contract'])}", f"{t['count']} scopes", "#1d1d1d", None),
                (f"Remaining as at {_dmy(r['as_on'])}", f"AED {_m(t['remaining'])}", f"received to date AED {_m(t['received'])}", RED, None),
                (f"Collected as at {_dmy(r['as_on'])}", f"{t['pct']:.1f}%", "", "#1d1d1d", t["pct"])]
    return [("Total contract value", f"AED {_m(t['contract'])}", f"{t['projects']} projects · {t['count']} scopes", "#1d1d1d", None),
            ("Received", f"AED {_m(t['received'])}", f"{t['completed']} scopes completed", "#1E7B34", None),
            ("Remaining", f"AED {_m(t['remaining'])}", f"{t['partial']} partly paid · {t['pending']} not started", RED, None),
            ("Collected", f"{t['pct']:.1f}%", "", "#1d1d1d", t["pct"])]


def _pays_of(r, x):
    return x["period_payments"] if r["period"] else x["payments"]


def _html(r, detail, pdf_url, excel_url):
    logo = export_web.logo_data_uri()
    cols = _cols(r, detail)

    def pill(s):
        bg, fg = STATUS_COLOURS[s]
        return f'<span class="pill" style="background:{bg};color:{fg}">{s}</span>'

    def bar(p):
        return f'<span class="bar"><i style="width:{max(0, min(100, p))}%"></i></span><span class="pct">{p:.0f}%</span>'

    def pays(x):
        ps = _pays_of(r, x)
        if not ps:
            return '<span class="none">No payment yet</span>'
        return '<div class="pays">' + "".join(f'<span><em>{_dmy(p["date"])}</em>{_m(p["amount"])}</span>' for p in ps) + "</div>"

    def cell(k, x, g=None, i=0, n=1):
        if k == "proj":
            return f'<td class="proj" rowspan="{n + (1 if n > 1 else 0)}">{escape(g["project"])}</td>' if i == 0 else ""
        if k == "scope":
            return f'<td class="l scope">{escape(x["scope"])}</td>'
        if k == "pays":
            return f'<td class="l">{pays(x)}</td>'
        if k == "inper":
            return f'<td class="n b g">{_m(x["in_period"])}</td>'
        if k == "received":
            return f'<td class="n{"" if r["period"] else " b"}">{_m(x["received"])}</td>'
        if k == "remaining":
            return f'<td class="n{" due" if x["remaining"] > 0.005 else ""}">{_m(x["remaining"])}</td>'
        if k == "contract":
            return f'<td class="n">{_m(x["contract"])}</td>'
        if k == "pct":
            return f'<td class="c">{bar(x["pct"])}</td>'
        if k == "last":
            return (f'<td class="n">{_m(x["last_amount"])}<small>{_dmy(x["last_date"])}</small></td>' if x["last_date"]
                    else '<td class="n"><span class="none">-</span></td>')
        if k == "status":
            return f'<td class="c">{pill(x["status"])}</td>'
        return f'<td class="l rem">{escape(x["remarks"])}</td>'

    def sumrow(cls, label, t, colspan_label):
        out = []
        for k, _, _ in cols:
            if k == "proj" and colspan_label == 2:
                out.append(f'<td colspan="2" class="l">{label}</td>')
            elif k == "proj":
                continue
            elif k == "scope":
                if colspan_label != 2:
                    out.append(f'<td class="l">{label}</td>')
            elif k in ("contract", "received", "remaining"):
                out.append(f'<td class="n">{_m(t[k])}</td>')
            elif k == "inper":
                out.append(f'<td class="n">{_m(t["in_period"])}</td>')
            elif k == "pct":
                out.append(f'<td class="c">{bar(t["pct"])}</td>')
            else:
                out.append("<td></td>")
        return f'<tr class="{cls}">' + "".join(out) + "</tr>"

    head = "".join(f'<th class="{"n" if k in NUMS else ("l" if k in ("scope", "remarks", "pays") else "c")}">{escape(h)}</th>' for k, h, _ in cols)
    body = []
    for g in r["groups"]:
        n = len(g["rows"])
        for i, x in enumerate(g["rows"]):
            body.append(f'<tr class="{"first" if i == 0 else ""}">' + "".join(cell(k, x, g, i, n) for k, _, _ in cols) + "</tr>")
        if n > 1:
            body.append(sumrow("sub", f"Project {escape(g['project'])} total", g, 1))
    s = r["shown"]
    foot = sumrow("", "TOTAL" + ("" if s["count"] == r["totals"]["count"] or r["period"] else " (shown)"), s, 2) if r["rows"] else ""
    empty = "" if r["rows"] else f'<p class="nothing">No payments in {escape(r["period_label"] or "this selection")}.</p>'
    kpi = '<div class="kpis">' + "".join(
        f'<div><span>{escape(lab)}</span><b style="color:{col}">{escape(val)}</b>'
        + (f'<span class="bar big"><i style="width:{min(100, b)}%"></i></span>' if b is not None else f"<small>{escape(note)}</small>") + "</div>"
        for lab, val, note, col, b in _kpis(r)) + "</div>"
    return f"""<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Project Payment Tracker - {_title(detail)}</title><style>
*{{box-sizing:border-box}} body{{margin:0;background:#ECEEF1;font:12.5px/1.4 Arial,Helvetica,sans-serif;color:#1d1d1d}}
.page{{max-width:1280px;margin:18px auto;background:#fff;padding:28px 32px;box-shadow:0 2px 10px rgba(0,0,0,.08)}}
.top{{display:flex;justify-content:space-between;align-items:flex-end;border-bottom:2px solid {RED};padding-bottom:10px}}
.top img{{height:32px}} h1{{margin:0;font-size:19px;letter-spacing:.08em;color:{RED};text-align:right}}
.top p{{margin:3px 0 0;text-align:right;color:{MUTED};font-size:12px}}
.kpis{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin:16px 0 14px}}
.kpis > div{{border:1px solid #E6E1DC;border-radius:8px;padding:10px 14px;background:#FCFAF8}}
.kpis > div > span:first-child{{display:block;font-size:10.5px;color:{MUTED};text-transform:uppercase;letter-spacing:.06em}}
.kpis b{{display:block;font-size:17px;margin:3px 0 2px;font-variant-numeric:tabular-nums}}
.kpis small{{color:{MUTED};font-size:11px}}
.wrap{{overflow-x:auto}} table{{width:100%;border-collapse:collapse;min-width:980px}}
th{{background:{DARK};color:#fff;font-size:10.5px;letter-spacing:.05em;text-transform:uppercase;padding:8px 8px;text-align:center}}
th.l{{text-align:left}} th.n{{text-align:right}}
td{{padding:6px 8px;border-bottom:1px solid #ECECEC;vertical-align:middle}} td.l{{text-align:left}} td.c{{text-align:center;white-space:nowrap}}
td.n{{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}} td.b{{font-weight:700}} td.g{{color:#1E7B34}} td.due{{color:{RED}}}
td small{{display:block;color:{MUTED};font-size:10.5px}}
td.proj{{font-weight:700;font-size:14px;text-align:center;background:#F7F4F1;border-right:1px solid #E6E1DC;border-bottom:2px solid #D9D3CD;vertical-align:top;padding-top:8px;width:58px}}
tr.first td{{border-top:2px solid #D9D3CD}} td.scope{{font-weight:600}} td.rem{{color:#444;font-size:11.5px;max-width:220px}}
tr.sub td{{background:#F7F4F1;font-weight:700;border-bottom:2px solid #D9D3CD;font-size:12px}}
tfoot td{{background:{DARK};color:#fff;font-weight:700;border:none;padding:8px}}
.bar{{display:inline-block;width:64px;height:7px;border-radius:4px;background:#EAE6E2;vertical-align:middle;overflow:hidden}}
.bar i{{display:block;height:100%;background:linear-gradient(90deg,#2E8B57,#3DA86B);border-radius:4px}}
.bar.big{{display:block;width:100%;height:8px;margin-top:8px}} tfoot .bar{{background:rgba(255,255,255,.25)}}
.pct{{display:inline-block;width:38px;text-align:right;font-size:11.5px;font-variant-numeric:tabular-nums;margin-left:6px}}
.pill{{display:inline-block;padding:2px 10px;border-radius:10px;font-size:11px;font-weight:700}}
.pays{{display:flex;flex-wrap:wrap;gap:4px 6px;max-width:420px}}
.pays span{{border:1px solid #E6E1DC;border-radius:5px;padding:2px 7px;font-variant-numeric:tabular-nums;font-size:11.5px;white-space:nowrap;background:#FCFAF8}}
.pays em{{font-style:normal;color:{MUTED};margin-right:6px}} .none{{color:#aaa;font-style:italic}}
.nothing{{text-align:center;color:{MUTED};padding:30px 0;font-size:14px}}
@media print{{body{{background:#fff}} .page{{box-shadow:none;margin:0;max-width:none;padding:0}} @page{{size:A4 landscape;margin:10mm}}}}
@media(max-width:760px){{.page{{padding:14px;margin:0}} .kpis{{grid-template-columns:1fr 1fr}}}}
</style></head><body>
{export_web.preview_bar("Project payment tracker", _sub(r, detail), pdf_url, excel_url)}
<div class="page">
 <div class="top"><div>{f'<img src="{logo}" alt="">' if logo else '<b>INFINIA</b>'}</div><div><h1>PROJECT PAYMENT TRACKER</h1><p>{escape(_sub(r, detail))}</p></div></div>
 {kpi}
 {empty or f'<div class="wrap"><table><thead><tr>{head}</tr></thead><tbody>{"".join(body)}</tbody><tfoot>{foot}</tfoot></table></div>'}
</div></body></html>"""


def _pdf(r, detail):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    C = colors.HexColor
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=7.8, leading=9.6)
    bold = ParagraphStyle("bb", parent=base, fontName="Helvetica-Bold")
    rt = ParagraphStyle("r", parent=base, alignment=TA_RIGHT)
    rtb = ParagraphStyle("rb", parent=bold, alignment=TA_RIGHT)
    rtg = ParagraphStyle("rg", parent=rtb, textColor=C("#1E7B34"))
    rtr = ParagraphStyle("rr", parent=rt, textColor=C(RED))
    rem = ParagraphStyle("rem", parent=base, fontSize=7.2, leading=8.8, textColor=C("#444444"))
    hd = ParagraphStyle("h", parent=bold, fontSize=7, leading=8.6, textColor=colors.white, alignment=TA_CENTER)
    hdl = ParagraphStyle("hl", parent=hd, alignment=0)
    hdr = ParagraphStyle("hr", parent=hd, alignment=TA_RIGHT)
    proj = ParagraphStyle("p", parent=bold, fontSize=11, leading=13, alignment=TA_CENTER)
    white = ParagraphStyle("w", parent=bold, textColor=colors.white)
    whiter = ParagraphStyle("wr", parent=white, alignment=TA_RIGHT)
    P = lambda t, st=base: Paragraph(escape(str(t)), st)

    class Bar(Flowable):
        """A collected-% bar with the figure beside it."""
        def __init__(self, pct, w=32, light=False):
            super().__init__(); self.pct, self.bw, self.light = pct, w, light
            self.width, self.height = w + 26, 9

        def draw(self):
            c = self.canv
            c.setFillColor(C("#5A5F64") if self.light else C("#EAE6E2")); c.roundRect(0, 2, self.bw, 5, 2.5, stroke=0, fill=1)
            c.setFillColor(C("#3DA86B")); w = self.bw * max(0, min(100, self.pct)) / 100
            if w > 0:
                c.roundRect(0, 2, max(w, 3), 5, 2.5, stroke=0, fill=1)
            c.setFillColor(colors.white if self.light else C("#1d1d1d")); c.setFont("Helvetica-Bold" if self.light else "Helvetica", 7.2)
            c.drawRightString(self.bw + 26, 1.6, f"{self.pct:.0f}%")

    def pill(s):
        bg, fg = STATUS_COLOURS[s]
        t = Table([[Paragraph(s, ParagraphStyle("pl", parent=bold, fontSize=7, leading=8.4, textColor=C(fg), alignment=TA_CENTER))]],
                  colWidths=[17 * mm], rowHeights=[4.2 * mm])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), C(bg)), ("ROUNDEDCORNERS", [5, 5, 5, 5]),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                               ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
        return t

    page = landscape(A4)
    W = page[0] - 20 * mm
    buf = io.BytesIO()

    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7); canvas.setFillColor(C(MUTED))
        canvas.drawString(10 * mm, 6 * mm, f"{COMPANY} · Project payment tracker · {_sub(r, detail)}")
        canvas.drawRightString(page[0] - 10 * mm, 6 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=9 * mm, bottomMargin=12 * mm,
                            title=f"Project Payment Tracker - {_title(detail)}")
    logo = export_web._logo_image(44)
    top = Table([[logo or P("INFINIA", bold),
                  [Paragraph("PROJECT PAYMENT TRACKER", ParagraphStyle("t", parent=bold, fontSize=15, leading=18, textColor=C(RED), alignment=TA_RIGHT)),
                   Paragraph(escape(_sub(r, detail)), ParagraphStyle("st", parent=base, fontSize=8.5, textColor=C(MUTED), alignment=TA_RIGHT))]]],
                colWidths=[W / 2, W / 2])
    top.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("LINEBELOW", (0, 0), (-1, 0), 1.4, C(RED)),
                             ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    lab = ParagraphStyle("k", parent=base, fontSize=6.8, textColor=C(MUTED))
    note = ParagraphStyle("kn", parent=base, fontSize=7, textColor=C(MUTED))
    kw = (W - 3 * 4 * mm) / 4
    kcells = []
    for lb, val, nt, col, b in _kpis(r):
        kcells.append([P(lb.upper(), lab), Paragraph(escape(val), ParagraphStyle("kb", parent=bold, fontSize=12.5, leading=15, textColor=C(col))),
                       Bar(b, w=kw - 8 - 14 - 30) if b is not None else P(nt, note)])
    kp = Table([[Table([[x] for x in k], colWidths=[kw - 8]) for k in kcells]], colWidths=[kw + 4 * mm] * 3 + [kw])
    kp.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 4 * mm)]))
    for cell in kp._cellvalues[0]:
        cell.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .6, C("#E6E1DC")), ("BACKGROUND", (0, 0), (-1, -1), C("#FCFAF8")),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (0, 0), 6), ("BOTTOMPADDING", (0, -1), (-1, -1), 6),
                                  ("TOPPADDING", (0, 1), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -2), 1)]))
    story = [top, Spacer(1, 7), kp, Spacer(1, 8)]
    if not r["rows"]:
        story.append(Paragraph(f"No payments in {escape(r['period_label'] or 'this selection')}.",
                               ParagraphStyle("no", parent=base, fontSize=11, textColor=C(MUTED), alignment=TA_CENTER, spaceBefore=30)))
        doc.build(story, onFirstPage=frame, onLaterPages=frame); buf.seek(0)
        return buf

    cols = _cols(r, detail)
    fixed = sum(w for _, _, w in cols) * mm
    widths = [w * mm if w else max(30 * mm, W - fixed) for _, _, w in cols]
    if sum(widths) > W:                       # never wider than the page
        k = W / sum(widths); widths = [w * k for w in widths]
    ix = {k: i for i, (k, _, _) in enumerate(cols)}
    hstyle = lambda k: hdr if k in NUMS else (hdl if k in ("scope", "remarks", "pays") else hd)
    data = [[Paragraph(escape(h.upper()), hstyle(k)) for k, h, _ in cols]]
    st = [("BACKGROUND", (0, 0), (-1, 0), C(DARK)), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
          ("TOPPADDING", (0, 0), (-1, -1), 2.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2),
          ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
    pw = widths[ix["pays"]] if "pays" in ix else 0

    def paygrid(x):
        ps = _pays_of(r, x)
        if not ps:
            return Paragraph("<i>No payment yet</i>", ParagraphStyle("np", parent=base, textColor=C("#AAAAAA")))
        per = 2
        chips = [Paragraph(f"<font color='{MUTED}'>{_dmy(p['date'])}</font>&nbsp;&nbsp;{_m(p['amount'])}", ParagraphStyle("pc", parent=base, fontSize=7.2))
                 for p in ps]
        rows = [chips[i:i + per] + [""] * (per - len(chips[i:i + per])) for i in range(0, len(chips), per)]
        g = Table(rows, colWidths=[(pw - 8) / per] * per)
        g.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                               ("TOPPADDING", (0, 0), (-1, -1), .5), ("BOTTOMPADDING", (0, 0), (-1, -1), .5)]))
        return g

    def line(x, first_of=None):
        row = []
        for k, _, _ in cols:
            if k == "proj":
                row.append(Paragraph(escape(first_of), proj) if first_of else "")
            elif k == "scope":
                row.append(P(x["scope"], bold))
            elif k == "contract":
                row.append(P(_m(x["contract"]), rt))
            elif k == "pays":
                row.append(paygrid(x))
            elif k == "inper":
                row.append(P(_m(x["in_period"]), rtg))
            elif k == "received":
                row.append(P(_m(x["received"]), rt if r["period"] else rtb))
            elif k == "remaining":
                row.append(P(_m(x["remaining"]), rtr if x["remaining"] > 0.005 else rt))
            elif k == "pct":
                row.append(Bar(x["pct"]))
            elif k == "last":
                row.append(Paragraph(f"<font color='{MUTED}' size='6.8'>{_dmy(x['last_date'])}</font>&nbsp;&nbsp;{_m(x['last_amount'])}", rt)
                           if x["last_date"] else P("-", ParagraphStyle("d", parent=rt, textColor=C("#AAAAAA"))))
            elif k == "status":
                row.append(pill(x["status"]))
            else:
                row.append(P(x["remarks"], rem))
        return row

    def total(label, t, light=False):
        row = []
        for k, _, _ in cols:
            if k == "proj":
                row.append(P(label, white) if light else "")
            elif k == "scope":
                row.append("" if light else P(label, bold))
            elif k in ("contract", "received", "remaining"):
                row.append(P(_m(t[k]), whiter if light else rtb))
            elif k == "inper":
                row.append(P(_m(t["in_period"]), whiter if light else rtg))
            elif k == "pct":
                row.append(Bar(t["pct"], light=light))
            else:
                row.append("")
        return row

    for g in r["groups"]:
        n = len(g["rows"])
        start = len(data)
        for i, x in enumerate(g["rows"]):
            data.append(line(x, g["project"] if i == 0 else None))
        if n > 1:
            data.append(total(f"Project {g['project']} total", g))
            st += [("BACKGROUND", (1, len(data) - 1), (-1, len(data) - 1), C("#F7F4F1"))]
        end = len(data) - 1
        st += [("BACKGROUND", (0, start), (0, end), C("#F7F4F1")), ("VALIGN", (0, start), (0, end), "TOP"),
               ("LINEABOVE", (0, start), (-1, start), .9, C("#CFC8C1")),
               ("LINEBELOW", (1, start), (-1, end - 1), .25, C("#E4E4E4")), ("LINEAFTER", (0, start), (0, end), .5, C("#E0DAD4"))]
    s = r["shown"]
    data.append(total("TOTAL" + ("" if s["count"] == r["totals"]["count"] or r["period"] else " (shown)"), s, light=True))
    st += [("SPAN", (0, -1), (1, -1)), ("BACKGROUND", (0, -1), (-1, -1), C(DARK)), ("TOPPADDING", (0, -1), (-1, -1), 5), ("BOTTOMPADDING", (0, -1), (-1, -1), 5)]
    tbl = Table(data, colWidths=widths, repeatRows=1)
    tbl.setStyle(TableStyle(st))
    story.append(tbl)
    doc.build(story, onFirstPage=frame, onLaterPages=frame)
    buf.seek(0)
    return buf


def _excel(r, detail):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    wb = Workbook(); ws = wb.active; ws.title = "Tracker"
    dark = PatternFill("solid", fgColor=DARK.lstrip("#")); tint = PatternFill("solid", fgColor="F7F4F1")
    thin = Side(style="thin", color="E4E4E4"); group = Side(style="medium", color="CFC8C1")
    per = r["period"]
    # (key, heading, width)
    spec = [("proj", "Project", 10), ("scope", "Scope", 30), ("contract", "Contract (AED)", 17)]
    if per:
        spec.append(("inper", "Received in period (AED)", 18))
    spec += [("received", "Received to date (AED)" if per else "Received (AED)", 18), ("remaining", "Remaining (AED)", 17),
             ("pct", "Collected", 11), ("last", "Last payment (AED)", 17), ("lastd", "Last payment date", 13),
             ("status", "Status", 12), ("remarks", "Remarks", 46)]
    col = {k: i for i, (k, _, _) in enumerate(spec, 1)}
    CC, RR, MM = L(col["contract"]), L(col["received"]), L(col["remaining"])
    ws["A1"] = "PROJECT PAYMENT TRACKER"; ws["A1"].font = Font(bold=True, size=15, color=RED.lstrip("#"))
    ws["A2"] = f"{COMPANY} · {_sub(r, detail)}"; ws["A2"].font = Font(color=MUTED.lstrip("#"))
    # The four figures sit over the wide columns, so they never show as ###.
    for c, (lab, val, note, colr, b) in zip((2, 3, 4, 5), _kpis(r)):
        ws.cell(3, c, lab).font = Font(color=MUTED.lstrip("#"), size=9)
        cell = ws.cell(4, c, val); cell.font = Font(bold=True, size=11, color=colr.lstrip("#"))
        cell.alignment = Alignment(horizontal="left")
    hr = 6
    for i, (k, h, w) in enumerate(spec, 1):
        c = ws.cell(hr, i, h.upper()); c.font = Font(bold=True, color="FFFFFF", size=9); c.fill = dark
        c.alignment = Alignment(horizontal="right" if k in NUMS else ("left" if k in ("scope", "remarks") else "center"), vertical="center", wrap_text=True)
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[hr].height = 30
    row = hr + 1
    fills = {k: PatternFill("solid", fgColor=v[0].lstrip("#")) for k, v in STATUS_COLOURS.items()}
    fonts = {k: Font(bold=True, color=v[1].lstrip("#"), size=9) for k, v in STATUS_COLOURS.items()}
    first_data = row
    money = [k for k in ("contract", "inper", "received", "remaining", "last") if k in col]
    for g in r["groups"]:
        start = row
        for x in g["rows"]:
            vals = {"proj": g["project"], "scope": x["scope"], "contract": x["contract"], "inper": x["in_period"],
                    "received": x["received"], "remaining": f"={CC}{row}-{RR}{row}", "pct": f"=IF({CC}{row}=0,0,{RR}{row}/{CC}{row})",
                    "last": x["last_amount"] if x["last_date"] else None, "lastd": M._as_date(x["last_date"]) if x["last_date"] else None,
                    "status": x["status"], "remarks": x["remarks"]}
            for k, i in col.items():
                ws.cell(row, i, vals[k])
            for k in money:
                ws.cell(row, col[k]).number_format = "#,##0.00"
            ws.cell(row, col["pct"]).number_format = "0%"; ws.cell(row, col["lastd"]).number_format = "dd-mmm-yy"
            for k in ("proj", "scope", "inper" if per else "received"):
                ws.cell(row, col[k]).font = Font(bold=True)
            ws.cell(row, col["status"]).fill = fills[x["status"]]; ws.cell(row, col["status"]).font = fonts[x["status"]]
            for i in range(1, len(spec) + 1):
                k = spec[i - 1][0]
                ws.cell(row, i).border = Border(bottom=thin, top=group if row == start else None)
                ws.cell(row, i).alignment = Alignment(vertical="center", wrap_text=k in ("scope", "remarks"),
                                                      horizontal="center" if k in ("proj", "pct", "lastd", "status") else None)
            row += 1
        if len(g["rows"]) > 1:
            ws.cell(row, col["scope"], f"Project {g['project']} total")
            for k in ("contract", "inper", "received"):
                if k in col:
                    ws.cell(row, col[k], f"=SUM({L(col[k])}{start}:{L(col[k])}{row - 1})")
            ws.cell(row, col["remaining"], f"={CC}{row}-{RR}{row}"); ws.cell(row, col["pct"], f"=IF({CC}{row}=0,0,{RR}{row}/{CC}{row})")
            for i in range(1, len(spec) + 1):
                c = ws.cell(row, i); c.fill = tint; c.font = Font(bold=True)
                c.number_format = "0%" if spec[i - 1][0] == "pct" else "#,##0.00"
                c.alignment = Alignment(horizontal="center" if spec[i - 1][0] == "pct" else None, vertical="center")
            row += 1
    last = row - 1
    if r["rows"]:
        # Totals add the scope lines only (not the project subtotal lines).
        ws.cell(row, 1, "TOTAL")
        for k in ("contract", "inper", "received"):
            if k in col:
                Lk = L(col[k])
                ws.cell(row, col[k], f'=SUMIFS({Lk}{first_data}:{Lk}{last},$A${first_data}:$A${last},"<>")')
        ws.cell(row, col["remaining"], f"={CC}{row}-{RR}{row}"); ws.cell(row, col["pct"], f"=IF({CC}{row}=0,0,{RR}{row}/{CC}{row})")
        for i in range(1, len(spec) + 1):
            c = ws.cell(row, i); c.fill = dark; c.font = Font(bold=True, color="FFFFFF")
            c.number_format = "0%" if spec[i - 1][0] == "pct" else "#,##0.00"
            c.alignment = Alignment(horizontal="center" if spec[i - 1][0] == "pct" else None, vertical="center")
        PC = L(col["pct"])
        ws.conditional_formatting.add(f"{PC}{first_data}:{PC}{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="3DA86B"))
    else:
        ws.cell(row, 1, f"No payments in {r['period_label'] or 'this selection'}.").font = Font(italic=True, color=MUTED.lstrip("#"))
    ws.freeze_panes = ws.cell(hr + 1, 3)
    export_web.print_ready(ws, "landscape", header_row=hr, title="Project Payment Tracker")
    if detail:
        p = wb.create_sheet("Payments")
        ph = ["Project", "Scope", "Payment no.", "Date", "Amount (AED)", "Note"]
        for i, h in enumerate(ph, 1):
            c = p.cell(1, i, h.upper()); c.font = Font(bold=True, color="FFFFFF", size=9); c.fill = dark
            c.alignment = Alignment(horizontal="right" if i == 5 else "left", vertical="center")
        for cl, w in zip("ABCDEF", (10, 30, 12, 13, 16, 40)):
            p.column_dimensions[cl].width = w
        pr = 2
        for g in r["groups"]:
            for x in g["rows"]:
                nums = {pp["id"]: k for k, pp in enumerate(x["payments"], 1)}       # the payment's number on the scope
                for pay in _pays_of(r, x):
                    p.append([g["project"], x["scope"], nums.get(pay["id"]), M._as_date(pay["date"]), pay["amount"], pay["note"]])
                    p.cell(pr, 4).number_format = "dd-mmm-yy"; p.cell(pr, 5).number_format = "#,##0.00"
                    for i in range(1, 7):
                        p.cell(pr, i).border = Border(bottom=thin)
                    pr += 1
        p.append(["TOTAL", None, None, None, f"=SUM(E2:E{pr - 1})" if pr > 2 else 0])
        for i in range(1, 7):
            c = p.cell(pr, i); c.fill = dark; c.font = Font(bold=True, color="FFFFFF")
        p.cell(pr, 5).number_format = "#,##0.00"
        p.freeze_panes = "A2"; p.auto_filter.ref = f"A1:F{max(pr - 1, 1)}"
        export_web.print_ready(p, "portrait", header_row=1, title="Project Payments")
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


def _args(u, report, hide_completed, q, start, end, db):
    if RIGHT not in M.effective_permissions(u):
        raise HTTPException(status_code=403, detail="The project payment tracker is for accounts only.")
    return report == "detail", tracker(db, bool(int(hide_completed or 0)), q or "", M._as_date(start), M._as_date(end))


@router.get("/export/accounts/projects")
def export_tracker(token: str, report: str = "summary", format: str = "pdf", hide_completed: int = 0, q: str = "",
                   start: str = "", end: str = "", db: Session = Depends(get_db)):
    u = auth.get_download_user_from_token(token, db)
    detail, r = _args(u, report, hide_completed, q, start, end, db)
    when = re.sub(r"[^A-Za-z0-9]+", "_", r["period_label"]).strip("_") or r["as_on"]
    name = f"Project_Payment_Tracker_{'With_Dates' if detail else 'Summary'}_{when}"
    if format == "excel":
        return StreamingResponse(_excel(r, detail), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f"attachment; filename={name}.xlsx"})
    return StreamingResponse(_pdf(r, detail), media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="{name}.pdf"'})


@router.get("/export/accounts/projects/view", response_class=HTMLResponse)
def view_tracker(token: str, report: str = "summary", hide_completed: int = 0, q: str = "", start: str = "", end: str = "",
                 db: Session = Depends(get_db)):
    u = auth.get_download_user_from_token(token, db)
    detail, r = _args(u, report, hide_completed, q, start, end, db)
    t = auth.create_view_token(u.username)
    base = "/export/accounts/projects?" + urlencode({"report": report, "hide_completed": int(r["hide_completed"]), "q": r["q"],
                                                     "start": r["from"], "end": r["to"], "token": t})
    return HTMLResponse(_html(r, detail, base + "&format=pdf", base + "&format=excel"))
