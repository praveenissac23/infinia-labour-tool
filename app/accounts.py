"""Accounts - its own page, for admin and whoever is given the rights on
Settings > Access (the chief accountant).

  * Tax invoices and proforma invoices: numbered, kept, printed in the
    company's invoice layout (invoice_pdf), cancelled rather than
    deleted, a proforma turned into a tax invoice in one step.
  * Cash register: cash received and cash paid out, two sides like the
    office's own cash file, with the balance. Behind a password asked
    every time the page is opened - on the server as well as the screen:
    without a fresh key from the password nothing in it can be read,
    added, printed or exported.
"""
import io
import json
import os
import time
from datetime import datetime, timedelta, timezone
from html import escape

from fastapi import APIRouter, Body, Depends, File, Header, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from jose import jwt, JWTError
from sqlalchemy import Boolean, Column, Date, DateTime, Float, ForeignKey, Integer, String, Text, func
from sqlalchemy.orm import Session

import auth
import export_web
import invoice_pdf
import main as M
import models
from database import Base, get_db

router = APIRouter()
INV_RIGHT, REG_RIGHT = "accounts_invoices", "accounts_register"
INV = Depends(M.require_screen(INV_RIGHT))
REG = Depends(M.require_screen(REG_RIGHT))
# The register's password, kept only as a one-way hash (never the word
# itself). Admin can change it from the register page.
DEFAULT_REGISTER_HASH = "$2b$12$ciocx9PMuiMJy2QVk8cbN.X3SexmFYZ97OvbWmSHN5iLM4nOXwEtm"
KEY_MINUTES = 30
MAX_TRIES, LOCK_MINUTES = 5, 5


class Invoice(Base):
    """A tax invoice or a proforma invoice, as issued."""
    __tablename__ = "invoices"
    id = Column(Integer, primary_key=True)
    kind = Column(String, nullable=False, index=True)          # tax | proforma
    number = Column(String, nullable=False, index=True)
    inv_date = Column(Date, nullable=False)
    company_id = Column(Integer, ForeignKey("companies.id"), nullable=True)
    client = Column(String, default="")
    client_trn = Column(String, default="")
    client_address = Column(String, default="")
    project = Column(String, default="")       # (B+G+1+R) Villa
    project_no = Column(String, default="")    # 906
    plot = Column(String, default="")
    location = Column(String, default="")
    work = Column(String, default="")          # Phase 2 Works
    lines = Column(Text, default="[]")         # [{description, amount, vat}]
    subtotal = Column(Float, default=0.0)
    vat = Column(Float, default=0.0)
    total = Column(Float, default=0.0)
    status = Column(String, default="issued")  # issued | cancelled
    cancel_reason = Column(String, default="")
    from_proforma_id = Column(Integer, nullable=True)
    converted_to_id = Column(Integer, nullable=True)
    created_by = Column(Integer, ForeignKey("users.id"), nullable=True)
    updated_by = Column(Integer, ForeignKey("users.id"), nullable=True)
    created_at = Column(DateTime(timezone=True), server_default=func.now())
    updated_at = Column(DateTime(timezone=True), onupdate=func.now())


class CashEntry(Base):
    """One line of the cash register: cash received, or cash paid out."""
    __tablename__ = "cash_register"
    id = Column(Integer, primary_key=True)
    kind = Column(String, nullable=False, index=True)          # received | paid
    on_date = Column(Date, nullable=False, index=True)
    description = Column(String, default="")
    amount = Column(Float, default=0.0)
    remarks = Column(String, default="")
    created_by = Column(Integer, ForeignKey("users.id"), nullable=True)
    updated_by = Column(Integer, ForeignKey("users.id"), nullable=True)
    created_at = Column(DateTime(timezone=True), server_default=func.now())
    updated_at = Column(DateTime(timezone=True), onupdate=func.now())


def create_tables(engine):
    for t in (Invoice.__table__, CashEntry.__table__):
        t.create(bind=engine, checkfirst=True)


# ---- invoices ---------------------------------------------------------------

def invoice_signature():
    """The invoice signature and stamp kept in Settings (the purchase-order
    one until an invoice one is uploaded). Was: the same picture the
    purchase orders print - so there is one to upload and one to change."""
    return export_web.invoice_signature_or_lpo()


def _company(db, cid):
    """The issuing company's name, address and TRN for the paper. The main
    company falls back on the details printed on its invoices today; any
    other company prints only what is recorded for it."""
    co = db.get(models.Company, cid) if cid else None
    if not co:
        return dict(invoice_pdf.DEFAULT_COMPANY)
    mine = {"name": co.name, "address": (co.address or "").replace("\n", ", "), "trn": co.trn or ""}
    if "INFINIA CONTRACTING" in (co.name or "").upper() and "PRIME" not in (co.name or "").upper():
        return {**invoice_pdf.DEFAULT_COMPANY, **{k: v for k, v in mine.items() if v}}
    return {**mine, "web": invoice_pdf.DEFAULT_COMPANY["web"]}


def _lines(raw):
    out = []
    for l in raw or []:
        d = " ".join(str(l.get("description") or "").split())
        try:
            a = round(float(str(l.get("amount") or 0).replace(",", "")), 2)
        except ValueError:
            raise HTTPException(status_code=400, detail=f"The amount for '{d or 'a line'}' is not a number.")
        try:
            v = float(l.get("vat") if l.get("vat") not in (None, "") else 5)
        except ValueError:
            v = 5.0
        if not d and not a:
            continue
        if not d:
            raise HTTPException(status_code=400, detail="Every line needs a description.")
        if v not in (0.0, 5.0):
            raise HTTPException(status_code=400, detail="VAT on a line is 5% or 0%.")
        out.append({"description": d[:300], "amount": a, "vat": v})
    if not out:
        raise HTTPException(status_code=400, detail="Add at least one line with a description and an amount.")
    return out


def _dict(x, users=None):
    return {"id": x.id, "kind": x.kind, "number": x.number, "date": x.inv_date.isoformat(),
            "company_id": x.company_id, "client": x.client or "", "client_trn": x.client_trn or "",
            "client_address": x.client_address or "", "project": x.project or "", "project_no": x.project_no or "",
            "plot": x.plot or "", "location": x.location or "", "work": x.work or "",
            "lines": json.loads(x.lines or "[]"), "subtotal": x.subtotal, "vat": x.vat, "total": x.total,
            "status": x.status, "cancel_reason": x.cancel_reason or "",
            "from_proforma_id": x.from_proforma_id, "converted_to_id": x.converted_to_id,
            "by": (users or {}).get(x.updated_by or x.created_by, "")}


def next_number(db, kind, project_no=""):
    """The next number in the company's pattern: IC/26/906/01, IC/26/906/02
    ... (PI/... for a proforma). Following the last one for the project:
    'IC/26/906/PHASE2-10' is followed by 'IC/26/906/PHASE2-11'."""
    import re
    yy = f"{M._dubai_today():%y}"
    pre = ("PI" if kind == "proforma" else "IC") + f"/{yy}/" + (f"{project_no.strip()}/" if project_no.strip() else "")
    # Without a project number only IC/26/NN counts, not IC/26/906/NN.
    for last in (db.query(Invoice).filter(Invoice.kind == kind, Invoice.number.like(pre + "%"))
                   .order_by(Invoice.id.desc()).all()):
        if not project_no.strip() and "/" in last.number[len(pre):]:
            continue
        m = re.match(r"^(.*?)(\d+)$", last.number)
        if m:
            return m.group(1) + str(int(m.group(2)) + 1).zfill(len(m.group(2)))
    return pre + "01"


def _clean(db, p, kind, keep_id=None):
    if kind not in ("tax", "proforma"):
        raise HTTPException(status_code=400, detail="Tax invoice or proforma invoice?")
    number = " ".join(str(p.get("number") or "").split())
    if not number:
        raise HTTPException(status_code=400, detail="Give the invoice a number.")
    clash = db.query(Invoice).filter(Invoice.kind == kind, Invoice.number == number, Invoice.id != (keep_id or 0)).first()
    if clash:
        raise HTTPException(status_code=400, detail=f"{number} is already used (dated {clash.inv_date:%d %b %Y}).")
    d = M._as_date(p.get("date"))
    if not d:
        raise HTTPException(status_code=400, detail="Put the invoice date.")
    client = " ".join(str(p.get("client") or "").split())
    if not client:
        raise HTTPException(status_code=400, detail="Who is the invoice to? Type the client.")
    lines = _lines(p.get("lines"))
    sub, vat, total = invoice_pdf.totals(lines)
    s = lambda k, n=200: " ".join(str(p.get(k) or "").split())[:n]
    return {"kind": kind, "number": number[:60], "inv_date": d, "company_id": p.get("company_id") or None,
            "client": client[:200], "client_trn": s("client_trn", 30), "client_address": s("client_address"),
            "project": s("project"), "project_no": s("project_no", 30), "plot": s("plot", 60),
            "location": s("location"), "work": s("work"), "lines": json.dumps(lines),
            "subtotal": sub, "vat": vat, "total": total}


@router.get("/employees/accounts/invoices")
def list_invoices(kind: str = "tax", db: Session = Depends(get_db), user: models.User = INV):
    users = {u.id: (u.full_name or u.username) for u in db.query(models.User).all()}
    rows = db.query(Invoice).filter(Invoice.kind == kind).order_by(Invoice.inv_date.desc(), Invoice.id.desc()).all()
    live = [r for r in rows if r.status != "cancelled"]
    companies = [{"id": c.id, "name": c.name} for c in db.query(models.Company).filter(models.Company.active == True).all()]  # noqa: E712
    return {"rows": [_dict(r, users) for r in rows], "companies": companies,
            "totals": {"count": len(live), "subtotal": round(sum(r.subtotal for r in live), 2),
                       "vat": round(sum(r.vat for r in live), 2), "total": round(sum(r.total for r in live), 2)},
            "clients": sorted({r.client for r in db.query(Invoice).all() if r.client}),
            "signature": bool(invoice_signature())}


@router.get("/employees/accounts/invoices/next")
def get_next(kind: str = "tax", project_no: str = "", db: Session = Depends(get_db), user: models.User = INV):
    return {"number": next_number(db, kind, project_no)}


@router.post("/employees/accounts/invoices")
def add_invoice(payload: dict = Body(...), db: Session = Depends(get_db), user: models.User = INV):
    x = Invoice(**_clean(db, payload, payload.get("kind") or "tax"), created_by=user.id, updated_by=user.id)
    db.add(x); db.commit()
    M.log_action(db, user.id, "invoice_add", f"{x.kind} {x.number} {x.client} AED {x.total:,.2f}")
    return {"ok": True, "id": x.id, "number": x.number}


@router.put("/employees/accounts/invoices/{iid}")
def edit_invoice(iid: int, payload: dict = Body(...), db: Session = Depends(get_db), user: models.User = INV):
    x = db.get(Invoice, iid)
    if not x:
        raise HTTPException(status_code=404, detail="That invoice is not on file.")
    if x.status == "cancelled":
        raise HTTPException(status_code=400, detail="A cancelled invoice cannot be changed.")
    for k, v in _clean(db, payload, x.kind, keep_id=x.id).items():
        setattr(x, k, v)
    x.updated_by = user.id
    db.commit()
    M.log_action(db, user.id, "invoice_edit", f"{x.kind} {x.number} AED {x.total:,.2f}")
    return {"ok": True}


@router.post("/employees/accounts/invoices/{iid}/cancel")
def cancel_invoice(iid: int, payload: dict = Body(default={}), db: Session = Depends(get_db), user: models.User = INV):
    """Cancelled, never deleted: the number stays used and the paper that
    went out stays on the register, stamped CANCELLED."""
    x = db.get(Invoice, iid)
    if not x:
        raise HTTPException(status_code=404, detail="That invoice is not on file.")
    x.status = "cancelled"
    x.cancel_reason = str((payload or {}).get("reason") or "").strip()[:200]
    x.updated_by = user.id
    db.commit()
    M.log_action(db, user.id, "invoice_cancel", f"{x.kind} {x.number} - {x.cancel_reason}")
    return {"ok": True}


@router.post("/employees/accounts/invoices/{iid}/convert")
def convert_proforma(iid: int, db: Session = Depends(get_db), user: models.User = INV):
    """The proforma's client, project and lines become a new tax invoice,
    numbered next in the tax series; the two are linked."""
    p = db.get(Invoice, iid)
    if not p or p.kind != "proforma":
        raise HTTPException(status_code=404, detail="That proforma is not on file.")
    if p.status == "cancelled":
        raise HTTPException(status_code=400, detail="A cancelled proforma cannot be turned into a tax invoice.")
    if p.converted_to_id and db.get(Invoice, p.converted_to_id):
        t = db.get(Invoice, p.converted_to_id)
        raise HTTPException(status_code=400, detail=f"Already turned into tax invoice {t.number}.")
    t = Invoice(kind="tax", number=next_number(db, "tax", p.project_no or ""), inv_date=M._dubai_today(),
                company_id=p.company_id, client=p.client, client_trn=p.client_trn, client_address=p.client_address,
                project=p.project, project_no=p.project_no, plot=p.plot, location=p.location, work=p.work,
                lines=p.lines, subtotal=p.subtotal, vat=p.vat, total=p.total, from_proforma_id=p.id,
                created_by=user.id, updated_by=user.id)
    db.add(t); db.flush()
    p.converted_to_id = t.id
    db.commit()
    M.log_action(db, user.id, "invoice_convert", f"proforma {p.number} -> tax invoice {t.number}")
    return {"ok": True, "id": t.id, "number": t.number}


def _pdf_response(db, x):
    d = _dict(x)
    d["date_text"] = f"{x.inv_date:%d %B %Y}"
    buf = invoice_pdf.build(d, company=_company(db, x.company_id), signature=invoice_signature())
    name = ("Tax_Invoice_" if x.kind == "tax" else "Proforma_") + x.number.replace("/", "-")
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f'inline; filename="{name}.pdf"'})


@router.get("/export/accounts/invoice/{iid}")
def export_invoice(iid: int, token: str, db: Session = Depends(get_db)):
    u = auth.get_download_user_from_token(token, db)
    if INV_RIGHT not in M.effective_permissions(u):
        raise HTTPException(status_code=403, detail="Invoices are for accounts only.")
    x = db.get(Invoice, iid)
    if not x:
        raise HTTPException(status_code=404, detail="That invoice is not on file.")
    return _pdf_response(db, x)


@router.get("/export/accounts/invoice/{iid}/view")
def view_invoice(iid: int, token: str, db: Session = Depends(get_db)):
    """The invoice on its own page with Download PDF, Print and Share above
    it - Share hands the PDF to WhatsApp, mail or any app on a phone."""
    from fastapi.responses import HTMLResponse
    from urllib.parse import quote
    import export_web
    u = auth.get_download_user_from_token(token, db)
    if INV_RIGHT not in M.effective_permissions(u):
        raise HTTPException(status_code=403, detail="Invoices are for accounts only.")
    x = db.get(Invoice, iid)
    if not x:
        raise HTTPException(status_code=404, detail="That invoice is not on file.")
    t = quote(auth.create_view_token(u.username), safe="")
    pdf = f"/export/accounts/invoice/{iid}?token={t}"
    title = ("Tax Invoice " if x.kind == "tax" else "Proforma Invoice ") + x.number
    d = _dict(x)
    sub = " - ".join(v for v in (d.get("client") or "", f"AED {float(d.get('total') or 0):,.2f}") if v)
    return HTMLResponse(f"""<!doctype html><html><head><meta charset="utf-8"><title>{escape(title)}</title>
<meta name="viewport" content="width=device-width, initial-scale=1">
<style>{export_web.PREVIEW_BAR_CSS}
  body {{ margin:0; background:#F1EFEA; font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif; }}
  .doc {{ width:min(210mm, calc(100% - 24px)); height:calc(100vh - 90px); margin:12px auto; background:white;
          box-shadow:0 1px 6px rgba(0,0,0,.14); }}
  .doc iframe {{ width:100%; height:100%; border:0; }}
  .phone {{ display:none; text-align:center; color:#666; font-size:14px; padding:30px 16px; }}
  @media (max-width: 700px) {{ .doc {{ display:none; }} .phone {{ display:block; }} }}
</style></head><body>
<div class="pvbar"><h1>{escape(title)}</h1><span class="who">{escape(sub)}</span><span class="sp"></span>
  <a class="btn dark" href="{escape(pdf)}" target="_blank">Download PDF</a>
  <a class="btn" href="{escape(pdf)}" target="_blank">Print</a>
  {export_web.share_button(pdf, title)}</div>
{export_web.SHARE_JS}
<div class="doc"><iframe src="{escape(pdf)}#view=FitH" title="{escape(title)}"></iframe></div>
<div class="phone">Tap <b>Share</b> to send it on WhatsApp or mail, or <b>Download PDF</b> to open it.</div>
</body></html>""")


@router.get("/export/accounts/invoices")
def export_invoice_register(token: str, kind: str = "tax", format: str = "excel", db: Session = Depends(get_db)):
    """The list of invoices in the app's standard report: Preview (on
    screen, with Download PDF / Download Excel / Print), PDF or Excel."""
    u = auth.get_download_user_from_token(token, db)
    if INV_RIGHT not in M.effective_permissions(u):
        raise HTTPException(status_code=403, detail="Invoices are for accounts only.")
    title = "Tax invoices" if kind == "tax" else "Proforma invoices"
    found = db.query(Invoice).filter(Invoice.kind == kind).order_by(Invoice.inv_date, Invoice.id).all()
    live = [x for x in found if x.status != "cancelled"]
    # A cancelled invoice stays on the list, its figures shown as text so
    # the totals line (which adds numbers only) leaves it out.
    amt = lambda x, v: round(v or 0, 2) if x.status != "cancelled" else f"{(v or 0):,.2f}"
    rows = [{"Number": x.number, "Date": f"{x.inv_date:%d-%b-%y}", "Client": x.client or "",
             "Client TRN": x.client_trn or "-", "Project": " - ".join(v for v in (x.project_no, x.project) if v) or "-",
             "Net (AED)": amt(x, x.subtotal), "VAT (AED)": amt(x, x.vat), "Total (AED)": amt(x, x.total),
             "Status": "Cancelled" if x.status == "cancelled" else ("Invoiced" if x.converted_to_id else "Issued")}
            for x in found]
    cx = len(found) - len(live)
    sub = (f"{len(live)} {'invoice' if kind == 'tax' else 'proforma'}{'' if len(live) == 1 else 's'}"
           + (f" (+{cx} cancelled, not in the totals)" if cx else "") + f" | As at {M._dubai_today():%d %b %Y}")
    money = ["Net (AED)", "VAT (AED)", "Total (AED)"]
    name = title.replace(" ", "_")
    if format == "view":
        from urllib.parse import quote
        url = f"/export/accounts/invoices?kind={kind}&token={quote(auth.create_view_token(u.username), safe='')}"
        return M._preview_page(title, sub, rows, url, url, money_cols=money)
    if format == "pdf":
        return StreamingResponse(export_web.build_store_report_pdf(title, rows, sub, money_cols=money),
                                 media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{name}.pdf"'})
    return StreamingResponse(export_web.build_store_report_excel(title, rows, sub, money_cols=money),
                             media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={name}.xlsx"})


# ---- the cash register (password) -------------------------------------------------

def _hash(db):
    return M.get_setting(db, "register_password_hash") or DEFAULT_REGISTER_HASH


def _tries(db, user):
    raw = M.get_setting(db, f"register_tries_{user.id}") or "0|0"
    n, until = raw.split("|")
    return int(n), float(until)


@router.post("/employees/accounts/register/unlock")
def unlock(payload: dict = Body(...), db: Session = Depends(get_db), user: models.User = REG):
    """The password, every time the register is opened. Five wrong in a row
    locks it for five minutes."""
    n, until = _tries(db, user)
    if until > time.time():
        mins = int((until - time.time()) // 60) + 1
        raise HTTPException(status_code=429, detail=f"Too many wrong passwords. Try again in {mins} minute{'s' if mins > 1 else ''}.")
    if not auth.verify_password(str(payload.get("password") or ""), _hash(db)):
        n += 1
        lock = time.time() + LOCK_MINUTES * 60 if n >= MAX_TRIES else 0
        M.put_setting(db, f"register_tries_{user.id}", f"{0 if lock else n}|{lock}")
        M.log_action(db, user.id, "register_wrong_password", f"attempt {n}" + (" - locked 5 minutes" if lock else ""))
        left = MAX_TRIES - n
        raise HTTPException(status_code=401, detail="Wrong password." + (f" {left} more tr{'ies' if left > 1 else 'y'} before it locks." if left > 0 else " Locked for 5 minutes."))
    M.put_setting(db, f"register_tries_{user.id}", "0|0")
    key = jwt.encode({"sub": user.username, "scope": "register",
                      "exp": datetime.now(timezone.utc) + timedelta(minutes=KEY_MINUTES)}, auth.SECRET_KEY, algorithm=auth.ALGORITHM)
    M.log_action(db, user.id, "register_open", "cash register opened")
    return {"key": key, "minutes": KEY_MINUTES}


def _check_key(key, user):
    try:
        p = jwt.decode(key or "", auth.SECRET_KEY, algorithms=[auth.ALGORITHM])
    except JWTError:
        p = {}
    if p.get("scope") != "register" or p.get("sub") != user.username:
        raise HTTPException(status_code=423, detail="The register is locked - enter the password.")


def register_key(user: models.User = REG, x_register_key: str = Header(default="")):
    _check_key(x_register_key, user)
    return user


RK = Depends(register_key)


def register(db, start=None, end=None):
    q = db.query(CashEntry)
    rows = q.order_by(CashEntry.on_date, CashEntry.id).all()
    before = [r for r in rows if start and r.on_date < start]
    inside = [r for r in rows if (not start or r.on_date >= start) and (not end or r.on_date <= end)]
    opening = round(sum(r.amount if r.kind == "received" else -r.amount for r in before), 2)
    d = lambda r: {"id": r.id, "kind": r.kind, "date": r.on_date.isoformat(), "description": r.description or "",
                   "amount": round(r.amount or 0, 2), "remarks": r.remarks or ""}
    rec = [d(r) for r in inside if r.kind == "received"]
    paid = [d(r) for r in inside if r.kind == "paid"]
    tr, tp = round(sum(r["amount"] for r in rec), 2), round(sum(r["amount"] for r in paid), 2)
    last = max((r.on_date for r in inside), default=None)
    # The same lines in date order with the running balance - the paper
    # reads as a bank statement, so it never splits into two uneven sides.
    bal, ledger = opening, []
    for r in inside:
        bal = round(bal + (r.amount if r.kind == "received" else -r.amount), 2)
        ledger.append({**d(r), "balance": bal})
    return {"from": start.isoformat() if start else "", "to": end.isoformat() if end else "",
            "opening": opening, "received": rec, "paid": paid, "ledger": ledger, "total_received": tr, "total_paid": tp,
            "balance": round(opening + tr - tp, 2), "as_on": last.isoformat() if last else ""}


@router.get("/employees/accounts/register")
def get_register(start: str = "", end: str = "", db: Session = Depends(get_db), user: models.User = RK):
    return register(db, M._as_date(start), M._as_date(end))


def _entry(p):
    kind = str(p.get("kind") or "")
    if kind not in ("received", "paid"):
        raise HTTPException(status_code=400, detail="Cash received or cash paid out?")
    d = M._as_date(p.get("date"))
    if not d:
        raise HTTPException(status_code=400, detail="Put the date.")
    desc = " ".join(str(p.get("description") or "").split())
    if not desc:
        raise HTTPException(status_code=400, detail="Type a description.")
    try:
        amt = round(float(str(p.get("amount") or 0).replace(",", "")), 2)
    except ValueError:
        raise HTTPException(status_code=400, detail="The amount must be a number.")
    if amt <= 0:
        raise HTTPException(status_code=400, detail="Type the amount.")
    return {"kind": kind, "on_date": d, "description": desc[:250], "amount": amt,
            "remarks": " ".join(str(p.get("remarks") or "").split())[:200]}


@router.post("/employees/accounts/register")
def add_entry(payload: dict = Body(...), db: Session = Depends(get_db), user: models.User = RK):
    x = CashEntry(**_entry(payload), created_by=user.id, updated_by=user.id)
    db.add(x); db.commit()
    M.log_action(db, user.id, "register_entry", f"entry #{x.id} added")      # no figures in the activity log
    return {"ok": True, "id": x.id}


@router.put("/employees/accounts/register/{eid}")
def edit_entry(eid: int, payload: dict = Body(...), db: Session = Depends(get_db), user: models.User = RK):
    x = db.get(CashEntry, eid)
    if not x:
        raise HTTPException(status_code=404, detail="That line is no longer there.")
    for k, v in _entry(payload).items():
        setattr(x, k, v)
    x.updated_by = user.id
    db.commit()
    M.log_action(db, user.id, "register_entry", f"entry #{eid} changed")
    return {"ok": True}


@router.delete("/employees/accounts/register/{eid}")
def delete_entry(eid: int, db: Session = Depends(get_db), user: models.User = RK):
    x = db.get(CashEntry, eid)
    if not x:
        raise HTTPException(status_code=404, detail="That line is no longer there.")
    db.delete(x); db.commit()
    M.log_action(db, user.id, "register_entry", f"entry #{eid} removed")
    return {"ok": True}


@router.post("/employees/accounts/register/password")
def change_password(payload: dict = Body(...), db: Session = Depends(get_db), user: models.User = RK):
    if user.role != "admin":
        raise HTTPException(status_code=403, detail="Only an admin can change the register password.")
    if not auth.verify_password(str(payload.get("current") or ""), _hash(db)):
        raise HTTPException(status_code=400, detail="The current password is wrong.")
    new = str(payload.get("new") or "")
    if len(new) < 8:
        raise HTTPException(status_code=400, detail="Use at least 8 characters.")
    M.put_setting(db, "register_password_hash", auth.hash_password(new))
    M.log_action(db, user.id, "register_password", "cash register password changed")
    return {"ok": True}


def _register_pdf(r):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT, TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    RED = colors.HexColor("#C0392B"); TINT = colors.HexColor("#F7F5F3"); GREY = colors.HexColor("#DDDDDD")
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=8.5, leading=10.5)
    bold = ParagraphStyle("bb", parent=base, fontName="Helvetica-Bold")
    rt = ParagraphStyle("r", parent=base, alignment=TA_RIGHT); rtb = ParagraphStyle("rb", parent=bold, alignment=TA_RIGHT)
    hd = ParagraphStyle("h", parent=bold, fontSize=8, textColor=colors.white)
    hdr = ParagraphStyle("hr", parent=hd, alignment=TA_RIGHT)
    muted = ParagraphStyle("m", parent=base, textColor=colors.HexColor("#6B7178"))
    P = lambda s_, st=base: Paragraph(escape(str(s_)), st)
    m = lambda v: f"{v:,.2f}" if v not in ("", None) else ""
    page = landscape(A4); W = page[0] - 24 * mm
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
                            title="Cash Register")
    logo = export_web._logo_image(42)
    period = (f"{_d(r['from']) or 'start'} to {_d(r['to']) or 'date'}" if (r["from"] or r["to"]) else "All entries")
    top = Table([[logo or P("INFINIA", bold), Paragraph("CASH REGISTER", ParagraphStyle("t", parent=bold, fontSize=15, leading=19, textColor=RED, alignment=TA_RIGHT))],
                 ["", Paragraph(escape(period), ParagraphStyle("p", parent=base, alignment=TA_RIGHT, textColor=colors.HexColor("#6B7178")))]],
                colWidths=[W / 2, W / 2])
    top.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("LINEBELOW", (0, 1), (-1, 1), 1.2, RED),
                             ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    # One ledger in date order: Date | Description | Remarks | Received | Paid out | Balance.
    data = [[P("DATE", hd), P("DESCRIPTION", hd), P("REMARKS", hd), P("RECEIVED (AED)", hdr), P("PAID OUT (AED)", hdr), P("BALANCE (AED)", hdr)]]
    data.append([P(""), P("Balance brought forward", ParagraphStyle("i", parent=muted, fontName="Helvetica-Oblique")), P(""), P(""), P(""), P(m(r["opening"]), rtb)])
    for x in r["ledger"]:
        data.append([P(_d(x["date"])), P(x["description"]), P(x["remarks"], muted),
                     P(m(x["amount"]) if x["kind"] == "received" else "", rt), P(m(x["amount"]) if x["kind"] == "paid" else "", rt),
                     P(m(x["balance"]), rtb)])
    data.append([P("TOTAL", bold), "", "", P(m(r["total_received"]), rtb), P(m(r["total_paid"]), rtb), P(m(r["balance"]), rtb)])
    data.append([P(f"Balance as on {_d(r['as_on']) or '-'}", bold), "", "", "", "", P(f"AED {r['balance']:,.2f}", rtb)])
    widths = [W * f for f in (.09, .42, .16, .11, .11, .11)]
    n = len(data)
    # The totals and the closing balance never sit alone on a page: the
    # last few lines go with them in a block that is kept together.
    cut = max(2, n - 5)
    head, tail = data[:cut], data[cut:]
    t = Table(head, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), RED), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                           ("LINEBELOW", (0, 1), (-1, -1), .3, GREY), ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#FBF6F4"))]
                          + [("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FAFAFA")) for i in range(3, cut, 2)]))
    k = len(tail)
    t2 = Table(tail, colWidths=widths)
    st = [("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, k - 3), .3, GREY),
          ("SPAN", (0, k - 2), (2, k - 2)), ("LINEABOVE", (0, k - 2), (-1, k - 2), 1, colors.black), ("BACKGROUND", (0, k - 2), (-1, k - 2), TINT),
          ("SPAN", (0, k - 1), (4, k - 1)), ("LINEABOVE", (0, k - 1), (-1, k - 1), 1, RED), ("BACKGROUND", (0, k - 1), (-1, k - 1), TINT),
          ("TOPPADDING", (0, k - 1), (-1, k - 1), 6), ("BOTTOMPADDING", (0, k - 1), (-1, k - 1), 6)]
    st += [("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FAFAFA")) for i in range(k - 2) if (cut + i) % 2 == 1]
    t2.setStyle(TableStyle(st))
    from reportlab.platypus import KeepTogether
    tail_block = KeepTogether([t2])
    doc.build([top, Spacer(1, 8), t, tail_block])
    buf.seek(0)
    return buf


def _d(iso):
    d = M._as_date(iso) if iso else None
    return f"{d:%d-%b-%y}" if d else ""


def _register_excel(r):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook(); ws = wb.active; ws.title = "Cash Register"
    fill = PatternFill("solid", fgColor="C0392B"); thin = Side(style="thin", color="DDDDDD")
    ws["A1"] = "CASH REGISTER"; ws["A1"].font = Font(bold=True, size=14, color="C0392B")
    period = (f"{_d(r['from']) or 'start'} to {_d(r['to']) or 'date'}" if (r["from"] or r["to"]) else "All entries")
    ws["F1"] = period; ws["F1"].font = Font(color="6B7178"); ws["F1"].alignment = Alignment(horizontal="right")
    ws.append(["DATE", "DESCRIPTION", "REMARKS", "RECEIVED (AED)", "PAID OUT (AED)", "BALANCE (AED)"])
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = fill
        c.alignment = Alignment(horizontal="right" if c.column >= 4 else "left", vertical="center")
    ws.append([None, "Balance brought forward", None, None, None, r["opening"]])
    first = ws.max_row
    for c in ws[first]:
        c.font = Font(italic=True, bold=(c.column == 6)); c.fill = PatternFill("solid", fgColor="FBF6F4")
    ws.cell(first, 6).number_format = "#,##0.00"
    for x in r["ledger"]:
        ws.append([M._as_date(x["date"]), x["description"], x["remarks"],
                   x["amount"] if x["kind"] == "received" else None, x["amount"] if x["kind"] == "paid" else None, None])
        row = ws.max_row
        ws.cell(row, 6, f"=F{row - 1}+D{row}-E{row}").font = Font(bold=True)        # the balance is a formula
        ws.cell(row, 1).number_format = "dd-mmm-yy"
        for col in (4, 5, 6):
            ws.cell(row, col).number_format = "#,##0.00"
        for col in range(1, 7):
            ws.cell(row, col).border = Border(bottom=thin)
    last = ws.max_row
    ws.append(["TOTAL", None, None, f"=SUM(D{first + 1}:D{last})", f"=SUM(E{first + 1}:E{last})", f"=F{first}+D{last + 1}-E{last + 1}"])
    tot = ws.max_row
    ws.merge_cells(start_row=tot, start_column=1, end_row=tot, end_column=3)
    for col in range(1, 7):
        c = ws.cell(tot, col); c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="F7F5F3")
        c.border = Border(top=Side(style="medium", color="000000"))
        if col >= 4: c.number_format = "#,##0.00"
    ws.append([f"Balance as on {_d(r['as_on']) or '-'}", None, None, None, None, f"=F{tot}"])
    ws.merge_cells(start_row=tot + 1, start_column=1, end_row=tot + 1, end_column=5)
    ws.cell(tot + 1, 1).font = Font(bold=True); ws.cell(tot + 1, 6).font = Font(bold=True)
    ws.cell(tot + 1, 6).number_format = '"AED "#,##0.00'
    for col in range(1, 7):
        ws.cell(tot + 1, col).border = Border(top=Side(style="thin", color="C0392B"))
    for col, w in zip("ABCDEF", (12, 52, 24, 16, 16, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"
    export_web.print_ready(ws, "landscape", header_row=2, title="Cash Register")
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


@router.get("/export/accounts/register")
def export_register(token: str, rk: str, format: str = "pdf", start: str = "", end: str = "", db: Session = Depends(get_db)):
    u = auth.get_download_user_from_token(token, db)
    if REG_RIGHT not in M.effective_permissions(u):
        raise HTTPException(status_code=403, detail="The cash register is for accounts only.")
    _check_key(rk, u)
    r = register(db, M._as_date(start), M._as_date(end))
    if format == "excel":
        return StreamingResponse(_register_excel(r), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": "attachment; filename=Cash_Register.xlsx"})
    return StreamingResponse(_register_pdf(r), media_type="application/pdf",
                             headers={"Content-Disposition": 'inline; filename="Cash_Register.pdf"'})
