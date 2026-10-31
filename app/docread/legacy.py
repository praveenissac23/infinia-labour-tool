"""Bridges to the existing readers.

  * to_lines(doc)        - the layout model as the browser's lines ({p, y, c, x, xe, t, ocr}), so a file read on the
                           server can also be given to the existing est_* readers (a fallback during roll-out).
  * browser_lines(data)  - what portal/est/agreement.js sends today for a file, emulated on the server (SheetJS
                           formatted cells for Excel, pdf.js-style item grouping for PDFs, Tesseract page mode 3 with
                           its own line grouping for scans). Used by the scoring harness to run the OLD readers on
                           exactly the kind of input they get in the portal.
"""
import io
import os
import re
import subprocess

from .excel import display


def to_lines(doc):
    out = []
    for reg in doc.regions:
        for r in reg.rows:
            cs = [c.text for c in r.cells]
            l = {"p": r.page, "y": round(r.y, 1) if not float(r.y).is_integer() else int(r.y), "c": cs, "t": "  ".join(x for x in cs if x)}
            if reg.kind not in ("sheet", "docx", "grid"):
                l["x"] = [round(c.x0) for c in r.cells]
                l["xe"] = [round(c.x1) for c in r.cells]
            else:
                l["sheet"] = reg.name
            if r.ocr:
                l["ocr"] = 1
            out.append(l)
    return out


# ---- the browser, emulated ------------------------------------------------------------------------------

def _xlsx_lines(data):
    """SheetJS sheet_to_json(header:1, raw:false): every row of every visible sheet (hidden rows too - SheetJS keeps
    them), formatted text, merged cells only in their first cell, empty cells as ''."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out = []
    for si, ws in enumerate(wb.worksheets):
        if ws.sheet_state != "visible":
            continue
        for row in ws.iter_rows():
            cells = [display(c.value, c.number_format) if c.value is not None else "" for c in row]
            cells = [re.sub(r"\s+", " ", x).strip() for x in cells]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                out.append({"p": si + 1, "c": cells, "t": "  ".join(x for x in cells if x), "sheet": ws.title})
    return out


def _xls_lines(data):
    import xlrd
    bk = xlrd.open_workbook(file_contents=data, formatting_info=True)
    out = []
    for si, sh in enumerate(bk.sheets()):
        if sh.visibility:
            continue
        for r in range(sh.nrows):
            cells = []
            for c in range(sh.ncols):
                x = sh.cell(r, c)
                fmt = ""
                try:
                    fmt = bk.format_map[bk.xf_list[x.xf_index].format_key].format_str
                except Exception:
                    pass
                v = x.value
                if x.ctype == 3:
                    try:
                        v = xlrd.xldate.xldate_as_datetime(x.value, bk.datemode)
                    except Exception:
                        pass
                cells.append(display(v, fmt) if x.ctype not in (0, 6) else "")
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                out.append({"p": si + 1, "c": cells, "t": "  ".join(x for x in cells if x), "sheet": sh.name})
    return out


def _pdf_lines(data):
    """pdf.js text items -> rows (|dy| <= max(2, 0.35 h)) -> cells (gap < max(4, 0.9 h)), as agreement.js does."""
    import pdfplumber
    out, scanned = [], []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for pn, page in enumerate(pdf.pages, 1):
            ws = page.extract_words(x_tolerance=1.5, y_tolerance=2.5, keep_blank_chars=True)
            if sum(len(w["text"]) for w in ws) < 25:
                scanned.append(pn)
                continue
            items = sorted(({"s": w["text"], "x": w["x0"], "y": w["bottom"], "w": w["x1"] - w["x0"], "h": w["bottom"] - w["top"]} for w in ws),
                           key=lambda i: (i["y"], i["x"]))
            rows = []
            for it in items:
                r = rows[-1] if rows else None
                if r and abs(r["y"] - it["y"]) <= max(2, it["h"] * 0.35):
                    r["items"].append(it)
                else:
                    rows.append({"y": it["y"], "items": [it]})
            for r in rows:
                r["items"].sort(key=lambda i: i["x"])
                cells, cur, end, size = [], None, -1, 0
                for it in r["items"]:
                    gap = it["x"] - end
                    if cur and gap < max(4, it["h"] * 0.9):
                        cur["t"] += (" " if gap > it["h"] * 0.15 else "") + it["s"]
                    else:
                        cur = {"t": it["s"], "x": round(it["x"])}
                        cells.append(cur)
                    end = it["x"] + it["w"]
                    cur["e"] = round(end)
                    size = max(size, it["h"])
                cells = [c for c in cells if c["t"].strip()]
                if cells:
                    out.append({"p": pn, "y": round(r["y"]), "h": round(size, 1), "c": [re.sub(r"\s+", " ", c["t"]).strip() for c in cells],
                                "x": [c["x"] for c in cells], "xe": [c["e"] for c in cells], "t": "  ".join(c["t"] for c in cells)})
    if scanned:
        from .ocr import render_pdf_pages
        imgs = render_pdf_pages(data, scanned, dpi=200)
        for pn in scanned:
            if pn in imgs:
                out += _ocr_lines(imgs[pn], pn, 200 / 72.0)
        out.sort(key=lambda l: (l["p"], l.get("y", 0)))
    return out


def _ocr_lines(img, pn, scale):
    """Tesseract (page mode 3, as the browser runs it) and agreement.js ocrLines grouping."""
    from PIL import Image
    if img.width < 1600:                     # the browser enlarges small pictures ("up")
        k = 2
        img = img.resize((img.width * k, img.height * k), Image.LANCZOS)
        scale *= k
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    p = subprocess.run(["tesseract", "stdin", "stdout", "--psm", "3", "-c", "preserve_interword_spaces=1", "tsv"], input=buf.getvalue(),
                       capture_output=True, env=dict(os.environ, OMP_THREAD_LIMIT="1"), timeout=120)
    lines = {}
    for ln in p.stdout.decode("utf-8", "replace").splitlines()[1:]:
        f = ln.split("\t")
        if len(f) < 12 or f[0] != "5" or not f[11].strip():
            continue
        key = (f[2], f[3], f[4])
        lines.setdefault(key, []).append((f[11].strip(), int(f[6]), int(f[7]), int(f[6]) + int(f[8]), int(f[7]) + int(f[9]), float(f[10])))
    out = []
    for ws in lines.values():
        ws = [w for w in ws if w[5] > 15]
        if not ws:
            continue
        y0 = min(w[2] for w in ws)
        h = (max(w[4] for w in ws) - y0) / scale
        cells, cur, end = [], None, -1
        cut = False
        for t, x0, _, x1, _, _ in sorted(ws, key=lambda w: w[1]):
            x0s, x1s = x0 / scale, x1 / scale
            if re.fullmatch(r"[|\[\]{}~=_—–\\/]+", t) and t != "/":
                cut = True
                continue
            t = t.strip("|[]")
            if not t:
                cut = True
                continue
            if cur and not cut and x0s - end < max(6, h * 1.1):
                cur["t"] += " " + t
            else:
                cur = {"t": t, "x": round(x0s)}
                cells.append(cur)
            cut = False
            end = x1s
            cur["e"] = round(x1s)
        out.append({"p": pn, "y": round(y0 / scale), "h": round(min(h, 14), 1), "c": [c["t"] for c in cells], "x": [c["x"] for c in cells],
                    "xe": [c["e"] for c in cells], "t": "  ".join(c["t"] for c in cells), "ocr": 1})
    out.sort(key=lambda l: l["y"])
    return out


def _docx_lines(data):
    from docx import Document
    d = Document(io.BytesIO(data))
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    out = []
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            t = re.sub(r"\s+", " ", Paragraph(el, d).text).strip()
            if t:
                out.append({"p": 1, "c": [t], "t": t})
        elif tag == "tbl":
            for tr in Table(el, d).rows:
                cells, seen = [], set()
                for tc in tr.cells:
                    if id(tc._tc) in seen:
                        continue
                    seen.add(id(tc._tc))
                    cells.append(re.sub(r"\s+", " ", tc.text).strip())
                if any(cells):
                    out.append({"p": 1, "c": cells, "t": "  ".join(c for c in cells if c), "tbl": 1})
    return out


def browser_lines(data, name=""):
    from .load import sniff
    k = sniff(data, name)
    if k == "pdf":
        return _pdf_lines(data)
    if k == "xls":
        return _xls_lines(data)
    if k in ("xlsx", "xlsm"):
        return _xlsx_lines(data)
    if k == "docx":
        return _docx_lines(data)
    if k == "image":
        from PIL import Image, ImageOps
        img = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("L")
        return _ocr_lines(img, 1, 2.0)
    raise ValueError(k)
