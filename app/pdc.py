"""PDC tracker - Expiry Reminder > PDCs.

Post-dated cheques the company has given out (and any expected payment
worth seeing beside them, such as the month's salaries): who, the
cheque, the date it can be presented, the amount. The accountant's own
sheet is a month-by-month grid - payees down the side, months across,
totals both ways - and that is the main view and the printed paper.

Only admin and a login holding the "pdc" right (the chief accountant)
can see any of it: the page, the figures, the papers, the bell line and
the badge. A cheque is reminded 14 and 7 days before its date, on the
day, and every day after while it is still not marked cleared.
"""
import io
from datetime import date, timedelta
from html import escape
from urllib.parse import urlencode

from fastapi import APIRouter, Depends, HTTPException, Body
from fastapi.responses import HTMLResponse, StreamingResponse
from sqlalchemy.orm import Session

import main as M
import models, auth, export_web
from database import get_db

router = APIRouter()
RIGHT = "pdc"
PDC = Depends(M.require_screen(RIGHT))
COMPANY = "INFINIA CONTRACTING L.L.C."
RED = "#B7322A"
REMIND = (14, 7)          # days before the cheque date


def _may(user):
    if RIGHT not in M.effective_permissions(user):
        raise HTTPException(status_code=403, detail="The PDC tracker is for admin and the chief accountant only.")


def _money(v):
    return f"{v:,.2f}"


def _month(s, default=None):
    try:
        y, m = (int(x) for x in str(s).split("-")[:2])
        return date(y, m, 1)
    except Exception:
        return default or M._dubai_today().replace(day=1)


def _add_months(d, n):
    y, m = divmod(d.month - 1 + n, 12)
    return date(d.year + y, m + 1, 1)


def _same_day(d, n):
    """The same day n months on - the 31st becomes the month's last day."""
    first = _add_months(d.replace(day=1), n)
    last = (_add_months(first, 1) - timedelta(days=1)).day
    return first.replace(day=min(d.day, last))


def _row(x, today, users):
    left = None if x.date_tbc else (x.cheque_date - today).days
    return {"id": x.id, "payee": x.payee, "cheque_no": x.cheque_no or "", "bank": x.bank or "",
            "date": x.cheque_date.isoformat(), "date_tbc": bool(x.date_tbc), "month": f"{x.cheque_date:%b-%y}", "amount": round(x.amount or 0, 2), "notes": x.notes or "",
            "status": x.status or "pending", "cleared_on": x.cleared_on.isoformat() if x.cleared_on else "",
            "days_left": left, "by": users.get(x.updated_by or x.created_by, "")}


def summary(db):
    today = M._dubai_today()
    pend = db.query(models.Pdc).filter(models.Pdc.status == "pending").all()
    nodate = [x for x in pend if x.date_tbc]
    days = [((x.cheque_date - today).days, x) for x in pend if not x.date_tbc]
    overdue = [x for d, x in days if d < 0]
    week = [x for d, x in days if 0 <= d <= 7]
    fortnight = [x for d, x in days if 8 <= d <= 14]
    s = lambda xs: round(sum(x.amount or 0 for x in xs), 2)
    return {"overdue": len(overdue), "overdue_amount": s(overdue), "week": len(week), "week_amount": s(week),
            "fortnight": len(fortnight), "fortnight_amount": s(fortnight),
            "pending": len(pend), "pending_amount": s(pend),
            "no_date": len(nodate), "no_date_amount": s(nodate),
            "milestone": any(d < 0 or d == 0 or d in REMIND for d, _ in days),
            "soon": sorted(overdue + week + fortnight, key=lambda x: x.cheque_date)}


@router.get("/employees/pdc")
def list_pdc(db: Session = Depends(get_db), user: models.User = PDC):
    today = M._dubai_today()
    users = {u.id: (u.full_name or u.username) for u in db.query(models.User).all()}
    rows = db.query(models.Pdc).order_by(models.Pdc.cheque_date, models.Pdc.id).all()
    s = summary(db)
    s.pop("soon")
    return {"rows": [_row(x, today, users) for x in rows], "summary": s,
            "payees": sorted({x.payee for x in rows} | {p.name for p in db.query(models.Supplier).filter(models.Supplier.active == True).all()}),  # noqa: E712
            "banks": sorted({x.bank for x in rows if x.bank})}


def _clean(p):
    payee = " ".join(str(p.get("payee") or "").split())
    if not payee:
        raise HTTPException(status_code=400, detail="Who is the cheque for? Type the payee.")
    d = M._as_date(p.get("date"))
    if not d and p.get("_keep_date"):
        d = p["_keep_date"]
    if not d:
        raise HTTPException(status_code=400, detail="Put the cheque date.")
    try:
        amt = round(float(str(p.get("amount") or 0).replace(",", "")), 2)
    except ValueError:
        raise HTTPException(status_code=400, detail="The amount must be a number.")
    if amt <= 0:
        raise HTTPException(status_code=400, detail="Type the amount.")
    st = str(p.get("status") or "pending").lower()
    if st not in ("pending", "cleared", "cancelled"):
        st = "pending"
    return {"payee": payee[:120], "cheque_no": str(p.get("cheque_no") or "").strip()[:40],
            "bank": str(p.get("bank") or "").strip()[:60], "cheque_date": d, "amount": amt,
            "notes": str(p.get("notes") or "").strip()[:200], "status": st}


@router.post("/employees/pdc")
def add_pdc(payload: dict = Body(...), db: Session = Depends(get_db), user: models.User = PDC):
    """One cheque - or a run of monthly cheques (rent, instalments,
    salaries): repeat = how many months, the cheque number counting up
    when it ends in digits."""
    base = _clean(payload)
    try:
        n = max(1, min(36, int(payload.get("repeat") or 1)))
    except ValueError:
        n = 1
    no = base["cheque_no"]
    digits = len(no) - len(no.rstrip("0123456789"))
    made = []
    for i in range(n):
        x = dict(base)
        x["cheque_date"] = _same_day(base["cheque_date"], i)
        if i and digits:
            x["cheque_no"] = no[:-digits] + str(int(no[-digits:]) + i).zfill(digits)
        if x["status"] == "cleared":
            x["cleared_on"] = M._dubai_today()
        r = models.Pdc(**x, created_by=user.id, updated_by=user.id)
        db.add(r); made.append(r)
    db.commit()
    M.log_action(db, user.id, "pdc_add", f"{base['payee']} {_money(base['amount'])}"
                 + (f" x {n} months from {base['cheque_date']}" if n > 1 else f" dated {base['cheque_date']}"))
    return {"ok": True, "ids": [r.id for r in made]}


@router.put("/employees/pdc/{pid}")
def edit_pdc(pid: int, payload: dict = Body(...), db: Session = Depends(get_db), user: models.User = PDC):
    x = db.get(models.Pdc, pid)
    if not x:
        raise HTTPException(status_code=404, detail="That cheque is no longer there.")
    was = x.status
    # A cheque still waiting for its date may be saved without one; putting
    # a date in settles it.
    tbc = bool(x.date_tbc) and not M._as_date(payload.get("date"))
    for k, v in _clean({**payload, "_keep_date": x.cheque_date if tbc else None}).items():
        setattr(x, k, v)
    x.date_tbc = tbc
    if x.status == "cleared" and was != "cleared":
        x.cleared_on = M._dubai_today()
    if x.status != "cleared":
        x.cleared_on = None
    x.updated_by = user.id
    db.commit()
    M.log_action(db, user.id, "pdc_edit", f"#{pid} {x.payee} {_money(x.amount)} {x.cheque_date} {x.status}")
    return {"ok": True}


@router.post("/employees/pdc/{pid}/clear")
def clear_pdc(pid: int, db: Session = Depends(get_db), user: models.User = PDC):
    """Cleared - or back to pending if pressed by mistake."""
    x = db.get(models.Pdc, pid)
    if not x:
        raise HTTPException(status_code=404, detail="That cheque is no longer there.")
    if x.status == "cleared":
        x.status, x.cleared_on = "pending", None
    else:
        x.status, x.cleared_on = "cleared", M._dubai_today()
    x.updated_by = user.id
    db.commit()
    M.log_action(db, user.id, "pdc_clear", f"#{pid} {x.payee} {_money(x.amount)} -> {x.status}")
    return {"ok": True, "status": x.status}


@router.delete("/employees/pdc/{pid}")
def delete_pdc(pid: int, db: Session = Depends(get_db), user: models.User = PDC):
    x = db.get(models.Pdc, pid)
    if not x:
        raise HTTPException(status_code=404, detail="That cheque is no longer there.")
    M.log_action(db, user.id, "pdc_delete", f"#{pid} {x.payee} {_money(x.amount)} {x.cheque_date}")
    db.delete(x); db.commit()
    return {"ok": True}


# ---- The month-by-month grid (the accountant's sheet) ----------------------

def tracker(db, start="", months=4, show="due"):
    """Payees down the side, months across. show='due': only what is
    still to be paid; 'all': paid and to be paid. Anything still pending
    from before the first month is gathered in an Overdue column."""
    a = _month(start)
    months = max(1, min(12, int(months or 4)))
    cols = [_add_months(a, i) for i in range(months)]
    end = _add_months(a, months)
    q = db.query(models.Pdc).filter(models.Pdc.status != "cancelled")
    if show != "all":
        q = q.filter(models.Pdc.status == "pending")
    rows_in = q.order_by(models.Pdc.cheque_date).all()
    over = [x for x in rows_in if x.cheque_date < a and x.status == "pending"]
    inside = [x for x in rows_in if a <= x.cheque_date < end]
    payees = {}
    for x in over + inside:
        p = payees.setdefault(x.payee, {"payee": x.payee, "notes": [], "overdue": 0.0,
                                        "cells": [0.0] * months, "cheques": [],
                                        # the cheques behind each figure, so the exact date shows under it
                                        "overdue_dates": [], "dates": [[] for _ in range(months)]})
        if x.notes and x.notes not in p["notes"]:
            p["notes"].append(x.notes)
        p["cheques"].append({"date": x.cheque_date.isoformat(), "no": x.cheque_no or "", "amount": x.amount, "status": x.status})
        q = {"id": x.id, "date": x.cheque_date.isoformat(), "label": "date to fill" if x.date_tbc else f"{x.cheque_date:%d-%b-%y}",
             "tbc": bool(x.date_tbc), "no": x.cheque_no or "",
             "amount": round(x.amount or 0, 2), "cleared": x.status == "cleared"}
        if x.cheque_date < a:
            p["overdue"] += x.amount or 0
            p["overdue_dates"].append(q)
        else:
            i = (x.cheque_date.year - a.year) * 12 + x.cheque_date.month - a.month
            p["cells"][i] += x.amount or 0
            p["dates"][i].append(q)
    rows = sorted(payees.values(), key=lambda r: (min(c["date"] for c in r["cheques"]), r["payee"].lower()))
    for r in rows:
        r["overdue"] = round(r["overdue"], 2)
        r["cells"] = [round(v, 2) for v in r["cells"]]
        r["total"] = round(r["overdue"] + sum(r["cells"]), 2)
        r["notes"] = "; ".join(r["notes"])
    tot = [round(sum(r["cells"][i] for r in rows), 2) for i in range(months)]
    od = round(sum(r["overdue"] for r in rows), 2)
    return {"from": f"{a:%Y-%m}", "months": months, "show": show,
            "columns": [{"key": f"{c:%Y-%m}", "label": f"{c:%b-%y}"} for c in cols],
            "overdue": od, "rows": rows, "totals": tot, "grand": round(od + sum(tot), 2),
            "label": f"{cols[0]:%B %Y}" + (f" to {cols[-1]:%B %Y}" if months > 1 else "")}


@router.get("/employees/pdc/tracker")
def get_tracker(start: str = "", months: int = 4, show: str = "due", db: Session = Depends(get_db), user: models.User = PDC):
    return tracker(db, start, months, show)


def _head(t):
    return (["Payee"] + (["Overdue"] if t["overdue"] else []) + [c["label"] for c in t["columns"]] + ["Total"])


def _vals(t, r):
    return ([r["overdue"]] if t["overdue"] else []) + r["cells"] + [r["total"]]


def _dates(t, r):
    """The cheques under each figure, same order as _vals (less Total)."""
    return ([r["overdue_dates"]] if t["overdue"] else []) + r["dates"]


def _when(qs):
    """'05-Oct-26 #000451' - one line per cheque; with its amount when
    one figure is made of several cheques."""
    out = []
    for q in qs:
        bit = q["label"] + (f" #{q['no']}" if q["no"] else "") + (" (cleared)" if q["cleared"] else "")
        if len(qs) > 1:
            bit += f" - {_money(q['amount'])}"
        out.append(bit)
    return out


def _html(t, pdf_url, excel_url):
    logo = export_web.logo_data_uri()
    m = lambda v: _money(v) if v else ""
    head = "".join(f"<th>{escape(h)}</th>" for h in _head(t))
    def line(r):
        note = f"<span>{escape(r['notes'])}</span>" if r["notes"] else ""
        vals, ds = _vals(t, r), _dates(t, r)
        cells = ""
        for i, v in enumerate(vals[:-1]):
            od = " od" if (t["overdue"] and i == 0 and v) else ""
            when = "".join(f"<small>{escape(w)}</small>" for w in _when(ds[i]))
            cells += f"<td class='n{od}'>{m(v)}{when}</td>"
        return f"<tr><td class='l'><b>{escape(r['payee'])}</b>{note}</td>{cells}<td class='n b'>{_money(vals[-1])}</td></tr>"
    body = "".join(line(r) for r in t["rows"])
    foot = ("<tr><td class='l'>TOTAL (Monthly)</td>" + "".join(f"<td class='n'>{_money(v)}</td>" for v in
            (([t["overdue"]] if t["overdue"] else []) + t["totals"] + [t["grand"]])) + "</tr>")
    return f"""<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=820, minimum-scale=0.3, maximum-scale=4, user-scalable=yes">
<title>PDC Tracker - {escape(t['label'])}</title><style>
*{{box-sizing:border-box}} body{{margin:0;background:#ECEEF1;font:12.5px/1.4 Arial,Helvetica,sans-serif;color:#1d1d1d}}
.bar{{position:sticky;top:0;background:#fff;border-bottom:1px solid #ddd;padding:10px 16px;display:flex;gap:8px;justify-content:flex-end}}
.bar a,.bar button{{font:600 13px Arial;padding:8px 14px;border-radius:6px;border:1px solid #ddd;background:#fff;color:#222;text-decoration:none;cursor:pointer}}
.bar .p{{background:{RED};border-color:{RED};color:#fff}}
.page{{max-width:1100px;margin:18px auto;background:#fff;padding:30px 36px;box-shadow:0 2px 10px rgba(0,0,0,.08)}}
.top{{display:flex;justify-content:space-between;align-items:flex-end;border-bottom:2px solid {RED};padding-bottom:10px}}
.top img{{height:30px}} h1{{margin:0;font-size:18px;letter-spacing:.08em;color:{RED};text-align:right}}
.meta{{display:flex;justify-content:space-between;gap:20px;margin:12px 0;font-size:12.5px}} .meta span{{color:#777}} .meta b{{margin-left:6px}}
.wrap{{overflow-x:auto}} table{{width:100%;border-collapse:collapse;min-width:640px}}
th{{background:{RED};color:#fff;font-size:11px;letter-spacing:.04em;padding:7px 8px;text-align:right}} th:first-child{{text-align:left}}
td{{padding:6px 8px;border-bottom:1px solid #e3e3e3}} td.l span{{display:block;font-size:11px;color:#777}}
td.n{{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap;vertical-align:top}} td.n small{{display:block;font-size:10.5px;color:#777;font-weight:normal}} td.b{{font-weight:700}} td.od{{color:{RED};font-weight:600}}
tbody tr:nth-child(even) td{{background:#FAFAFA}}
tfoot td{{font-weight:700;background:#F3F1EF;border-top:1.5px solid #222;border-bottom:none}}
.sign{{display:grid;grid-template-columns:repeat(3,1fr);gap:40px;margin-top:56px}} .sign div{{border-top:1px solid #222;padding-top:6px;text-align:center;font-size:12px}}
@media print{{.bar{{display:none}} body{{background:#fff}} .page{{box-shadow:none;margin:0;max-width:none;padding:0}}}}
@media(max-width:640px){{.page{{padding:16px;margin:0}}}}
</style></head><body>
{export_web.preview_bar("PDC Tracker", t['label'], pdf_url, excel_url)}
<div class="page">
 <div class="top"><div>{f'<img src="{logo}" alt="">' if logo else '<b>INFINIA</b>'}</div><h1>PDC TRACKER</h1></div>
 <div class="meta"><div><span>Company</span><b>{COMPANY}</b></div><div><span>Period</span><b>{escape(t['label'])}</b></div>
  <div><span>Showing</span><b>{'To be paid' if t['show'] != 'all' else 'Paid and to be paid'}</b></div></div>
 <div class="wrap"><table><thead><tr>{head}</tr></thead><tbody>{body or '<tr><td colspan="9" style="color:#999">No cheques in this period.</td></tr>'}</tbody><tfoot>{foot}</tfoot></table></div>
 <div class="sign"><div>Prepared By</div><div>Checked By</div><div>Approved By</div></div>
</div></body></html>"""


def _pdf(t):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT, TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, KeepTogether
    R = colors.HexColor(RED)
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=8.5, leading=10.5)
    bold = ParagraphStyle("bb", parent=base, fontName="Helvetica-Bold")
    small = ParagraphStyle("s", parent=base, fontSize=7, leading=8.5, textColor=colors.HexColor("#777777"))
    smr = ParagraphStyle("sr", parent=small, alignment=TA_RIGHT)
    rt = ParagraphStyle("r", parent=base, alignment=TA_RIGHT)
    rtb = ParagraphStyle("rb", parent=bold, alignment=TA_RIGHT)
    hd = ParagraphStyle("h", parent=bold, fontSize=8, textColor=colors.white, alignment=TA_RIGHT)
    hdl = ParagraphStyle("hl", parent=hd, alignment=0)
    cen = ParagraphStyle("c", parent=base, alignment=TA_CENTER)
    P = lambda s, st=base: Paragraph(escape(str(s)), st)
    page = landscape(A4)
    W = page[0] - 24 * mm
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
                            title=f"PDC Tracker - {t['label']}")
    logo = export_web._logo_image(42)
    top = Table([[logo or P("INFINIA", bold), Paragraph("PDC TRACKER", ParagraphStyle("t", parent=bold, fontSize=15, leading=19, textColor=R, alignment=TA_RIGHT))]],
                colWidths=[W / 2, W / 2])
    top.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("LINEBELOW", (0, 0), (-1, 0), 1.5, R),
                             ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    meta = Table([[Paragraph(f"<font color='#777777'>Company</font>&nbsp;&nbsp;<b>{COMPANY}</b>", base),
                   Paragraph(f"<font color='#777777'>Period</font>&nbsp;&nbsp;<b>{escape(t['label'])}</b>", rt)]], colWidths=[W / 2, W / 2])
    meta.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 6)]))
    heads = _head(t)
    m = lambda v: _money(v) if v else ""
    data = [[P(heads[0].upper(), hdl)] + [P(h.upper(), hd) for h in heads[1:]]]
    for r in t["rows"]:
        name = [P(r["payee"], bold)] + ([P(r["notes"], small)] if r["notes"] else [])
        vals, ds = _vals(t, r), _dates(t, r)
        cells = []
        for i, v in enumerate(vals[:-1]):
            cells.append([P(m(v), rt)] + [Paragraph(escape(w), smr) for w in _when(ds[i])] if v else "")
        data.append([name] + cells + [P(_money(vals[-1]), rtb)])
    if not t["rows"]:
        data.append([P("No cheques in this period.")] + [""] * (len(heads) - 1))
    data.append([P("TOTAL (Monthly)", bold)] + [P(_money(v), rtb) for v in (([t["overdue"]] if t["overdue"] else []) + t["totals"] + [t["grand"]])])
    n = len(heads) - 1
    first = W * .34
    widths = [first] + [(W - first) / n] * n
    tb = Table(data, colWidths=widths, repeatRows=1)
    st = [("BACKGROUND", (0, 0), (-1, 0), R), ("VALIGN", (0, 0), (-1, -1), "TOP"),
          ("LINEBELOW", (0, 1), (-1, -2), .3, colors.HexColor("#DDDDDD")),
          ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F3F1EF")), ("LINEABOVE", (0, -1), (-1, -1), 1, colors.black)]
    for i in range(2, len(data) - 1, 2):
        st.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FAFAFA")))
    if t["overdue"]:
        st.append(("TEXTCOLOR", (1, 1), (1, -2), R))
    tb.setStyle(TableStyle(st))
    sign = Table([[P("Prepared By", cen), "", P("Checked By", cen), "", P("Approved By", cen)]],
                 colWidths=[W * .26, W * .11, W * .26, W * .11, W * .26])
    sign.setStyle(TableStyle([("LINEABOVE", (0, 0), (0, 0), .8, colors.black), ("LINEABOVE", (2, 0), (2, 0), .8, colors.black),
                              ("LINEABOVE", (4, 0), (4, 0), .8, colors.black)]))
    doc.build([top, meta, Spacer(1, 6), tb, KeepTogether([Spacer(1, 40), sign])])
    buf.seek(0)
    return buf


def _excel(t):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as L
    wb = Workbook(); ws = wb.active; ws.title = "PDC Tracker"
    R = RED.lstrip("#")
    heads = _head(t)
    n = len(heads)
    ws.column_dimensions["A"].width = 44
    for i in range(2, n + 1):
        ws.column_dimensions[L(i)].width = 16
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws["A1"] = "PDC TRACKER"; ws["A1"].font = Font(bold=True, size=15, color=R); ws["A1"].alignment = Alignment(horizontal="right")
    ws.row_dimensions[1].height = 30
    ws["A2"] = f"{COMPANY}   -   {t['label']}"; ws["A2"].font = Font(bold=True)
    for i, h in enumerate(heads, 1):
        c = ws.cell(4, i, h.upper()); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=R)
        c.alignment = Alignment(horizontal="left" if i == 1 else "right")
    row = 5
    thin = Side(style="thin", color="DDDDDD")
    for r in t["rows"]:
        ws.cell(row, 1, r["payee"] + (f"  ({r['notes']})" if r["notes"] else ""))
        vals = _vals(t, r)[:-1]
        for i, v in enumerate(vals, 2):
            if v:
                ws.cell(row, i, v).number_format = "#,##0.00"
        ws.cell(row, n, f"=SUM({L(2)}{row}:{L(n - 1)}{row})").number_format = "#,##0.00"
        ws.cell(row, n).font = Font(bold=True)
        # The exact cheque dates on a line of their own under the figures
        # (text, so the totals' SUMs pass over it).
        row += 1
        ds = _dates(t, r)
        ws.cell(row, 1, "   cheque date").font = Font(italic=True, size=9, color="777777")
        for i, qs in enumerate(ds, 2):
            if qs:
                c = ws.cell(row, i, "\n".join(_when(qs)))
                c.font = Font(italic=True, size=9, color="777777")
                c.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 13 * max([1] + [len(q) for q in ds])
        for i in range(1, n + 1):
            ws.cell(row, i).border = Border(bottom=thin)
        row += 1
    ws.cell(row, 1, "TOTAL (Monthly)").font = Font(bold=True)
    for i in range(2, n + 1):
        c = ws.cell(row, i, f"=SUM({L(i)}5:{L(i)}{row - 1})" if row > 5 else 0)
        c.number_format = "#,##0.00"; c.font = Font(bold=True)
    for i in range(1, n + 1):
        ws.cell(row, i).fill = PatternFill("solid", fgColor="F3F1EF"); ws.cell(row, i).border = Border(top=Side(style="medium"))
    ws.freeze_panes = "B5"
    export_web.print_ready(ws, "landscape", header_row=4, title="PDC Tracker")
    # The cheques themselves, one line each, on a second sheet.
    w2 = wb.create_sheet("Cheques")
    for i, (h, wd) in enumerate((("Cheque date", 13), ("Payee", 36), ("Cheque no", 14), ("Amount", 15), ("Status", 11), ("Notes", 34)), 1):
        c = w2.cell(1, i, h.upper()); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=R)
        w2.column_dimensions[L(i)].width = wd
    rr = 2
    for r in t["rows"]:
        for q in sorted(r["cheques"], key=lambda q: q["date"]):
            w2.cell(rr, 1, M._as_date(q["date"])).number_format = "dd-mmm-yy"
            w2.cell(rr, 2, r["payee"]); w2.cell(rr, 3, q["no"])
            w2.cell(rr, 4, q["amount"]).number_format = "#,##0.00"; w2.cell(rr, 5, q["status"].title()); w2.cell(rr, 6, r["notes"])
            rr += 1
    w2.freeze_panes = "A2"
    export_web.print_ready(w2, "portrait", header_row=1, title="PDC Cheques")
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


@router.get("/export/pdc")
def export_pdc(token: str, start: str = "", months: int = 4, show: str = "due", format: str = "pdf", db: Session = Depends(get_db)):
    _may(auth.get_download_user_from_token(token, db))
    t = tracker(db, start, months, show)
    name = f"PDC_Tracker_{t['from']}"
    if format == "excel":
        return StreamingResponse(_excel(t), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f"attachment; filename={name}.xlsx"})
    return StreamingResponse(_pdf(t), media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={name}.pdf"})


@router.get("/export/pdc/view", response_class=HTMLResponse)
def view_pdc(token: str, start: str = "", months: int = 4, show: str = "due", db: Session = Depends(get_db)):
    user = auth.get_download_user_from_token(token, db)
    _may(user)
    t = tracker(db, start, months, show)
    v = auth.create_view_token(user.username)
    base = "/export/pdc?" + urlencode({"start": t["from"], "months": t["months"], "show": show, "token": v})
    return HTMLResponse(_html(t, base + "&format=pdf", base + "&format=excel"))


def notification(db, today):
    """The bell line - for admin and the pdc right only (the caller
    checks). It pops up on a reminder day (14 or 7 days before, the day
    itself, or overdue); otherwise it sits quietly in the bell."""
    s = summary(db)
    if not (s["overdue"] or s["week"] or s["fortnight"] or s["no_date"]):
        return None
    bits = [f"{s['overdue']} overdue" if s["overdue"] else "",
            f"{s['week']} due within 7 days" if s["week"] else "",
            f"{s['fortnight']} within 14 days" if s["fortnight"] else "",
            f"{s['no_date']} with no date yet" if s["no_date"] else ""]
    first = ", ".join(f"{x.payee} {_money(x.amount)} on {x.cheque_date:%d %b}" for x in s["soon"][:3]) \
        or "Put the cheque dates in so the reminders can work"
    more = len(s["soon"]) - 3
    return {"id": f"pdc-{today}" if s["milestone"] else f"pdc-wk{today.isocalendar()[1]}", "kind": "pdc",
            "title": "PDCs: " + ", ".join(b for b in bits if b),
            "detail": first + (f" and {more} more" if more > 0 else ""),
            "screen": "pdc", "when": today.isoformat(), "count": s["overdue"] + s["week"],
            "level": "warn" if s["overdue"] else "info"}
